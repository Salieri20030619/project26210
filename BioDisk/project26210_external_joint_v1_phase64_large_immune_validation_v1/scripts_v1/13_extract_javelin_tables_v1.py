#!/usr/bin/env python3
import argparse
import csv
import gzip
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value).replace("\t", " ").replace("\r", " ").replace("\n", " ")


def read_signature(path):
    rows = []
    with open(path, encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle, delimiter="\t"):
            if row["signature_id"] == "f1_top100_weighted":
                rows.append(row)
    if len(rows) != 100:
        raise RuntimeError(f"Expected 100 frozen F1 genes, observed {len(rows)}")
    return rows


def write_sheet(sheet, output_path, header_row=2, first_data_row=3, compressed=False):
    opener = gzip.open if compressed else open
    mode = "wt" if compressed else "w"
    with opener(output_path, mode, encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        header = [clean(cell.value) for cell in sheet[header_row]]
        writer.writerow(header)
        count = 0
        for row in sheet.iter_rows(min_row=first_data_row, values_only=True):
            writer.writerow([clean(value) for value in row])
            count += 1
    return header, count


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("signature")
    parser.add_argument("output_dir")
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    signature = read_signature(args.signature)
    target_genes = {row["gene"] for row in signature} | {"TYMP"}

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)

    clinical_sheet = workbook["S11_Clinical_data"]
    clinical_header, clinical_rows = write_sheet(
        clinical_sheet, output_dir / "javelin_clinical_v1.tsv", header_row=2, first_data_row=3
    )

    expression_sheet = workbook["S13_Gene_expression_TPM"]
    expression_header = [clean(cell.value) for cell in expression_sheet[2]]
    found = {}
    expression_output = output_dir / "javelin_f1_tymp_expression_v1.tsv.gz"
    with gzip.open(expression_output, "wt", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(expression_header)
        for row in expression_sheet.iter_rows(min_row=3, values_only=True):
            gene = clean(row[0])
            if gene in target_genes:
                writer.writerow([clean(value) for value in row])
                found[gene] = found.get(gene, 0) + 1

    deconv_header, deconv_rows = write_sheet(
        workbook["S14_ImmuneNet_decon_data"],
        output_dir / "javelin_immunenet_deconvolution_v1.tsv.gz",
        header_row=2,
        first_data_row=3,
        compressed=True,
    )

    geneset_sheet = workbook["S15_Pathway_genesets"]
    geneset_rows = 0
    with (output_dir / "javelin_pathway_genesets_long_v1.tsv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["pathway", "collection", "gene_order", "gene"])
        for row in geneset_sheet.iter_rows(min_row=2, values_only=True):
            pathway = clean(row[0])
            collection = clean(row[1])
            for gene_order, value in enumerate(row[2:], start=1):
                gene = clean(value)
                if gene:
                    writer.writerow([pathway, collection, gene_order, gene])
                    geneset_rows += 1

    pathway_header, pathway_rows = write_sheet(
        workbook["S16_Pathway_scores"],
        output_dir / "javelin_pathway_scores_v1.tsv.gz",
        header_row=2,
        first_data_row=3,
        compressed=True,
    )

    clinical_ids = set()
    with (output_dir / "javelin_clinical_v1.tsv").open(encoding="utf-8", newline="") as handle:
        clinical_data = list(csv.DictReader(handle, delimiter="\t"))
        clinical_ids = {row["ID"] for row in clinical_data}
    expression_ids = set(expression_header[1:])
    pathway_ids = set()
    with gzip.open(output_dir / "javelin_pathway_scores_v1.tsv.gz", "rt", encoding="utf-8", newline="") as handle:
        pathway_data = list(csv.DictReader(handle, delimiter="\t"))
        pathway_ids = {row["row_id"] for row in pathway_data}

    coverage_path = output_dir / "javelin_f1_gene_coverage_v1.tsv"
    with coverage_path.open("w", encoding="utf-8", newline="") as handle:
        fields = ["gene_order", "gene", "raw_weight", "present", "row_count"]
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in signature:
            writer.writerow(
                {
                    "gene_order": row["gene_order"],
                    "gene": row["gene"],
                    "raw_weight": row["raw_weight"],
                    "present": str(row["gene"] in found).upper(),
                    "row_count": found.get(row["gene"], 0),
                }
            )

    pfs_complete = sum(row.get("PFS_P", "") not in ("", "NA") for row in clinical_data)
    arm_complete = sum(row.get("TRT01P", "") not in ("", "NA") for row in clinical_data)
    gates = [
        ("clinical_rows", clinical_rows, 886, clinical_rows == 886),
        ("clinical_unique_ids", len(clinical_ids), 886, len(clinical_ids) == 886),
        ("expression_samples", len(expression_ids), 726, len(expression_ids) == 726),
        ("expression_ids_in_clinical", len(expression_ids & clinical_ids), len(expression_ids), expression_ids <= clinical_ids),
        ("pathway_samples", pathway_rows, 726, pathway_rows == 726),
        ("pathway_ids_match_expression", len(pathway_ids & expression_ids), 726, pathway_ids == expression_ids),
        ("deconvolution_rows", deconv_rows, 727, deconv_rows == 727),
        ("f1_genes_present", sum(gene in found for gene in target_genes if gene != "TYMP"), 70,
         sum(gene in found for gene in target_genes if gene != "TYMP") >= 70),
        ("tymp_present", int("TYMP" in found), 1, "TYMP" in found),
        ("clinical_pfs_complete", pfs_complete, 886, pfs_complete == 886),
        ("clinical_arm_complete", arm_complete, 886, arm_complete == 886),
        ("pathway_count", len(pathway_header) - 1, 58, len(pathway_header) - 1 == 58),
        ("geneset_memberships", geneset_rows, 1, geneset_rows > 0),
    ]
    with (output_dir / "javelin_data_gate_v1.tsv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["gate", "observed", "expected_or_threshold", "status"])
        for gate, observed, expected, passed in gates:
            writer.writerow([gate, observed, expected, "PASS" if passed else "FAIL"])

    print(f"clinical_rows={clinical_rows}")
    print(f"expression_samples={len(expression_ids)}")
    print(f"pathway_samples={pathway_rows}")
    print(f"deconvolution_rows={deconv_rows}")
    print(f"f1_genes_present={sum(row['gene'] in found for row in signature)}")
    print(f"tymp_present={'TYMP' in found}")
    print(f"pathways={len(pathway_header) - 1}")
    print(f"all_gates_pass={all(item[3] for item in gates)}")


if __name__ == "__main__":
    main()
