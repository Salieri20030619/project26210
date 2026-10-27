#!/usr/bin/env python3
import csv
import gzip
import os
import sys
from collections import Counter

from openpyxl import load_workbook


NA_VALUES = {None, "", "NA", "N/A", "na", "NaN"}


def clean(value):
    if value in NA_VALUES:
        return ""
    return str(value).strip()


def write_tsv(path, header, rows, gzip_output=False):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    opener = gzip.open if gzip_output else open
    kwargs = {"encoding": "utf-8", "newline": ""}
    if gzip_output:
        kwargs["compresslevel"] = 6
    with opener(path, "wt" if gzip_output else "w", **kwargs) as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(header)
        for row in rows:
            writer.writerow([clean(value) for value in row])


def extract_sheet(workbook, sheet_name, header_row, output_path):
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(iterator)
    header = [clean(value) for value in next(iterator)]
    rows = []
    for row in iterator:
        values = list(row[: len(header)])
        if not any(clean(value) for value in values):
            continue
        rows.append(values)
    write_tsv(output_path, header, rows, gzip_output=output_path.endswith(".gz"))
    return header, rows


def counter_text(values):
    counts = Counter(value if value else "MISSING" for value in values)
    return ";".join(f"{key}:{counts[key]}" for key in sorted(counts))


def main(s1_path, s4_path, signature_path, output_dir):
    os.makedirs(output_dir, exist_ok=True)

    s1 = load_workbook(s1_path, read_only=True, data_only=True)
    clinical_header, clinical_rows = extract_sheet(
        s1, "Clinical_and_Immune_Data", 2,
        os.path.join(output_dir, "checkmate_clinical_all_v1.tsv.gz"),
    )
    if len(clinical_header) != len(set(clinical_header)):
        raise SystemExit("Clinical workbook has duplicate column names")
    clinical = [dict(zip(clinical_header, (clean(value) for value in row))) for row in clinical_rows]

    s4 = load_workbook(s4_path, read_only=True, data_only=True)
    rna_header, rna_rows = extract_sheet(
        s4, "RNA_Expression", 2,
        os.path.join(output_dir, "checkmate_rna_expression_v1.tsv.gz"),
    )
    ssgsea_header, ssgsea_rows = extract_sheet(
        s4, "ssGSEA", 2,
        os.path.join(output_dir, "checkmate_ssgsea_v1.tsv.gz"),
    )
    cibersort_header, cibersort_rows = extract_sheet(
        s4, "CIBERSORTx", 2,
        os.path.join(output_dir, "checkmate_cibersortx_v1.tsv.gz"),
    )

    rna_samples = rna_header[1:]
    ssgsea_samples = ssgsea_header[2:]
    cibersort_samples = cibersort_header[1:]
    genes = [clean(row[0]) for row in rna_rows]
    pathways = [clean(row[1]) for row in ssgsea_rows]
    immune_cell_types = [clean(row[0]) for row in cibersort_rows]

    rna_clinical = [row for row in clinical if row.get("RNA_ID", "") in set(rna_samples)]
    write_tsv(
        os.path.join(output_dir, "checkmate_clinical_rna_subset_v1.tsv"),
        clinical_header,
        ([row.get(column, "") for column in clinical_header] for row in rna_clinical),
    )

    with open(signature_path, encoding="utf-8", newline="") as handle:
        signature_rows = list(csv.DictReader(handle, delimiter="\t"))
    f1_genes = [row["gene"] for row in signature_rows if row["signature_id"] == "f1_top100_weighted"]
    f1_present = sorted(set(f1_genes) & set(genes))
    f1_missing = sorted(set(f1_genes) - set(genes))
    write_tsv(
        os.path.join(output_dir, "checkmate_f1_gene_coverage_v1.tsv"),
        ["gene", "present"],
        ((gene, int(gene in set(genes))) for gene in f1_genes),
    )

    clinical_rna_ids = [row.get("RNA_ID", "") for row in rna_clinical]
    nivolumab_rows = [row for row in rna_clinical if "NIVOLUMAB" in row.get("Arm", "").upper()]
    everolimus_rows = [row for row in rna_clinical if "EVEROLIMUS" in row.get("Arm", "").upper()]
    t_reg_names = [name for name in immune_cell_types if "reg" in name.lower() or "treg" in name.lower()]

    gates = [
        ("clinical_rows", len(clinical), 1000, len(clinical) >= 1000),
        ("rna_samples", len(rna_samples), 311, len(rna_samples) == 311),
        ("rna_sample_ids_unique", len(set(rna_samples)), 311, len(set(rna_samples)) == 311),
        ("rna_clinical_rows", len(rna_clinical), 311, len(rna_clinical) == 311),
        ("rna_ids_mapped", len(set(clinical_rna_ids) & set(rna_samples)), 311, len(set(clinical_rna_ids) & set(rna_samples)) == 311),
        ("rna_clinical_ids_unique", len(set(clinical_rna_ids)), 311, len(set(clinical_rna_ids)) == 311),
        ("rna_genes", len(genes), 43893, len(genes) == 43893),
        ("rna_gene_names_unique", len(set(genes)), len(genes), len(set(genes)) == len(genes)),
        ("f1_signature_genes", len(f1_genes), 100, len(f1_genes) == 100),
        ("f1_genes_present", len(f1_present), 70, len(f1_present) >= 70),
        ("tymp_present", int("TYMP" in set(genes)), 1, "TYMP" in set(genes)),
        ("ssgsea_samples_match", len(set(ssgsea_samples) ^ set(rna_samples)), 0, set(ssgsea_samples) == set(rna_samples)),
        ("cibersort_samples_match", len(set(cibersort_samples) ^ set(rna_samples)), 0, set(cibersort_samples) == set(rna_samples)),
        ("cibersort_treg_feature", len(t_reg_names), 1, len(t_reg_names) >= 1),
        ("nivolumab_rna_patients", len(nivolumab_rows), 150, len(nivolumab_rows) >= 150),
        ("everolimus_rna_patients", len(everolimus_rows), 100, len(everolimus_rows) >= 100),
    ]
    gate_rows = [
        {"gate": gate, "observed": observed, "expected_or_threshold": expected, "status": "PASS" if passed else "FAIL"}
        for gate, observed, expected, passed in gates
    ]
    gate_rows.extend([
        {"gate": "cohort_distribution", "observed": counter_text(row.get("Cohort", "") for row in rna_clinical), "expected_or_threshold": "descriptive", "status": "INFO"},
        {"gate": "arm_distribution", "observed": counter_text(row.get("Arm", "") for row in rna_clinical), "expected_or_threshold": "descriptive", "status": "INFO"},
        {"gate": "orr_distribution_nivolumab", "observed": counter_text(row.get("ORR", "") for row in nivolumab_rows), "expected_or_threshold": "descriptive", "status": "INFO"},
        {"gate": "benefit_distribution_nivolumab", "observed": counter_text(row.get("Benefit", "") for row in nivolumab_rows), "expected_or_threshold": "descriptive", "status": "INFO"},
        {"gate": "ssgsea_pathways", "observed": len(pathways), "expected_or_threshold": "descriptive", "status": "INFO"},
        {"gate": "cibersort_cell_types", "observed": len(immune_cell_types), "expected_or_threshold": "descriptive", "status": "INFO"},
        {"gate": "cibersort_treg_names", "observed": ";".join(t_reg_names), "expected_or_threshold": "descriptive", "status": "INFO"},
        {"gate": "f1_missing_genes", "observed": ";".join(f1_missing), "expected_or_threshold": "descriptive", "status": "INFO"},
    ])
    write_tsv(
        os.path.join(output_dir, "checkmate_data_gate_v1.tsv"),
        ["gate", "observed", "expected_or_threshold", "status"],
        ([row["gate"], row["observed"], row["expected_or_threshold"], row["status"]] for row in gate_rows),
    )
    for row in gate_rows:
        print(f'{row["gate"]}\t{row["observed"]}\t{row["expected_or_threshold"]}\t{row["status"]}')
    if any(row["status"] == "FAIL" for row in gate_rows):
        raise SystemExit(2)


if __name__ == "__main__":
    main(*sys.argv[1:5])
