#!/usr/bin/env python3
import argparse
import csv
import json
from pathlib import Path

from openpyxl import load_workbook


def compact(value, limit=180):
    if value is None:
        return ""
    text = str(value).replace("\t", " ").replace("\r", " ").replace("\n", " ")
    return text if len(text) <= limit else text[: limit - 3] + "..."


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("output_dir")
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)

    inventory = []
    previews = {}
    for index, sheet in enumerate(workbook.worksheets, start=1):
        inventory.append(
            {
                "sheet_index": index,
                "sheet_name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
            }
        )
        if 11 <= index <= 16 or any(f"{number}" in sheet.title for number in range(11, 17)):
            rows = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                rows.append([compact(value) for value in row[: min(sheet.max_column, 30)]])
                if row_index >= 12:
                    break
            previews[sheet.title] = rows

    with (output_dir / "javelin_workbook_inventory_v1.tsv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(inventory[0]), delimiter="\t")
        writer.writeheader()
        writer.writerows(inventory)
    with (output_dir / "javelin_workbook_previews_v1.json").open("w", encoding="utf-8") as handle:
        json.dump(previews, handle, ensure_ascii=False, indent=2)

    print(f"sheets={len(inventory)}")
    for row in inventory:
        print(f"{row['sheet_index']:02d}\t{row['sheet_name']}\t{row['max_row']}x{row['max_column']}")
    print("PREVIEW_KEYS=" + ",".join(previews))


if __name__ == "__main__":
    main()
