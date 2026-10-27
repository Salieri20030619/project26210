#!/usr/bin/env python3
import sys
from openpyxl import load_workbook


def render(value):
    if value is None:
        return ""
    text = str(value).replace("\t", " ").replace("\n", " ")
    return text[:160]


def main(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    print(f"WORKBOOK\t{path}")
    print("SHEETS\t" + "\t".join(workbook.sheetnames))
    for sheet in workbook.worksheets:
        print(f"SHEET\t{sheet.title}\tROWS\t{sheet.max_row}\tCOLS\t{sheet.max_column}")
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            print(f"ROW\t{sheet.title}\t{row_index}\t" + "\t".join(render(value) for value in row[:40]))
            if row_index >= 8:
                break


if __name__ == "__main__":
    main(sys.argv[1])
