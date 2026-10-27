from __future__ import annotations

import argparse
import csv
import gzip
import json
from pathlib import Path

from openpyxl import load_workbook


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value).replace("\t", " ").replace("\r", " ").replace("\n", " ")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("signature_members")
    parser.add_argument("output_dir")
    args = parser.parse_args()

    workbook_path = Path(args.workbook)
    members_path = Path(args.signature_members)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    expression_output = output_dir / "javelin_r2_core_expression_v1.tsv.gz"
    coverage_output = output_dir / "javelin_r2_core_gene_coverage_v1.tsv"
    status_output = output_dir / "javelin_r2_extraction_status_v1.json"
    for output in [expression_output, coverage_output, status_output]:
        if output.exists():
            raise FileExistsError(f"Refusing to overwrite existing output: {output}")

    with members_path.open(encoding="utf-8-sig", newline="") as handle:
        members = list(csv.DictReader(handle, delimiter="\t"))
    target_genes = sorted({row["gene"].upper() for row in members})
    if len(target_genes) != 1770 or "TYMP" in target_genes:
        raise RuntimeError(f"Unexpected frozen target set: n={len(target_genes)} TYMP={'TYMP' in target_genes}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["S13_Gene_expression_TPM"]
    header = [clean(cell.value) for cell in sheet[2]]
    sample_ids = header[1:]
    found_counts: dict[str, int] = {}
    written_rows = 0
    with gzip.open(expression_output, "wt", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(header)
        for values in sheet.iter_rows(min_row=3, values_only=True):
            gene = clean(values[0]).upper()
            if gene in target_genes:
                writer.writerow([gene, *[clean(value) for value in values[1:]]])
                found_counts[gene] = found_counts.get(gene, 0) + 1
                written_rows += 1
    workbook.close()

    with coverage_output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["gene", "present", "row_count"])
        for gene in target_genes:
            writer.writerow([gene, str(gene in found_counts).upper(), found_counts.get(gene, 0)])

    status = {
        "status": "PASS" if len(sample_ids) == 726 and len(found_counts) >= 1400 else "FAIL",
        "workbook": str(workbook_path),
        "frozen_unique_genes": len(target_genes),
        "expression_samples": len(sample_ids),
        "present_unique_genes": len(found_counts),
        "coverage_fraction": len(found_counts) / len(target_genes),
        "expression_rows_written": written_rows,
        "duplicated_target_gene_rows": sum(count - 1 for count in found_counts.values()),
        "tymp_extracted": "TYMP" in found_counts,
    }
    status_output.write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(status, ensure_ascii=False, indent=2))
    if status["status"] != "PASS" or status["tymp_extracted"]:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
