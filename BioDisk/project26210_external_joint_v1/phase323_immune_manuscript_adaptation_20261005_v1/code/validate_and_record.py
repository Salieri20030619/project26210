from pathlib import Path
import json,hashlib,csv,sys,re
from datetime import datetime,timezone,timedelta
P=Path(__file__).resolve().parents[1];R=P.parent
sys.path.insert(0,str(R/'phase322_integrated_immune_report_20261005_v1/code/vendor'))
from docx import Document
blocks=json.loads((P/'inputs/DOCUMENT_CONTENT.json').read_text())
out=next((P/'delivery').glob('*.docx'));doc=Document(out)
tables=[b for b in blocks if b['type']=='table']
alltext='\n'.join(p.text for p in doc.paragraphs)
checks={
'all_text_present':all(b['text'] in alltext for b in blocks if b['type'] in ['title','h','p']),
'seven_tables_exact':len(doc.tables)==7 and all([[c.text for c in row.cells] for row in t.rows]==[b['headers']]+[[str(v) for v in row] for row in b['rows']] for t,b in zip(doc.tables,tables)),
'source_hashes_unchanged':all(hashlib.sha256(Path(x['path']).read_bytes()).hexdigest()==x['sha256'] for x in json.loads((P/'inputs/SOURCE_MANIFEST.json').read_text())),
'ten_rendered_pages':len(list((P/'render/v1').glob('page-*.png')))==10,
'render_exit_zero':(P/'logs/render.exit').read_text().strip()=='0',
'repeating_table_headers':all(t.rows[0]._tr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}trPr/{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader') is not None for t in doc.tables)
}
assert all(checks.values()),checks
now=datetime.now(timezone(timedelta(hours=8))).isoformat(timespec='seconds')
qc={'status':'PASS','time':now,'pages':10,'tables':7,'checks':checks,'visual_review':'All 10 rendered pages individually inspected; readable Chinese and symbols, no clipped content, blank pages, orphan headings or split numeric values; continued table repeats header.','no_new_models':True,'full_manuscript_unchanged':True,'delivery_type':'immune chapter prepared for integration, not merged full manuscript'}
(P/'qc/COMPLETE.json').write_text(json.dumps(qc,ensure_ascii=False,indent=2))
runtime={'dependency_loader':'Unavailable in this tool session; no bundled runtime exposed by available tools or package search','python':sys.executable,'packages':str(R/'phase322_integrated_immune_report_20261005_v1/code/vendor'),'renderer':'Existing unchanged packaged render_docx.py','libreoffice':'Existing /usr/bin/libreoffice on sgs10 headless, no Windows desktop process','operation_marker':'Packaged document edit marker ran successfully once before authoring using available local Node; no bundled Node available'}
(P/'qc/DOCUMENT_RUNTIME.json').write_text(json.dumps(runtime,indent=2))
delivery=[{'name':f.name,'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in sorted((P/'delivery').iterdir()) if f.is_file()]
(P/'qc/DELIVERY_MANIFEST.json').write_text(json.dumps(delivery,ensure_ascii=False,indent=2))
record=f"""## Phase323（2026-10-05）：免疫章节适配完整报告，背景与方法已优化
用户纠正恢复入口：免疫报告的用途是适配并入完整报告，重点优化背景和方法；本轮按此完成章节稿，不将Phase322独立报告视为该任务终点。
以Phase273正式全文和Phase322修订v2为源，新增承接TYMP相关代谢、MTA、SPP1富集巨噬状态的三段背景；方法展开数据与独立样本、宽松/严格成员、归一化和等权评分、资格、患者秩残差及精确符号检验、相互控制与删基因、跨细胞/组织/空间/临床分析及分别BH家族。
去除独立摘要和独立结论，保留6项结果、7张表及主要阴性；结果段原值与七表逐项核对不变，无新模型。标题层级不预占全文编号，表免疫1–7为临时标签；衔接说明明确与已有临床/空间内容去重位置。
交付：{P}/delivery/26210_免疫章节_全文适配_v3.docx；另有全文衔接说明.txt。10页逐页图像核验及内容/来源检查PASS，证据qc/COMPLETE.json。
本次完成的是可并入全文的章节稿，尚未实际合并或覆盖45页完整报告；完整主稿仍Phase273。旧Phase322及v2保留。当前断点为本章节审阅及后续全文整合，不追加生物学计算或下载。
"""
for name in ['CODEX_HANDOFF.md','PROJECT_STATUS.md']:
 f=R/name;f.write_text(record+'\n'+f.read_text())
with (R/'DECISIONS.md').open('a') as f:f.write('\n'+record)
manifest=R/'RESULTS_MANIFEST.tsv'
with manifest.open() as f:header=next(csv.reader(f,delimiter='\t'))
print('MANIFEST_SCHEMA',header)
artifacts=[('phase323_chapter_docx','immune_manuscript_chapter',out),('phase323_integration_guide','integration_instructions',P/'delivery/全文衔接说明.txt'),('phase323_validation','content_and_render_validation',P/'qc/COMPLETE.json'),('phase323_source_manifest','input_provenance',P/'inputs/SOURCE_MANIFEST.json')]
expected=['result_id','stage','artifact_type','absolute_path','generated_at','validation_status','validation_evidence','notes']
assert header==expected,header
with manifest.open('a',newline='') as f:
 w=csv.writer(f,delimiter='\t',lineterminator='\n')
 for rid,kind,path in artifacts:w.writerow([rid,'phase323_immune_manuscript_adaptation',kind,str(path),now,'PASS',str(P/'qc/COMPLETE.json'),'Background/method adaptation; all original table values retained; not merged into full manuscript'])
print(json.dumps({'status':'PASS','pages':10,'delivery':delivery},ensure_ascii=False))
