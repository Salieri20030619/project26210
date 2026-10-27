from __future__ import annotations

import hashlib
import importlib.util
import json
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent
OUTPUT = ROOT / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_PROCESS_DETAIL_REVIEW_v7.docx"
PART_OUTPUT = ROOT / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_PROCESS_DETAIL_REVIEW_v7.part.docx"
STATUS = ROOT / "build_status_v2.json"
PART_STATUS = ROOT / "build_status_v2.part.json"

source = PROJECT / "phase41_part1_client_report_narrative_local_v1" / "03_build_part1_narrative_client_report_v1.py"
spec = importlib.util.spec_from_file_location("phase41_report", source)
base = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(base)


PROCESS = {
    "B01": "首先逐队列核对原始对象、文库名称、患者编号和组织来源，在每个数据集内部统一细胞标识，再建立跨队列的患者—文库—细胞对应表。原始计数仍按数据集分别保存，不把五个研究直接拼成一个未经区分的生物学批次；跨队列比较只使用经过核对的公共元数据和统一分析字段。",
    "B02": "各队列先独立完成细胞和基因层面的质量控制，随后在共同基因空间中进行scVI整合，以降低研究来源造成的表达偏移。细胞类型依据整合后的表达结构和经典标志基因分层标注，并在每个细胞上保留原始队列、患者、文库和组织来源，供后续患者层面统计使用。",
    "M01": "恶性细胞分析先从肿瘤来源样本中提取统一注释为上皮细胞的细胞，排除正常组织上皮和所有免疫、内皮及基质细胞。这个步骤只确定需要接受CNV检验的候选池，不把“来自肿瘤组织”或“属于上皮谱系”直接等同于恶性。",
    "M02": "候选上皮细胞分别进入inferCNV和CopyKAT。正常参考按“同患者正常上皮—同研究正常上皮池—同患者肿瘤淋巴细胞”的顺序选择；inferCNV先根据染色体位置平滑表达并形成HMM CNV分组，只有分组的CNV基因负荷高于参考细胞中位数加3倍稳健离差与参考95分位数两者中的较高阈值，才记为高CNV。CopyKAT独立保留非整倍体、二倍体和未定义状态。最终按预设真值表合并两套结果：高CNV且非整倍体列为双方法一致恶性；高CNV而另一方法不可用时可列为inferCNV单方法支持；高CNV/二倍体或低CNV/非整倍体均保留为方法不一致，不纳入保守恶性集合。",
    "M03": "在每名合格患者内部读取inferCNV HMM分组的染色体区段状态，先把具有相同或近似CNV轮廓的细胞汇总为主要克隆节点，再按节点间CNV距离连接为确定性树。根节点仅选择CNV负荷较低的展示锚点；随后把每个节点中占优势的cNMF程序映射到树上，用于比较克隆结构和转录状态，而不把连边直接解释为真实时间谱系。",
    "M04": "以保守恶性集合的原始表达为输入，选取2,500个高变基因，在cNMF秩4至8之间分别重复分解。秩的选择同时考虑重复运行的一致性、程序稳定性和细胞使用度分布，确定秩7后再计算每个细胞的F1—F7使用度；F1与TYMP的关系先在患者内计算，再汇总患者方向并进行留一队列敏感性复核。",
    "M05": "为避免把单细胞数量当作生物学重复，每名合格患者按程序状态汇总为伪批量表达，共形成28个伪批量样本。各程序分别进行差异表达，再把排序基因表用于功能富集，由富集通路、代表基因和程序间差异共同确定F1—F7的功能标签。",
    "M06": "从GWAS Catalog整理肾癌易感位点及其映射基因，构建scDRS基因集并为每个恶性细胞计算遗传易感表达分数。随后在每名患者内部计算F1使用度与scDRS分数的Spearman相关，患者相关系数作为独立观测，以符号检验评估方向一致性，并用Bootstrap给出中位效应区间。",
    "M07": "稳健性复核把总体效应拆回队列和患者：逐队列查看患者相关方向，检查细胞层面的scDRS显著性是否经多重校正保留，并比较单队列结果与全体患者汇总是否一致。只有跨队列方向能够复现的信号才会提升总体证据等级。",
    "M08": "在同一批11,950个恶性细胞上分别以GENIE3和GRNBoost2推断转录因子—靶基因关系。两套初始网络各自经过cisTarget基序富集筛选，保留具有基序支持的regulon，再用AUCell计算每个细胞的regulon活性；两条算法链独立运行，最后才按调控因子名称对齐比较。",
    "M09": "对两套pySCENIC活性矩阵分别估计三类患者内效应：F1连续值与regulon活性的相关、F1高低组差异，以及最高与最低四分位差异。每个端点先在患者内计算，再汇总患者方向并校正多重检验；连续、分组、四分位和双算法复现构成4项证据，4/4列为高置信，3/4保留为次级支持，2/4仅作探索。",
    "M10": "metaVIPER使用独立的调控网络资源，把恶性细胞表达转换为5,566个调控因子的蛋白活性，11,921个输入完整的细胞进入计算。修复后不按矩阵行号直接比较，而是先按调控因子键和细胞键重排，再在20个患者检查点逐项核对旧、新活性值，从而区分“顺序错位”与“数值被改变”。",
    "M11": "先把8个MR候选与metaVIPER可推断蛋白取交集，得到CASP9、CHMP1A、PRSS8和PXN。对每名患者提取这些蛋白的活性效应，并与对应MR效应方向配对计算相关；整体相关和20名患者各自的相关方向同时汇总，用来判断遗传效应与单细胞蛋白活性是否同向。",
    "M12": "每个候选蛋白均在患者内部接受三种互补检验：与连续F1的相关、F1高低组差异和最高—最低四分位差异。患者内效应转换为方向后进行患者层面符号检验，并在同一候选家族内校正多重比较；患者少于10名时，只有达到名义显著的结果才报告统计支持，否则只保留方向描述。",
    "M13": "终末共同激活先用Monocle3、Slingshot和scTour各自定义患者内轨迹末端窗口，再比较候选蛋白、TYMP和RRM2B在末端与非末端的变化。由于TYMP和RRM2B缺少合格肾癌regulon，补充验证不把RNA表达冒充VIPER活性，而是分别构建严格酶轴和包含通路、scFEA模块的广义功能轴，并按连续F1、分组、四分位和轨迹四类证据进行4/4—3/4—2/4分级。",
    "M14": "在每个恶性细胞的标准化表达上，对4条预设代谢通路计算AUCell富集分数；同时把表达矩阵输入scFEA，在代谢反应图上估计168个模块的相对通量。两类评分按细胞键回填到同一恶性细胞对象，并检查细胞覆盖和缺失情况，形成后续酶轴比较的统一代谢矩阵。",
    "M15": "代谢程序沿用与调控分析相同的患者内框架，分别估计代谢轴与连续F1的关系、F1高低组差异和最高—最低四分位差异，再加入多轨迹末端变化作为第四项证据。横断面和轨迹端点分开计算，避免用细胞数放大显著性。",
    "M16": "TYMP和RRM2B分别建立两层指标：严格轴仅保留与该酶直接对应的表达或功能代理，广义轴则纳入相邻的嘧啶代谢、RNR相关通路和scFEA模块。四类患者内端点使用相同方向和显著性规则分别评级，因此两种酶可以得到同向、反向或不支持的独立结论。",
    "M17": "将TYMP、TK1、ODC1、SRM、SMS、PCYT2等12个候选基因放入同一矩阵，逐一计算其与连续F1的患者内关系以及四分位和分组差异。基因对关系则在每名患者内部计算相关，再汇总患者方向并进行BH校正，以筛选跨患者可复现的核苷转运和多胺代谢协同。",
    "M18": "Monocle3、Slingshot、CytoTRACE2、scTour和CellRank分别从同一恶性细胞集合独立重建状态变化：前两者基于低维结构拟合伪时间，CytoTRACE2估计分化潜能，scTour学习连续轨迹表示，CellRank估计终末状态概率。各方法的方向在患者内锚定并转换到可比较尺度，不把不同算法输出直接混成一条伪时间。",
    "M19": "对Monocle3、Slingshot和scTour，在每名患者内分别划定终末窗口，并比较终末与其余细胞的候选基因和蛋白变化；CytoTRACE2和CellRank用于检查状态方向与终末概率是否相容。最终按方法统计同向患者数，并判断信号能否跨轨迹复现。",
    "C01": "在14名细胞数满足要求的患者中分别运行CellChat，不先把患者细胞混合。每名患者均以恶性细胞、内皮细胞和巨噬细胞作为发送或接收群体，估计配体—受体通信概率，再比较F1/TYMP相关恶性状态下各发送—接收方向的变化；患者内差值汇总后进行方向检验和FDR校正。",
    "C02": "对CellChat和后续NicheNet涉及的受体逐名患者检查目标接收细胞中的表达比例，并对ITGAL/ITGB2、ITGAM/ITGB2等复合体同时要求组成亚基具备表达基础。稀疏或缺失受体不以零值强行补齐，而是记录可评估患者数，用来限定后续通讯结果的解释范围。",
    "C03": "NicheNet把患者内恶性细胞的响应基因作为目标集，同时结合候选发送细胞中的配体表达和既有配体—靶基因先验网络，为每个患者计算配体活性。随后以患者为单位汇总活性方向并进行多重校正，筛选能够重复解释恶性细胞响应的上游信号。",
    "C04": "受体关联以连续TYMP和连续F1为主要端点，在每名患者内估计受体特征随恶性状态的变化，再对患者效应进行汇总和FDR校正。F1高低组只作为小样本敏感性分析；当可评估患者少于10名时，名义P值达到显著可报告，否则仅描述方向。",
    "C05": "从TCGA-KIRC提取510例患者的RNA表达和OS、PFS随访信息，各候选与TYMP均使用预设中位数切点划分双低、单高和双高三组，不搜索最优阈值。先用全局生存检验比较三组，再以双低为参照拟合Cox模型；多变量模型统一调整年龄、性别和分期，并同时报告OS与PFS。",
}


def detailed_analysis(doc, code, title, why, method, result, conclusion=None):
    base.add_reference_heading(doc, f"{code}  {title}", 2)
    process = PROCESS.get(code, method)
    first = base.add_reference_body(doc, process, keep=True)
    first.paragraph_format.keep_together = True
    result_text = result.strip()
    if conclusion:
        result_text += " " + conclusion.strip()
    second = base.add_reference_body(doc, result_text)
    second.paragraph_format.keep_together = True


def update_cover(doc):
    note = (
        "客户正文审阅稿：已纳入全部26项分析，正文补充了样本进入、算法处理、患者层面汇总和判定规则；"
        "本版仅放分析流程图，数据结果图待正文确认后统一编排。"
    )
    for paragraph in doc.paragraphs:
        if "客户正文审阅稿" in paragraph.text:
            base.clear_runs(paragraph)
            base.report.h.set_run(paragraph.add_run("版本说明  "), bold=True, color=base.report.h.TEAL)
            base.report.h.set_run(paragraph.add_run(note), color="000000")
            return


def keep_appendix_together(doc):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "附录  分析完成清单":
            paragraph.paragraph_format.page_break_before = True
            return


def build():
    for path in (OUTPUT, PART_OUTPUT, STATUS, PART_STATUS):
        if path.exists():
            raise FileExistsError(f"Refusing to overwrite existing file: {path}")

    base.ROOT = ROOT
    base.OUTPUT = OUTPUT
    base.PART_OUTPUT = PART_OUTPUT
    base.STATUS = STATUS
    base.PART_STATUS = PART_STATUS
    base.ASSETS = PROJECT / "phase41_part1_client_report_narrative_local_v1" / "assets_v1"
    base.FLOWCHART = base.ASSETS / "part1_workflow_narrative_v2.png"
    base.narrative_analysis = detailed_analysis
    base.build()

    doc = Document(OUTPUT)
    update_cover(doc)
    keep_appendix_together(doc)
    doc.save(OUTPUT)

    status = json.loads(STATUS.read_text(encoding="utf-8"))
    status.update(
        {
            "status": "PASS",
            "editorial_revision": "expanded process narrative without repeated result opener",
            "process_overrides": len(PROCESS),
            "analysis_shows_phrase_count": sum(p.text.count("分析显示") for p in doc.paragraphs),
            "output": str(OUTPUT),
            "size_bytes": OUTPUT.stat().st_size,
            "sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        }
    )
    STATUS.write_text(json.dumps(status, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(status, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
