from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_PROCESS_DETAIL_REVIEW_v7.docx"
RENDER = ROOT / "rendered_v2"
STATUS = ROOT / "validation_status_v3.json"

correct_cohorts = ["GSE156632", "GSE159115", "GSE178481", "GSE207493", "GSE242299"]
excluded_cohorts = ["GSE202374", "GSE264586", "GSE294109"]
analysis_ids = ["B01", "B02"] + [f"M{i:02d}" for i in range(1, 20)] + [f"C{i:02d}" for i in range(1, 6)]
forbidden = ["为什么做", "怎么做", "得到什么", "本项结论", "目的、方法和结果", "分析显示"]

doc = Document(DOCX)
paragraph_text = "\n".join(p.text for p in doc.paragraphs)
all_text = paragraph_text
for table in doc.tables:
    for row in table.rows:
        all_text += "\n" + "\t".join(cell.text for cell in row.cells)

checks = []


def add(name, passed, detail):
    checks.append({"check": name, "pass": bool(passed), "detail": detail})


for cohort in correct_cohorts:
    add(f"correct_cohort_{cohort}", cohort in all_text, "present" if cohort in all_text else "missing")
for cohort in excluded_cohorts:
    add(f"excluded_cohort_{cohort}", cohort not in all_text, "absent" if cohort not in all_text else "unexpectedly present")
for analysis_id in analysis_ids:
    add(f"analysis_{analysis_id}", analysis_id in all_text, "present" if analysis_id in all_text else "missing")
for phrase in forbidden:
    add(f"forbidden_{phrase}", phrase not in all_text, "absent" if phrase not in all_text else "present")

process_markers = [
    "同患者正常上皮—同研究正常上皮池—同患者肿瘤淋巴细胞",
    "中位数加3倍稳健离差",
    "参考95分位数",
    "高CNV且非整倍体",
    "方法不一致",
    "患者相关系数作为独立观测",
    "4/4列为高置信",
    "不搜索最优阈值",
]
for marker in process_markers:
    add(f"process_{marker}", marker in all_text, "present" if marker in all_text else "missing")

analysis_headings = []
for i, paragraph in enumerate(doc.paragraphs):
    token = paragraph.text.split("  ", 1)[0].strip()
    if token in analysis_ids:
        analysis_headings.append((token, i))
        next_text = doc.paragraphs[i + 1].text.strip() if i + 1 < len(doc.paragraphs) else ""
        add(f"process_length_{token}", len(next_text) >= 65, f"{len(next_text)} chars")

add("analysis_heading_count", len(analysis_headings) == 26, str(len(analysis_headings)))
add("workflow_caption", "图1  第一部分总体分析流程" in all_text, "process flowchart caption")
add("statistical_policy", "患者是独立统计单位" in all_text and "患者数3–9" in all_text, "patient-level and small-n policy")
add("enzyme_axis", "广义TYMP/嘧啶轴" in all_text and "广义RNR相关轴" in all_text, "positive and reverse enzyme axes")
add("result_figure_policy", "未插入数据结果图" in all_text, "text-review edition")

pngs = sorted(RENDER.glob("page-*.png"))
add("render_status", (RENDER / "render_status.json").exists(), "render status exists")
add("rendered_pages", len(pngs) > 0, f"{len(pngs)} PNG pages")
add("pdf_exists", (RENDER / f"{DOCX.stem}.pdf").exists(), "rendered PDF")

with zipfile.ZipFile(DOCX) as zf:
    media = [name for name in zf.namelist() if name.startswith("word/media/")]
add("single_process_figure", len(media) == 1, f"{len(media)} embedded media files")

status = "PASS" if all(item["pass"] for item in checks) else "FAIL"
result = {
    "status": status,
    "docx": str(DOCX),
    "sha256": hashlib.sha256(DOCX.read_bytes()).hexdigest(),
    "analysis_count": len(analysis_headings),
    "rendered_pages": len(pngs),
    "checks_total": len(checks),
    "checks_passed": sum(item["pass"] for item in checks),
    "checks": checks,
}
STATUS.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
raise SystemExit(0 if status == "PASS" else 1)
