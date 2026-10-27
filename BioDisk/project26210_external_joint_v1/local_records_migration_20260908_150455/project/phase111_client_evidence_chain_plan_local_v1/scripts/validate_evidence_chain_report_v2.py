from pathlib import Path
import hashlib, json, re
from PIL import Image, ImageChops
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(r'C:\Users\31937\project26210_phase111_client_next_stage_plan_v1')
DOCX = ROOT / 'output_v2' / '项目26210_双路线证据与补充分析计划_v2.docx'
RENDER = ROOT / 'render_v9_word'
PDF = RENDER / '项目26210_双路线证据与补充分析计划_v2.pdf'
OUT = ROOT / 'output_v2' / 'REPORT_VALIDATION_v2.json'

def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as fh:
        for chunk in iter(lambda: fh.read(1048576), b''):
            h.update(chunk)
    return h.hexdigest().upper()

def main():
    checks = []
    def add(name, ok, evidence):
        checks.append({'check': name, 'status': 'PASS' if ok else 'FAIL', 'evidence': evidence})

    add('docx_exists', DOCX.exists() and DOCX.stat().st_size > 0, str(DOCX))
    add('pdf_exists', PDF.exists() and PDF.stat().st_size > 0, str(PDF))
    status = json.loads((RENDER / 'render_status.json').read_text(encoding='utf-8'))
    add('word_render', status.get('status') == 'PASS', status)
    add('page_count', status.get('document_pages') == 10 and status.get('rendered_pngs') == 10, status)

    pngs = sorted(RENDER.glob('page-*.png'))
    metrics, dims = [], []
    for path in pngs:
        im = Image.open(path).convert('RGB')
        dims.append(im.size)
        diff = ImageChops.difference(im, Image.new('RGB', im.size, 'white')).convert('L')
        bbox = diff.point(lambda x: 255 if x > 12 else 0).getbbox()
        if bbox is None:
            metrics.append({'page': path.name, 'blank': True})
        else:
            l, t, r, b = bbox
            metrics.append({'page': path.name, 'blank': False, 'bbox': bbox, 'margins': [l, t, im.width-r, im.height-b]})
    add('png_count', len(pngs) == 10, len(pngs))
    add('png_dimensions', len(set(dims)) == 1, dims)
    add('no_blank_pages', all(not x['blank'] for x in metrics), metrics)
    add('content_inside_page', all(x['blank'] or min(x['margins']) >= 8 for x in metrics), metrics)

    doc = Document(DOCX)
    paragraph_text = [p.text for p in doc.paragraphs]
    table_text = [p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
    text = '\n'.join(paragraph_text + table_text)
    required = ['路线A：F1/TYMP → MTA → SPP1 → 空间结构', '路线B：F1/TYMP → MTA → SPP1 → 预后与治疗', 'UnitedMet 255个肿瘤样本', '14名患者中14/14方向为正', '0/7通过FDR', '现在可以明确说', '现在不能说']
    missing = [x for x in required if x not in text]
    add('required_evidence', not missing, missing)
    forbidden = re.findall(r'Phase\s*\d+', text, flags=re.I)
    add('no_internal_phase_labels', not forbidden, forbidden)
    tokens = [x for x in ['TODO', 'TBD', '[[', ']]', ':codex-file-citation'] if x in text]
    add('no_placeholders', not tokens, tokens)

    levels = [int(p.style.name.split()[-1]) for p in doc.paragraphs if p.style and p.style.name.startswith('Heading ')]
    add('heading_hierarchy', bool(levels) and all(b <= a + 1 for a, b in zip(levels, levels[1:])), levels)
    add('table_count', len(doc.tables) == 7, len(doc.tables))
    flags = [t.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is not None for t in doc.tables]
    add('table_headers', all(flags), flags)
    drawings = doc._element.body.findall('.//' + qn('w:drawing'))
    descrs = [x.get('descr') for d in drawings for x in d.findall('.//' + qn('wp:docPr'))]
    add('figure_alt_text', len(drawings) == 1 and all(descrs), {'drawings': len(drawings), 'descrs': descrs})

    failed = [x for x in checks if x['status'] == 'FAIL']
    result = {'status': 'PASS' if not failed else 'FAIL', 'checks_passed': len(checks)-len(failed), 'checks_total': len(checks), 'docx_sha256': digest(DOCX), 'pdf_sha256': digest(PDF), 'page_metrics': metrics, 'checks': checks}
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if failed:
        raise SystemExit(1)

if __name__ == '__main__':
    main()
