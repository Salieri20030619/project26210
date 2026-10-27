from __future__ import annotations

import json
import math
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.io import loadmat


if len(sys.argv) != 4:
    raise SystemExit("Usage: 01_preprocess_ccl180_ccle_v1.py <CCL180_dir> <CCLE_dir> <output_dir>")

ccl_dir = Path(sys.argv[1]).resolve()
ccle_dir = Path(sys.argv[2]).resolve()
out_dir = Path(sys.argv[3]).resolve()
results_dir = out_dir / "preprocessed_v1"
qc_dir = out_dir / "qc_v1"
results_dir.mkdir(parents=True, exist_ok=True)
qc_dir.mkdir(parents=True, exist_ok=True)


def matlab_strings(array: np.ndarray) -> list[str]:
    values: list[str] = []
    for value in np.ravel(array):
        while isinstance(value, np.ndarray) and value.size == 1:
            value = value.item()
        values.append(str(value))
    return values


def cell_key(value: object) -> str:
    key = re.sub(r"[^A-Z0-9]", "", str(value).upper())
    aliases = {"786O": "7860", "NCIH322": "NCIH322M", "RPMI8826": "RPMI8226"}
    return aliases.get(key, key)


cells = matlab_strings(loadmat(ccl_dir / "Transcripts_cellline.mat", squeeze_me=True)["CellLineSelectedTrans"])
genes = matlab_strings(loadmat(ccl_dir / "Transcripts_list.mat", squeeze_me=True)["GeneSelectedTrans"])
expression = np.asarray(loadmat(ccl_dir / "Transcripts_Astra.mat", squeeze_me=True)["transcriptomicsAstra"], dtype=float)
if expression.shape != (len(cells), len(genes)):
    raise RuntimeError(f"Transcript matrix mismatch: {expression.shape}, cells={len(cells)}, genes={len(genes)}")

gene_index = {gene.upper(): index for index, gene in enumerate(genes)}
if "TYMP" not in gene_index or "MTAP" not in gene_index:
    raise RuntimeError("TYMP or MTAP missing from CCL180 transcript list")

expr = pd.DataFrame(
    {
        "cell_line": cells,
        "cell_key": [cell_key(value) for value in cells],
        "tymp_expression": expression[:, gene_index["TYMP"]],
        "mtap_expression": expression[:, gene_index["MTAP"]],
    }
)
if expr["cell_key"].duplicated().any():
    raise RuntimeError("Duplicate normalized transcript cell-line key")

meta = pd.read_csv(ccl_dir / "CellLinesMeta.csv", header=None)
meta.columns = ["meta_cell_line", "meta_field2", "meta_field3", "tissue", "medium_group", "meta_field6"]
meta["cell_key"] = meta["meta_cell_line"].map(cell_key)
meta = meta.drop_duplicates("cell_key")

workbook_path = ccl_dir / "metabolomicsData.xlsx"
workbook = load_workbook(workbook_path, read_only=True, data_only=True)

annotation_sheet = workbook["annotation"]
annotation_rows = annotation_sheet.iter_rows(values_only=True)
annotation_header = [str(value) if value is not None else "" for value in next(annotation_rows)]
annotation = pd.DataFrame(annotation_rows, columns=annotation_header)
annotation["ionIdx"] = pd.to_numeric(annotation["ionIdx"], errors="coerce").astype("Int64")
annotation["ionMz"] = pd.to_numeric(annotation["ionMz"], errors="coerce")

exact_terms = {
    "5'-s-methyl-5'-thioadenosine",
    "hmdb0001173",
    "c00170",
    "c11h15n5o3s",
}
annotation_blob = annotation.fillna("").astype(str).agg("|".join, axis=1).str.lower()
exact_mta_mask = annotation_blob.map(lambda value: any(term in value for term in exact_terms))

near_targets = [296.0813, 297.0891, 298.0969]
near_mask = annotation["ionMz"].map(lambda value: False if pd.isna(value) else min(abs(value - target) for target in near_targets) <= 0.03)
near = annotation.loc[near_mask, [
    "ionIdx", "ionMz", "formula", "ion", "score", "label (bona fide)", "HMDB ids", "KEGG ids"
]].copy()
near["accepted_as_mta"] = False
near["rejection_reason"] = "formula/database annotation conflicts with MTA identity"
near.to_csv(results_dir / "ccl180_near_mta_mass_features_v1.tsv", sep="\t", index=False)

panel = {
    134: "methionine",
    1103: "sah",
    89: "n_acetylputrescine",
}
for ion_id, name in panel.items():
    rows = annotation[annotation["ionIdx"] == ion_id]
    if rows.empty:
        raise RuntimeError(f"Required CCL180 panel ion missing: {ion_id} {name}")

injection_sheet = workbook["injections"]
injection_rows = injection_sheet.iter_rows(values_only=True)
injection_header = [str(value) if value is not None else "" for value in next(injection_rows)]
injections = pd.DataFrame(injection_rows, columns=injection_header)
injections = injections[injections["dsActive"].astype(bool)].copy()
injections["dsIdx"] = pd.to_numeric(injections["dsIdx"], errors="raise").astype(int)
injections["cell_line_injection"] = injections["dsUser2"].astype(str)
injections["cell_key"] = injections["cell_line_injection"].map(cell_key)
injections["batch"] = injections["dsUser1"].astype(str)
injections["injection_tissue"] = injections["dsUser4"].astype(str)

intensity_sheet = workbook["intensities1"]
intensity_rows = intensity_sheet.iter_rows(values_only=True)
intensity_header = list(next(intensity_rows))
injection_ids = [int(str(value).split("/")[0].strip()) for value in intensity_header[3:]]
selected_rows: dict[int, tuple[object, ...]] = {}
for row in intensity_rows:
    try:
        ion_id = int(row[0])
    except (TypeError, ValueError):
        continue
    if ion_id in panel:
        selected_rows[ion_id] = row
    if len(selected_rows) == len(panel):
        break
if set(selected_rows) != set(panel):
    raise RuntimeError(f"Missing panel intensity rows: {set(panel) - set(selected_rows)}")

long_parts: list[pd.DataFrame] = []
for ion_id, metabolite in panel.items():
    values = pd.to_numeric(pd.Series(selected_rows[ion_id][3:]), errors="coerce")
    if len(values) != len(injection_ids):
        raise RuntimeError("Intensity/injection header length mismatch")
    frame = pd.DataFrame({"dsIdx": injection_ids, "metabolite": metabolite, "intensity": values})
    long_parts.append(frame)
long = pd.concat(long_parts, ignore_index=True).merge(
    injections[["dsIdx", "cell_line_injection", "cell_key", "batch", "injection_tissue"]],
    on="dsIdx",
    how="inner",
    validate="many_to_one",
)
if long["intensity"].isna().any() or (long["intensity"] <= 0).any():
    raise RuntimeError("Missing or non-positive selected CCL180 intensity")
long.to_csv(results_dir / "ccl180_selected_injections_long_v1.tsv", sep="\t", index=False)

agg = (
    long.groupby(["cell_key", "metabolite"], as_index=False)
    .agg(
        n_injections=("dsIdx", "nunique"),
        n_batches=("batch", "nunique"),
        median_intensity=("intensity", "median"),
        mean_intensity=("intensity", "mean"),
        sd_intensity=("intensity", "std"),
        min_intensity=("intensity", "min"),
        max_intensity=("intensity", "max"),
    )
)
agg.to_csv(results_dir / "ccl180_selected_metabolite_aggregation_v1.tsv", sep="\t", index=False)

medians = agg.pivot(index="cell_key", columns="metabolite", values="median_intensity").reset_index()
counts = agg.pivot(index="cell_key", columns="metabolite", values="n_injections").reset_index()
counts = counts.rename(columns={name: f"{name}_n_injections" for name in panel.values()})
ccl = expr.merge(meta[["cell_key", "meta_cell_line", "tissue", "medium_group"]], on="cell_key", how="left", validate="one_to_one")
ccl = ccl.merge(medians, on="cell_key", how="inner", validate="one_to_one")
ccl = ccl.merge(counts, on="cell_key", how="left", validate="one_to_one")
ccl["is_renal"] = ccl["tissue"].astype(str).str.upper().eq("KIDNEY")
ccl.to_csv(results_dir / "ccl180_preprocessed_all_v1.tsv", sep="\t", index=False)
ccl[ccl["is_renal"]].to_csv(results_dir / "ccl180_preprocessed_renal_v1.tsv", sep="\t", index=False)

ccle = pd.read_csv(ccle_dir / "CCLE_metabolomics_20190502.csv")
ccle_mta_present = any(
    term in str(column).lower()
    for column in ccle.columns
    for term in ("methylthioadenosine", "5'-mta", "5-mta")
)
mapping = pd.read_csv(ccl_dir / "CellLinesMeta_CCLE.csv")
mapping = mapping[mapping["NameAstra"].notna() & mapping["CCLE_ID"].notna()].copy()
mapping["cell_key"] = mapping["NameAstra"].map(cell_key)
mapping = mapping[mapping["cell_key"].isin(set(expr["cell_key"]))].drop_duplicates("CCLE_ID")

ccle_join = mapping.merge(ccle, on="CCLE_ID", how="inner", validate="one_to_one")
ccle_join = ccle_join.merge(expr, on="cell_key", how="inner", validate="one_to_one")
ccle_join = ccle_join.merge(meta[["cell_key", "tissue", "medium_group"]], on="cell_key", how="left", validate="one_to_one")
ccle_out = ccle_join[[
    "cell_key", "NameAstra", "CCLE_ID", "DepMap_ID", "tissue", "medium_group",
    "tymp_expression", "mtap_expression", "methionine", "5-adenosylhomocysteine", "putrescine"
]].copy()
ccle_out = ccle_out.rename(columns={"NameAstra": "cell_line", "5-adenosylhomocysteine": "sah"})
ccle_out["is_renal"] = ccle_out["tissue"].astype(str).str.upper().eq("KIDNEY")
ccle_out.to_csv(results_dir / "ccle_preprocessed_all_v1.tsv", sep="\t", index=False)
ccle_out[ccle_out["is_renal"]].to_csv(results_dir / "ccle_preprocessed_renal_v1.tsv", sep="\t", index=False)

shared = ccl[["cell_key", "cell_line", "methionine", "sah"]].merge(
    ccle_out[["cell_key", "methionine", "sah"]],
    on="cell_key",
    how="inner",
    suffixes=("_ccl180", "_ccle"),
    validate="one_to_one",
)
shared.to_csv(results_dir / "shared_platform_metabolites_v1.tsv", sep="\t", index=False)

identity_audit = pd.DataFrame(
    [
        {
            "dataset": "CCL180",
            "exact_mta_evaluable": False,
            "exact_mta_matches": int(exact_mta_mask.sum()),
            "near_mass_annotation_rows": int(len(near)),
            "decision": "DIRECT_MTA_NOT_EVALUABLE",
            "reason": "No exact name/HMDB/KEGG/formula match; near-mass ions have conflicting identities",
        },
        {
            "dataset": "CCLE",
            "exact_mta_evaluable": False,
            "exact_mta_matches": int(ccle_mta_present),
            "near_mass_annotation_rows": 0,
            "decision": "DIRECT_MTA_NOT_EVALUABLE",
            "reason": "No MTA feature in the released metabolite matrix",
        },
    ]
)
identity_audit.to_csv(results_dir / "direct_mta_feasibility_v1.tsv", sep="\t", index=False)

checks = [
    ("ccl180_transcript_155", len(expr) == 155, f"observed={len(expr)}"),
    ("ccl180_tymp_mtap_complete", not expr[["tymp_expression", "mtap_expression"]].isna().any().any(), "TYMP/MTAP complete"),
    ("ccl180_direct_mta_absent", int(exact_mta_mask.sum()) == 0, f"matches={int(exact_mta_mask.sum())}"),
    ("ccl180_near_mass_not_mta", len(near) > 0 and not near["accepted_as_mta"].any(), f"rows={len(near)}"),
    ("ccl180_selected_ions_complete", set(selected_rows) == set(panel), str(sorted(selected_rows))),
    ("ccl180_injections_1153", len(injections) == 1153, f"observed={len(injections)}"),
    ("ccl180_paired_155", len(ccl) == 155, f"observed={len(ccl)}"),
    ("ccl180_panel_complete", not ccl[["methionine", "sah", "n_acetylputrescine"]].isna().any().any(), "complete"),
    ("ccl180_renal_below_10", 0 < int(ccl["is_renal"].sum()) < 10, f"renal={int(ccl['is_renal'].sum())}"),
    ("ccle_direct_mta_absent", not ccle_mta_present, f"present={ccle_mta_present}"),
    ("ccle_paired_148", len(ccle_out) == 148, f"observed={len(ccle_out)}"),
    ("ccle_panel_complete", not ccle_out[["methionine", "sah", "putrescine"]].isna().any().any(), "complete"),
    ("ccle_renal_below_10", 0 < int(ccle_out["is_renal"].sum()) < 10, f"renal={int(ccle_out['is_renal'].sum())}"),
    ("shared_platform_148", len(shared) == 148, f"observed={len(shared)}"),
    ("both_include_786", "7860" in set(ccl["cell_key"]) and "7860" in set(ccle_out["cell_key"]), "7860 matched"),
]
checks_frame = pd.DataFrame(checks, columns=["check_id", "pass", "detail"])
checks_frame.to_csv(qc_dir / "preprocessing_checks_v1.tsv", sep="\t", index=False)

status = {
    "status": "PASS_PREPROCESSING" if bool(checks_frame["pass"].all()) else "FAIL_PREPROCESSING",
    "ccl180_all_n": int(len(ccl)),
    "ccl180_renal_n": int(ccl["is_renal"].sum()),
    "ccle_all_n": int(len(ccle_out)),
    "ccle_renal_n": int(ccle_out["is_renal"].sum()),
    "shared_platform_n": int(len(shared)),
    "direct_mta_evaluable": False,
    "checks_passed": int(checks_frame["pass"].sum()),
    "checks_total": int(len(checks_frame)),
}
(qc_dir / "preprocessing_status_v1.json").write_text(json.dumps(status, indent=2) + "\n", encoding="utf-8")
if status["status"].startswith("FAIL"):
    failed = checks_frame.loc[~checks_frame["pass"], "check_id"].tolist()
    raise RuntimeError(f"Preprocessing failed: {failed}")

summary = [
    "# Phase 104 preprocessing summary",
    "",
    "- Direct MTA decision: **DIRECT_MTA_NOT_EVALUABLE** in both CCL180 and CCLE.",
    f"- CCL180: {len(ccl)} paired cell lines; renal n={int(ccl['is_renal'].sum())}; selected adjacent panel is methionine, SAH and N-acetylputrescine.",
    f"- CCLE: {len(ccle_out)} cell lines paired to CCL180 TYMP/MTAP expression; renal n={int(ccle_out['is_renal'].sum())}; selected adjacent panel is methionine, SAH and putrescine.",
    f"- Shared CCL180/CCLE platform comparison: {len(shared)} overlapping cell lines.",
    "- Adjacent metabolites are not MTA substitutes and cannot support a direct TYMP-to-MTA claim.",
]
(out_dir / "PHASE104_PREPROCESSING_SUMMARY_v1.md").write_text("\n".join(summary) + "\n", encoding="ascii")
print("PASS_PREPROCESSING")
print(json.dumps(status, sort_keys=True))
