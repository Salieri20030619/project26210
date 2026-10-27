from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


def main() -> int:
    if len(sys.argv) != 3:
        raise SystemExit("usage: extract_docx_reference_local_v1.py INPUT.docx OUTPUT_DIR")
    source = Path(sys.argv[1]).resolve()
    output_dir = Path(sys.argv[2]).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    doc = Document(str(source))
    paragraph_rows = []
    text_lines = []
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text
        paragraph_rows.append(
            {"index": index, "style": paragraph.style.name if paragraph.style else "", "text": text}
        )
        if text.strip():
            text_lines.append(f"[P{index}][{paragraph_rows[-1]['style']}] {text}")

    tables = []
    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.replace("\r", " ").replace("\n", " | ") for cell in row.cells]
            rows.append(cells)
            text_lines.append(f"[TABLE {table_index} ROW {row_index}] " + "\t".join(cells))
        tables.append({"index": table_index, "rows": rows, "columns_max": max((len(x) for x in rows), default=0)})

    headers = []
    footers = []
    sections = []
    for section_index, section in enumerate(doc.sections, start=1):
        header_text = "\n".join(p.text for p in section.header.paragraphs if p.text.strip())
        footer_text = "\n".join(p.text for p in section.footer.paragraphs if p.text.strip())
        headers.append(header_text)
        footers.append(footer_text)
        sections.append(
            {
                "index": section_index,
                "page_width": section.page_width,
                "page_height": section.page_height,
                "top_margin": section.top_margin,
                "bottom_margin": section.bottom_margin,
                "left_margin": section.left_margin,
                "right_margin": section.right_margin,
                "orientation": str(section.orientation),
            }
        )

    with zipfile.ZipFile(source) as archive:
        names = set(archive.namelist())
        document_xml = archive.read("word/document.xml")
        root = ET.fromstring(document_xml)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        inserted = len(root.findall(".//w:ins", ns))
        deleted = len(root.findall(".//w:del", ns))
        comment_count = 0
        if "word/comments.xml" in names:
            comments_root = ET.fromstring(archive.read("word/comments.xml"))
            comment_count = len(comments_root.findall(".//w:comment", ns))

    props = doc.core_properties
    structure = {
        "source": str(source),
        "paragraph_count": len(doc.paragraphs),
        "nonempty_paragraph_count": sum(bool(p.text.strip()) for p in doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "section_count": len(doc.sections),
        "tracked_insertions": inserted,
        "tracked_deletions": deleted,
        "comments": comment_count,
        "headers": headers,
        "footers": footers,
        "sections": sections,
        "core_properties": {
            "title": props.title,
            "subject": props.subject,
            "author": props.author,
            "last_modified_by": props.last_modified_by,
            "created": props.created.isoformat() if props.created else None,
            "modified": props.modified.isoformat() if props.modified else None,
        },
        "paragraphs": paragraph_rows,
        "tables": tables,
    }
    (output_dir / "docx_structure.json").write_text(
        json.dumps(structure, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (output_dir / "docx_extracted_text.txt").write_text("\n".join(text_lines) + "\n", encoding="utf-8")
    print(json.dumps({key: structure[key] for key in (
        "paragraph_count", "nonempty_paragraph_count", "table_count", "inline_shape_count",
        "section_count", "tracked_insertions", "tracked_deletions", "comments"
    )}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
