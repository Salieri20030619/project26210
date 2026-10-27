from __future__ import annotations

import hashlib
import importlib.util
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent
OUTPUT = ROOT / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_NARRATIVE_REVIEW_v5.docx"
STATUS = ROOT / "build_status_v2.json"
PART_OUTPUT = ROOT / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_NARRATIVE_REVIEW_v5.part.docx"
PART_STATUS = ROOT / "build_status_v2.part.json"
ASSETS = ROOT / "assets_v1"
FLOWCHART = ASSETS / "part1_workflow_narrative_v2.png"

source = PROJECT / "phase40_part1_client_report_local_v1" / "03_build_part1_complete_client_report_v1.py"
spec = importlib.util.spec_from_file_location("phase40_report", source)
report = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(report)


original_configure_styles = report.h.configure_styles
original_add_title_page = report.add_title_page


def configure_reference_body_styles(doc):
    """Keep the project cover, but use the reference report's restrained body typography."""
    original_configure_styles(doc)
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(22)
    normal.paragraph_format.space_after = Pt(4)

    for name, size, before, after in (
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 11.5, 8, 4),
    ):
        style = styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = report.h.RGBColor.from_string("000000")
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    caption = styles["Caption"]
    caption.font.name = "宋体"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(9.5)


def add_reference_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.first_line_indent = Pt(0)
    for run in p.runs:
        report.h.set_run(run, size={1: 16, 2: 13, 3: 11.5}[level], bold=True, color="000000", font="黑体")
    return p


def add_reference_body(doc, text, *, bold_lead=None, keep=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(22)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if bold_lead and text.startswith(bold_lead):
        report.h.set_run(p.add_run(bold_lead), bold=True, color="000000", font="宋体")
        report.h.set_run(p.add_run(text[len(bold_lead):]), color="000000", font="宋体")
    else:
        report.h.set_run(p.add_run(text), color="000000", font="宋体")
    p.paragraph_format.keep_with_next = keep
    return p


def narrative_analysis(doc, code, title, why, method, result, conclusion=None):
    """Render one analysis as two connected research-report paragraphs."""
    add_reference_heading(doc, f"{code}  {title}", 2)
    purpose = why.rstrip("。")
    method_text = method[0].lower() + method[1:] if method[:1].isascii() else method
    first = add_reference_body(doc, f"为{purpose}，本研究{method_text}", keep=True)
    first.paragraph_format.keep_together = True
    interpretation = result
    if conclusion:
        interpretation += conclusion
    second = add_reference_body(doc, f"分析显示，{interpretation}")
    second.paragraph_format.keep_together = True


def narrative_title_page(doc):
    original_add_title_page(doc)
    replacements = {
        "客户正文审阅稿：已纳入全部26项分析的目的、方法和结果；本版仅放分析流程图，数据结果图待正文确认后统一编排。":
        "客户正文审阅稿：已纳入全部26项分析及其研究逻辑、实施路径和主要结果；本版仅放分析流程图，数据结果图待正文确认后统一编排。"
    }
    for p in doc.paragraphs:
        for needle, replacement in replacements.items():
            if needle in p.text:
                for run in p.runs:
                    run.text = ""
                report.h.set_run(p.add_run("版本说明  "), bold=True, color=report.h.TEAL)
                report.h.set_run(p.add_run(replacement), color="000000")


def clear_runs(paragraph):
    for run in paragraph.runs:
        run.text = ""


def remove_page_break_before(doc, heading_text):
    paragraphs = list(doc.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text == heading_text and index > 0:
            previous = paragraphs[index - 1]
            if previous._p.xpath(".//w:br[@w:type='page']"):
                previous._element.getparent().remove(previous._element)
            return


def postprocess_layout(doc):
    for paragraph in doc.paragraphs:
        if paragraph.text == "第一部分 外部单细胞验证与机制分析报告":
            clear_runs(paragraph)
            run = paragraph.add_run("第一部分 外部单细胞验证\n与机制分析报告")
            report.h.set_run(run, size=23, bold=True, color=report.h.NAVY)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if "客户正文审阅稿：已纳入全部26项分析的目的、方法和结果" in paragraph.text:
            clear_runs(paragraph)
            report.h.set_run(paragraph.add_run("版本说明  "), bold=True, color=report.h.TEAL)
            report.h.set_run(
                paragraph.add_run(
                    "客户正文审阅稿：已纳入全部26项分析及其研究逻辑、实施路径和主要结果；"
                    "本版仅放分析流程图，数据结果图待正文确认后统一编排。"
                ),
                color="000000",
            )

    for heading in (
        "3  恶性细胞、CNV克隆与cNMF程序",
        "6  肿瘤微环境通讯与临床验证",
        "7  综合解释",
    ):
        remove_page_break_before(doc, heading)

    completion = doc.tables[-1]
    target_row = None
    c05_row = None
    for row in completion.rows:
        key = row.cells[0].text.strip()
        if key == "C01–C04":
            target_row = row
        elif key == "C05":
            c05_row = row
    if target_row is not None and c05_row is not None:
        target_row.cells[0].text = "C01–C05"
        target_row.cells[1].text = "CellChat、受体、NicheNet与TCGA-KIRC联合生存"
        completion._tbl.remove(c05_row._tr)


def build():
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing deliverable: {OUTPUT}")
    if PART_OUTPUT.exists() or PART_STATUS.exists():
        raise FileExistsError("Refusing to overwrite an existing partial build")

    report.ROOT = ROOT
    report.OUTPUT = PART_OUTPUT
    report.STATUS = PART_STATUS
    report.ASSETS = ASSETS
    report.FLOWCHART = FLOWCHART
    report.h.configure_styles = configure_reference_body_styles
    report.h.add_heading = add_reference_heading
    report.h.add_body = add_reference_body
    report.analysis = narrative_analysis
    report.add_title_page = narrative_title_page
    report.build_document()

    doc = Document(PART_OUTPUT)
    postprocess_layout(doc)
    doc.save(OUTPUT)

    status = json.loads(PART_STATUS.read_text(encoding="utf-8"))
    status.update(
        {
            "status": "PASS",
            "editorial_revision": "reference-derived continuous narrative",
            "mechanical_labels_removed": ["为什么做", "怎么做", "得到什么", "本项结论"],
            "output": str(OUTPUT),
            "size_bytes": OUTPUT.stat().st_size,
            "sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        }
    )
    STATUS.write_text(json.dumps(status, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    PART_OUTPUT.unlink()
    PART_STATUS.unlink()
    print(json.dumps(status, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
