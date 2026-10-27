from __future__ import annotations

import csv
import hashlib
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


PHASE = Path(__file__).resolve().parent
SOURCE = Path(
    r"D:\project26210_external_joint_v1\phase56_part1_incremental_report_update_local_v1"
) / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_WITH_FIGURES_REVIEW_v11.docx"
EVIDENCE_V2 = Path(
    r"D:\project26210_external_joint_v1\phase65_part1_immune_mechanistic_convergence_local_v1"
) / "results_v2" / "IMMUNE_MECHANISTIC_EVIDENCE_REGISTER_v2.tsv"
PHASE67 = Path(
    r"D:\project26210_external_joint_v1\phase67_part1_direct_program_replication_local_v1"
) / "results_v5"
OUTPUT = PHASE / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_WITH_FIGURES_REVIEW_v12.docx"
EVIDENCE_V3 = PHASE / "IMMUNE_MECHANISTIC_EVIDENCE_REGISTER_v3.tsv"
GRADE_SUMMARY = PHASE / "PART1_IMMUNE_EVIDENCE_GRADE_UPDATE_v1.md"
QC = PHASE / "qc_local_v1"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(archive.namelist())
            if name.startswith("word/media/")
        }


def find_exact(document: Document, text: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one paragraph {text!r}; found {len(matches)}")
    return matches[0]


def find_starting(document: Document, prefix: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one paragraph starting {prefix!r}; found {len(matches)}")
    return matches[0]


def replace_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_before(anchor: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    paragraph = anchor.insert_paragraph_before(text, style=style)
    paragraph.paragraph_format.keep_with_next = style.startswith("Heading")
    return paragraph


def add_labeled_before(anchor: Paragraph, label: str, body: str) -> Paragraph:
    paragraph = anchor.insert_paragraph_before(style="Normal")
    lead = paragraph.add_run(label)
    lead.bold = True
    paragraph.add_run(body)
    return paragraph


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(round(width_cm / 2.54 * 1440))))
    tc_w.set(qn("w:type"), "dxa")


def format_cell(cell, *, bold: bool = False, font_size: float = 8.5) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.bold = bold
            run.font.name = "宋体"
            run.font.size = Pt(font_size)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def insert_evidence_table(document: Document, anchor: Paragraph) -> None:
    headers = ["证据轴", "外部复现", "等级", "可写结论"]
    rows = [
        [
            "髓系/DC–TYMP代谢与调控核心",
            "GSE207493发现；GSE264586与GSE294109均有18个同方向FDR签名",
            "A",
            "氧化/缺氧、红氧、糖代谢及IRF/SPI1/NF-κB核心获得双外部队列患者级复现",
        ],
        [
            "Treg–TYMP代谢与调控核心",
            "GSE207493发现；GSE294109有13个同方向FDR签名；GSE264586无冻结Treg注释",
            "B",
            "存在单外部队列复现，但不能写成双队列确认",
        ],
        [
            "F1细胞来源与Treg情境关系",
            "大队列bulk关联可复现；细胞分辨率显示F1主要位于肿瘤细胞",
            "A/B",
            "bulk F1–Treg关系主要是组成或共享微环境情境，不代表Treg内源F1",
        ],
        [
            "F1直接临床预测",
            "三套外部临床队列未形成稳健一致验证",
            "D",
            "不能作为已验证的独立免疫治疗疗效指标",
        ],
        [
            "TYMP直接临床预测",
            "既有response/PFS/OS模型未获直接支持",
            "D",
            "机制复现不等同于疗效或生存预测验证",
        ],
    ]
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = [3.5, 5.2, 1.4, 6.4]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        set_cell_width(cell, widths[idx])
        shade_cell(cell, "D9EAF7")
        format_cell(cell, bold=True, font_size=9.0)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(values):
            cells[idx].text = text
            set_cell_width(cells[idx], widths[idx])
            format_cell(cells[idx])
    anchor._p.addprevious(table._tbl)


def update_evidence_register() -> dict[str, int]:
    with EVIDENCE_V2.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    fields = list(rows[0].keys())
    updates = {
        "I08": {
            "finding": "TYMP-high myeloid/DC externally replicated metabolic core",
            "n": "GSE207493 discovery plus GSE264586 and GSE294109 external replication",
            "effect": "Nine metabolic/stress signatures replicate in both external cohorts; the broader pyrimidine-remodeling package remains discovery-cohort evidence",
            "statistic": "18 total R2 signatures across two external cohorts, including 9 metabolic/stress signatures; patient/timepoint rho, cycle adjusted, within-dataset family BH FDR",
            "global_grade": "A",
            "external_grade": "R2_TWO_EXTERNAL_FDR_CORE",
            "interpretation": "TYMP marks a replicated oxidative/hypoxic, redox and glycolytic myeloid/DC state; only the replicated core receives grade A.",
            "limitation": "Expression-derived scores are not matched metabolomics; broad pyrimidine remodeling outside the replicated core remains C1-level discovery evidence.",
            "source": "gse207493_metabolic_*_v1.tsv;phase67_cross_dataset_replication_v1.tsv",
        },
        "I09": {
            "finding": "TYMP-high Treg externally replicated metabolic core",
            "n": "GSE207493 discovery plus GSE294109 external replication; GSE264586 not evaluable for Treg",
            "effect": "Six hypoxia, respiration, redox and glycolytic signatures replicate in one external cohort",
            "statistic": "13 total Treg R1 signatures in GSE294109, including 6 metabolic/stress signatures; patient/timepoint rho, cycle adjusted, family BH FDR",
            "global_grade": "B",
            "external_grade": "R1_SINGLE_EXTERNAL_FDR_CORE",
            "interpretation": "Treg TYMP adaptation has independent single-cohort support but is not a two-external-cohort replication.",
            "limitation": "GSE264586 lacks the frozen Treg annotation; the full discovery signature is not externally replicated.",
            "source": "gse207493_metabolic_*_v1.tsv;phase67_cross_dataset_replication_v1.tsv",
        },
        "I11": {
            "finding": "TYMP-high myeloid/DC externally replicated TF core",
            "n": "GSE207493 discovery plus GSE264586 and GSE294109 external replication",
            "effect": "Nine IRF/SPI1/NF-kB-centered positive regulons replicate in both external cohorts",
            "statistic": "IRF1,SPI1,RELB,STAT6,IRF7,STAT2,ZBTB7A,CEBPB,RELA are R2; patient/timepoint rho, cycle adjusted, family BH FDR",
            "global_grade": "A",
            "external_grade": "R2_TWO_EXTERNAL_FDR_CORE",
            "interpretation": "A lineage-specific regulatory core now has two independent external cohort replications.",
            "limitation": "The grade applies to the nine replicated regulons, not all 40 discovery-cohort TFs, and does not establish causality.",
            "source": "de_novo_tf_cross_method_FDR_v1.tsv;phase67_cross_dataset_replication_v1.tsv",
        },
        "I12": {
            "finding": "TYMP-high Treg externally replicated TF core",
            "n": "GSE207493 discovery plus GSE294109 external replication; GSE264586 not evaluable for Treg",
            "effect": "Seven GATA3/NF-kB/STAT-centered positive regulons replicate in one external cohort",
            "statistic": "GATA3,IRF2,NFKB2,REL,TCF3,RELB,STAT1 are R1 in GSE294109; patient/timepoint rho, cycle adjusted, family BH FDR",
            "global_grade": "B",
            "external_grade": "R1_SINGLE_EXTERNAL_FDR_CORE",
            "interpretation": "The Treg regulatory core has independent single-cohort support, not two-cohort confirmation.",
            "limitation": "GSE264586 lacks the frozen Treg annotation; the grade applies only to the seven replicated regulons and is not causal evidence.",
            "source": "de_novo_tf_cross_method_FDR_v1.tsv;phase67_cross_dataset_replication_v1.tsv",
        },
    }
    seen = set()
    for row in rows:
        if row["evidence_id"] in updates:
            row.update(updates[row["evidence_id"]])
            seen.add(row["evidence_id"])
    if seen != set(updates):
        raise RuntimeError(f"Missing evidence rows: {sorted(set(updates) - seen)}")
    with EVIDENCE_V3.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    counts: dict[str, int] = {}
    for row in rows:
        counts[row["global_grade"]] = counts.get(row["global_grade"], 0) + 1
    expected = {"A": 5, "B": 6, "C1": 2, "D": 2}
    if counts != expected:
        raise RuntimeError(f"Unexpected evidence grade counts: {counts}")
    return counts


def write_grade_summary(counts: dict[str, int]) -> None:
    GRADE_SUMMARY.write_text(
        """# 第一部分免疫系列证据等级更新（Phase68）

## 更新结果

- 统一15项证据等级由 A=3、B=4、C1=6、D=2 更新为 A=5、B=6、C1=2、D=2。
- I08 与 I11 仅将两套外部队列均同方向且通过FDR的髓系/DC代谢和调控核心升级为 A。
- I09 与 I12 仅将 GSE294109 单外部队列复现的Treg代谢和调控核心升级为 B；GSE264586无冻结Treg细注释，不能称为双外部队列复现。
- F1肿瘤代谢与调控发现仍为 C1；F1和TYMP直接临床预测仍为 D。

## 外部复现基础

- 直接分析覆盖50,531个Myeloid/DC细胞和6,341个Treg细胞。
- 18个Myeloid/DC签名在GSE264586与GSE294109中均与发现队列同方向并通过FDR。
- 13个Treg签名在GSE294109中同方向并通过FDR。
- 评分移除TYMP锚定基因，按患者/时间点汇总，进行S/G2M校正，并在数据集、谱系和签名家族内做BH校正。

## 解释边界

升级只适用于外部复现的冻结核心，不外推到整个发现签名集；表达评分不是配对代谢组，调控相关不是因果验证，机制复现也不是response、PFS或OS预测验证。
""",
        encoding="utf-8",
    )


def main() -> None:
    for path in [OUTPUT, EVIDENCE_V3, GRADE_SUMMARY]:
        if path.exists():
            raise FileExistsError(f"Refusing to overwrite existing output: {path}")
    for path in [SOURCE, EVIDENCE_V2, PHASE67 / "phase67_cross_dataset_replication_v1.tsv"]:
        if not path.exists():
            raise FileNotFoundError(path)

    QC.mkdir(parents=True, exist_ok=True)
    source_hash = sha256(SOURCE)
    source_media = media_hashes(SOURCE)
    counts = update_evidence_register()
    write_grade_summary(counts)

    document = Document(SOURCE)
    replace_text(
        find_starting(document, "版本说明  客户图文审阅稿："),
        "版本说明  客户图文审阅稿：保留全部26项分析与15组正文结果图，在新增患者三层复核基础上，进一步纳入免疫系列外部直接复现与统一证据等级更新；正文继续采用患者层统计和分级判定。",
    )

    summary_anchor = find_exact(document, "1  数据范围与分析设计")
    add_labeled_before(
        summary_anchor,
        "免疫系列升级  ",
        "冻结的TYMP去锚代谢与调控程序在GSE264586和GSE294109中完成患者级直接复算。共分析50,531个Myeloid/DC细胞和6,341个Treg细胞；18个髓系/DC签名在两个外部队列均同方向并通过FDR，13个Treg签名在GSE294109单外部队列同方向并通过FDR。统一15项证据更新为A=5、B=6、C1=2、D=2；机制证据升级不改变F1或TYMP直接临床预测仍为D级的结论。",
    )

    stats_anchor = find_exact(document, "1.3  总体流程")
    insert_before(
        stats_anchor,
        "免疫程序外部复现沿用冻结签名并移除TYMP锚定基因；在患者、时间点和谱系内计算相关，进行S/G2M校正，以患者跨时间点中位效应推断，并在每个数据集、谱系和签名家族内进行BH校正。",
    )

    section7 = find_exact(document, "7  综合解释")
    insert_before(section7, "C07  免疫程序外部直接复现与证据升级", "Heading 2")
    insert_before(
        section7,
        "Phase67使用Phase64冻结的50个TYMP去锚签名，在两套外部免疫治疗单细胞队列中直接重算TYMP连续表达与程序活性的患者级关系。GSE264586和GSE294109合计纳入50,531个Myeloid/DC细胞；GSE294109另纳入6,341个Treg细胞。每个患者/时间点/谱系至少50个细胞才进入分析，相关性经过S/G2M残差化，FDR按数据集、谱系和签名家族控制。",
    )
    insert_evidence_table(document, section7)
    add_labeled_before(
        section7,
        "髓系/DC双外部复现  ",
        "GSE264586有21个、GSE294109有23个Myeloid/DC签名与发现队列同方向并通过FDR；其中18个在两个外部队列均达到R2。复现核心包括缺氧反应、氧化呼吸/电子传递、ROS解毒、糖酵解、糖异生、磷酸戊糖途径和嘌呤分解，PDH复合体调节项保持预设负向方向；同时包括IRF1、SPI1、RELB、STAT6、IRF7、STAT2、ZBTB7A、CEBPB和RELA正调控子。该核心由原来的单队列内部跨方法收敛升级为A级外部机制复现。",
    )
    add_labeled_before(
        section7,
        "Treg单外部复现  ",
        "GSE294109有13个Treg签名同方向并通过FDR，包括6个缺氧/呼吸/红氧/糖代谢签名和GATA3、IRF2、NFKB2、REL、TCF3、RELB、STAT1七个正调控子。由于GSE264586没有冻结Treg细注释，该核心只能升为B级单外部队列支持，不能写作双队列确认。",
    )
    add_labeled_before(
        section7,
        "统一等级与边界  ",
        "15项免疫证据由A=3、B=4、C1=6、D=2更新为A=5、B=6、C1=2、D=2。升级只适用于外部复现的冻结核心：原始髓系/DC广泛嘧啶重塑中未进入R2核心的部分仍是发现证据，F1肿瘤代谢与调控程序仍为C1。Phase66的42项情境转移FDR结果提示临床背景依赖；Phase67的直接相关复现更接近机制验证，但两者都不等同于response、PFS或OS预测。F1和TYMP作为独立临床预测指标均继续保持D级。",
    )

    find_starting(document, "第一，cNMF F1与TYMP在患者内总体同向").add_run(
        " 第七，TYMP去锚的髓系/DC代谢—调控核心已在两套外部队列获得患者级FDR复现，Treg核心在一套可评价外部队列获得支持。"
    )
    find_starting(document, "F1不对应稳定的RCC遗传易感性富集。").add_run(
        " 免疫程序复现不提供疗效方向或生存效应，不能据此升级F1或TYMP的临床预测等级。"
    )
    find_starting(document, "综合结论  第一部分26项分析已经全部完成。").add_run(
        " Phase67进一步把TYMP相关髓系/DC代谢—调控核心推进到双外部队列机制复现，并为Treg核心提供单外部队列支持；该升级限定在机制层，不改变F1/TYMP直接临床预测仍为D级。"
    )

    document.core_properties.modified = datetime.now(timezone.utc)
    document.save(OUTPUT)

    reopened = Document(OUTPUT)
    text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    table_text = "\n".join(cell.text for table in reopened.tables for row in table.rows for cell in row.cells)
    required = [
        "C07  免疫程序外部直接复现与证据升级",
        "50,531个Myeloid/DC细胞和6,341个Treg细胞",
        "18个在两个外部队列均达到R2",
        "A=5、B=6、C1=2、D=2",
        "F1和TYMP作为独立临床预测指标均继续保持D级",
    ]
    missing = [item for item in required if item not in text and item not in table_text]
    if missing:
        raise RuntimeError(f"Required report text missing after save: {missing}")
    if "•" in text or "•" in table_text:
        raise RuntimeError("Solid bullet character found in report")
    output_media = media_hashes(OUTPUT)
    if source_media != output_media:
        raise RuntimeError("Embedded media changed during report update")

    status = {
        "status": "REPORT_UPDATED_PENDING_RENDER_QC",
        "source": str(SOURCE),
        "source_sha256_before": source_hash,
        "source_sha256_after": sha256(SOURCE),
        "output": str(OUTPUT),
        "output_bytes": OUTPUT.stat().st_size,
        "output_sha256": sha256(OUTPUT),
        "paragraphs": len(reopened.paragraphs),
        "tables": len(reopened.tables),
        "media_parts_preserved_byte_for_byte": len(output_media),
        "evidence_grade_counts": counts,
        "required_text_checks": {item: item in text or item in table_text for item in required},
        "edits": [
            "updated version note, executive summary and statistical rule",
            "added C07 external direct immune-program replication section",
            "added one bordered evidence-grade table",
            "updated supported, unsupported and final-conclusion paragraphs",
            "created 15-row immune evidence register v3",
        ],
    }
    (QC / "report_update_status_v1.json").write_text(
        json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
