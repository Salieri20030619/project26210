from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "reference" / "Project0106_report_20260430_reference.docx"
OUT = ROOT / "reference_inspection_v1"


def run_props(run):
    font = run.font.name
    if not font and run._element.rPr is not None and run._element.rPr.rFonts is not None:
        font = run._element.rPr.rFonts.get(qn("w:eastAsia")) or run._element.rPr.rFonts.get(qn("w:ascii"))
    size = run.font.size.pt if run.font.size else None
    color = run.font.color.rgb if run.font.color and run.font.color.rgb else None
    return font, size, str(color) if color else None, bool(run.bold), bool(run.italic)


def main():
    OUT.mkdir(exist_ok=True)
    doc = Document(SOURCE)
    paras = []
    styles = Counter()
    fonts = Counter()
    for i, p in enumerate(doc.paragraphs, start=1):
        text = p.text.strip()
        if not text:
            continue
        style = p.style.name if p.style else ""
        styles[style] += 1
        run_summary = []
        for run in p.runs:
            if not run.text.strip():
                continue
            props = run_props(run)
            fonts[props] += len(run.text)
            run_summary.append({
                "text": run.text,
                "font": props[0],
                "size_pt": props[1],
                "color": props[2],
                "bold": props[3],
                "italic": props[4],
            })
        paras.append({"index": i, "style": style, "text": text, "runs": run_summary})

    tables = []
    for ti, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": ti, "rows": len(table.rows), "cols": len(table.columns), "content": rows})

    sections = []
    for si, section in enumerate(doc.sections, start=1):
        sections.append({
            "index": si,
            "page_width_in": section.page_width.inches,
            "page_height_in": section.page_height.inches,
            "top_margin_in": section.top_margin.inches,
            "right_margin_in": section.right_margin.inches,
            "bottom_margin_in": section.bottom_margin.inches,
            "left_margin_in": section.left_margin.inches,
        })

    summary = {
        "source": str(SOURCE),
        "paragraphs_nonempty": len(paras),
        "tables": len(tables),
        "inline_shapes": len(doc.inline_shapes),
        "sections": sections,
        "paragraph_styles": styles.most_common(),
        "dominant_run_formats_by_characters": [
            {
                "font": k[0], "size_pt": k[1], "color": k[2], "bold": k[3], "italic": k[4], "characters": v
            }
            for k, v in fonts.most_common(20)
        ],
    }
    (OUT / "reference_structure_v1.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    with (OUT / "reference_paragraphs_v1.tsv").open("w", encoding="utf-8", newline="") as handle:
        handle.write("index\tstyle\ttext\n")
        for p in paras:
            text = p["text"].replace("\t", " ").replace("\r", " ").replace("\n", " ")
            handle.write(f'{p["index"]}\t{p["style"]}\t{text}\n')
    (OUT / "reference_tables_v1.json").write_text(
        json.dumps(tables, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
