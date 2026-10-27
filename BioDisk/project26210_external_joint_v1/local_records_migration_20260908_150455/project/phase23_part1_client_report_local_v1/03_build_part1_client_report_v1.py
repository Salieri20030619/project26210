from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets_v1"
OUTPUT = ROOT / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_v1.docx"
STATUS = ROOT / "build_status_v1.json"

FONT = "Microsoft YaHei"
FONT_SERIF = "SimSun"
NAVY = "18324A"
BLUE = "2E6F9E"
TEAL = "2D8C87"
GOLD = "B88A2A"
INK = "202B33"
MUTED = "64727D"
LIGHT_BLUE = "EAF3F8"
LIGHT_TEAL = "E9F5F3"
LIGHT_GOLD = "FBF4E4"
LIGHT_GRAY = "F3F5F7"
WHITE = "FFFFFF"
BORDER = "CCD6DD"
CONTENT_DXA = 9360


def set_run(run, *, size=None, bold=None, italic=None, color=INK, font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def paragraph_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run(run, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, NAVY, 8, 4),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def configure_section(section):
    # Named override: A4 customer report with exactly 6.5 inches usable width.
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(0.90)
    section.bottom_margin = Inches(0.90)
    side_margin = (21.0 / 2.54 - 6.5) / 2
    section.left_margin = Inches(side_margin)
    section.right_margin = Inches(side_margin)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.45)
    section.different_first_page_header_footer = True

    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    set_run(hp.add_run("project26210  |  外部单细胞数据验证"), size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    set_run(fp.add_run("—  "), size=9, color=MUTED)
    add_page_field(fp)
    set_run(fp.add_run("  —"), size=9, color=MUTED)


def add_title_page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run("\n\n\n")

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    set_run(kicker.add_run("EXTERNAL SINGLE-CELL VALIDATION"), size=10, bold=True, color=GOLD)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    set_run(title.add_run("外部单细胞数据验证报告"), size=28, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(26)
    set_run(subtitle.add_run("F1、TYMP、SPP1 髓系状态及肿瘤微环境关联分析"), size=14, color=BLUE)

    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.space_after = Pt(36)
    set_run(rule.add_run("━━━━━━━━━━━━━━━━━━━━"), size=10, color=GOLD)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(6)
    set_run(meta.add_run("项目：project26210"), size=11, bold=True, color=INK)
    meta2 = doc.add_paragraph()
    meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta2.paragraph_format.space_after = Pt(0)
    set_run(meta2.add_run("报告日期：2026年8月"), size=10.5, color=MUTED)
    doc.add_page_break()


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, size={1: 16, 2: 13, 3: 11.5}[level], bold=True,
                color={1: NAVY, 2: BLUE, 3: NAVY}[level])
    return p


def add_body(doc, text, *, bold_lead=None, keep=False):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True, color=NAVY)
        set_run(p.add_run(text[len(bold_lead):]))
    else:
        set_run(p.add_run(text))
    p.paragraph_format.keep_with_next = keep
    return p


def add_labeled(doc, label, text, *, result=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.04)
    p.paragraph_format.space_after = Pt(6)
    if result:
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), LIGHT_BLUE)
        p_pr.append(shd)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(9)
    set_run(p.add_run(f"{label}："), bold=True, color=TEAL if result else NAVY)
    set_run(p.add_run(text))
    return p


def add_callout(doc, label, text, fill=LIGHT_GOLD):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.10)
    p.paragraph_format.right_indent = Inches(0.10)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    set_run(p.add_run(f"{label}  "), bold=True, color=GOLD if fill == LIGHT_GOLD else TEAL)
    set_run(p.add_run(text), color=INK)
    return p


def add_table(doc, headers, rows, widths_dxa, numeric_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.keep_with_next = True
        set_run(p.add_run(str(value)), size=9, bold=True, color=WHITE)
    repeat_table_header(hdr)
    cant_split(hdr)
    for ridx, values in enumerate(rows):
        row = table.add_row()
        cant_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            if ridx % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            set_run(p.add_run(str(value)), size=8.8, color=INK)
    set_table_geometry(table, widths_dxa)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def font_path(bold=False):
    base = Path(r"C:\Windows\Fonts")
    candidates = ["msyhbd.ttc", "msyhbd.ttf"] if bold else ["msyh.ttc", "msyh.ttf"]
    for name in candidates:
        path = base / name
        if path.exists():
            return str(path)
    return str(base / "simhei.ttf")


def wrap_text(draw, text, font, max_width):
    lines, current = [], ""
    for char in text:
        trial = current + char
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def add_arrow(draw, start, end, color=BLUE, width=8):
    if not color.startswith("#"):
        color = f"#{color}"
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 22
    left = (end[0] - length * math.cos(angle - math.pi / 6), end[1] - length * math.sin(angle - math.pi / 6))
    right = (end[0] - length * math.cos(angle + math.pi / 6), end[1] - length * math.sin(angle + math.pi / 6))
    draw.polygon([end, left, right], fill=color)


def build_flowchart(path):
    def pc(value):
        return value if value.startswith("#") else f"#{value}"

    img = Image.new("RGB", (1800, 1120), pc(WHITE))
    draw = ImageDraw.Draw(img)
    title_font = ImageFont.truetype(font_path(True), 48)
    step_font = ImageFont.truetype(font_path(True), 30)
    text_font = ImageFont.truetype(font_path(False), 25)
    note_font = ImageFont.truetype(font_path(False), 24)
    draw.text((900, 58), "外部单细胞验证总体流程", font=title_font, fill=pc(NAVY), anchor="mm")
    draw.text((900, 112), "五个数据集全程独立处理，患者层面完成统计汇总", font=note_font, fill=pc(MUTED), anchor="mm")

    xs = [100, 650, 1200]
    ys = [180, 480, 780]
    box_w, box_h = 500, 210
    steps = [
        (0, 0, "01  数据输入", "GSE178481、GSE202374、GSE207493、GSE264586、GSE294109", LIGHT_BLUE, BLUE),
        (1, 0, "02  原始计数与元数据", "按数据集合并95个对象，统一患者、样本和细胞标识", LIGHT_BLUE, BLUE),
        (2, 0, "03  质控与聚类", "细胞质控、标准化、PCA、UMAP及独立聚类", LIGHT_BLUE, BLUE),
        (2, 1, "04  细胞注释", "建立广义类型、精细类型与置信度分层", LIGHT_TEAL, TEAL),
        (1, 1, "05  状态映射", "识别SPP1/C1QC髓系状态及F1高低状态", LIGHT_TEAL, TEAL),
        (0, 1, "06  CNV复核", "inferCNV与CopyKAT筛选保守恶性细胞集合", LIGHT_TEAL, TEAL),
        (0, 2, "07  程序评分", "投射18个预设转录程序并采用两种计分方法", LIGHT_GOLD, GOLD),
        (1, 2, "08  患者层面分析", "横断面关联、F1特异性、纵向变化与疗效分组", LIGHT_GOLD, GOLD),
        (2, 2, "09  证据综合", "按样本量、跨方法和跨队列一致性形成结论", LIGHT_GOLD, GOLD),
    ]
    coords = {}
    for col, row, title, body, fill, accent in steps:
        x, y = xs[col], ys[row]
        coords[(col, row)] = (x, y)
        draw.rounded_rectangle([x, y, x + box_w, y + box_h], radius=24, fill=pc(fill), outline=pc(accent), width=4)
        draw.rounded_rectangle([x, y, x + box_w, y + 58], radius=22, fill=pc(accent))
        draw.rectangle([x, y + 34, x + box_w, y + 58], fill=pc(accent))
        draw.text((x + 24, y + 29), title, font=step_font, fill=pc(WHITE), anchor="lm")
        lines = wrap_text(draw, body, text_font, box_w - 48)
        line_h = 39
        total_h = len(lines) * line_h
        yy = y + 74 + (box_h - 74 - total_h) / 2
        for line in lines:
            draw.text((x + box_w / 2, yy), line, font=text_font, fill=pc(INK), anchor="ma")
            yy += line_h

    gap = 25
    # Snake path: top row left->right, down, middle row right->left, down, bottom row left->right.
    add_arrow(draw, (600, 285), (650, 285))
    add_arrow(draw, (1150, 285), (1200, 285))
    add_arrow(draw, (1450, 390), (1450, 480))
    add_arrow(draw, (1200, 585), (1150, 585))
    add_arrow(draw, (650, 585), (600, 585))
    add_arrow(draw, (350, 690), (350, 780))
    add_arrow(draw, (600, 885), (650, 885))
    add_arrow(draw, (1150, 885), (1200, 885))

    draw.text((900, 1055), "输出：面向客户的跨队列证据排序与核心假设判断", font=note_font, fill=pc(NAVY), anchor="mm")
    img.save(path, dpi=(220, 220))


def set_image_alt(paragraph, title, description):
    for inline in paragraph._p.xpath(".//wp:inline"):
        doc_pr = inline.find(qn("wp:docPr"))
        if doc_pr is not None:
            doc_pr.set("title", title)
            doc_pr.set("descr", description)


def build_document():
    ASSETS.mkdir(exist_ok=True)
    flowchart = ASSETS / "part1_analysis_flowchart_v1.png"
    build_flowchart(flowchart)

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    core = doc.core_properties
    core.title = "外部单细胞数据验证报告"
    core.subject = "F1、TYMP、SPP1髓系状态及肿瘤微环境关联分析"
    core.author = "project26210 analysis team"
    core.keywords = "single-cell RNA-seq, F1, TYMP, SPP1, C1QC, CNV"
    core.comments = "Client-facing report generated from validated project26210 external single-cell outputs."

    add_title_page(doc)

    add_heading(doc, "报告摘要", 1)
    add_body(doc, "本研究整合五个外部单细胞RNA测序队列，在统一的数据处理框架下完成细胞质控、聚类注释、空间参考状态映射、CNV恶性细胞复核、转录程序评分和患者层面验证。五个队列始终独立处理，最终在患者层面比较不同生物学程序之间的关系。")
    add_callout(doc, "核心发现 1", "SPP1髓系状态与血管重塑的正向关系最一致。四个达到10例的队列在两种计分方法下均为正向，GSE202374的6例方向也相同。", LIGHT_TEAL)
    add_callout(doc, "核心发现 2", "F1与C1QC髓系状态呈一致正向关系。GSE207493的11例患者中，两种计分方法的相关系数分别为0.618和0.645；GSE294109的6例结果方向相同。", LIGHT_BLUE)
    add_callout(doc, "核心发现 3", "TYMP—SPP1关系不支持原定正向假设。GSE207493的11例结果在两种方法下均为负向，GSE294109的小样本结果在两种方法间方向不一致。", LIGHT_GOLD)
    add_body(doc, "统计以患者为独立单位。有效患者数达到10例时报告相关系数并结合FDR和方法一致性判断；少于10例时仅报告方向。达到10例的关联结果中，没有一项达到FDR<0.05。")

    add_heading(doc, "1. 项目概述", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_body(doc, "空间转录组结果提示，肿瘤细胞内F1/TYMP相关程序、SPP1或C1QC髓系状态以及血管和T细胞功能之间可能存在协同变化。外部单细胞数据能够把这些信号拆分到恶性细胞、髓系细胞、T/NK细胞和血管相关细胞中，从独立患者队列检验主要关系是否重复出现。")
    add_heading(doc, "1.2 研究目标", 2)
    add_body(doc, "本部分围绕三个目标展开：建立可追溯的外部单细胞参考；在CNV复核后的恶性细胞中量化F1和TYMP程序；在患者层面评估这些程序与SPP1/C1QC髓系状态、血管重塑、T细胞功能障碍、细胞毒性及代谢背景之间的关系。")
    data_heading = add_heading(doc, "1.3 数据来源", 2)
    data_heading.paragraph_format.page_break_before = True
    add_table(
        doc,
        ["数据集", "患者数", "样本数", "质控后细胞数", "主要设计信息"],
        [
            ["GSE178481", "11", "26", "156,310", "多样本肾癌队列"],
            ["GSE202374", "6", "6", "39,332", "横断面肾癌队列"],
            ["GSE207493", "19", "19", "160,710", "横断面肾癌队列"],
            ["GSE264586", "12", "19", "88,152", "含Cycle 2和治疗结束随访"],
            ["GSE294109", "15", "25", "140,615", "含疗效分组信息"],
        ],
        [1600, 1100, 1100, 1700, 3860],
        numeric_cols=(1, 2, 3),
    )
    add_body(doc, "五个队列共纳入63名患者、95个样本。原始对象合并后共有623,431个细胞，质控后保留585,119个细胞；用于程序评分的目标区室RDS包含505,639个细胞。")
    add_heading(doc, "1.4 统计口径", 2)
    add_table(
        doc,
        ["项目", "报告规则"],
        [
            ["独立单位", "患者；同一患者的多个样本或细胞不重复作为统计单位"],
            ["有效患者数 ≥10", "报告相关系数，并结合FDR、两种计分方法及跨队列一致性判断"],
            ["有效患者数 <10", "只报告方向，不展示或解读P值、FDR、置信区间和效应量"],
            ["多重检验", "在每个数据集和计分方法内进行BH-FDR校正"],
        ],
        [2100, 7260],
    )

    add_heading(doc, "2. 分析流程", 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(flowchart), width=Inches(6.45))
    set_image_alt(p, "外部单细胞验证总体流程", "五个外部数据集经独立合并、质控聚类、细胞注释、状态映射、CNV复核、程序评分和患者层面分析后形成证据综合。")
    cap = doc.add_paragraph("图1 外部单细胞数据验证总体流程", style="Caption")
    add_body(doc, "分析从原始计数和患者元数据出发，先建立稳定的细胞类型与状态框架，再利用CNV结果筛选恶性细胞，随后将预设转录程序投射到相应细胞区室，并在患者层面完成横断面、配对和疗效分组分析。")

    add_heading(doc, "3. 分析结果", 1)

    add_heading(doc, "3.1 原始计数合并与身份追踪", 2)
    add_labeled(doc, "为什么做", "五个公开队列由95个独立对象组成，首先需要统一细胞、样本和患者标识，形成可重复追踪的分析输入。")
    add_labeled(doc, "怎么做", "仅使用确认的原始计数，按数据集分别合并；细胞ID由数据集、对象名和原始条形码共同构成，并保留原始元数据字段。")
    add_labeled(doc, "得到什么", "95个对象共623,431个细胞全部完成计数守恒核对，生成5个彼此独立的原始计数对象。计数均为非负、有限整数，细胞ID全局唯一。", result=True)

    add_heading(doc, "3.2 患者映射与细胞质量控制", 2)
    add_labeled(doc, "为什么做", "后续检验以患者为统计单位，因此必须先确认每个对象对应的患者、样本及随访时间点，并去除低质量细胞。")
    add_labeled(doc, "怎么做", "完成95个对象的患者与样本交叉表；在各数据集内按检测基因数、总计数和线粒体比例进行质控，随后分别完成标准化和特征选择。")
    add_labeled(doc, "得到什么", "63名患者和95个样本全部获得明确映射。质控后保留585,119个细胞，去除38,312个低质量细胞；五个数据集分别保留156,310、39,332、160,710、88,152和140,615个细胞。", result=True)

    add_heading(doc, "3.3 降维、聚类与聚类质量评估", 2)
    add_labeled(doc, "为什么做", "无监督聚类用于发现数据集中主要细胞群和状态，为后续细胞类型注释及程序评分建立结构基础。")
    add_labeled(doc, "怎么做", "各数据集分别完成PCA、邻近图构建、0.4/0.8双分辨率聚类和UMAP展示，并结合细胞数量、患者覆盖和样本来源评估聚类结构。")
    add_labeled(doc, "得到什么", "主分辨率0.8下共获得178个聚类。十组数据集—分辨率组合均通过细胞数守恒检查，UMAP结果完整展示了各队列的主要细胞结构。", result=True)

    add_heading(doc, "3.4 细胞类型注释", 2)
    add_labeled(doc, "为什么做", "F1/TYMP信号、髓系状态和T细胞功能需要在正确的细胞区室内解释，因此需要建立统一且可追溯的细胞类型体系。")
    add_labeled(doc, "怎么做", "综合每个聚类的标志基因、经典marker评分、患者覆盖和组织来源，建立广义类型、精细类型及注释置信度。")
    add_labeled(doc, "得到什么", "585,119个细胞全部完成注释，共得到12类广义细胞类型和31类精细标签。用于本次程序验证的目标区室包含505,639个细胞，其中T/NK细胞270,774个、髓系细胞120,022个、候选恶性细胞40,302个、正常上皮细胞34,691个、内皮细胞30,480个、成纤维细胞9,370个。", result=True)

    add_heading(doc, "3.5 空间参考状态映射", 2)
    add_labeled(doc, "为什么做", "空间分析中的F1高低状态、SPP1髓系状态和C1QC髓系状态需要在外部单细胞数据中找到对应的细胞级表达模式。")
    add_labeled(doc, "怎么做", "使用冻结的空间参考定义投射14类参考状态；对髓系细胞重新计算SPP1和C1QC状态分数，对候选恶性细胞计算F1 top-100加权分数。")
    add_labeled(doc, "得到什么", "71,045个巨噬细胞相关细胞被分为SPP1状态15,322个、C1QC状态15,294个和其他状态40,429个。40,302个候选恶性细胞中，10,079个为F1高状态、30,143个为F1低状态、80个未分型。", result=True)

    add_heading(doc, "3.6 CNV恶性细胞复核", 2)
    add_labeled(doc, "为什么做", "上皮样表达可提示肿瘤表型，但恶性细胞分析需要进一步结合拷贝数变化证据。")
    add_labeled(doc, "怎么做", "对适合的患者和样本分别运行inferCNV和CopyKAT，并将既有精确细胞CNV结果与新计算结果按固定优先级合并，形成保守恶性细胞集合。")
    add_labeled(doc, "得到什么", "40,302个候选恶性细胞中，16,645个进入最终恶性细胞集合。GSE207493和GSE294109提供的恶性细胞数量最多，是恶性细胞程序关联分析的主要信息来源。", result=True)
    add_table(
        doc,
        ["数据集", "候选恶性细胞", "最终恶性细胞", "覆盖患者数"],
        [
            ["GSE178481", "3,907", "238", "3"],
            ["GSE202374", "1,835", "1,557", "4"],
            ["GSE207493", "21,448", "8,589", "17"],
            ["GSE264586", "6,698", "1,817", "5"],
            ["GSE294109", "6,414", "4,444", "8"],
            ["合计", "40,302", "16,645", "—"],
        ],
        [2200, 2300, 2300, 2560],
        numeric_cols=(1, 2, 3),
    )

    add_heading(doc, "3.7 转录程序构建与评分", 2)
    add_labeled(doc, "为什么做", "客户假设涉及多个细胞区室和生物学过程，需要把每个假设转化为可跨队列比较的患者层面数值。")
    add_labeled(doc, "怎么做", "在接触外部结局前冻结18个转录程序，共213条基因记录；分别采用基因Z分数均值和log-normalized均值进行主分析与敏感性分析，再在指定细胞区室内汇总到患者。")
    add_labeled(doc, "得到什么", "五个队列的目标区室均完成程序投射。F1 top-100加权分数与原始参考分数的最大绝对误差不超过1×10⁻⁸，两种计分方法共同进入后续结果判断。", result=True)

    add_heading(doc, "3.8 患者层面核心关联分析", 2)
    add_labeled(doc, "为什么做", "核心目标是判断恶性细胞F1/TYMP程序是否与髓系状态和代谢背景协同变化，以及SPP1髓系状态是否伴随血管重塑和T细胞功能改变。")
    add_labeled(doc, "怎么做", "在每个数据集内计算患者层面的Spearman相关，并使用两种计分方法复核。有效患者数达到10例时报告相关系数；少于10例时仅报告方向。")
    add_labeled(doc, "得到什么", "120个“数据集×方法×关系”组合中，42个达到10例、24个按方向报告、54个因有效患者不足未检验。达到10例的结果中没有FDR<0.05。", result=True)

    add_heading(doc, "3.8.1 恶性细胞程序与髓系状态", 3)
    add_body(doc, "该组关系只有GSE207493达到10例，共11名患者。F1—C1QC为最清楚的正向关系；F1—SPP1和TYMP—SPP1均为负向。")
    add_table(
        doc,
        ["关系", "主分析 rho", "敏感性分析 rho", "结果方向"],
        [
            ["F1 → SPP1髓系状态", "-0.555", "-0.582", "负向"],
            ["TYMP → SPP1髓系状态", "-0.409", "-0.336", "负向"],
            ["F1 → APOE/C1QC髓系状态", "0.327", "0.273", "正向"],
            ["F1 → C1QC髓系状态", "0.618", "0.645", "正向，幅度最大"],
            ["F1 → HLA-II髓系状态", "-0.373", "-0.364", "负向"],
            ["F1 → 缺氧/血管生成", "0.318", "0.209", "正向"],
            ["TYMP → 缺氧/血管生成", "-0.345", "-0.273", "负向"],
            ["F1 → 多胺/MTA背景", "0.100", "0.255", "正向"],
            ["TYMP → 多胺/MTA背景", "0.318", "0.273", "正向"],
        ],
        [3500, 1600, 1900, 2360],
        numeric_cols=(1, 2),
    )
    add_body(doc, "GSE294109中有6名患者可用于同组关系的计算。方向上，F1—C1QC、F1—缺氧/血管生成和TYMP—多胺/MTA均为正向；TYMP—缺氧/血管生成为负向；TYMP—SPP1在两种方法间方向不一致。")

    add_heading(doc, "3.8.2 SPP1髓系状态与微环境", 3)
    add_table(
        doc,
        ["数据集", "患者数", "血管重塑 rho\n主/敏感性", "T细胞功能障碍 rho\n主/敏感性", "细胞毒性 rho\n主/敏感性"],
        [
            ["GSE178481", "10", "0.430 / 0.127", "0.612 / 0.564", "-0.127 / -0.188"],
            ["GSE207493", "19", "0.174 / 0.165", "0.089 / -0.023", "0.037 / 0.009"],
            ["GSE264586", "12", "0.434 / 0.455", "0.441 / 0.427", "-0.126 / -0.119"],
            ["GSE294109", "14", "0.081 / 0.152", "0.103 / 0.143", "-0.385 / -0.407"],
        ],
        [1450, 950, 2300, 2460, 2200],
        numeric_cols=(1, 2, 3, 4),
    )
    add_body(doc, "SPP1—血管重塑在四个达到10例的队列及两种方法下全部为正向；GSE202374的6例结果方向也相同。SPP1—T细胞功能障碍在GSE178481、GSE264586和GSE294109中为正向，在GSE207493中接近零。SPP1—细胞毒性在三个达到10例的队列中为负向，在GSE207493中接近零且略偏正。")

    add_heading(doc, "3.9 F1恶性细胞特异性复核", 2)
    add_labeled(doc, "为什么做", "如果F1代表恶性细胞相关程序，同一患者的恶性细胞F1水平应整体高于正常上皮细胞。")
    add_labeled(doc, "怎么做", "在同时具有足量恶性细胞和正常上皮细胞的患者中进行配对比较，并分别用加权F1和未加权F1评分复核。")
    add_labeled(doc, "得到什么", "各队列可配对患者数均少于10，因此按方向读取。GSE207493的9对患者、GSE294109的3对患者均显示恶性细胞F1偏高；GSE178481的2对患者总体也偏正向。", result=True)

    add_heading(doc, "3.10 治疗过程中的纵向变化", 2)
    add_labeled(doc, "为什么做", "纵向样本可观察同一患者在治疗过程中髓系状态、T细胞功能和血管相关程序是否同步变化。")
    add_labeled(doc, "怎么做", "在GSE264586中将基线分别与Cycle 2和治疗结束样本配对，比较患者内程序分数变化。")
    add_labeled(doc, "得到什么", "Cycle 2有4对患者，治疗结束有3对患者，均按方向读取。Cycle 2时SPP1髓系状态上升，C1QC状态和细胞毒性下降；治疗结束时SPP1与C1QC状态下降，细胞毒性继续下降，IFN反应和单核细胞程序上升。", result=True)

    add_heading(doc, "3.11 疗效分组比较", 2)
    add_labeled(doc, "为什么做", "疗效分组可用于判断基线或整体微环境程序是否与临床获益方向相联系。")
    add_labeled(doc, "怎么做", "在GSE294109中比较获益组与无效组的患者层面程序分数，并用两种计分方法核对方向。")
    add_labeled(doc, "得到什么", "无效组最多为3名患者，因此按方向读取。获益组的SPP1髓系状态、单核细胞程序和侵袭/EMT程序偏高；C1QC、HLA-II、IFN反应、T细胞功能障碍和细胞毒性偏低。", result=True)

    add_heading(doc, "4. 结果综合", 1)
    add_table(
        doc,
        ["客户假设", "综合判断", "主要依据"],
        [
            ["SPP1髓系状态伴随血管重塑增强", "支持度最高", "五个队列方向一致，其中四个队列达到10例"],
            ["F1关联C1QC/APOE型髓系状态", "方向支持", "GSE207493达到10例且两种方法一致；GSE294109小样本方向相同"],
            ["F1关联缺氧/血管生成", "方向支持", "GSE207493为正向；GSE294109小样本方向相同"],
            ["TYMP关联多胺/MTA背景", "方向支持", "GSE207493为正向；GSE294109小样本方向相同"],
            ["SPP1髓系状态伴随T细胞功能障碍", "多数队列支持", "三个达到10例的队列为正向，GSE207493接近零"],
            ["SPP1髓系状态伴随细胞毒性降低", "队列间不一致", "三个达到10例的队列为负向，GSE207493接近零且略偏正"],
            ["F1关联SPP1髓系状态", "未形成稳定正向关系", "GSE207493为负向，GSE294109小样本为正向"],
            ["TYMP关联SPP1髓系状态", "不支持原假设", "GSE207493为负向，GSE294109小样本在方法间方向不一致"],
        ],
        [2700, 1900, 4760],
    )

    add_heading(doc, "5. 项目结论", 1)
    add_body(doc, "外部单细胞结果最集中地支持“SPP1髓系状态—血管重塑”轴，并给出“F1—C1QC髓系状态”的一致正向信号。SPP1与T细胞功能障碍在多数队列中呈正向，但在最大队列中接近零；SPP1与细胞毒性的关系存在队列差异。")
    add_body(doc, "与客户核心假设直接相关的F1—SPP1和TYMP—SPP1没有形成稳定正向关系。GSE207493中二者均为负向，其中TYMP—SPP1在GSE294109的小样本结果中还表现出方法间方向不一致。综合现有外部单细胞证据，TYMP—SPP1轴的支持度下降，而SPP1—血管重塑和F1—C1QC轴更值得作为后续空间验证和机制讨论的重点。")

    add_heading(doc, "附录：分析版本与数据依据", 1)
    add_body(doc, "本报告基于五个Phase 14细胞评分RDS、最终CNV v6汇总表及患者层面关联结果重新整理。RDS逐一读取并完成身份、行数、列数、MD5、患者覆盖、样本覆盖和细胞ID唯一性检查。")
    add_table(
        doc,
        ["项目", "版本/结果"],
        [
            ["R", "4.5.2"],
            ["细胞评分对象", "Phase 14 run 20260812_082000_provisional_v3"],
            ["最终CNV门控", "20260812_final_cnv_gated_v6"],
            ["患者层面分析", "Final v6；报告口径按患者数10例重新分层"],
            ["计分方法", "主分析：基因Z分数均值；敏感性分析：log-normalized均值"],
        ],
        [2600, 6760],
    )

    doc.save(OUTPUT)
    sha = hashlib.sha256(OUTPUT.read_bytes()).hexdigest()
    status = {
        "status": "PASS_BUILD_COMPLETE",
        "output": str(OUTPUT),
        "sha256": sha,
        "flowchart": str(flowchart),
        "design_preset": "narrative_proposal",
        "named_overrides": [
            "A4_CLIENT_REPORT_OVERRIDE: A4 portrait with 6.5-inch usable width",
            "CJK_FONT_OVERRIDE: Microsoft YaHei for all Chinese text",
            "EDITORIAL_COVER_OVERRIDE: centered research-report cover",
        ],
        "figures": 1,
        "data_result_figures": 0,
    }
    STATUS.write_text(json.dumps(status, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(status, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build_document()
