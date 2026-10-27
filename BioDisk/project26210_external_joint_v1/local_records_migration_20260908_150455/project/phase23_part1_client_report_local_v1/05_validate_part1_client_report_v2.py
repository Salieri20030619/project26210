from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_v2.docx"
BUILD = ROOT / "build_status_v2.json"
RENDER = ROOT / "render_v2" / "render_status.json"
OUT = ROOT / "validation_v2"
SOURCE_STATUS = ROOT.parent / "phase22_part1_rds_report_local_v1" / "reporting_tables_v2" / "reporting_tables_status_v2.json"


def table_geometry_ok(table):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    grid = table._tbl.tblGrid
    grid_widths = [int(c.get(qn("w:w"))) for c in grid]
    if tbl_w is None or tbl_ind is None:
        return False
    if int(tbl_w.get(qn("w:w"))) != 9360 or int(tbl_ind.get(qn("w:w"))) != 120:
        return False
    if sum(grid_widths) != 9360:
        return False
    for row in table.rows:
        widths = []
        for cell in row.cells:
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                return False
            widths.append(int(tc_w.get(qn("w:w"))))
        if widths != grid_widths:
            return False
    return True


def main():
    OUT.mkdir(exist_ok=True)
    doc = Document(DOCX)
    text_parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            text_parts.extend(cell.text for cell in row.cells)
    text = "\n".join(text_parts)
    build = json.loads(BUILD.read_text(encoding="utf-8"))
    render = json.loads(RENDER.read_text(encoding="utf-8"))
    source = json.loads(SOURCE_STATUS.read_text(encoding="utf-8"))

    heading_counts = {
        name: sum(1 for p in doc.paragraphs if p.style and p.style.name == name)
        for name in ("Heading 1", "Heading 2", "Heading 3")
    }
    image_alt = []
    for inline in doc.element.body.xpath(".//wp:inline"):
        doc_pr = inline.find(qn("wp:docPr"))
        image_alt.append((doc_pr.get("title", ""), doc_pr.get("descr", "")) if doc_pr is not None else ("", ""))

    with zipfile.ZipFile(DOCX) as zf:
        names = set(zf.namelist())
        document_xml = zf.read("word/document.xml").decode("utf-8")

    pngs = sorted((ROOT / "render_v2").glob("page-*.png"))
    checks = {
        "docx_exists_nonempty": DOCX.exists() and DOCX.stat().st_size > 100_000,
        "build_status_pass": build["status"] == "PASS_BUILD_COMPLETE",
        "source_reporting_tables_pass": source["status"] == "PASS_REPORTING_TABLES_COMPLETE",
        "patient_threshold_10": source["minimum_units_for_inferential_reporting"] == 10,
        "render_status_pass": render["status"] == "PASS" and render["document_pages"] == 10,
        "all_rendered_pages_present": len(pngs) == 10 and all(p.stat().st_size > 10_000 for p in pngs),
        "flowchart_embedded_with_alt_text": len(image_alt) == 1 and all(image_alt[0]),
        "real_heading_styles_present": heading_counts == {"Heading 1": 7, "Heading 2": 15, "Heading 3": 2},
        "all_table_geometry_exact": len(doc.tables) == 7 and all(table_geometry_ok(t) for t in doc.tables),
        "required_analysis_labels": all(text.count(label) >= 11 for label in ("为什么做：", "怎么做：", "得到什么：")),
        "required_conclusions": all(
            phrase in text
            for phrase in ("SPP1髓系状态与血管重塑", "F1与C1QC髓系状态", "TYMP—SPP1关系不支持")
        ),
        "small_sample_rule_stated": "少于10例时仅报告方向" in text and "有效患者数 <10" in text,
        "no_small_n_inferential_values": all(token not in text for token in ("0.0188", "0.0416", "0.02797")),
        "no_defensive_phrases": all(
            phrase not in text for phrase in ("不能证明", "不能替代", "不能表述", "尚不足以", "证据边界", "局限性")
        ),
        "no_comments_or_tracked_changes": "word/comments.xml" not in names and "<w:ins" not in document_xml and "<w:del" not in document_xml,
        "no_placeholder_text": all(token not in text for token in ("TODO", "TBD", "待补", "此处插入")),
    }
    overall = "PASS" if all(checks.values()) else "FAIL"
    result = {
        "schema_version": "part1_client_report_validation_v2",
        "status": overall,
        "docx": str(DOCX),
        "sha256": hashlib.sha256(DOCX.read_bytes()).hexdigest(),
        "checks_passed": sum(checks.values()),
        "checks_total": len(checks),
        "pages_visually_inspected": 10,
        "flowchart_count": len(image_alt),
        "table_count": len(doc.tables),
        "heading_counts": heading_counts,
        "review_state": "AWAITING_USER_REVIEW",
        "checks": checks,
    }
    (OUT / "PART1_CLIENT_REPORT_VALIDATION_STATUS_v2.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if overall != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
