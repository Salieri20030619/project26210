from __future__ import annotations

import importlib.util
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE_SCRIPT = Path(r"C:\Users\31937\project26210_immune_report_rebuild_v1\scripts\02_build_report_v1.py")
spec = importlib.util.spec_from_file_location("immune_report_base", BASE_SCRIPT)
base = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(base)


def add_callout(doc, title, text, fill=base.CALLOUT, accent=base.BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "22")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    r1 = p.add_run(title)
    base.set_run_font(r1, size=10.5, bold=True, color=accent)
    r1.add_break()
    r2 = p.add_run(text)
    base.set_run_font(r2, size=10.2, color=base.INK)
    return p


def add_labelled(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_together = True
    r1 = p.add_run(label)
    base.set_run_font(r1, size=10.7, bold=True, color=base.DARK_BLUE)
    r2 = p.add_run(text)
    base.set_run_font(r2, size=10.7, color=base.INK)
    return p


def add_figure_guide_item(doc, label, text):
    """Compact panel-level legend text for a dedicated caption page."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_together = True
    r1 = p.add_run(label)
    base.set_run_font(r1, size=8.8, bold=True, color=base.DARK_BLUE)
    r2 = p.add_run(text)
    base.set_run_font(r2, size=8.8, color=base.INK)
    return p

def add_module(doc, title, why, data, result, judgement):
    base.add_heading(doc, title, 2)
    p1 = add_labelled(doc, "为什么做：", why)
    p2 = add_labelled(doc, "使用数据与统计单位：", data)
    p3 = add_labelled(doc, "结果：", result)
    p1.paragraph_format.keep_with_next = True
    p2.paragraph_format.keep_with_next = True
    p3.paragraph_format.keep_with_next = True
    add_callout(doc, "证据判定", judgement, fill=base.CALLOUT, accent=base.TEAL)
    return p3


def add_cover(doc, now):
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(14)
    base.add_paragraph(
        doc,
        "外部免疫多组学专题报告",
        size=11,
        color=base.GOLD,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=18,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("肾透明细胞癌")
    base.set_run_font(r, size=29, bold=True, color=base.NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(10)
    r2 = p2.add_run("免疫机制与转化证据整合分析报告")
    base.set_run_font(r2, size=21, bold=True, color=base.NAVY)
    base.add_paragraph(
        doc,
        "从免疫细胞来源与程序复现，到空间、扰动、bulk可辨识性与临床验证",
        size=12.5,
        color=base.TEAL,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=36,
    )
    add_callout(
        doc,
        "报告范围",
        "本报告仅覆盖免疫专题分析，不纳入旧版恶性细胞、CNV、cNMF等非免疫章节。正文只使用分析名称，不展示内部工作编号。不重复拆分主图Panel；另加入4张主图未覆盖的代表性图件，并避开条形图和箱线图。",
        fill=base.SOFT_TEAL,
        accent=base.TEAL,
    )
    base.add_paragraph(
        doc,
        f"报告日期：{now.strftime('%Y年%m月%d日')}  |  客户交付版",
        size=10,
        color=base.MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=6,
    )
    base.add_paragraph(
        doc,
        "所有推断以患者、配对患者或空间切片为独立统计单位。",
        size=9.5,
        color=base.MUTED,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    base.page_break(doc)


def configure_report_styles(doc):
    base.configure_styles(doc)
    normal = doc.styles["Normal"]
    base.set_style_font(normal, 11, base.INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name in ["List Bullet", "List Bullet 2", "List Number"]:
        style = doc.styles[name]
        base.set_style_font(style, 11, base.INK)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.10


def assert_sources(project_root: Path, main_figure: Path, representative: Path):
    required = [
        main_figure,
        representative / "immune_pathway_remodeling_heatmap.png",
        representative / "clinical_validation_forest.png",
        representative / "virtual_knockout_enrichment.png",
        representative / "bulk_identifiability_diagnostic.png",
        project_root / "phase63_part1_immune_evidence_convergence_local_v1" / "IMMUNE_EVIDENCE_REGISTER_v1.tsv",
        project_root / "phase65_part1_immune_mechanistic_convergence_local_v1" / "results_v2" / "METABOLIC_ENDPOINT_SUMMARY_v1.tsv",
        project_root / "phase67_part1_direct_program_replication_local_v1" / "results_v5" / "phase67_external_replication_summary_v1.tsv",
        project_root / "phase69_85_server_status_sync_local_v1" / "phase70" / "results_v1" / "composition_adjusted_bridge_meta_v1.tsv",
        project_root / "phase69_85_server_status_sync_local_v1" / "phase71" / "results_v2" / "phase71_relation_summary_all_v1.tsv",
        project_root / "phase69_85_server_status_sync_local_v1" / "phase81" / "results_v1" / "phase81_primary_models_v1.tsv",
        project_root / "phase86_macrophage_state_identifiability_local_v1" / "results_v1" / "PHASE86_IDENTIFIABILITY_SUMMARY_v1.md",
        project_root / "phase87_three_cohort_myeloid_trajectory_local_v1" / "results_v1" / "phase87_dataset_method_endpoint_summary_v1.tsv",
        project_root / "phase88_three_cohort_fine_state_cellchat_local_v1" / "results_v1" / "phase88_dataset_endpoint_summary_v1.tsv",
    ]
    missing = [str(path) for path in required if not path.exists() or path.stat().st_size == 0]
    if missing:
        raise SystemExit("Missing required source files:\n" + "\n".join(missing))


def main():
    if len(sys.argv) != 5:
        raise SystemExit("Usage: 01_build_full_immune_report_v1.py <project_root> <main_figure_png> <representative_figures_dir> <output_docx>")
    project_root = Path(sys.argv[1])
    main_figure = Path(sys.argv[2])
    representative = Path(sys.argv[3])
    output = Path(sys.argv[4])
    if output.exists():
        raise SystemExit(f"Refusing existing output: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    assert_sources(project_root, main_figure, representative)
    now = datetime.now().astimezone()

    doc = Document()
    configure_report_styles(doc)
    base.configure_section(doc.sections[0])
    core = doc.core_properties
    core.title = "肾透明细胞癌免疫机制与转化证据整合分析报告"
    core.subject = "免疫细胞来源、外部复现、空间生态、扰动、bulk可辨识性与临床验证"
    core.author = "项目分析团队"
    core.keywords = "ccRCC; TYMP; immune; myeloid; macrophage; spatial; clinical validation"
    core.comments = "免疫专题完整重构版；正文不使用内部工作编号。"

    add_cover(doc, now)

    base.add_heading(doc, "执行摘要", 1)
    add_callout(
        doc,
        "核心结论",
        "最稳健的正向证据是TYMP在Myeloid/DC中的细胞来源、去除TYMP锚点后的代谢与调控程序在两套外部单细胞数据中的患者级复现，以及Monocyte到多个巨噬细胞状态的横断面排序。现有临床、空间、扰动、bulk分解和通讯证据均不足以把这一广义髓系背景升级为选择性SPP1极化、5-MTA因果介导或response/PFS/OS预测。",
        fill=base.SOFT_TEAL,
        accent=base.TEAL,
    )
    for text in [
        "免疫初筛和患者级伪批量表明，治疗时间与临床获益伴随广泛但高度细胞类型依赖的通路重塑；单基因差异稀疏，且低样本比较不能写成稳定疗效标志物。",
        "GSE207493的19位患者显示F1主要来自肿瘤细胞，TYMP在Myeloid/DC最强，并在Treg中呈较弱但可评估的富集。Myeloid/DC的代谢与转录调控核心随后在GSE264586和GSE294109直接复现。",
        "空间组成校正仅支持SPP1状态与血管重塑在FFPE和冰冻制备中同时出现；更精细的肿瘤TYMP邻近SPP1状态检验在两种制备中均不支持。",
        "虚拟敲除、MTAP自然扰动和外源MTA签名转移未建立两外部队列一致的TYMP-MTA-SPP1因果链；客户现有文件也缺少可重算的5-MTA定量矩阵和干预-救援设计。",
        "bulk中可以恢复粗粒度巨噬细胞总量，但无法可靠区分SPP1、C1QC和Other macrophage。细亚状态临床建模因此被明确禁止。",
    ]:
        base.add_paragraph(doc, text)

    base.add_heading(doc, "统计报告规则", 2)
    base.add_table(
        doc,
        ["独立样本数", "正文报告方式", "证据使用"],
        [
            ["n < 10", "报告效应方向；仅当名义P < 0.05时报告P值", "不以FDR或单队列结果升级为稳健证据"],
            ["n ≥ 10", "同时报告效应、P值和BH校正FDR（q）", "按预设家族和跨队列规则判定支持"],
            ["细胞数/spot数", "只描述数据规模，不作为独立重复", "患者、配对患者或切片才是推断单位"],
        ],
        [1500, 3820, 4040],
        font_size=9.2,
    )
    add_callout(
        doc,
        "解释原则",
        "方向一致不等于显著；单制备、单外部队列、低样本或代理签名只能支持方向性与可检验性，不能替代独立临床终点、直接代谢测量或因果实验。",
        fill=base.SOFT_MAGENTA,
        accent=base.MAGENTA,
    )
    base.page_break(doc)

    base.add_heading(doc, "数据与证据架构", 1)
    base.add_table(
        doc,
        ["数据来源", "规模", "在报告中的作用"],
        [
            ["GSE207493", "19位患者；全图谱160,710个细胞", "免疫细胞来源、代谢/调控程序发现、跨谱系与排序发现"],
            ["GSE264586", "12位患者；全图谱88,152个细胞", "免疫治疗纵向、外部程序复现、空间外的机制压力测试"],
            ["GSE294109", "15位患者；全图谱140,615个细胞", "免疫治疗获益比较、外部程序复现、机制与通讯复核"],
            ["GSE175540", "24张空间切片：12 FFPE + 12冰冻", "组成校正、代理链转移和细亚状态邻域检验"],
            ["E-MTAB-3218 / CheckMate / JAVELIN", "独立临床转录组；模型样本量按终点报告", "response/PFS和治疗交互验证"],
            ["TCGA-KIRC", "496例可用于MTAP自然扰动交互", "检验MTAP缺失是否改变TYMP相关免疫状态"],
        ],
        [1900, 2400, 5060],
        font_size=8.8,
    )
    base.add_heading(doc, "完整免疫分析主图", 2)
    base.add_figure(
        doc,
        main_figure,
        5.95,
        "图1 免疫专题分析总览。A至F的视觉编码、阅读顺序、关键数值和证据边界见下一页完整图注。",
        "免疫专题完整主图",
        "A至F复合主图：按数据架构、队列分离图谱、细胞来源、程序复现、髓系排序、空间和CellChat压力测试顺序阅读；关键统计量在相邻图注页逐项解释。",
    )
    base.add_heading(doc, "图1完整图注与读图指南", 2)
    add_figure_guide_item(doc, "阅读顺序：", "按A→B→C→D→E→F阅读。A-B交代证据层和队列，C-D从发现推进到外部复现，E-F再用状态排序、空间邻域和通讯特异性做压力测试。箭头表示分析流程，不表示因果；不同队列从未合并做细胞级推断。")
    add_figure_guide_item(doc, "A｜证据架构：", "四个方框分别是发现单细胞GSE207493（19位患者、160,710个细胞）、外部单细胞GSE264586+GSE294109（27位患者、228,767个细胞）、空间GSE175540（12张FFPE+12张冰冻切片）及临床CheckMate+JAVELIN（311+726份RNA样本）。方框间箭头只表示验证层级由近及远。")
    add_figure_guide_item(doc, "B｜三队列免疫图谱：", "颜色区分广义细胞谱系；三幅UMAP保留各队列原始坐标，不能按点密度直接比较队列丰度。GSE207493、GSE264586和GSE294109分别含19/12/15位患者和160,710/88,152/140,615个细胞，总计389,477个细胞。")
    add_figure_guide_item(doc, "C｜TYMP来源与发现程序：", "左侧特征图的颜色仅表示各自特征内的相对强度；右侧气泡填色为谱系内标准化水平，面积为细胞数。下方圆环表示31条代谢通路中通过周期校正BH FDR的比例，文字给出跨SCENIC/CollecTRI-VIPER一致TF数：F1-high肿瘤为22/31通路、49个TF，TYMP-high Myeloid/DC为17/31、40个TF，TYMP-high Treg为10/31、8个TF。")
    add_figure_guide_item(doc, "D｜外部程序复现：", "每行一个去锚代谢程序或TF regulon；两列分别为GSE264586（n=12）与GSE294109（n=13）。填色为残差Spearman rho（青色负、洋红正），圆面积为−log10(q)，同色连线表示两队列方向一致。18个Myeloid/DC签名在两队列均通过FDR，13个Treg签名仅在可评价的单一外部队列支持。例：缺氧反应rho=0.2083/0.2394、P=0.000488/0.000244、q=0.000868/0.000977；IRF1 regulon rho=0.3486/0.2981、P=0.000488/0.000244、q=0.00114/0.000570。")
    add_figure_guide_item(doc, "E｜髓系状态排序：", "上排颜色表示Monocyte、SPP1、C1QC和Other macrophage，黑色箭头为Slingshot方向；下排曲线是各程序沿拟时序的缩放中位数。共68,711个髓系细胞、45位患者、62个患者-样本单元。Slingshot与独立PCA-kNN测地拟时序的患者级中位rho为0.841、0.559和0.430；SPP1相对Monocyte的测地/Slingshot中位差分别为0.333/0.418、0.247/0.266、0.269/0.286，六项q均≤0.00684。该panel支持横断面排序，不证明真实谱系转换。")
    add_figure_guide_item(doc, "F｜空间与通讯压力测试：", "空间图两行分别为代表性FFPE和冰冻切片，列为肿瘤TYMP、SPP1、C1QC和内皮信号；每幅图独立缩放颜色，不能跨图比较绝对强度。24张切片中，SPP1状态-血管重塑在FFPE/冰冻均支持（rho=0.2121/0.2083，P均=0.000488，q=0.00220/0.00439），但肿瘤TYMP-邻近SPP1不支持（rho=0.00894/0.00985，P=0.3050/0.2930，q=0.3928/0.3802）。弦图显示GSE207493患者中位数最高的14条非自环显著CellChat边：弦宽∝√中位概率、颜色表示发送者。右侧气泡填色为特异性rho、面积为患者数、黑色外框表示正式支持；6个SPP1相对C1QC/Other端点均无正式支持，GSE207493中rho范围−0.255至0.244、q均=0.9526。")
    add_callout(doc, "主图结论", "最强正向证据集中在D的双外部Myeloid/DC程序复现和E的髓系横断面排序；F仅支持SPP1状态与血管背景共定位，不支持肿瘤TYMP邻近SPP1或SPP1特异通讯。因此主图不能被解读为TYMP已选择性驱动SPP1巨噬细胞。", fill=base.SOFT_TEAL, accent=base.TEAL)
    base.page_break(doc)

    base.add_heading(doc, "1 免疫治疗情境初筛与患者级差异", 1)
    add_module(
        doc,
        "1.1 F1-TYMP与免疫治疗结局初筛",
        "先判断肿瘤F1-TYMP关系是否随免疫治疗获益或治疗时间发生稳定变化，避免在缺少临床关联时直接把机制轴写成疗效标志物。",
        "GSE264586和GSE294109；严格CNV与扩展恶性定义并列。GSE294109可评估比较仅为扩展层5比2、严格层4比1；GSE264586展示对比仅1对合格患者。",
        "GSE294109中获益组相对无效组的F1-TYMP差异在两套恶性定义下均为负向；因独立样本数不足10且名义P未显著，正文仅保留方向。GSE264586纵向仅能描述方向，不能进行患者层推断。",
        "不支持F1-TYMP作为当前免疫治疗疗效标志物；该模块的作用是给后续免疫机制分析设定临床边界。",
    )
    add_module(
        doc,
        "1.2 免疫组成与状态特征",
        "检验免疫细胞比例、T/NK功能状态和髓系状态是否与获益或F1-TYMP强度相伴。",
        "GSE294109按患者汇总多区域后比较9名获益与3名无效患者，共27项组成/状态特征；GSE264586为3至4对纵向患者。F1-TYMP与组成的关联按可评估患者单独计算。",
        "GSE294109的27项特征无一通过FDR，最小P=0.04196、q=0.46158。扩展恶性定义下，F1-TYMP与Treg比例在7名患者中同向，rho=0.9286且P=0.00252；严格层5名患者中与T/NK dysfunction同向，rho=0.900且P=0.03739。GSE264586低样本纵向比较仅报告方向，未形成稳定支持。",
        "小样本结果提示Treg/免疫抑制背景值得继续验证，但不能称为稳定组成机制或疗效预测。",
    )
    add_module(
        doc,
        "1.3 患者级伪批量基因与通路分析",
        "把单细胞表达汇总到患者或配对患者层面，区分稀疏单基因差异与更广泛的通路重塑。",
        "GSE264586和GSE294109；12个edgeR quasi-likelihood模型、133,463条基因统计和594条Hallmark通路统计。纵向模型为3至4对患者；获益模型为9至12名患者且无效组均为3名。",
        "12个模型合计仅8个基因达到FDR，说明没有稳定单基因标志物集合；通路层每模型有2至38条FDR信号。GSE264586的TNFA/NFKB、IFN-gamma、mTORC1等随治疗阶段和细胞类型改变方向；GSE294109获益组出现CD4/单核IFN-gamma降低、CD8/NK/Treg氧化磷酸化降低和C1QC巨噬细胞EMT升高等模式。低样本纵向结果仅作方向性描述。",
        "证据支持情境依赖的免疫通路重塑，不支持单一普适疗效程序。",
    )
    base.add_figure(
        doc,
        representative / "immune_pathway_remodeling_heatmap.png",
        6.20,
        "图2 患者级免疫通路重塑热图。列为12个患者级伪批量模型，行为12条代表性通路；红/蓝分别表示正/负NES，颜色越深表示幅度越大，圆点表示该模型内BH FDR<0.05。NES应在同一模型内读方向和相对强度，不应跨不同细胞类型直接比较绝对值。12个模型的594条Hallmark统计中，每模型有2至38条FDR信号，而133,463条基因统计合计仅8个基因达到FDR，说明通路重塑强于稳定单基因标志物；3至4对患者的纵向模型仅作方向性解释。",
        "患者级免疫通路重塑热图",
        "十二个患者级伪批量模型中十二条代表性免疫与代谢通路的NES热图，显著单元格以圆点标记。",
    )
    # Natural section flow avoids sparse half-pages.

    base.add_heading(doc, "2 免疫细胞来源与谱系内程序发现", 1)
    add_module(
        doc,
        "2.1 F1与TYMP的细胞来源",
        "明确F1和TYMP主要由哪些细胞类型贡献，避免把bulk或跨细胞相关误写成同一细胞内机制。",
        "GSE207493，19位患者、102,723个用于细胞来源核对的细胞；对比经S/G2M细胞周期校正。",
        "F1在肿瘤细胞中显著高于非肿瘤细胞，校正后中位差0.2092，P=7.63×10^-6、q=1.53×10^-5；F1在Treg中反而低于其他T细胞，中位差-0.0398，P=1.53×10^-5、q=1.53×10^-5。TYMP在Myeloid/DC中最强，中位差1.3049，P=3.81×10^-6、q=1.53×10^-5；Treg相对其他T细胞也有较弱富集，中位差0.3148，P=1.53×10^-5、q=1.53×10^-5。",
        "F1应解释为肿瘤细胞程序；TYMP机制主线优先落在Myeloid/DC，Treg为次级谱系。",
    )
    add_module(
        doc,
        "2.2 谱系内代谢程序",
        "检验F1高肿瘤、TYMP高Myeloid/DC和TYMP高Treg是否各自伴随可重复的代谢状态，并排除细胞周期解释。",
        "GSE207493；肿瘤与Myeloid/DC各19位患者，Treg 16位患者；31条预设Reactome代谢通路。",
        "F1高肿瘤有22/31条通路通过FDR；缺氧反应校正后P=7.63×10^-6、q=4.73×10^-5。TYMP高Myeloid/DC有17/31条通路通过FDR；糖酵解P=3.81×10^-6、q=2.96×10^-5，嘧啶salvage同为P=3.81×10^-6、q=2.96×10^-5。TYMP高Treg有10/31条通路通过FDR；缺氧反应P=3.05×10^-5、q=2.37×10^-4。",
        "三类端点都存在代谢重塑，但Myeloid/DC的核苷酸、糖代谢和氧化还原核心更完整；表达通路不是直接代谢物测量。",
    )
    add_module(
        doc,
        "2.3 谱系内转录调控程序",
        "使用两种独立调控推断方法筛选与F1或TYMP同向的TF-regulon，减少单一算法偏差。",
        "GSE207493；SCENIC和CollecTRI-VIPER在患者层分别评估连续、分组和四分位端点。肿瘤与Myeloid/DC均为19位患者，Treg为13位患者。",
        "两方法同方向且通过FDR的结果共176条：肿瘤F1 95条、Myeloid/DC-TYMP 67条、Treg-TYMP 14条。Myeloid/DC核心包括IRF1、SPI1、RELB、STAT6、IRF7、STAT2、ZBTB7A、CEBPB和RELA；Treg核心包括GATA3、IRF2、NFKB2、REL、TCF3、RELB和STAT1。",
        "这是发现队列内的跨方法收敛，证明程序可被两种算法识别，但尚不是外部复现或因果调控。",
    )
    # Natural section flow avoids sparse half-pages.

    base.add_heading(doc, "3 外部情境转移与直接程序复现", 1)
    add_module(
        doc,
        "3.1 临床/时间情境中的基因集转移",
        "检验发现队列中的冻结Myeloid/DC和Treg程序是否在免疫治疗的时间或获益对比中呈基因集富集。",
        "GSE264586和GSE294109的5个合格患者级伪批量模型；50个冻结签名中123个模型-签名组合通过基因覆盖门槛。GSE264586模型仅3至4对患者，GSE294109 Treg为9名患者。",
        "共42个模型-签名组合通过模型内FDR，另有10个仅名义结果。GSE264586治疗末期C1QC巨噬细胞和FCN1单核细胞呈广泛IRF/NFKB/STAT调控转移；GSE294109获益对比方向混合。由于所有可解释模型均少于10个独立患者，本报告只保留方向，不用这些FDR结果升级证据。",
        "说明程序能进入免疫治疗情境，但不提供统一疗效方向，也不能替代直接外部相关复算。",
    )
    add_module(
        doc,
        "3.2 Myeloid/DC程序的双外部直接复现",
        "直接在外部单细胞数据中重算TYMP连续表达与去除TYMP锚点后的程序分数关系，验证机制是否跨数据集保持。",
        "GSE264586：12位患者、19个患者-时间点、16,594个Myeloid细胞；GSE294109：13位患者、33,895个Myeloid细胞。相关在患者/时间点内经S/G2M残差化后汇总。",
        "18个Myeloid/DC签名在两个外部数据集均与发现方向一致并通过FDR。代表性结果：缺氧反应在GSE264586中rho=0.2083、P=0.000488、q=0.000868，在GSE294109中rho=0.2394、P=0.000244、q=0.000977；IRF1 regulon分别为rho=0.3486、P=0.000488、q=0.00114和rho=0.2981、P=0.000244、q=0.000570。复现核心还包括呼吸链、ROS解毒、糖酵解/糖异生、磷酸戊糖、嘌呤分解及SPI1、RELB、STAT6、IRF7、STAT2、ZBTB7A、CEBPB、RELA。",
        "这是当前最高等级的免疫机制证据：广义TYMP相关Myeloid/DC代谢-调控核心可跨两个外部单细胞数据集复现。",
    )
    add_module(
        doc,
        "3.3 Treg程序的单外部直接复现",
        "评估TYMP高Treg程序是否也能脱离发现队列复现，并明确缺失注释带来的证据上限。",
        "仅GSE294109可评价冻结Treg细注释：11位患者、6,229个Treg细胞；GSE264586无法评价该端点。",
        "13个Treg签名同方向并通过FDR。ROS解毒rho=0.1125、P=0.000977、q=0.00244；GATA3 regulon rho=0.1641、P=0.000977、q=0.00195。其余包括缺氧、呼吸链、糖酵解/糖异生以及IRF2、NFKB2、REL、TCF3、RELB、STAT1。",
        "属于单外部数据集支持，不能写作双外部确认；Treg机制证据低于Myeloid/DC核心。",
    )
    # Natural section flow avoids sparse half-pages.

    base.add_heading(doc, "4 临床转化与bulk终点校验", 1)
    add_module(
        doc,
        "4.1 双外部复现核心的临床复合评分",
        "检验18个复现签名构成的等家族复合分数能否预测独立临床队列的response、PFS或治疗获益。",
        "E-MTAB-3218、CheckMate和JAVELIN；7个预设主模型，均使用独立临床样本并按模型报告效应。",
        "7个主模型中0个通过FDR。E-MTAB PFS HR=0.639，P=0.1335、q=0.6652；CheckMate nivolumab PFS HR=0.861，P=0.5482、q=0.6652；JAVELIN联合治疗PFS HR=1.045，P=0.5702、q=0.6652。",
        "机制复现没有转化为独立临床预测，不能据此升级response/PFS证据。",
    )
    base.add_figure(
        doc,
        representative / "clinical_validation_forest.png",
        6.20,
        "图3 外部复现髓系核心的独立临床验证森林图。点为效应比，横线为95%CI，竖线1表示无效应；点在1左/右分别表示较低/较高事件风险或相应结局方向，必须结合端点解释。右侧同时列出P和BH FDR，七个预设模型均未通过FDR。代表性PFS结果为E-MTAB HR=0.639、P=0.1335、q=0.6652；CheckMate nivolumab HR=0.861、P=0.5482、q=0.6652；JAVELIN联合治疗HR=1.045、P=0.5702、q=0.6652，因此外部程序复现未转化为稳定临床预测。",
        "临床验证森林图",
        "E-MTAB-3218、CheckMate和JAVELIN七个预设临床模型的效应比、置信区间、P值和FDR。",
    )
    add_module(
        doc,
        "4.2 修正临床终点后的完整链与TYMP-SPP1模型",
        "重新核对PFS事件编码，避免旧终点方向错误影响临床结论，并复测完整链代理和TYMP-SPP1联合特征。",
        "CheckMate nivolumab PFS 165例、146个事件；JAVELIN联合治疗PFS 354例、157个事件。模型按患者临床协变量调整。",
        "完整链代理在CheckMate为HR=1.017，P=0.8444、q=0.8444，在JAVELIN为HR=1.114，P=0.1770、q=0.5464；TYMP-SPP1联合特征在CheckMate为HR=0.928，P=0.4098、q=0.5464，在JAVELIN为HR=1.082，P=0.3024、q=0.5464，方向冲突。",
        "修正后的临床终点仍不支持完整机制链或TYMP-SPP1联合特征；旧的错误事件编码结果不得继续引用。",
    )
    add_module(
        doc,
        "4.3 文献状态桥接",
        "检验公开论文中的APOE-TAM、POSTN-CAF和T-cell inflamed状态能否在两个bulk ICI队列中形成可重复的TYMP桥接。",
        "CheckMate与JAVELIN；TYMP-APOE、TYMP-POSTN、TYMP-TIB三类连续桥接特征。",
        "TYMP-APOE方向冲突；TYMP-POSTN在JAVELIN中名义P=0.0409但q=0.2456，CheckMate中P=0.6654、q=0.8076；TYMP-TIB在CheckMate中P=0.9887、q=0.9887，在JAVELIN中P=0.3935、q=0.8076。公开仓库也未提供可复现的SPP1-TAM状态矩阵。",
        "论文状态可以提供背景，但没有形成跨队列临床桥接。",
    )
    # Natural section flow avoids sparse half-pages.

    base.add_heading(doc, "5 空间、跨谱系与跨区室证据", 1)
    add_module(
        doc,
        "5.1 组成校正的空间桥接",
        "检验TYMP/F1与髓系状态、T细胞功能和血管重塑的空间关系是否超出粗粒度细胞组成解释。",
        "GSE175540，12张FFPE和12张冰冻切片；每种制备均以12张切片为独立单位，9条预设桥接分别校正相关细胞组成与测序深度。",
        "仅SPP1状态到血管重塑在两种制备中同时支持：FFPE中位rho=0.2121、P=0.000488、q=0.00220；冰冻中位rho=0.2083、P=0.000488、q=0.00439。TYMP到SPP1在FFPE中P=0.06396、q=0.06396，在冰冻中P=0.8501、q=0.9697，均不支持。其余7条只在FFPE支持。",
        "空间证据支持SPP1状态与血管背景共定位，不支持肿瘤TYMP到SPP1状态的双制备桥接。",
    )
    add_module(
        doc,
        "5.2 MTA-response/SPP1转录桥接",
        "检验文献来源的MTA-A2B-STAT3响应代理能否连接TYMP、SPP1状态和免疫抑制程序。",
        "GSE207493、GSE264586和GSE294109；分别19、12、13位患者，共67,825个Myeloid/DC细胞；患者/时间点内经S/G2M残差化。",
        "外部数据中，TYMP到MTA-response在GSE264586为rho=0.2034、P=0.000488、q=0.00146，在GSE294109为rho=0.1165、P=0.000244、q=0.00146；MTA-response到去SPP1锚点状态分别为rho=0.0962、P=0.01611、q=0.02417和rho=0.1200、P=0.00806、q=0.01050。第三条TYMP到去锚SPP1状态只在GSE294109支持，GSE264586为P=0.09229、q=0.09229。",
        "三条必需边仅2/3跨两个外部数据集支持，因此完整转录桥未闭合；代理分数不是5-MTA测量。",
    )
    add_module(
        doc,
        "5.3 跨谱系免疫生态",
        "评估髓系MTA-response、SPP1状态和免疫抑制程序是否与T细胞dysfunction、cytotoxicity或Treg比例在患者层同步变化。",
        "GSE207493、GSE264586、GSE294109分别19、12、13位患者，共229,326个冻结免疫细胞；9条预设跨谱系关系。",
        "9条关系无一满足两外部数据集规则。代表性地，髓系免疫抑制到T dysfunction在GSE264586为rho=0.573、P=0.0513、q=0.1903，在GSE294109为rho=0.648、P=0.0165、q=0.1488；生态复合分数到T抑制复合分数分别为rho=0.503、P=0.0952、q=0.1903和rho=0.538、P=0.0576、q=0.2594。",
        "存在同向背景，但未形成跨谱系正式支持；患者共变不等于直接细胞通讯。",
    )
    add_module(
        doc,
        "5.4 MTA-response代理的空间转移",
        "检验MTA-response代理在空间切片中是否连接TYMP、SPP1状态、血管和T细胞功能。",
        "GSE175540，12张FFPE和12张冰冻切片；6条主要关系及4个方向一致性链产品。",
        "6条主要关系中0条获得双制备支持，4个链产品也均未通过。MTA-response到血管程序在FFPE为rho=0.0857、P=0.000488、q=0.00293，但冰冻为rho=0.0212、P=0.03418、q=0.2051；其余TYMP-MTA-response-SPP1关系均不支持。",
        "空间结果不支持跨制备的MTA-response机制链；链产品不是因果中介效应。",
    )
    add_module(
        doc,
        "5.5 患者配对的跨区室代理链",
        "用同一患者的肿瘤区室和髓系区室检验肿瘤TYMP/MTA-production与髓系MTA-response、SPP1和免疫抑制的关系。",
        "GSE207493 19位、GSE264586 4位、GSE294109 9位患者；GSE264586按预注册仅作低样本描述。",
        "仅GSE294109的肿瘤MTA-production到髓系MTA-response呈正向，n=9、rho=0.833且P=0.00527；其余主要关系不显著或方向不稳。GSE207493对应关系为rho=0.286、P=0.2353、q=0.4706；GSE264586仅描述为负向。",
        "单外部小样本结果不能闭合患者配对代理链，也不能证明代谢物跨区室运输。",
    )
    add_module(
        doc,
        "5.6 细亚状态空间邻域检验",
        "在统一细胞状态参考下，直接检验肿瘤区室TYMP是否邻近SPP1状态巨噬细胞，减少早期粗粒度代理的歧义。",
        "GSE175540，24张切片；13状态参考分别映射到12张FFPE和12张冰冻切片，连续肿瘤TYMP经肿瘤区室校正。",
        "肿瘤TYMP到邻近SPP1状态在FFPE为rho=0.00894、P=0.3050、q=0.3928，在冰冻为rho=0.00985、P=0.2930、q=0.3802，均不支持。C1QC邻域控制仅在FFPE支持：rho=0.02682、P=0.00505、q=0.02523；冰冻P=0.02942、q=0.14710。",
        "更精细的空间检验进一步否定双制备TYMP邻近SPP1状态的结论。",
    )
    # Natural section flow avoids sparse half-pages.

    base.add_heading(doc, "6 扰动、自然实验与湿实验审计", 1)
    add_module(
        doc,
        "6.1 四基因虚拟敲除",
        "用网络扰动方法检验STAT3、SPP1、TYMP和MTAP是否对预设髓系程序产生跨数据集稳定影响。",
        "GSE207493、GSE264586、GSE294109各自重建髓系GRN，共12次scTenifoldKnk分析；10个主要KO-程序假设。",
        "0/10主要假设达到两外部数据集支持；仅SPP1 KO和MTAP KO到去SPP1锚点状态在GSE264586获得单外部支持，其余8项无外部支持。",
        "网络嵌入距离是无方向扰动量，不能写成基因上调/下调，也不能证明5-MTA或SPP1蛋白功能。",
    )
    add_module(
        doc,
        "6.2 虚拟敲除受影响基因的通路富集",
        "把无方向网络扰动扩展到通路层，判断跨数据集稳定受影响的是SPP1特异轴还是更广泛免疫网络。",
        "三个单细胞数据集、四个KO；每次约3,007至3,009个可评估基因。",
        "TYMP KO跨外部数据集主要富集MHC-II抗原加工与呈递，例如MHC-II复合体q分别为2.36×10^-7和1.38×10^-5；MTAP KO主要富集补体、凝血、基质和囊泡/分泌相关过程，例如补体激活q分别为1.52×10^-9和9.58×10^-11。STAT3 KO和SPP1 KO没有精确条目达到两外部富集。",
        "结果指向广义抗原呈递、补体-溶酶体和基质网络，而不是可重复的MTA/SPP1特异因果路线。",
    )
    base.add_figure(
        doc,
        representative / "virtual_knockout_enrichment.png",
        6.20,
        "图4 虚拟敲除受影响基因的跨数据集富集。每条通路的两个点分别为GSE264586和GSE294109，横轴为−log10(BH FDR)，越靠右表示富集证据越强；连线仅帮助比较同一通路的两队列强度，不表示方向或因果。TYMP KO的MHC-II复合体q=2.36×10^-7和1.38×10^-5；MTAP KO的补体激活q=1.52×10^-9和9.58×10^-11。STAT3 KO和SPP1 KO没有精确通路条目在两个外部数据集同时富集，结果更符合广义抗原呈递及补体-基质网络扰动。",
        "虚拟敲除跨数据集富集图",
        "TYMP和MTAP虚拟敲除后，两个外部数据集的代表性通路FDR以成对点图展示。",
    )
    add_module(
        doc,
        "6.3 外源MTA扰动签名的转移",
        "尝试用公开的人MTA暴露、MTAP抑制和Mtap敲除数据学习外源扰动签名，再投射到肾癌髓系和空间数据。",
        "GSE249082、GSE298469、GSE276571和GSE166213用于训练；得到89基因直接MTA签名和26基因跨情境共识。随后投射到三个单细胞数据集和24张空间切片。",
        "26基因共识未达到预设双向基因数门槛，因此所有转移均为探索性。直接MTA签名到免疫抑制仅在GSE294109支持：13位患者、rho=0.0433、P=0.00610、q=0.01221；空间SPP1状态只在FFPE支持，冰冻为rho=0.0216、P=0.0342、q=0.0684。",
        "训练来源主要是成纤维细胞而非巨噬细胞，且共识签名未合格，不能升级TYMP-MTA-SPP1机制。",
    )
    add_module(
        doc,
        "6.4 MTAP自然扰动",
        "利用MTAP拷贝缺失或低表达作为自然扰动，检验其是否增强TYMP与SPP1/MTA-response/免疫抑制的耦合。",
        "TCGA-KIRC 496例，其中MTAP loss 148例；外部单细胞正式合并28位患者。",
        "TCGA中TYMP×MTAP loss对SPP1状态P=0.2545、q=0.3818，对MTA-response P=0.1076、q=0.3229，对免疫抑制P=0.5518、q=0.5518。外部单细胞交互同样不支持：SPP1状态P=0.7665、q=0.9029，MTA-response P=0.8612、q=0.9029，免疫抑制P=0.9029、q=0.9029。",
        "MTAP缺失确实降低MTAP表达，但没有形成预设机制交互；自然扰动不支持完整链。",
    )
    add_module(
        doc,
        "6.5 客户湿实验数据审计",
        "核对客户所述5-MTA升高和机制验证是否存在可重算的样本级原始表、分组、重复和统计结果。",
        "审计项目服务器、本地项目、下载目录以及两个客户工作簿：候选基因实验计划和532行空间代谢物注释表。",
        "未找到样本ID、实验组、生物学重复、MTA峰面积/浓度、内标、批次、误差或统计检验；代谢物注释表也没有精确5-MTA条目。报告中的44倍MTA升高目前只能视为客户陈述，无法由现有文件重算。",
        "因果升级至少需要样本清单、MTA定量、TYMP干预、巨噬细胞SPP1读出和phenocopy/rescue链条。",
    )
    # Natural section flow avoids sparse half-pages.

    base.add_heading(doc, "7 bulk巨噬细胞状态与可辨识性", 1)
    add_module(
        doc,
        "7.1 统一细胞状态向bulk投影",
        "评估单细胞定义的SPP1、C1QC和Other macrophage状态能否在CheckMate和JAVELIN bulk转录组中恢复并用于PFS分析。",
        "CheckMate 311份RNA样本、JAVELIN 726份RNA样本；共同478基因平台和同一13状态参考。主要PFS模型分别为165和354例。",
        "投影总巨噬细胞与官方巨噬细胞分数相关rho分别为0.459和0.540，但SPP1状态比例与PFS在CheckMate为HR=0.970、P=0.7537、q=0.9186，在JAVELIN为HR=1.130、P=0.1318、q=0.5272；TYMP×SPP1状态交互也均不支持。",
        "投影可追踪粗粒度巨噬细胞背景，但临床模型不支持SPP1状态或TYMP交互。",
    )
    add_module(
        doc,
        "7.2 TYMP与巨噬细胞共扩增",
        "区分TYMP是否选择性偏向SPP1状态，还是与整体巨噬细胞负荷共同升高。",
        "CheckMate 309例、JAVELIN 723例；模型调整官方巨噬细胞分数。",
        "条件模型中，SPP1巨噬细胞比例随TYMP在CheckMate呈负向beta=-0.1196、P=0.03862、q=0.10386，在JAVELIN为beta=-0.0726、P=0.05193、q=0.10386；SPP1/C1QC比值同样不支持。未条件化的SPP1和C1QC原始权重则在两个队列均随TYMP升高，例如SPP1权重在CheckMate P=0.00978、q=0.01304，在JAVELIN P=1.82×10^-7、q=7.30×10^-7。",
        "支持的是广义巨噬细胞富集生态共扩增，不是选择性SPP1极化。",
    )
    add_module(
        doc,
        "7.3 条件分解与失败诊断",
        "在总巨噬细胞负荷条件下进一步测试TYMP是否与三种细亚状态具有独立关系。",
        "CheckMate 309例、JAVELIN 723例；SPP1、C1QC和Other macrophage分别进入条件模型。",
        "Other macrophage在两个队列均被赋予零权重，导致完整三状态模型不可估计。可估计部分中，SPP1残差在CheckMate为beta=-0.1052、P=0.05122、q=0.10244，在JAVELIN为beta=-0.0222、P=0.3617、q=0.3617；均未出现预期正向支持。",
        "完整分解验证失败；部分结果只能作为诊断，不能写成三状态已被可靠分离。",
    )
    add_module(
        doc,
        "7.4 巨噬细胞细亚状态可辨识性基准",
        "用已知真值的伪bulk混合物直接测试三种细亚状态是否能在478基因平台上被可靠恢复。",
        "38位患者、1,824个确定性500,000-read伪bulk混合物；四种冻结/去锚/留一数据集方法，同时测试纯巨噬和组织背景混合物。",
        "三个预设门禁全部失败。冻结参考在三个纯巨噬测试中的最小状态rho仅0.245至0.295，最大MAE 0.633至0.689，主导状态准确率0.299至0.399；Other macrophage零权重比例达到1.000。相反，组织背景中的总巨噬细胞量仍可恢复，rho=0.745至0.894。",
        "bulk只允许使用粗粒度巨噬细胞总量；SPP1/C1QC/Other细亚状态权重不得用于临床或机制结论。",
    )
    base.add_figure(
        doc,
        representative / "bulk_identifiability_diagnostic.png",
        6.20,
        "图5 bulk巨噬细胞细亚状态可辨识性诊断。列为四种冻结/去锚/留一投影策略及测试数据集；左侧相关和主导状态准确率越高越好，右侧MAE与零权重比例越低越好。基准包含38位患者和1,824个500,000-read伪bulk混合物。三个纯巨噬测试中最小状态rho仅0.245至0.295、最大MAE为0.633至0.689、主导状态准确率为0.299至0.399，Other macrophage零权重比例=1.000；相反，总巨噬细胞量rho=0.745至0.894。故该图只支持bulk粗粒度巨噬细胞总量，不支持SPP1/C1QC/Other细亚状态权重。",
        "bulk细亚状态可辨识性诊断图",
        "四种投影策略在三个数据集中的细亚状态相关、主导状态准确率、总巨噬细胞相关、误差和零权重比例。",
    )
    # Natural section flow avoids sparse half-pages.

    base.add_heading(doc, "8 髓系状态排序与细亚状态通讯", 1)
    add_module(
        doc,
        "8.1 三数据集髓系状态排序",
        "在不依赖bulk投影的前提下，检验Monocyte、SPP1-state、C1QC-state和Other macrophage是否呈可重复的横断面顺序。",
        "GSE207493、GSE264586、GSE294109共68,711个髓系细胞、45位患者、62个患者-样本单元；Slingshot与独立PCA-kNN测地拟时序并行。",
        "两种方法的患者级一致性中位rho分别为0.841、0.559和0.430。SPP1-state均晚于Monocyte：GSE207493测地/Slingshot中位差0.333/0.418，均P=3.81×10^-6、q=4.45×10^-6；GSE264586为0.247/0.266，均P=0.00195、q=0.00684/0.00456；GSE294109为0.269/0.286，P=0.000244/0.000488、q=0.000244/0.000684。",
        "支持可重复的横断面状态排序，不证明真实谱系转换、治疗诱导或时间先后。",
    )
    add_module(
        doc,
        "8.2 去锚状态程序的轨迹一致性",
        "检查排序是否仅由SPP1/C1QC锚定基因驱动，还是去除锚点后仍有程序趋势。",
        "同一三个单细胞数据集和两种拟时序方法；每个数据集按患者汇总。",
        "SPP1去锚程序在GSE207493两方法均支持；GSE294109两方法也支持。GSE264586测地法rho=0.553、P=0.00488、q=0.00684，但Slingshot仅rho=0.082、P=0.3394、q=0.4751。Monocyte程序在三个数据集、两种方法中均下降，例如GSE264586 Slingshot rho=-0.713、P=0.000488、q=0.00342。",
        "状态顺序稳定，但去锚SPP1程序未达到两个外部数据集双方法复现门槛。",
    )
    add_module(
        doc,
        "8.3 三数据集细亚状态CellChat",
        "检验SPP1-state相对C1QC/Other macrophage是否具有特异的恶性、T细胞或内皮通讯，而不是仅表现为整体通讯更强。",
        "三个数据集分开运行，36个患者-样本单元、28位患者、46,964个抽样细胞和161,865条配体-受体记录；6个预设特异性端点。",
        "GSE207493的16位患者中，六个特异端点rho范围-0.255至0.244，P范围0.3413至0.9526，q均为0.9526；GSE294109仅6至8位患者，所有端点仅报告方向且名义P均不显著；GSE264586仅4位患者且TYMP变异不足。GSE207493的绝对Malignant-to-SPP1和SPP1-to-Malignant强度在次级层呈名义关联，但该层未进行正式FDR判定，且未形成相对C1QC/Other的特异性。",
        "六个主要端点无一获得正式支持；绝对通讯强度更符合组成或总体活性，不能写成SPP1特异通讯。",
    )
    # Natural section flow avoids sparse half-pages.

    h9 = base.add_heading(doc, "9 综合证据与下一步验证", 1)
    h9.paragraph_format.page_break_before = True
    base.add_table(
        doc,
        ["证据层级", "可以写入结论", "不能外推"],
        [
            ["跨两个外部单细胞数据集复现", "TYMP相关Myeloid/DC代谢与TF调控核心", "response/PFS/OS预测或因果驱动"],
            ["单外部数据集支持", "Treg-TYMP程序及部分患者配对方向", "双外部确认或普适机制"],
            ["空间双制备", "SPP1状态与血管背景共定位", "肿瘤TYMP邻近SPP1状态"],
            ["细胞层排序", "Monocyte到多个巨噬状态的横断面顺序", "真实谱系转换或治疗诱导"],
            ["bulk", "粗粒度巨噬细胞总量", "SPP1/C1QC/Other细亚状态分解"],
            ["临床与扰动", "用于限定负向边界", "已验证疗效标志物或5-MTA因果闭环"],
        ],
        [2050, 3650, 3660],
        font_size=8.8,
    )
    base.add_heading(doc, "9.1 当前最稳健的结论", 2)
    for text in [
        "TYMP的最稳定免疫细胞来源是Myeloid/DC；其代谢和调控程序在两个外部单细胞数据集中直接复现。",
        "TYMP与巨噬细胞富集生态共同升高，但没有证据支持bulk中选择性SPP1比例上升。",
        "Monocyte到SPP1/C1QC/Other macrophage的横断面顺序可复现，但去锚SPP1程序在GSE264586没有双方法通过。",
        "空间和通讯结果均没有建立肿瘤TYMP到SPP1状态的双制备/跨数据集特异联系。",
        "临床和扰动压力测试共同把TYMP-5-MTA-SPP1保留为待验证假说，而不是已闭合机制或临床标志物。",
    ]:
        base.add_paragraph(doc, text)
    base.add_heading(doc, "9.2 建议的验证优先级", 2)
    for text in [
        "建立匹配的样本清单与LC-MS/LC-MS/MS 5-MTA定量，报告原始峰面积/浓度、内标、单位、LOD/LOQ、批次和生物学重复。",
        "在巨噬细胞或肿瘤-巨噬共培养中进行TYMP干预，同时测量5-MTA和SPP1/C1QC/抗原呈递/补体等状态读出。",
        "加入MTA phenocopy以及MTAP/代谢清除或受体通路rescue，区分TYMP表达效应与5-MTA介导效应。",
        "空间验证应使用独立FFPE与冰冻样本，并预先固定肿瘤TYMP区室、邻域半径和SPP1状态定义。",
        "临床转化应在独立样本中直接测量可辨识的粗粒度髓系轴；在新平台验证前不要使用bulk三状态权重。",
    ]:
        base.add_paragraph(doc, text)
    add_callout(
        doc,
        "最终报告口径",
        "可以表述为“TYMP相关广义髓系免疫代谢与调控程序具有外部复现，并与巨噬细胞富集生态和状态排序相伴”；不能表述为“TYMP已选择性驱动SPP1巨噬细胞、5-MTA介导已被直接证明，或该轴已成为免疫治疗疗效/生存标志物”。",
        fill=base.SOFT_TEAL,
        accent=base.TEAL,
    )

    appendix_heading = base.add_heading(doc, "附录：免疫分析覆盖清单", 1)
    appendix_heading.paragraph_format.page_break_before = True
    coverage_rows = [
        ["免疫治疗初筛", "F1-TYMP结局、免疫组成/状态、患者级伪批量", "已覆盖"],
        ["细胞来源与程序发现", "F1/TYMP来源、代谢程序、调控程序", "已覆盖"],
        ["外部复现", "临床情境转移、Myeloid/DC双外部、Treg单外部", "已覆盖"],
        ["临床转化", "复合评分、修正PFS、论文状态桥接", "已覆盖"],
        ["空间与生态", "组成校正、MTA桥接、跨谱系、空间转移、跨区室、细亚状态空间", "已覆盖"],
        ["扰动与因果", "虚拟KO、KO富集、MTA签名、MTAP自然扰动、湿实验审计", "已覆盖"],
        ["bulk状态", "投影、共扩增、条件分解、可辨识性基准", "已覆盖"],
        ["细胞层后续", "髓系排序、去锚程序、细亚状态CellChat", "已覆盖"],
    ]
    base.add_table(doc, ["主题", "分析模块", "状态"], coverage_rows, [2100, 5960, 1300], font_size=8.8)
    base.add_paragraph(
        doc,
        "报告结束",
        size=8.5,
        color=base.MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=0,
    )

    visible_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    forbidden = []
    if re.search(r"(?i)\bphase\s*\d+\b", visible_text):
        forbidden.append("internal work identifier")
    if re.search(r"五\s*队列|5\s*队列", visible_text):
        forbidden.append("five-cohort narrative")
    if forbidden:
        raise SystemExit("Forbidden visible text: " + ", ".join(forbidden))
    if len(doc.inline_shapes) != 5:
        raise SystemExit(f"Expected one composite plus four representative figures, found {len(doc.inline_shapes)}")

    doc.save(output)
    print(
        json.dumps(
            {
                "status": "PASS_BUILD",
                "output": str(output),
                "paragraphs": len(doc.paragraphs),
                "tables": len(doc.tables),
                "inline_shapes": len(doc.inline_shapes),
                "visible_phase_identifiers": 0,
                "visible_five_cohort_phrases": 0,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()







