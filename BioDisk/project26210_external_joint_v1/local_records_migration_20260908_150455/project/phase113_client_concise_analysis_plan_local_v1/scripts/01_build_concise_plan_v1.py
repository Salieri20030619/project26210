from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(r'D:\project26210_external_joint_v1\phase113_client_concise_analysis_plan_local_v1')
REFERENCE = ROOT / 'PCa_reference.docx'
OUT = ROOT / 'output_v1' / '项目26210_无权重七特征公共数据迁移分析思路_概括版_v1.docx'


def clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)


def set_run_font(run, bold=None) -> None:
    run.font.name = 'Microsoft YaHei'
    run.font.size = Pt(12)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn(f'w:{key}'), 'Microsoft YaHei')


def set_paragraph_layout(paragraph, *, first_indent=True, center=False,
                         space_before=0, space_after=0, keep_with_next=False) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.keep_with_next = keep_with_next
    fmt.widow_control = True
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.first_line_indent = None
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.first_line_indent = Cm(0.85) if first_indent else None


def add_text_paragraph(document, text, *, first_indent=True, space_before=0,
                       space_after=0, keep_with_next=False):
    paragraph = document.add_paragraph()
    set_paragraph_layout(
        paragraph,
        first_indent=first_indent,
        space_before=space_before,
        space_after=space_after,
        keep_with_next=keep_with_next,
    )
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_heading(document, text):
    paragraph = document.add_paragraph()
    set_paragraph_layout(
        paragraph,
        first_indent=False,
        space_before=6,
        keep_with_next=True,
    )
    run = paragraph.add_run(text)
    set_run_font(run, bold=True)
    return paragraph


def add_step(document, number, title, text):
    paragraph = document.add_paragraph()
    set_paragraph_layout(paragraph, first_indent=False, space_before=2)
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(0.85)
    fmt.first_line_indent = Cm(-0.85)
    lead = paragraph.add_run(f'{number}. {title}：')
    set_run_font(lead, bold=True)
    body = paragraph.add_run(text)
    set_run_font(body)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph()
    set_paragraph_layout(paragraph, first_indent=False)
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(0.85)
    fmt.first_line_indent = Cm(-0.43)
    run = paragraph.add_run(f'•  {text}')
    set_run_font(run)
    return paragraph


def add_no_split(paragraph):
    ppr = paragraph._element.get_or_add_pPr()
    keep_lines = OxmlElement('w:keepLines')
    ppr.append(keep_lines)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document(str(REFERENCE))
    clear_document_body(document)

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.175)
    section.right_margin = Cm(3.175)

    normal = document.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(12)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn('w:eastAsia'), 'Microsoft YaHei'
    )

    title = document.add_paragraph()
    set_paragraph_layout(title, first_indent=False, center=True, space_after=8, keep_with_next=True)
    title_run = title.add_run('基于无权重七特征代谢指纹的公共数据迁移分析思路')
    set_run_font(title_run, bold=True)

    add_heading(document, '一、分析目标与总体逻辑')
    add_text_paragraph(
        document,
        '客户现有实验结果显示，TYMP敲低后细胞代谢谱整体改变，并形成一组方向明确的核苷及甲硫氨酸代谢特征。本项目不以公共数据库替代该实验，也不再追加湿实验，而是先从客户数据中提取可迁移的代谢指纹，再在独立公共队列中检验其能否复现、能否映射到转录调控，以及能否连接到SPP1+巨噬细胞、空间免疫结构、临床预后和治疗反应。由此形成“F1/TYMP—MTA代谢—SPP1+巨噬细胞—空间结构/临床结局”的分层证据链。'
    )

    add_heading(document, '二、无权重七特征的定义')
    add_text_paragraph(
        document,
        '以客户图中方向一致、公共数据中具备较高可映射性的7种代谢物作为固定特征：5′-MTA下降，Thymidine和2′-Deoxyuridine升高，Methionine、S-adenosylmethionine（SAM）、Thymine和Spermidine下降。所有特征先在各队列内进行标准化，再按实验方向统一正负号并取算术平均，得到无权重七特征得分。得分越高，表示样本越接近客户实验中的shTYMP样代谢状态。该定义在正式分析前冻结，不根据某个队列的显著性临时增删特征或调整权重。'
    )
    add_text_paragraph(
        document,
        '主得分为［−Z(5′-MTA)+Z(Thymidine)+Z(2′-Deoxyuridine)−Z(Methionine)−Z(SAM)−Z(Thymine)−Z(Spermidine)］/7。单个5′-MTA分析仍作为独立结果保留，用于说明单代谢物证据与多特征联合证据是否一致，但不以七特征得分覆盖既往的阴性结果。'
    )

    add_heading(document, '三、具体分析步骤')
    add_step(
        document, '1', '特征映射与数据质控',
        '统一不同数据库中的代谢物名称、同义词和检测标识，核对样本、平台、批次及临床信息。完整覆盖7种代谢物的队列进入主分析；部分覆盖队列只作为敏感性分析。'
    )
    add_step(
        document, '2', '公共代谢组主验证',
        '在UnitedMet肾癌队列RC18（约97例）和RC20（约52例）中分别计算固定七特征得分，观察样本分布、七个分量的方向贡献和队列间一致性；分析得分与TYMP表达、肿瘤分期、分级及可用临床结局的关系。两个队列独立分析后再做方向一致性汇总，不先合并样本制造显著性。'
    )
    add_step(
        document, '3', '跨平台敏感性分析',
        '在CPTAC及其验证队列中计算3特征子得分，在NCI-60中计算5特征子得分，并与完整七特征结果比较，判断信号是否由单一平台、单一队列或个别代谢物驱动；部分特征得分不作为主结论。'
    )
    add_step(
        document, '4', '从代谢特征迁移到转录特征',
        '在同时具有代谢组和转录组的数据中，将七特征得分作为连续变量，分析其与TYMP、F1相关基因、甲硫氨酸/SAM循环、嘧啶代谢及多胺代谢通路的相关性和差异通路。优先保留跨队列方向一致且具有明确代谢解释的基因与通路，形成后续可投射到转录组数据的“七特征转录代理”。'
    )
    add_step(
        document, '5', 'SPP1+巨噬细胞与单细胞定位',
        '仅在转录代理通过前一步验证后，将其投射到肾癌单细胞数据。按患者为独立样本，比较高、低代理状态下SPP1+巨噬细胞比例、SPP1程序活性、巨噬细胞状态转换及其与肿瘤细胞的配体—受体通讯。重点判断该代谢状态是否稳定指向SPP1+免疫抑制生态位，而不是仅报告单个细胞层面的显著性。'
    )
    add_step(
        document, '6', '空间结构验证',
        '在空间转录组中定位七特征转录代理、TYMP/F1相关程序和SPP1+巨噬细胞，分析其空间共定位、邻域富集及与肿瘤区、缺氧区、血管区和T细胞排斥区的关系。若样本数有限，以每张切片和每位患者的效应方向为主，不以spot数量替代生物学重复。'
    )
    add_step(
        document, '7', '预后与治疗关联',
        '在肾癌转录组队列中计算冻结后的转录代理，分析其与分期、分级、总生存或无进展生存的关系，并在免疫治疗队列中评估应答差异。报告连续效应、分组曲线和多变量模型，并观察TYMP/F1或SPP1巨噬细胞分层后的结果。'
    )

    add_heading(document, '四、结果判定与停止规则')
    add_text_paragraph(
        document,
        '证据按层级分别判定，不用下游阳性替代上游失败。若完整七特征在RC18和RC20中方向一致并连接到预设代谢通路，则继续构建转录代理；若得分不稳定，只保留可复现的单特征或子模块。若转录代理不能在独立队列复现，则不进入单细胞、空间和临床外推；仅支持一条下游分支时，作为独立候选线索呈现。客户实验与公共数据证据分别陈述。'
    )

    add_heading(document, '五、预计交付内容')
    add_text_paragraph(
        document,
        '预计交付包括：七特征名称映射、数据覆盖与质控表；RC18和RC20完整七特征主结果及CPTAC、NCI-60跨平台敏感性结果；代谢特征对应的转录代理和通路解释；满足前置条件后形成的SPP1+巨噬细胞、空间生态位、预后及治疗关联结果；一张完整证据链总图，并明确标注各环节为“支持、部分支持或当前数据不支持”。'
    )

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            set_run_font(run, bold=run.bold)

    document.core_properties.title = '项目26210无权重七特征公共数据迁移分析思路（概括版）'
    document.core_properties.subject = '客户沟通用概括性分析方案'
    document.core_properties.author = 'Codex'
    document.save(str(OUT))
    print(str(OUT))


if __name__ == '__main__':
    build()
