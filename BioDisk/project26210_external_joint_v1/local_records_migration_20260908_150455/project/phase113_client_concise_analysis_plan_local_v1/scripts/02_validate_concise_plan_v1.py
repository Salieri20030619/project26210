import hashlib
import json
from pathlib import Path

import pdfplumber
from PIL import Image, ImageChops
from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r'D:\project26210_external_joint_v1\phase113_client_concise_analysis_plan_local_v1')
DOCX = ROOT / 'output_v1' / '项目26210_无权重七特征公共数据迁移分析思路_概括版_v1.docx'
PDF = ROOT / 'output_v1' / '项目26210_无权重七特征公共数据迁移分析思路_概括版_v1.pdf'
RENDER = ROOT / 'output_v1' / 'render_final'
OUT = ROOT / 'validation_v1' / 'validation_summary_v1.json'


def sha256(path):
    h = hashlib.sha256()
    with path.open('rb') as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest().upper()


def main():
    errors = []
    document = Document(str(DOCX))
    visible_runs = 0
    for p_index, paragraph in enumerate(document.paragraphs):
        for r_index, run in enumerate(paragraph.runs):
            if not run.text.strip():
                continue
            visible_runs += 1
            size = run.font.size.pt if run.font.size else None
            rpr = run._element.rPr
            rfonts = rpr.rFonts if rpr is not None else None
            east_asia = rfonts.get(qn('w:eastAsia')) if rfonts is not None else None
            if size != 12.0:
                errors.append(f'run_size:{p_index}:{r_index}:{size}')
            if east_asia != 'Microsoft YaHei':
                errors.append(f'run_font:{p_index}:{r_index}:{east_asia}')

    section = document.sections[0]
    geometry = {
        'page_width_cm': round(section.page_width.cm, 3),
        'page_height_cm': round(section.page_height.cm, 3),
        'top_margin_cm': round(section.top_margin.cm, 3),
        'bottom_margin_cm': round(section.bottom_margin.cm, 3),
        'left_margin_cm': round(section.left_margin.cm, 3),
        'right_margin_cm': round(section.right_margin.cm, 3),
    }
    expected_geometry = {
        'page_width_cm': 21.0,
        'page_height_cm': 29.7,
        'top_margin_cm': 2.54,
        'bottom_margin_cm': 2.54,
        'left_margin_cm': 3.175,
        'right_margin_cm': 3.175,
    }
    for key, expected in expected_geometry.items():
        if abs(geometry[key] - expected) > 0.02:
            errors.append(f'geometry:{key}:{geometry[key]}')

    full_text = '\n'.join(p.text for p in document.paragraphs)
    required = [
        'F1/TYMP', '5′-MTA', 'Thymidine', '2′-Deoxyuridine', 'Methionine',
        'S-adenosylmethionine', 'Thymine', 'Spermidine', 'RC18', 'RC20',
        'CPTAC', 'NCI-60', 'SPP1+巨噬细胞', '空间结构验证', '预后与治疗关联',
        '当前数据不支持',
    ]
    for term in required:
        if term not in full_text:
            errors.append(f'missing_term:{term}')
    if 'Phase' in full_text or 'phase' in full_text:
        errors.append('internal_phase_label_present')

    with pdfplumber.open(str(PDF)) as pdf:
        page_count = len(pdf.pages)
        page_text_lengths = [len((page.extract_text() or '').strip()) for page in pdf.pages]
    if page_count != 3:
        errors.append(f'pdf_page_count:{page_count}')
    if any(length < 100 for length in page_text_lengths):
        errors.append(f'sparse_page:{page_text_lengths}')

    pixel_audit = []
    for png in sorted(RENDER.glob('page-*.png')):
        image = Image.open(png).convert('RGB')
        white = Image.new('RGB', image.size, 'white')
        diff = ImageChops.difference(image, white).convert('L')
        mask = diff.point(lambda value: 255 if value > 12 else 0)
        bbox = mask.getbbox()
        if bbox is None:
            errors.append(f'blank_page:{png.name}')
            continue
        left, top, right, bottom = bbox
        margins = {
            'left_px': left,
            'top_px': top,
            'right_px': image.width - right,
            'bottom_px': image.height - bottom,
        }
        if min(margins.values()) < 35:
            errors.append(f'boundary_risk:{png.name}:{margins}')
        pixel_audit.append({'page': png.name, 'size': image.size, 'ink_bbox': bbox, **margins})

    summary = {
        'status': 'PASS' if not errors else 'FAIL',
        'docx': str(DOCX),
        'pdf': str(PDF),
        'sha256_docx': sha256(DOCX),
        'sha256_pdf': sha256(PDF),
        'visible_runs_checked': visible_runs,
        'font_rule': 'Microsoft YaHei 12 pt for every visible run',
        'geometry': geometry,
        'pdf_pages': page_count,
        'page_text_lengths': page_text_lengths,
        'pixel_audit': pixel_audit,
        'errors': errors,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    raise SystemExit(0 if not errors else 1)


if __name__ == '__main__':
    main()
