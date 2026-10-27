from __future__ import annotations

import csv
import hashlib
import json
import os
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


WORK = Path(__file__).resolve().parent
PROJECT = Path(r"D:\project26210_external_joint_v1")
PHASE = PROJECT / "phase89_part1_immune_evidence_freeze_report_update_local_v1"
SOURCE_REPORT = WORK / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_WITH_FIGURES_REVIEW_v12.docx"
SOURCE_EVIDENCE = WORK / "IMMUNE_MECHANISTIC_EVIDENCE_REGISTER_v3.tsv"
OUTPUT_REPORT = WORK / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_WITH_FIGURES_REVIEW_v13.docx"
OUTPUT_EVIDENCE = WORK / "IMMUNE_MECHANISTIC_EVIDENCE_REGISTER_v4.tsv"
OUTPUT_FREEZE = WORK / "PHASE89_IMMUNE_EVIDENCE_FREEZE_v1.tsv"
OUTPUT_SUMMARY = WORK / "PHASE89_INTEGRATION_SUMMARY_v1.md"
SOURCE_MANIFEST = WORK / "SOURCE_ARTIFACT_MANIFEST_SHA256_v1.tsv"
QC = WORK / "qc_local_v1"


SOURCE_ARTIFACTS = [
    PROJECT / "phase68_part1_immune_report_update_local_v1/PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_WITH_FIGURES_REVIEW_v12.docx",
    PROJECT / "phase68_part1_immune_report_update_local_v1/IMMUNE_MECHANISTIC_EVIDENCE_REGISTER_v3.tsv",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase69/results_v1/PHASE69_R2_COMPOSITE_CLINICAL_SUMMARY_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase70/results_v1/PHASE70_COMPOSITION_AWARE_SPATIAL_SUMMARY_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase71/PROJECT_STATUS_PHASE71_ADDENDUM_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase72/PROJECT_STATUS_PHASE72_ADDENDUM_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase73/PROJECT_STATUS_PHASE73_ADDENDUM_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase74/PROJECT_STATUS_PHASE74_ADDENDUM_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase75/results_v1/PHASE75_SPATIAL_TRANSFER_SUMMARY_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase76/results_v1/PHASE76_CROSS_COMPARTMENT_SUMMARY_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase77/PROJECT_STATUS_PHASE77_ADDENDUM_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase78/analysis_v1/PHASE78_DECISION_v1.json",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase79/results_v1/PHASE79_DECISION_v1.json",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase80/project26210_phase80_client_wetlab_audit_v1/PHASE80_CLIENT_WETLAB_AUDIT_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase81/PROJECT_STATUS_PHASE81_82_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase83/PROJECT_STATUS_PHASE83_84_v1.md",
    PROJECT / "phase69_85_server_status_sync_local_v1/phase85/results_v1/decomposition_v1/PHASE85_SUMMARY_v1.md",
    PROJECT / "phase86_macrophage_state_identifiability_local_v1/results_v1/PHASE86_DECISION_v1.json",
    PROJECT / "phase87_three_cohort_myeloid_trajectory_local_v1/results_v1/PHASE87_DECISION_v1.json",
    PROJECT / "phase88_three_cohort_fine_state_cellchat_local_v1/results_v1/PHASE88_DECISION_v1.json",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(archive.namelist())
            if name.startswith("word/media/")
        }


def find_exact(document: Document, text: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one paragraph {text!r}; found {len(matches)}")
    return matches[0]


def find_starting(document: Document, prefix: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one paragraph starting {prefix!r}; found {len(matches)}")
    return matches[0]


def replace_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_before(anchor: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    paragraph = anchor.insert_paragraph_before(text, style=style)
    paragraph.paragraph_format.keep_with_next = style.startswith("Heading")
    return paragraph


def add_labeled_before(anchor: Paragraph, label: str, body: str) -> Paragraph:
    paragraph = anchor.insert_paragraph_before(style="Normal")
    lead = paragraph.add_run(label)
    lead.bold = True
    paragraph.add_run(body)
    return paragraph


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    cell.width = Cm(width_dxa / 1440 * 2.54)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)


def format_cell(cell, *, bold: bool = False, font_size: float = 8.0, center: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        if center:
            paragraph.alignment = 1
        for run in paragraph.runs:
            run.bold = bold
            run.font.name = "宋体"
            run.font.size = Pt(font_size)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def write_source_manifest() -> None:
    rows = []
    for path in SOURCE_ARTIFACTS:
        if not path.is_file():
            raise FileNotFoundError(path)
        rows.append({
            "source_id": f"S{len(rows) + 1:02d}",
            "absolute_path": str(path),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        })
    with SOURCE_MANIFEST.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def update_evidence_register() -> Counter:
    with SOURCE_EVIDENCE.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    fields = list(rows[0])
    by_id = {row["evidence_id"]: row for row in rows}
    by_id["I01"].update({
        "finding": "F1-related/full-chain mechanism signatures as response/PFS/OS predictors",
        "n": "corrected CheckMate RNA n=311; JAVELIN RNA n=726; Phase69 also tested E-MTAB-3218",
        "effect": "No robust replicated clinical validation after correcting CheckMate event coding",
        "statistic": "All prior Phase64 CheckMate PFS/OS estimates are superseded. Phase69: 0/7 primary models passed FDR. Corrected Phase81 FULL_CHAIN_PROXY: CheckMate HR=1.0173 q=0.8444; JAVELIN HR=1.1138 q=0.5464.",
        "global_grade": "D",
        "external_grade": "CORRECTED_ENDPOINT_NOT_SUPPORTED",
        "interpretation": "F1-related or full-chain scores cannot be reported as validated stand-alone immunotherapy predictors.",
        "limitation": "Bulk signatures mix compartments; corrected cohorts, endpoints and treatment settings remain heterogeneous.",
        "source": "PHASE69_R2_COMPOSITE_CLINICAL_SUMMARY_v1.md;PROJECT_STATUS_PHASE81_82_v1.md",
    })
    by_id["I02"].update({
        "finding": "TYMP and TYMP-SPP1 bridge features as response/PFS/OS predictors",
        "n": "corrected CheckMate RNA n=311 and JAVELIN RNA n=726",
        "effect": "No robust cross-cohort direct or interaction validation after endpoint correction",
        "statistic": "Corrected Phase81 TYMP_SPP1_JOINT: CheckMate HR=0.9276 q=0.5464; JAVELIN HR=1.0821 q=0.5464, with direction conflict. Phase82 TYMP-APOE/POSTN/TIB bridges and Phase83 TYMP x SPP1-state interaction were unsupported.",
        "global_grade": "D",
        "external_grade": "CORRECTED_ENDPOINT_NOT_SUPPORTED",
        "interpretation": "TYMP cannot be reported as a validated stand-alone immunotherapy predictor or as a validated TYMP-SPP1 clinical bridge.",
        "limitation": "Bulk clinical cohorts do not resolve TYMP cell source, macrophage fine state or 5-MTA mediation.",
        "source": "PROJECT_STATUS_PHASE81_82_v1.md;PROJECT_STATUS_PHASE83_84_v1.md",
    })
    for evidence_id in ("I08", "I11"):
        by_id[evidence_id]["limitation"] = (
            by_id[evidence_id]["limitation"].rstrip(".")
            + "; Phase69-88 did not convert this replicated mechanism core into a clinical predictor, selective SPP1-state signal or causal 5-MTA bridge."
        )
    counts = Counter(row["global_grade"] for row in rows)
    if counts != Counter({"A": 5, "B": 6, "C1": 2, "D": 2}):
        raise RuntimeError(f"Unexpected grade distribution: {counts}")
    with OUTPUT_EVIDENCE.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    return counts


def write_freeze_table() -> list[dict[str, str]]:
    rows = [
        {
            "freeze_id": "F01",
            "evidence_axis": "Replicated myeloid mechanism core",
            "phases": "67,69",
            "result": "Eighteen TYMP-excluded myeloid/DC signatures replicated in two external cohorts, but 0/7 Phase69 clinical primary models passed FDR.",
            "freeze_status": "RETAIN_SUPPORTED_CONTEXT",
            "reportable_claim": "A broad TYMP-associated myeloid immunometabolic/regulatory core is externally reproducible.",
            "prohibited_overreach": "Do not call it a response/PFS predictor or proof of 5-MTA mediation.",
        },
        {
            "freeze_id": "F02",
            "evidence_axis": "Corrected clinical prediction",
            "phases": "81-83",
            "result": "Corrected CheckMate and JAVELIN FULL_CHAIN, TYMP-SPP1, paper-state and projected-state PFS models were unsupported or directionally conflicting.",
            "freeze_status": "NOT_CAUSAL_OR_CLINICAL",
            "reportable_claim": "Direct F1/TYMP clinical prediction remains grade D.",
            "prohibited_overreach": "Do not reuse superseded Phase64 CheckMate survival estimates.",
        },
        {
            "freeze_id": "F03",
            "evidence_axis": "Spatial and cross-compartment specificity",
            "phases": "70,72,75,76,78",
            "result": "SPP1-to-vascular context was reproducible, but continuous tumour-adjusted TYMP-to-neighbouring SPP1 state failed in both FFPE and frozen preparations; proxy chains were incomplete.",
            "freeze_status": "UNSUPPORTED_FOR_SPECIFICITY",
            "reportable_claim": "Vascular/macrophage context is compatible with a broad ecosystem signal.",
            "prohibited_overreach": "Do not claim a replicated TYMP-to-SPP1 spatial bridge or causal mediation.",
        },
        {
            "freeze_id": "F04",
            "evidence_axis": "Perturbation and causal closure",
            "phases": "73,74,77,79,80",
            "result": "Primary virtual-KO hypotheses had 0 two-external support; pathway enrichment was broad; the perturbation signature failed qualification; MTAP natural perturbation interactions failed; no recalculable wet-lab MTA package was found.",
            "freeze_status": "NOT_CAUSAL_OR_CLINICAL",
            "reportable_claim": "TYMP is embedded in broad antigen-presentation/myeloid network architecture.",
            "prohibited_overreach": "Do not claim TYMP causes 5-MTA accumulation, SPP1 secretion or macrophage reprogramming.",
        },
        {
            "freeze_id": "F05",
            "evidence_axis": "Bulk macrophage fine states",
            "phases": "83-86",
            "result": "Total macrophage load was partly recoverable, but Other Macrophage was zero in every bulk projection and all Phase86 operational/generalization/anchor gates failed across 1,824 pseudobulk mixtures.",
            "freeze_status": "PROHIBITED_FOR_FINE_BULK_USE",
            "reportable_claim": "Higher TYMP is associated with a broad macrophage-enriched tumour ecosystem.",
            "prohibited_overreach": "Do not quantify or clinically model SPP1/C1QC/Other fractions as identifiable bulk states.",
        },
        {
            "freeze_id": "F06",
            "evidence_axis": "Cell-level myeloid ordering",
            "phases": "87",
            "result": "SPP1-state cells were later than Monocytes in all three cohorts under Slingshot and independent PCA-kNN geodesic pseudotime; the anchor-excluded SPP1 program failed dual-method support in GSE264586.",
            "freeze_status": "RETAIN_SUPPORTED_CONTEXT",
            "reportable_claim": "Cross-sectional Monocyte-to-macrophage state ordering is reproducible.",
            "prohibited_overreach": "Do not infer lineage conversion, treatment induction, temporal ancestry or causality.",
        },
        {
            "freeze_id": "F07",
            "evidence_axis": "Fine-state communication",
            "phases": "88",
            "result": "None of six SPP1-versus-C1QC/Other CellChat specificity endpoints obtained formal support; GSE264586 reduced to four patients and was descriptive only.",
            "freeze_status": "UNSUPPORTED_FOR_SPECIFICITY",
            "reportable_claim": "Absolute communication intensity may track broad composition in GSE207493.",
            "prohibited_overreach": "Do not claim replicated SPP1-state-specific communication; NicheNet was not rerun.",
        },
        {
            "freeze_id": "F08",
            "evidence_axis": "Integrated thesis",
            "phases": "69-88",
            "result": "The evidence converges on broad myeloid/macrophage ecosystem biology but not on a selective SPP1-state, clinical-predictive or causal 5-MTA mechanism.",
            "freeze_status": "FINAL_EVIDENCE_FREEZE",
            "reportable_claim": "The TYMP-5-MTA-SPP1 model remains a testable biological hypothesis with supported broad context.",
            "prohibited_overreach": "Do not present the full chain as established or clinically validated.",
        },
    ]
    with OUTPUT_FREEZE.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    return rows


def insert_freeze_table(document: Document, anchor: Paragraph, rows: list[dict[str, str]]) -> None:
    headers = ["证据轴", "核心结果", "冻结状态", "报告口径"]
    widths = [1500, 3260, 1780, 2820]
    display = {
        "F01": ("髓系机制", "18个TYMP去锚髓系/DC签名在两套外部队列复现，但Phase69七个临床主模型均未通过FDR。", "保留：背景支持", "可写为广义TYMP相关髓系免疫代谢/调控核心可重复。"),
        "F02": ("临床预测", "修正后的CheckMate与JAVELIN全链、TYMP-SPP1及状态投影PFS模型不支持或方向冲突。", "不支持：临床", "F1/TYMP直接临床预测仍为D级。"),
        "F03": ("空间/跨区室", "SPP1到血管背景可重复；连续肿瘤校正TYMP到邻近SPP1状态在FFPE和冰冻制备均失败。", "不支持：SPP1特异", "可保留血管/巨噬细胞广义生态背景。"),
        "F04": ("扰动/因果", "虚拟敲除、富集、扰动签名、MTAP自然扰动及客户湿实验审计均未闭合5-MTA-SPP1因果链。", "不支持：因果", "可写为TYMP嵌入广义抗原呈递/髓系网络。"),
        "F05": ("bulk细亚状态", "巨噬细胞总量部分可恢复；Other Macrophage恒为零，Phase86三类可辨识性门禁全部失败。", "禁止：bulk细亚状态", "仅可报告TYMP与广义巨噬细胞富集生态相关。"),
        "F06": ("细胞状态排序", "三队列两种方法均显示SPP1状态晚于Monocyte；GSE264586去锚SPP1程序未通过双方法门禁。", "保留：横断面排序", "可报告可重复的Monocyte到巨噬细胞状态顺序。"),
        "F07": ("精细通讯", "六个SPP1相对C1QC/Other的CellChat特异终点均未获正式支持；GSE264586仅4位患者。", "不支持：SPP1特异", "绝对通讯强度只可作为广义组成背景。"),
        "F08": ("综合结论", "证据收敛于广义髓系/巨噬细胞生态，不收敛于选择性SPP1、临床预测或5-MTA因果机制。", "最终证据冻结", "TYMP-5-MTA-SPP1保留为待验证生物学假说。"),
    }
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        table.rows[0].cells[index].text = text
        shade_cell(table.rows[0].cells[index], "D9EAF7")
        format_cell(table.rows[0].cells[index], bold=True, font_size=8.5, center=True)
    set_repeat_table_header(table.rows[0])
    for row in rows:
        values = list(display[row["freeze_id"]])
        cells = table.add_row().cells
        for index, text in enumerate(values):
            cells[index].text = text
            if index == 2:
                shade_cell(cells[index], "EEF4F8")
            format_cell(cells[index], font_size=8.0, center=index == 2)
    set_table_geometry(table, widths)
    anchor._p.addprevious(table._tbl)


def write_summary() -> None:
    OUTPUT_SUMMARY.write_text(
        "# Phase 89 Part 1 immune evidence freeze and report integration\n\n"
        "- Status: `REPORT_GENERATED_PENDING_RENDER_AND_VALIDATION`.\n"
        "- Source: Phase 68 v12 report, preserved unchanged.\n"
        "- Output: non-overwriting v13 report.\n"
        "- Evidence register: v4 retains 15 rows and A=5, B=6, C1=2, D=2; I01/I02 now use corrected Phase81 survival-event coding.\n"
        "- Integrated freeze: broad TYMP-associated myeloid/macrophage context is retained; bulk three-state use is prohibited; cell-level state ordering is retained with a cross-sectional boundary; SPP1-specific communication and the full TYMP-to-5-MTA-to-SPP1 clinical/causal chain are unsupported.\n"
        "- NicheNet was not rerun in Phase88.\n",
        encoding="utf-8",
    )


def main() -> None:
    outputs = [OUTPUT_REPORT, OUTPUT_EVIDENCE, OUTPUT_FREEZE, OUTPUT_SUMMARY, SOURCE_MANIFEST]
    for path in outputs:
        if path.exists():
            raise FileExistsError(f"Refusing to overwrite: {path}")
    QC.mkdir(parents=True, exist_ok=True)
    write_source_manifest()
    counts = update_evidence_register()
    freeze_rows = write_freeze_table()
    write_summary()

    source_hash = sha256(SOURCE_REPORT)
    source_media = media_hashes(SOURCE_REPORT)
    document = Document(SOURCE_REPORT)

    replace_text(
        find_starting(document, "版本说明  客户图文审阅稿："),
        "版本说明  客户图文审阅稿：保留全部26项分析、15组正文结果图和Phase68证据升级，新增Phase69-88机制闭环压力测试与最终证据冻结；采用修正后的临床终点、患者层统计和分级判定。",
    )
    executive_anchor = find_exact(document, "1  数据范围与分析设计")
    add_labeled_before(
        executive_anchor,
        "最终证据冻结  ",
        "Phase69-88的临床、空间、扰动、bulk可辨识性、拟时序和细胞通讯压力测试共同限定了主线：可保留的是TYMP相关广义髓系免疫代谢核心、巨噬细胞富集生态共扩增和跨截面状态排序；不能升级的是选择性SPP1极化、SPP1特异通讯、5-MTA因果介导以及response/PFS/OS预测。",
    )
    stats_anchor = find_exact(document, "1.3  总体流程")
    insert_before(
        stats_anchor,
        "Phase69-88整合采用冻结证据口径：修正后的临床模型优先于被替代结果；bulk细亚状态必须先通过可辨识性门禁；拟时序仅解释横断面顺序；CellChat仅解释表达相容的通讯概率；一队列、小样本、单制备、次要或代理结果不得替代两外部队列正式支持。",
    )

    section7 = find_exact(document, "7  综合解释")
    c08 = insert_before(section7, "C08  机制闭环压力测试与最终证据冻结", "Heading 2")
    c08.paragraph_format.page_break_before = True
    insert_before(
        section7,
        "Phase69-85依次检验了临床转化、空间跨区室、跨谱系生态、虚拟敲除、扰动签名、MTAP自然扰动、客户湿实验数据、修正后的bulk ICI终点以及统一巨噬细胞状态投影。结果没有形成response/PFS/OS、空间、扰动和因果层面的闭环；重复出现的稳定成分是广义髓系程序、抗原呈递/补体-溶酶体网络以及巨噬细胞总体负荷，而不是选择性SPP1终末状态。",
    )
    insert_freeze_table(document, section7, freeze_rows)
    add_labeled_before(
        section7,
        "临床与因果边界  ",
        "Phase81确认旧Phase64 CheckMate生存事件编码颠倒，旧PFS/OS Cox及治疗交互结果全部被替代。修正后的CheckMate与JAVELIN FULL_CHAIN_PROXY、TYMP-SPP1联合特征、论文状态桥接和投影状态模型均未获得跨队列支持。Phase73-80也未得到两外部队列虚拟敲除主终点、合格的跨情境MTA扰动签名、MTAP自然扰动交互或可重算的客户5-MTA湿实验数据。",
    )
    add_labeled_before(
        section7,
        "Phase86-88闭环  ",
        "Phase86在38位患者、1,824个伪bulk混合物中判定SPP1/C1QC/Other三状态对bulk不可辨识，禁止继续用于细亚状态临床推断。Phase87在68,711个细胞中用Slingshot和独立PCA-kNN测地拟时序复现Monocyte到SPP1状态的横断面排序，但GSE264586的去锚SPP1程序未通过双方法门禁。Phase88在36个患者-样本单元中检验六个SPP1相对C1QC/Other的CellChat特异终点，无一获得正式支持。",
    )
    add_labeled_before(
        section7,
        "冻结后的写作规则  ",
        "可写为TYMP与广义髓系免疫代谢和巨噬细胞富集生态相关，并存在可重复的细胞状态顺序；不可写为bulk三状态已被可靠分解、TYMP选择性驱动SPP1极化、SPP1特异通讯已跨队列复现、5-MTA介导已被直接测量，或F1/TYMP已成为免疫治疗疗效及生存预测指标。",
    )

    supported = find_starting(document, "第一，cNMF F1与TYMP在患者内总体同向")
    supported.add_run(
        " 第八，后续压力测试保留了广义巨噬细胞富集生态和三队列Monocyte到巨噬细胞状态的横断面排序，但不支持把这些结果解释为选择性SPP1极化。"
    )
    unsupported = find_starting(document, "F1不对应稳定的RCC遗传易感性富集。")
    unsupported.add_run(
        " 此外，修正后的CheckMate终点未恢复F1/TYMP临床预测；bulk SPP1/C1QC/Other三状态未通过可辨识性门禁；三队列CellChat也未支持SPP1相对C1QC/Other的特异通讯。因此，旧生存结果、bulk细亚状态权重和绝对通讯强度均不能用于闭合TYMP-5-MTA-SPP1因果链。"
    )
    find_exact(document, "7.3  最终结论").paragraph_format.page_break_before = True
    final = find_starting(document, "综合结论  第一部分26项分析已经全部完成。")
    replace_text(
        final,
        "综合结论  第一部分26项分析及Phase52-88外部复核、机制压力测试已经完成。当前最稳健的证据支持TYMP相关F1为肿瘤细胞程序、TYMP高髓系/DC代谢-调控核心在两套外部队列复现，以及TYMP与广义巨噬细胞富集生态共同上升；Phase87进一步支持三队列横断面的Monocyte到巨噬细胞状态排序。与此同时，修正后的临床终点、空间跨区室、自然扰动、虚拟敲除、bulk可辨识性和细胞通讯结果均未建立选择性SPP1极化、SPP1特异通讯、直接5-MTA介导或response/PFS/OS预测。Phase86明确禁止把现有bulk三状态权重作为细亚状态真值，Phase88明确不给出两外部队列的SPP1特异通讯支持。因此，TYMP-5-MTA-SPP1轴应保留为具有广义髓系背景支持的待验证生物学假说，而不是已经闭合的因果机制或临床标志物。",
    )

    document.core_properties.modified = datetime.now(timezone.utc)
    document.save(OUTPUT_REPORT)

    reopened = Document(OUTPUT_REPORT)
    text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    table_text = "\n".join(cell.text for table in reopened.tables for row in table.rows for cell in row.cells)
    required = [
        "C08  机制闭环压力测试与最终证据冻结",
        "Phase86在38位患者、1,824个伪bulk混合物中",
        "Phase87在68,711个细胞中",
        "Phase88在36个患者-样本单元中",
        "禁止：bulk细亚状态",
        "TYMP-5-MTA-SPP1轴应保留为具有广义髓系背景支持的待验证生物学假说",
    ]
    missing = [phrase for phrase in required if phrase not in text and phrase not in table_text]
    if missing:
        raise RuntimeError(f"Required report text missing: {missing}")
    if media_hashes(OUTPUT_REPORT) != source_media:
        raise RuntimeError("Embedded media changed")
    if sha256(SOURCE_REPORT) != source_hash:
        raise RuntimeError("Source report changed")

    status = {
        "status": "REPORT_UPDATED_PENDING_RENDER_QC",
        "source_report": str(SOURCE_ARTIFACTS[0]),
        "source_report_sha256": source_hash,
        "output_report": OUTPUT_REPORT.name,
        "output_report_sha256": sha256(OUTPUT_REPORT),
        "paragraphs": len(reopened.paragraphs),
        "tables": len(reopened.tables),
        "media_parts_preserved_byte_for_byte": len(source_media),
        "evidence_rows": 15,
        "evidence_grade_counts": dict(counts),
        "freeze_rows": len(freeze_rows),
        "source_artifacts": len(SOURCE_ARTIFACTS),
        "document_skill_operation_marker": "UNAVAILABLE_IN_LOCAL_SKILL_RUNTIME",
        "edits": [
            "updated version note, executive summary and evidence rule",
            "added C08 with one fixed-geometry evidence-freeze table",
            "corrected I01/I02 to the Phase81 endpoint audit",
            "updated supported, unsupported and final-conclusion synthesis",
            "preserved all embedded media byte-for-byte",
        ],
    }
    (QC / "report_update_status_v1.json").write_text(json.dumps(status, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(status, ensure_ascii=False))


if __name__ == "__main__":
    main()
