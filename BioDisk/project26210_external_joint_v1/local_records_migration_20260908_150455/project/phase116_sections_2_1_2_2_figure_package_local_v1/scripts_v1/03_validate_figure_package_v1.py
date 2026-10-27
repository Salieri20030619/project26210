from __future__ import annotations

import csv
import hashlib
import json
import sys
import zipfile
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageFont
from docx import Document
from docx.oxml.ns import qn


BASE = Path(__file__).resolve().parent
PACKAGE = BASE / "package_v1"
DOCX = PACKAGE / "项目26210_2.1-2.2图件证据说明_v1.docx"
INDEX = PACKAGE / "图件索引与证据边界_v1.tsv"
RENDER = BASE / "render_word_v2"
OUT = BASE / "validation_v1"
OUT.mkdir(exist_ok=True)


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def check_table_geometry(doc: Document):
    issues = []
    for ti, table in enumerate(doc.tables, start=1):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        grid = table._tbl.tblGrid.findall(qn("w:gridCol"))
        widths = [int(x.get(qn("w:w"))) for x in grid]
        if tbl_w is None or int(tbl_w.get(qn("w:w"))) != sum(widths):
            issues.append(f"table{ti}:tblW")
        if tbl_ind is None or int(tbl_ind.get(qn("w:w"))) != 120:
            issues.append(f"table{ti}:tblInd")
        for ri, row in enumerate(table.rows, start=1):
            if len(row.cells) != len(widths):
                issues.append(f"table{ti}:row{ri}:ncols")
                continue
            for ci, cell in enumerate(row.cells):
                tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
                if tc_w is None or int(tc_w.get(qn("w:w"))) != widths[ci]:
                    issues.append(f"table{ti}:row{ri}:cell{ci+1}:tcW")
    return issues


def build_contact_sheet(pages):
    thumb_w = 360
    gap = 18
    label_h = 24
    cols = 3
    thumbs = []
    for page in pages:
        with Image.open(page) as img:
            img = img.convert("RGB")
            h = round(img.height * thumb_w / img.width)
            thumbs.append((page.name, img.resize((thumb_w, h), Image.Resampling.LANCZOS)))
    thumb_h = max(img.height for _, img in thumbs)
    rows = (len(thumbs) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * thumb_w + (cols + 1) * gap, rows * (thumb_h + label_h) + (rows + 1) * gap), "#D8DCE2")
    draw = ImageDraw.Draw(sheet)
    for idx, (name, img) in enumerate(thumbs):
        row, col = divmod(idx, cols)
        x = gap + col * (thumb_w + gap)
        y = gap + row * (thumb_h + label_h + gap)
        sheet.paste(img, (x, y + label_h))
        draw.text((x + 4, y + 4), name, fill="#1F2937")
    out = OUT / "render_contact_sheet_v1.jpg"
    sheet.save(out, quality=82, optimize=True)
    return out


def main():
    checks = []
    def record(name, passed, detail=""):
        checks.append({"check": name, "passed": bool(passed), "detail": str(detail)})

    record("docx_exists", DOCX.is_file(), DOCX)
    record("index_exists", INDEX.is_file(), INDEX)
    rows = list(csv.DictReader(INDEX.open(encoding="utf-8-sig"), delimiter="\t"))
    record("figure_index_rows_41", len(rows) == 41, len(rows))
    record("figure_index_names_unique", len({r["group"] + "/" + r["file_name"] for r in rows}) == 41)
    record("figure_formats_23_png_18_pdf", sum(r["format"] == "png" for r in rows) == 23 and sum(r["format"] == "pdf" for r in rows) == 18)
    record("no_zero_byte_figure", all(int(r["bytes"]) > 0 for r in rows))
    hash_ok = True
    source_filled = True
    for row in rows:
        path = PACKAGE / row["group"] / row["file_name"]
        hash_ok &= path.is_file() and digest(path) == row["sha256"]
        source_filled &= bool(row["original_source"].strip())
    record("all_figure_hashes_match_index", hash_ok)
    record("all_original_sources_recorded", source_filled)

    doc = Document(DOCX)
    full_text_parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text_parts.append(cell.text)
    full_text = "\n".join(full_text_parts)
    for phrase in (
        "当前数据不支持写成“TYMP促使肿瘤进入更恶性状态”",
        "F1–TYMP 17/20患者正向",
        "F1与CNV burden",
        "五方法恶性轨迹",
        "患者为主要推断单位",
    ):
        record(f"required_text::{phrase[:18]}", phrase in full_text)
    record("inline_images_4", len(doc.inline_shapes) == 4, len(doc.inline_shapes))
    alt_texts = [shape._inline.docPr.get("descr", "") for shape in doc.inline_shapes]
    record("all_images_have_meaningful_alt", all(len(x) >= 20 for x in alt_texts), alt_texts)
    record("table_geometry_exact", not check_table_geometry(doc), check_table_geometry(doc))
    record("heading_hierarchy_present", sum(p.style.name == "Heading 1" for p in doc.paragraphs) >= 6 and sum(p.style.name == "Heading 2" for p in doc.paragraphs) >= 6)

    render_status = json.loads((RENDER / "render_status.json").read_text(encoding="utf-8"))
    pages = sorted(RENDER.glob("page-*.png"))
    record("word_render_pass", render_status.get("status") == "PASS", render_status)
    record("rendered_pages_9", len(pages) == 9, len(pages))
    page_audit = []
    all_nonblank = True
    all_safe = True
    for page in pages:
        with Image.open(page) as img:
            rgb = img.convert("RGB")
            bg = Image.new("RGB", rgb.size, "white")
            diff = ImageChops.difference(rgb, bg)
            bbox = diff.getbbox()
            nonwhite = sum(1 for px in diff.convert("L").getdata() if px > 8)
            fraction = nonwhite / (rgb.width * rgb.height)
            safe = bool(bbox and bbox[0] >= 12 and bbox[1] >= 10 and bbox[2] <= rgb.width - 12 and bbox[3] <= rgb.height - 10)
            nonblank = bool(bbox and fraction > 0.003)
            all_nonblank &= nonblank
            all_safe &= safe
            page_audit.append({"page": page.name, "size": rgb.size, "bbox": bbox, "nonwhite_fraction": fraction, "nonblank": nonblank, "safe_pixel_boundary": safe})
    record("no_blank_pages", all_nonblank, page_audit)
    record("safe_page_pixel_boundaries", all_safe, page_audit)
    contact = build_contact_sheet(pages)
    record("contact_sheet_created", contact.is_file() and contact.stat().st_size > 0, contact)

    with zipfile.ZipFile(DOCX) as zf:
        names = set(zf.namelist())
    record("docx_ooxml_complete", {"word/document.xml", "[Content_Types].xml", "word/_rels/document.xml.rels"}.issubset(names))

    passed = all(c["passed"] for c in checks)
    summary = {
        "status": "PASS" if passed else "FAIL",
        "n_checks": len(checks),
        "n_failed": sum(not c["passed"] for c in checks),
        "docx_sha256": digest(DOCX),
        "index_sha256": digest(INDEX),
        "package_figure_files": len(rows),
        "rendered_pages": len(pages),
        "checks": checks,
        "page_audit": page_audit,
    }
    (OUT / "PACKAGE_VALIDATION_v1.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    with (OUT / "validation_checks_v1.tsv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["check", "passed", "detail"], delimiter="\t")
        writer.writeheader(); writer.writerows(checks)
    print(json.dumps({k: summary[k] for k in ("status", "n_checks", "n_failed", "package_figure_files", "rendered_pages")}, ensure_ascii=False))
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
