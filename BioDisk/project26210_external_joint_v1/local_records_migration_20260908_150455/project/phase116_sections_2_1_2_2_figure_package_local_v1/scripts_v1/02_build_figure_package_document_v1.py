from __future__ import annotations

import csv
import hashlib
import json
import os
import sys
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
PACKAGE = BASE / "package_v1"
OUTPUT_DOCX = PACKAGE / "项目26210_2.1-2.2图件证据说明_v1.docx"
INDEX_TSV = PACKAGE / "图件索引与证据边界_v1.tsv"
README = PACKAGE / "README_先看这里_v1.txt"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172B4D"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GOLD = "FFF4D6"
PALE_RED = "FDECEC"
PALE_GREEN = "EAF5EE"
GRID = "C9D2DC"
RISK = "9B1C1C"
GOLD = "7A5A00"
GREEN = "246B43"
FONT_CJK = "Microsoft YaHei"
FONT_LATIN = "Calibri"


CONCEPTS = {
    "Fig02_malignancy_CNV_composite_v1": (
        "2.1",
        "良恶性判定与CNV综合主图",
        "身份/描述性基础",
        "展示代表性inferCNV、共识恶性与非恶性细胞CNV burden，以及inferCNV/CopyKAT交叉判定。",
        "可证明恶性细胞识别和CNV分层已完成；不能证明TYMP导致恶性进展。",
    ),
    "Fig03_CNV_clone_trees_by_cNMF_program_v4": (
        "2.1",
        "CNV克隆树与cNMF程序主图（正式v4）",
        "结构性支持",
        "四个代表患者的中性根锚定CNV距离树，叠加主导cNMF程序和主要染色体增益/缺失。",
        "树表示CNV相似性组织，不是突变验证的系统发育或真实时间演化。",
    ),
    "Fig04_cNMF_rank7_TYMP_program_v1": (
        "2.2",
        "cNMF rank-7 F1/TYMP程序综合主图",
        "主要程序证据",
        "汇总rank选择、F1与TYMP患者内相关、程序通路注释和患者层F1高克隆比较。",
        "支持TYMP与F1转录程序相关；不支持将F1直接等同于最高恶性克隆或因果驱动。",
    ),
    "Fig09_malignant_trajectory_five_methods_v1": (
        "2.1",
        "五方法恶性细胞状态轨迹综合图",
        "关键边界",
        "并列展示Monocle3、Slingshot、CytoTRACE2、scTour和CellRank状态结果。",
        "TYMP/F1方向跨方法不一致；不能据此声称TYMP处于统一晚期状态或推动状态转移。",
    ),
    "Fig02B_CNV_burden_violin_v1": (
        "2.1",
        "共识恶性与非恶性细胞CNV burden",
        "描述性支持",
        "以患者阈值标准化CNV burden比较共识恶性与非恶性细胞。",
        "细胞数仅作描述，不能作为独立生物学重复。",
    ),
    "infercnv_GSE156632_RCC2": (
        "2.1",
        "代表病例inferCNV染色体热图",
        "代表性展示",
        "展示GSE156632 RCC2肿瘤上皮细胞的染色体分块CNV模式。",
        "单一病例图用于展示方法与结构，不代表跨患者普遍性。",
    ),
    "method_cross_table_v1": (
        "2.1",
        "inferCNV与CopyKAT交叉判定",
        "方法一致性",
        "展示两种CNV方法在细胞层的组合计数。",
        "两方法明确判定细胞中的完全一致率有限，需保留method-discordant类别。",
    ),
    "patient_malignancy_composition_v1": (
        "2.1",
        "患者层恶性状态组成",
        "描述性支持",
        "逐患者展示恶性共识、单方法恶性、非恶性和未决细胞比例。",
        "患者间组成差异明显，不得把细胞比例差异解释为TYMP效应。",
    ),
    "formal_primary_clone_tymp_landscape_v1": (
        "2.1/2.2",
        "正式CNV克隆TYMP/cNMF景观",
        "探索性桥接",
        "展示患者内主要CNV克隆、主导cNMF因子与TYMP水平。",
        "cNMF F1预先由TYMP证据选择，属于一致性而非独立验证。",
    ),
    "selected_F1_patient_tymp_deltas_v1": (
        "2.1/2.2",
        "F1高克隆与其他克隆TYMP差值",
        "探索性支持",
        "按患者展示F1高主要克隆相对其他主要克隆的TYMP中位数差。",
        "主阈值多重校正未通过，且统计强度受GSE207493影响。",
    ),
    "selected_F1_threshold_patient_sensitivity_v1": (
        "2.1/2.2",
        "F1–TYMP克隆阈值敏感性",
        "敏感性分析",
        "比较0.225、0.25和0.275三个CNV克隆距离阈值下的患者差值。",
        "阈值间显著性变化，不能挑选最有利阈值作为唯一结论。",
    ),
    "factor_contrast_threshold_sensitivity_v1": (
        "2.1/2.2",
        "各cNMF因子克隆差值阈值敏感性",
        "敏感性分析",
        "展示7个rank-7因子在多个克隆阈值下的TYMP差值。",
        "F1并非所有阈值下最强；需与因子家族多重校正一起解释。",
    ),
    "Monocle3_pseudotime_v1": (
        "2.1",
        "Monocle3恶性细胞拟时序",
        "轨迹描述",
        "患者内计算的Monocle3状态顺序展示。",
        "拟时序不是实际肿瘤时间、克隆祖先或因果进展。",
    ),
    "Slingshot_pseudotime_v2": (
        "2.1",
        "Slingshot恶性细胞拟时序（正式v2）",
        "轨迹描述",
        "正式人工复核通过的Slingshot v2状态曲线。",
        "F1呈正向候选但FDR未通过；旧v1未纳入包。",
    ),
    "CytoTRACE2_potency_v1": (
        "2.1",
        "CytoTRACE2细胞潜能",
        "轨迹/潜能边界",
        "展示恶性细胞潜能评分及与预设状态锚点的关系。",
        "F1方向与Slingshot/scTour相反；不能声称统一晚期F1状态。",
    ),
    "scTour_pseudotime_v1": (
        "2.1",
        "scTour恶性细胞拟时序",
        "轨迹描述",
        "展示scTour潜在状态顺序。",
        "F1/TYMP仅名义方向信号，轨迹家族FDR未通过。",
    ),
    "CellRank_fate_TYMP_state_v1": (
        "2.1",
        "CellRank TYMP状态命运概率",
        "状态概率描述",
        "展示由冻结TYMP状态定义得到的命运概率。",
        "状态标签含TYMP信息，不能作为TYMP作用的独立验证。",
    ),
    "CellRank_fate_Hypoxia_v1": (
        "2.1",
        "CellRank缺氧状态命运概率",
        "状态概率描述",
        "展示缺氧状态命运概率。",
        "概率为模型化状态兼容性，不是细胞真实迁移证据。",
    ),
    "CellRank_fate_Cycle_v1": (
        "2.1",
        "CellRank增殖状态命运概率",
        "状态概率描述",
        "展示细胞周期状态命运概率。",
        "不能由增殖状态概率推出TYMP驱动恶性进展。",
    ),
    "CellRank_fate_confidence_v1": (
        "2.1",
        "CellRank命运置信度",
        "质量/不确定性",
        "展示CellRank命运分配置信度。",
        "用于说明模型不确定性，不是生物效应量。",
    ),
    "F1_high_clone_pseudobulk_DE_volcano_v1": (
        "2.2",
        "F1高克隆配对伪批量差异表达火山图",
        "患者层差异表达",
        "14名患者、28个伪批量样本的患者阻断edgeR比较。",
        "基因层差异不等同于F1驱动或TYMP因果效应。",
    ),
    "F1_high_clone_top_DE_patient_heatmap_v1": (
        "2.2",
        "F1高克隆差异基因患者热图",
        "患者一致性",
        "展示主要差异基因在配对患者中的方向。",
        "患者方向存在异质性，不能以单个基因或单个患者概括。",
    ),
    "rank7_program_pathway_ORA_v1": (
        "2.2",
        "rank-7程序通路富集",
        "程序功能注释",
        "对各cNMF程序top-loading基因进行Hallmark/KEGG/Reactome ORA。",
        "这是程序装载基因注释；F1高克隆差异基因通路ORA本身未通过FDR。",
    ),
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT_LATIN
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)



def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        flag = OxmlElement("w:cantSplit")
        flag.set(qn("w:val"), "true")
        tr_pr.append(flag)


def set_alt_text(inline_shape, descr):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", descr)
    doc_pr.set("title", descr[:80])


def shade_paragraph(paragraph, fill, border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if border:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "6")
        left.set(qn("w:color"), border)
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_text(doc, text, *, size=10.5, bold=False, color="222222", after=6, italic=False, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_callout(doc, label, text, fill, accent):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.25
    shade_paragraph(p, fill, accent)
    r = p.add_run(label + "  ")
    set_run_font(r, size=10.5, bold=True, color=accent)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color="222222")
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_table(doc, headers, rows, widths, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    mark_header_row(hdr)
    prevent_row_split(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=font_size, bold=True, color=INK)
    for row_idx, values in enumerate(rows):
        added_row = table.add_row()
        prevent_row_split(added_row)
        cells = added_row.cells
        for i, value in enumerate(values):
            if row_idx % 2:
                set_cell_shading(cells[i], "F8FAFC")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color="222222")
        set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def source_path(group, stem, suffix):
    if group == "01_主管汇报主图":
        if stem == "Fig03_CNV_clone_trees_by_cNMF_program_v4":
            return str(Path(r"D:\project26210_external_joint_v1\phase46_part1_clone_tree_and_fig14_revision_local_v1\figures_local_v1") / f"{stem}{suffix}")
        return str(Path(r"D:\project26210_external_joint_v1\phase44_part1_client_figures_final_local_v1\figures_complete_v2") / f"{stem}{suffix}")
    if group == "02_2.1_CNV与亚克隆支持图":
        if stem in {"method_cross_table_v1", "patient_malignancy_composition_v1"}:
            return f"/mnt/BioDisk/ludingyu/project26210/results/part1_rebuild_v1/03_malignancy_cnv/malignancy_cnv_summary_v1/{stem}{suffix}"
        if stem.startswith("Fig02B"):
            return str(Path(r"D:\project26210_external_joint_v1\phase44_part1_client_figures_final_local_v1\figures_complete_v2") / f"{stem}{suffix}")
        if stem == "infercnv_GSE156632_RCC2":
            return str(Path(r"D:\project26210_external_joint_v1\phase43_part1_client_figures_local_v1\source_remote_v1") / f"{stem}{suffix}")
        return f"/mnt/BioDisk/ludingyu/project26210/results/part1_rebuild_v1/04_subclones_tymp/cnv_clone_tymp_cnmf_association_v1/{stem}{suffix}"
    if group == "03_2.1_恶性轨迹边界图":
        return f"/mnt/BioDisk/ludingyu/project26210/results/part1_rebuild_v1/09_metabolism_trajectory/trajectory_v1/figures/{stem}{suffix}"
    return f"/mnt/BioDisk/ludingyu/project26210/results/part1_rebuild_v1/05_cnmf/cnv_clone_program_function_v1/figures/{stem}{suffix}"


def build_index():
    records = []
    for group in ("01_主管汇报主图", "02_2.1_CNV与亚克隆支持图", "03_2.1_恶性轨迹边界图", "04_2.2_cNMF支持图"):
        for path in sorted((PACKAGE / group).iterdir()):
            if path.suffix.lower() not in {".png", ".pdf"}:
                continue
            stem = path.stem
            if stem not in CONCEPTS:
                raise ValueError(f"unmapped figure: {path}")
            section, concept, evidence_class, use, boundary = CONCEPTS[stem]
            width = height = ""
            if path.suffix.lower() == ".png":
                with Image.open(path) as img:
                    width, height = img.size
            records.append(
                {
                    "group": group,
                    "section": section,
                    "file_name": path.name,
                    "format": path.suffix.lower().lstrip("."),
                    "concept": concept,
                    "evidence_class": evidence_class,
                    "intended_use": use,
                    "interpretation_boundary": boundary,
                    "original_source": source_path(group, stem, path.suffix),
                    "bytes": path.stat().st_size,
                    "sha256": sha256(path),
                    "width_px": width,
                    "height_px": height,
                }
            )
    with INDEX_TSV.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]), delimiter="\t")
        writer.writeheader()
        writer.writerows(records)
    return records


def add_main_figure(doc, path, title, explanation, boundary):
    add_heading(doc, title, 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    shape = p.add_run().add_picture(str(path), width=Inches(6.18))
    set_alt_text(shape, title + "。" + explanation)
    add_text(doc, explanation, size=9.5, color="333333", after=3)
    add_callout(doc, "解释边界", boundary, PALE_GOLD, GOLD)


def build_doc(records):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("项目26210｜主管讨论证据包｜2.1–2.2")
    set_run_font(r, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("仅供内部讨论  •  ")
    set_run_font(r, size=8, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("项目26210｜2.1–2.2 图件证据说明")
    set_run_font(r, size=24, bold=True, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("基于《分析待做（新）》恶性上皮章节｜主管讨论版")
    set_run_font(r, size=12.5, color=MUTED)

    meta = add_table(
        doc,
        ["项目", "内容"],
        [
            ["参考文件", "分析待做（新）.docx；SHA-256 A32766F2…7753"],
            ["覆盖范围", "前半部分恶性上皮 2.1 与 2.2；不含后文髓系章节重复编号的2.1"],
            ["图件组成", f"{sum(1 for r in records if r['format']=='png')}张PNG；{sum(1 for r in records if r['format']=='pdf')}份对应PDF"],
            ["统计原则", "患者为主要推断单位；细胞数仅作描述；所有因果与时间演化表述均按现有证据降级"],
        ],
        [1900, 7460],
        9,
    )

    add_callout(
        doc,
        "主管讨论结论",
        "现有数据表明TYMP RNA与cNMF rank-7 F1恶性细胞转录程序相关，并在患者层显示方向支持；但当前数据不支持写成“TYMP促使肿瘤进入更恶性状态”。CNV负荷、拟时序、scDRS及跨队列结果没有形成一致的因果/进展证据。",
        PALE_RED,
        RISK,
    )

    add_heading(doc, "1. 原始任务要求与本包口径", 1)
    add_heading(doc, "1.1 任务2.1", 2)
    add_text(doc, "基于inferCNV与CopyKAT进行良恶性判定和CNV亚克隆划分，展示按NMF/cNMF程序着色的克隆树，并检验TYMP是否位于更晚或更恶性的状态。")
    add_heading(doc, "1.2 任务2.2", 2)
    add_text(doc, "运行cNMF提取恶性细胞程序，选择与TYMP表达同步的候选Factor-TYMP，并展示程序差异基因与功能注释。")
    add_callout(doc, "口径修正", "任务原文使用“证明”措辞。本包按已验证结果改为“检验并分级证据”：身份判定、程序相关、克隆关联、轨迹状态和因果促进是不同证据层，不能互相替代。", PALE_GOLD, GOLD)

    add_heading(doc, "2. 证据结论总表", 1)
    evidence_rows = [
        ["良恶性/CNV分层", "支持", "62,652个上皮细胞；38,670个肿瘤上皮细胞；10,844个malignant_consensus", "证明恶性细胞识别与CNV分层已完成；不是TYMP效应"],
        ["cNMF rank-7 F1程序", "支持但异质", "rank-7稳定门通过；F1–TYMP 17/20患者正向，中位rho=0.115，BH q=0.0197", "可称TYMP相关转录状态；GSE207493贡献主要统计强度"],
        ["GSE207493内部稳健性", "支持", "10/10患者正向；中位rho=0.147；留一患者方向稳定率1.0；平衡重抽样同向率1.0", "单队列内部稳健不等于跨队列普遍性"],
        ["F1高CNV克隆TYMP升高", "探索性", "主阈值0.25：中位差0.200，9/14正向，P=0.0210，7因子BH q=0.0735", "未通过主阈值多重校正；去掉GSE207493后方向不稳定"],
        ["F1与CNV burden", "当前数据不支持", "19名患者中位rho=-0.033，BH q=0.923", "F1/TYMP相关程序不能等同更高CNV恶性度"],
        ["五方法恶性轨迹", "方向矛盾/不一致", "Monocle3无F1/TYMP支持；Slingshot/scTour F1正向但FDR阴性；CytoTRACE2 F1为负向FDR信号", "不支持统一晚期状态或TYMP推动状态转移"],
        ["F1程序功能", "程序注释支持", "top-loading基因富集TNFA–NF-kB、IFNγ、糖酵解、缺氧等；多项FDR<0.001", "这是cNMF程序注释，不是患者配对差异通路的确认"],
        ["F1高克隆差异通路", "当前数据不支持", "14名患者配对伪批量DE可见基因差异，但DE基因ORA全部FDR=1", "不得声称F1高克隆具有经FDR确认的整体恶性通路增强"],
    ]
    add_table(doc, ["证据问题", "结论", "核心统计", "解释边界"], evidence_rows, [1700, 1450, 2900, 3310], 8.0)

    doc.add_page_break()
    add_heading(doc, "3. 主管汇报主图", 1)
    main_dir = PACKAGE / "01_主管汇报主图"
    add_main_figure(doc, main_dir / "Fig02_malignancy_CNV_composite_v1.png", "3.1 2.1 良恶性判定与CNV综合", "该图回答“哪些细胞可被保守地视为恶性、两种CNV方法如何一致或冲突”。正式恶性共识为10,844个细胞；方法明确判定细胞中的一致率为33.8%。", "用于建立恶性细胞分析对象，不用于证明TYMP驱动恶性。")
    doc.add_page_break()
    add_main_figure(doc, main_dir / "Fig03_CNV_clone_trees_by_cNMF_program_v4.png", "3.2 2.1 CNV克隆树与程序叠加", "正式v4展示四个代表患者的中性根锚定CNV距离树，节点/分支附主要CNV变化和主导cNMF程序。全体正式分析含14名患者、43个主要克隆节点。", "树是表达推断CNV距离组织，不是突变验证系统发育；不能按分支深度推断真实早晚期。")
    doc.add_page_break()
    add_main_figure(doc, main_dir / "Fig04_cNMF_rank7_TYMP_program_v1.png", "3.3 2.2 cNMF rank-7 F1/TYMP程序", "rank 4–8均完成，rank-7综合稳定性最高并通过稳定门。F1与TYMP在20名可评价患者中17名同向，中位rho=0.115，18端点BH q=0.0197。", "F1主导细胞99.66%来自GSE207493；正式结论必须保留队列权重与跨队列异质性。")
    doc.add_page_break()
    add_main_figure(doc, main_dir / "Fig09_malignant_trajectory_five_methods_v1.png", "3.4 2.1 五方法恶性状态轨迹", "该图并列展示五种状态/拟时序方法，用于判断TYMP/F1是否稳定定位于更晚状态。正式患者层结果显示方向依方法而变。", "这是反对过度表述的关键图：当前不能声称TYMP统一位于晚期状态或促进肿瘤状态转移。")

    doc.add_page_break()
    add_heading(doc, "4. 全部图件索引", 1)
    add_text(doc, "同一图同时提供PNG和PDF时，下表按图的概念合并列示；逐文件格式、字节数、SHA-256和原始路径见《图件索引与证据边界_v1.tsv》。", size=9.5, color=MUTED)
    seen = set()
    rows = []
    for rec in records:
        key = (rec["group"], Path(rec["file_name"]).stem)
        if key in seen:
            continue
        seen.add(key)
        _, concept, evidence_class, use, boundary = CONCEPTS[key[1]]
        formats = "/".join(sorted({r["format"].upper() for r in records if r["group"] == key[0] and Path(r["file_name"]).stem == key[1]}))
        rows.append([rec["section"], concept + f"（{formats}）", use, boundary])
    add_table(doc, ["任务", "图件", "用途", "边界"], rows, [850, 2450, 2700, 3360], 7.8)

    add_heading(doc, "5. 主管会上建议使用的表述", 1)
    wording_rows = [
        ["推荐", "现有五队列恶性细胞分析显示，TYMP RNA与cNMF rank-7 F1转录程序相关；该关系在患者层有方向支持，但强度存在队列异质性。"],
        ["推荐", "F1程序的装载基因富集炎症、缺氧和糖酵解等恶性相关转录背景，为TYMP相关状态提供功能注释。"],
        ["推荐", "CNV、克隆和多方法轨迹结果未形成TYMP推动恶性进展的一致证据，因此因果促进仍需扰动和纵向验证。"],
        ["禁止", "TYMP已被证明促使肿瘤进入更恶性状态。"],
        ["禁止", "cNMF F1就是恶性程度最高的遗传克隆，或克隆树证明F1/TYMP位于晚期分支。"],
        ["禁止", "细胞数、树分支数或单个队列的显著性等同于跨患者、跨队列的因果证据。"],
    ]
    add_table(doc, ["口径", "建议文本"], wording_rows, [1100, 8260], 9.0)

    add_heading(doc, "6. 数据与复核入口", 1)
    add_text(doc, "关键统计表位于本包“05_关键统计表”目录，包括CNV判定、克隆树、F1–TYMP患者层及留一研究审计、cNMF rank/因子/通路结果和五方法轨迹汇总。所有图件的原始来源与哈希均记录在图件索引中。")
    add_callout(doc, "最终提醒", "本包适合主管讨论和决定下一步证据路线。若主管要求真正回答“TYMP是否促使肿瘤进入更恶性状态”，下一阶段应先冻结扰动或纵向分析设计，而不是继续从同一批横断面细胞中挑选更有利的算法结果。", PALE_GREEN, GREEN)

    props = doc.core_properties
    props.title = "项目26210 2.1–2.2图件证据说明"
    props.subject = "主管讨论用恶性上皮CNV、cNMF与轨迹证据包"
    props.author = "Project 26210 analysis team"
    props.comments = "Generated from validated project artifacts; causal boundaries preserved."
    doc.save(OUTPUT_DOCX)


def main():
    records = build_index()
    if not records:
        raise SystemExit("no figure records")
    if any(r["bytes"] <= 0 for r in records):
        raise SystemExit("zero-byte figure found")
    build_doc(records)
    readme = f"""项目26210《分析待做（新）》2.1–2.2图件包 v1

先看：项目26210_2.1-2.2图件证据说明_v1.docx
逐文件索引：图件索引与证据边界_v1.tsv

目录：
01_主管汇报主图：4张推荐汇报主图及PDF。
02_2.1_CNV与亚克隆支持图：良恶性判定、患者组成、克隆TYMP差值与阈值敏感性。
03_2.1_恶性轨迹边界图：8张已通过人工图检的正式轨迹图；Slingshot v1未纳入。
04_2.2_cNMF支持图：患者配对差异表达、热图和rank-7程序通路注释。
05_关键统计表：主管复核用的患者层与验证统计。

核心口径：现有数据支持TYMP与F1/cNMF转录程序相关，但不支持“TYMP促使肿瘤进入更恶性状态”的因果结论。
图件文件数：{len(records)}；PNG={sum(r['format']=='png' for r in records)}；PDF={sum(r['format']=='pdf' for r in records)}。
"""
    README.write_text(readme, encoding="utf-8-sig")
    print(json.dumps({"docx": str(OUTPUT_DOCX), "index": str(INDEX_TSV), "figures": len(records)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
