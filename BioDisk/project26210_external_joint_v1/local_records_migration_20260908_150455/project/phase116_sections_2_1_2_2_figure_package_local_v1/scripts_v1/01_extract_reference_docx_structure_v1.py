from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.document import Document as _Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_block_items(parent: _Document):
    parent_elm = parent.element.body
    for child in parent_elm.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def main() -> int:
    if len(sys.argv) != 3:
        raise SystemExit("usage: script.py input.docx output.json")
    src = Path(sys.argv[1])
    out = Path(sys.argv[2])
    doc = Document(src)
    blocks = []
    for index, block in enumerate(iter_block_items(doc), start=1):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            blocks.append(
                {
                    "index": index,
                    "type": "paragraph",
                    "style": block.style.name if block.style else "",
                    "text": text,
                }
            )
        else:
            rows = []
            for row in block.rows:
                rows.append([cell.text.replace("\n", " | ").strip() for cell in row.cells])
            blocks.append({"index": index, "type": "table", "rows": rows})
    payload = {
        "source": str(src),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "blocks": blocks,
    }
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: payload[k] for k in ("source", "paragraph_count", "table_count", "inline_shape_count")}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
