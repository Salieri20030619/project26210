from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SCRIPT_DIR = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("immune_report_v1", SCRIPT_DIR / "02_build_report_v1.py")
base = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(base)
original_page_break = base.page_break
page_break_count = 0


def add_callout_v2(doc, title, text, fill=base.CALLOUT, accent=base.BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = base.Inches(0.12)
    p.paragraph_format.right_indent = base.Inches(0.08)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "22")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    r1 = p.add_run(title)
    base.set_run_font(r1, size=10.5, bold=True, color=accent)
    r1.add_break()
    r2 = p.add_run(text)
    base.set_run_font(r2, size=10.0, color=base.INK)
    return p


def page_break_v2(doc):
    global page_break_count
    page_break_count += 1
    # Suppress: after full-page overview, after program-replication callout,
    # and after state-ordering callout. Word can flow directly without blanks.
    if page_break_count in {3, 5, 6}:
        return None
    return original_page_break(doc)


base.add_callout = add_callout_v2
base.page_break = page_break_v2


if __name__ == "__main__":
    base.main()
