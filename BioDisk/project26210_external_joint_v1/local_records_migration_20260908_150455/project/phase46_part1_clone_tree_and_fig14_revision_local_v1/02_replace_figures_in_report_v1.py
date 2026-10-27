from __future__ import annotations

import hashlib
import json
from pathlib import Path

from PIL import Image
from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
PHASE = Path(__file__).resolve().parent
INPUT = (
    ROOT
    / "phase45_part1_client_report_with_figures_local_v1"
    / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_WITH_FIGURES_REVIEW_v9.docx"
)
OUTPUT = PHASE / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_WITH_FIGURES_REVIEW_v10.docx"
FIGURES = PHASE / "figures_local_v1"
QC = PHASE / "qc_local_v1"
QC.mkdir(parents=True, exist_ok=True)


def replace_text_preserve_runs(paragraph, old: str, new: str) -> None:
    full = "".join(run.text for run in paragraph.runs)
    if old not in full:
        raise RuntimeError(f"Expected caption text not found: {old}")
    replaced = full.replace(old, new, 1)
    paragraph.runs[0].text = replaced
    for run in paragraph.runs[1:]:
        run.text = ""


def next_drawing_paragraph(document: Document, caption_index: int):
    for paragraph in document.paragraphs[caption_index + 1 :]:
        if paragraph._p.xpath(".//a:blip"):
            return paragraph
        if paragraph.text.strip():
            break
    raise RuntimeError(f"No inline image immediately following caption paragraph {caption_index}")


def replace_inline_image(document: Document, caption_key: str, image_path: Path) -> dict:
    caption_index = next(
        index for index, paragraph in enumerate(document.paragraphs) if caption_key in paragraph.text
    )
    paragraph = next_drawing_paragraph(document, caption_index)
    blips = paragraph._p.xpath(".//a:blip")
    if len(blips) != 1:
        raise RuntimeError(f"Expected one image after {caption_key}, found {len(blips)}")
    relationship_id = blips[0].get(qn("r:embed"))
    image_part = document.part.related_parts[relationship_id]
    image_part._blob = image_path.read_bytes()

    with Image.open(image_path) as image:
        pixel_width, pixel_height = image.size
    aspect = pixel_width / pixel_height

    extents = paragraph._p.xpath(".//wp:inline/wp:extent")
    if len(extents) != 1:
        raise RuntimeError(f"Expected one wp:extent after {caption_key}, found {len(extents)}")
    cx = int(extents[0].get("cx"))
    cy = int(round(cx / aspect))
    extents[0].set("cy", str(cy))
    for shape_extent in paragraph._p.xpath(".//pic:spPr/a:xfrm/a:ext"):
        shape_extent.set("cx", str(cx))
        shape_extent.set("cy", str(cy))

    for doc_property in paragraph._p.xpath(".//wp:docPr"):
        doc_property.set("name", image_path.stem)
        doc_property.set("descr", image_path.stem.replace("_", " "))

    return {
        "caption_key": caption_key,
        "caption_paragraph_index": caption_index,
        "relationship_id": relationship_id,
        "target_part": str(image_part.partname),
        "source_image": str(image_path.relative_to(ROOT)),
        "pixel_width": pixel_width,
        "pixel_height": pixel_height,
        "display_cx": cx,
        "display_cy": cy,
    }


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {OUTPUT}")
    document = Document(INPUT)

    fig03_caption = next(paragraph for paragraph in document.paragraphs if "Fig03" in paragraph.text)
    replace_text_preserve_runs(
        fig03_caption,
        "节点与分支按主导 cNMF 程序着色。",
        "按亚克隆分支着色，并标注每条分支最主要的染色体 CNV 变化。",
    )

    replacements = [
        replace_inline_image(
            document,
            "Fig03",
            FIGURES / "Fig03_CNV_clone_trees_by_cNMF_program_v4.png",
        ),
        replace_inline_image(
            document,
            "Fig14",
            FIGURES / "Fig14_NicheNet_ligand_receptor_evidence_v2.png",
        ),
    ]
    document.save(OUTPUT)

    status = {
        "status": "REPORT_REBUILT_PENDING_RENDER_QC",
        "input": str(INPUT.relative_to(ROOT)),
        "output": str(OUTPUT.relative_to(ROOT)),
        "output_bytes": OUTPUT.stat().st_size,
        "output_sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        "replacements": replacements,
    }
    (QC / "report_replacement_status_v1.json").write_text(
        json.dumps(status, indent=2, ensure_ascii=False), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
