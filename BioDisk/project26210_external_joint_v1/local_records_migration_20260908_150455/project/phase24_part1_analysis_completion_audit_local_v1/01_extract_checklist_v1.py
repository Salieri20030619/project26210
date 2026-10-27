#!/usr/bin/env python
"""Extract all reviewable text and structural metadata from the analysis checklist DOCX."""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs).strip()


def table_to_rows(table: Table) -> list[list[str]]:
    return [[cell.text.strip().replace("\r", " ") for cell in row.cells] for row in table.rows]


def iter_blocks(document: Document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield "paragraph", Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield "table", Table(child, document)


def extract_xml_text(docx_path: Path) -> dict[str, list[str]]:
    output: dict[str, list[str]] = {}
    with zipfile.ZipFile(docx_path) as archive:
        xml_names = sorted(
            name
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )
        for name in xml_names:
            from lxml import etree

            root = etree.fromstring(archive.read(name))
            texts = [t for t in root.xpath("//*[local-name()='t']/text()") if t.strip()]
            if texts:
                output[name] = texts
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    document = Document(args.input_docx)

    blocks: list[dict] = []
    table_count = 0
    paragraph_count = 0
    for block_type, block in iter_blocks(document):
        if block_type == "paragraph":
            paragraph_count += 1
            text = paragraph_text(block)
            num_pr = block._p.pPr.numPr if block._p.pPr is not None else None
            num_id = num_pr.numId.val if num_pr is not None and num_pr.numId is not None else None
            ilvl = num_pr.ilvl.val if num_pr is not None and num_pr.ilvl is not None else None
            blocks.append(
                {
                    "block": "paragraph",
                    "index": paragraph_count,
                    "style": block.style.name if block.style else "",
                    "num_id": num_id,
                    "level": ilvl,
                    "text": text,
                }
            )
        else:
            table_count += 1
            blocks.append(
                {
                    "block": "table",
                    "index": table_count,
                    "rows": table_to_rows(block),
                }
            )

    package_names: list[str]
    with zipfile.ZipFile(args.input_docx) as archive:
        package_names = sorted(archive.namelist())

    result = {
        "input": str(args.input_docx.resolve()),
        "sha256": sha256(args.input_docx),
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "sections": len(document.sections),
        "blocks": blocks,
        "xml_text": extract_xml_text(args.input_docx),
        "package_members": package_names,
    }

    json_path = args.output_dir / "checklist_structure_v1.json"
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = [
        f"INPUT\t{result['input']}",
        f"SHA256\t{result['sha256']}",
        f"PARAGRAPHS\t{result['paragraphs']}",
        f"TABLES\t{result['tables']}",
        f"INLINE_SHAPES\t{result['inline_shapes']}",
        f"SECTIONS\t{result['sections']}",
        "",
        "DOCUMENT_ORDER",
    ]
    for item in blocks:
        if item["block"] == "paragraph":
            lines.append(
                f"P{item['index']:03d}\tstyle={item['style']}\tnum={item['num_id']}\tlevel={item['level']}\t{item['text']}"
            )
        else:
            lines.append(f"TABLE{item['index']:03d}")
            for row_index, row in enumerate(item["rows"], start=1):
                lines.append(f"T{item['index']:03d}R{row_index:03d}\t" + "\t".join(row))

    lines.extend(["", "RAW_XML_TEXT"])
    for name, texts in result["xml_text"].items():
        lines.append(f"[{name}]")
        lines.extend(texts)
    (args.output_dir / "checklist_structure_v1.txt").write_text(
        "\n".join(lines) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
