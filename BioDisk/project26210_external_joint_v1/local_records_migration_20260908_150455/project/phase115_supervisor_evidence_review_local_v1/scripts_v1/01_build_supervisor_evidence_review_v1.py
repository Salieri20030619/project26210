from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r'D:\project26210_external_joint_v1\phase115_supervisor_evidence_review_local_v1')
ASSETS = ROOT / 'assets_v1'
OUTPUT = ROOT / 'output_v1'
DOCX = OUTPUT / '项目26210_现有证据总览与下一步决策简报_v1.docx'
FIGURE = ASSETS / 'current_evidence_chain_status_v1.png'

FONT = 'Microsoft YaHei'
NAVY = '17365D'
BLUE = '1F4E79'
MID_BLUE = '2F5597'
TEXT = '263238'
MUTED = '5B6573'
GRID = 'B7C5D5'
HEADER_FILL = 'D9EAF7'
ALT_FILL = 'F7F9FB'
GREEN = 'D9EAD3'
AMBER = 'FFF2CC'
RED = 'F4CCCC'
GRAY = 'E7E6E6'
WHITE = 'FFFFFF'
TABLE_WIDTH = 9400


def set_run_font(run, size=10.5, bold=None, color=TEXT, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn(f'w:{key}'), FONT)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcpr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=70, bottom=70, start=100, end=100):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tcpr.append(tc_mar)
    for name, value in [('top', top), ('bottom', bottom), ('start', start), ('end', end)]:
        node = tc_mar.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def mark_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement('w:tblHeader')
    node.set(qn('w:val'), 'true')
    trpr.append(node)


def set_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblpr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = borders.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), GRID)


def remove_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblpr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = borders.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'nil')


def set_table_geometry(table, widths):
    if sum(widths) != TABLE_WIDTH:
        raise ValueError(f'table widths must sum to {TABLE_WIDTH}: {widths}')
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblpr = table._tbl.tblPr
    layout = tblpr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tblpr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    tblw = tblpr.find(qn('w:tblW'))
    if tblw is None:
        tblw = OxmlElement('w:tblW')
        tblpr.append(tblw)
    tblw.set(qn('w:w'), str(TABLE_WIDTH))
    tblw.set(qn('w:type'), 'dxa')
    tblind = tblpr.find(qn('w:tblInd'))
    if tblind is None:
        tblind = OxmlElement('w:tblInd')
        tblpr.append(tblind)
    tblind.set(qn('w:w'), '120')
    tblind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width / 1440 * 2.54)
            tcw = cell._tc.get_or_add_tcPr().find(qn('w:tcW'))
            if tcw is None:
                tcw = OxmlElement('w:tcW')
                cell._tc.get_or_add_tcPr().append(tcw)
            tcw.set(qn('w:w'), str(width))
            tcw.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
    set_table_borders(table)


def add_cell_text(cell, text, size=8.7, bold=False, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, widths, status_column=None, font_size=8.5):
    table = document.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    mark_header(table.rows[0])
    for index, header in enumerate(headers):
        set_cell_shading(table.cell(0, index), HEADER_FILL)
        add_cell_text(table.cell(0, index), header, size=8.7, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
    status_fill = {'支持': GREEN, '部分支持': AMBER, '情境依赖': AMBER, '不支持': RED, '未完成': GRAY, '不可估计': GRAY, '待检验': GRAY}
    for row_index, row_values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for column_index, value in enumerate(row_values):
            if row_index % 2 == 0:
                set_cell_shading(cells[column_index], ALT_FILL)
            if status_column is not None and column_index == status_column:
                key = str(value)
                fill = next((color for label, color in status_fill.items() if label in key), None)
                if fill:
                    set_cell_shading(cells[column_index], fill)
            add_cell_text(cells[column_index], value, size=font_size, bold=(status_column == column_index), align=WD_ALIGN_PARAGRAPH.LEFT)
        set_table_geometry(table, widths)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    return table


def add_body(document, text, bold_prefix=None, italic=False, color=TEXT, space_after=5):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(space_after)
    if bold_prefix and text.startswith(bold_prefix):
        lead = paragraph.add_run(bold_prefix)
        set_run_font(lead, bold=True, color=BLUE)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, italic=italic, color=color)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, italic=italic, color=color)
    return paragraph


def add_bullet(document, text, level=0):
    paragraph = document.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    paragraph.paragraph_format.left_indent = Cm(1.27 if level == 0 else 1.9)
    paragraph.paragraph_format.first_line_indent = Cm(-0.63)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f'Heading {level}')
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size={1: 15, 2: 12, 3: 10.5}[level], bold=True, color={1: BLUE, 2: MID_BLUE, 3: BLUE}[level])
    return paragraph


def add_callout(document, title, text, fill=HEADER_FILL, accent=BLUE):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ''
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    title_run = paragraph.add_run(title + '  ')
    set_run_font(title_run, size=10.5, bold=True, color=accent)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10.2, color=TEXT)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MUTED)


def set_alt_text(shape, description):
    docpr = shape._inline.docPr
    docpr.set('descr', description)


def create_evidence_figure():
    ASSETS.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(13.2, 3.2))
    ax.set_xlim(0, 13.2)
    ax.set_ylim(0, 3.2)
    ax.axis('off')
    nodes = [
        (0.2, 2.0, 1.7, 'F1 - TYMP', 'Supported\nheterogeneous', '#D9EAD3'),
        (2.25, 2.0, 1.7, 'TYMP - MTA', 'Not supported', '#F4CCCC'),
        (4.3, 2.0, 1.7, 'MTA - SPP1', 'Not closed', '#FFF2CC'),
        (6.35, 2.0, 1.7, 'TYMP/F1 - SPP1', 'Not supported', '#F4CCCC'),
        (8.4, 2.0, 1.8, 'SPP1 - Endothelium', 'Supported', '#D9EAD3'),
        (10.6, 2.0, 2.2, 'Immune / Clinical', 'Not supported', '#F4CCCC'),
    ]
    for x, y, width, title, status, fill in nodes:
        box = FancyBboxPatch((x, y), width, 0.82, boxstyle='round,pad=0.04,rounding_size=0.07', facecolor=fill, edgecolor='#5B6573', linewidth=1.2)
        ax.add_patch(box)
        ax.text(x + width / 2, y + 0.53, title, ha='center', va='center', fontsize=10.2, weight='bold', color='#17365D')
        ax.text(x + width / 2, y + 0.22, status, ha='center', va='center', fontsize=8.8, color='#333333')
    for start, end in [(1.9, 2.25), (3.95, 4.3), (6.0, 6.35), (8.05, 8.4), (10.2, 10.6)]:
        ax.annotate('', xy=(end, 2.41), xytext=(start, 2.41), arrowprops=dict(arrowstyle='->', color='#7A8793', lw=1.4))
    ax.text(6.6, 1.35, 'The only reproducible downstream anchor is the SPP1 - endothelial / vascular interface.', ha='center', fontsize=11, weight='bold', color='#1F4E79')
    ax.text(6.6, 0.85, 'Positive downstream evidence cannot repair unsupported upstream edges.', ha='center', fontsize=10, color='#9B1C1C')
    ax.text(6.6, 0.3, 'Current complete chain: NOT ESTABLISHED', ha='center', fontsize=12, weight='bold', color='#9B1C1C')
    fig.savefig(FIGURE, dpi=220, bbox_inches='tight', facecolor='white')
    plt.close(fig)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)

    styles = document.styles
    normal = styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for level, size, color, before, after in [(1, 15, BLUE, 14, 7), (2, 12, MID_BLUE, 10, 5), (3, 10.5, BLUE, 7, 3)]:
        style = styles[f'Heading {level}']
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = hp.add_run('项目26210 | 主管会议证据与决策简报')
    set_run_font(run, size=8.5, color=MUTED)
    ppr = hp._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '3')
    bottom.set(qn('w:color'), GRID)
    pbdr.append(bottom)
    ppr.append(pbdr)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Cm(16.6))
    set_table_geometry(table, [4700, 4700])
    remove_table_borders(table)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
    left = table.cell(0, 0).paragraphs[0]
    lr = left.add_run('内部讨论稿 | 2026-09-01')
    set_run_font(lr, size=8.5, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = right.add_run('第 ')
    set_run_font(label, size=8.5, color=MUTED)
    add_page_field(right)
    suffix = right.add_run(' 页')
    set_run_font(suffix, size=8.5, color=MUTED)


def build():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    create_evidence_figure()
    document = Document()
    configure_document(document)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(28)
    title.paragraph_format.space_after = Pt(10)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run('项目26210：现有证据总览与下一步决策简报')
    set_run_font(run, size=22, bold=True, color=NAVY)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    subrun = subtitle.add_run('供主管会议内部讨论 | 完整证据、当前卡点、客户需求变化与可选路线')
    set_run_font(subrun, size=11, color=MUTED)

    add_callout(
        document,
        '一句话结论',
        '项目并非“没有结果”，而是完整 TYMP-MTA-SPP1-血管-免疫/疗效链条的多个上游和临床边未获支持；目前最稳定的阳性锚点位于下游 SPP1 巨噬细胞-内皮/血管界面。客户最新要求相当于把主线改为治疗反应和免疫生态，而不是继续补代谢闭环。',
        fill=HEADER_FILL,
        accent=NAVY,
    )

    add_heading(document, '会议需要解决的三个问题', 1)
    add_table(
        document,
        ['问题', '当前事实', '会议需拍板'],
        [
            ('1. 主线是否改写', '代谢闭环与完整七特征均未获公共数据支持；客户明确要求代谢降为补充。', '是否转为“治疗反应分组下的 TYMP/F1-免疫生态”主线。'),
            ('2. TYMP-SPP1如何表述', '原始单细胞、细胞通讯和空间结果均不支持一般性 TYMP/F1-SPP1 边。', '把它作为待检验桥接，还是允许作为客户实验前提单独放置。'),
            ('3. 下一轮做到哪里停', '现有阳性集中在 SPP1-内皮/血管；免疫排斥和临床转化未闭合。', '批准最小反应分组分析，还是只围绕已支持的下游空间生态收束。'),
        ],
        [1850, 3650, 3900],
        font_size=9.0,
    )
    add_body(document, '建议阅读顺序：先看第1节的结论与卡点，再看第3节证据矩阵；第6-8节用于会议决策。', italic=True, color=MUTED)
    document.add_page_break()

    add_heading(document, '1. 执行摘要：项目为什么“进行不下去”', 1)
    add_body(document, '当前卡点不是统计方法不足，而是拟议主线要求多个连续证据边同时成立，但现有独立数据在关键上游边和临床终点上反复出现阴性或不可估计结果。继续增加同类敏感性分析不会自动形成闭环。')
    add_table(
        document,
        ['层级', '已经完成', '结果', '对主线的影响'],
        [
            ('F1-TYMP肿瘤状态', '五研究单细胞患者级相关及稳定性审计', '支持，但存在研究异质性和GSE207493权重', '可作为肿瘤状态起点，不能直接解释为TYMP功能或因果。'),
            ('代谢桥', 'UnitedMet、NCI60、CCL180、CCLE及七特征迁移', '完整桥不支持；部分缩减特征有候选信号', '不能再把代谢公共验证设为免疫分析成功前提，也不能写成已闭环。'),
            ('TYMP/F1-SPP1', '患者级丰度/极化、细状态通讯、空间邻域', '一般性边不支持', '“TYMP通过SPP1巨噬细胞”不能作为当前结论。'),
            ('SPP1-内皮/血管', '14患者细状态通讯及24切片空间验证', '稳定阳性', '这是当前最可用的下游机制锚点。'),
            ('T细胞排斥/疗效', '空间、免疫治疗单细胞及bulk临床模型', '异质或不支持', '尚不能形成免疫排斥或治疗预测结论。'),
        ],
        [1500, 2750, 2400, 2750],
        status_column=2,
        font_size=8.5,
    )
    add_callout(document, '核心矛盾', '客户希望得到“TYMP通过SPP1巨噬细胞造成T细胞功能下降/免疫屏障”的直观故事；项目现有证据只稳定支持故事的中下游一段，即SPP1巨噬细胞与内皮/血管生态的联系。', fill=RED, accent='9B1C1C')

    add_heading(document, '2. 数据必须分开：现有结果来自不同统计宇宙', 1)
    add_body(document, '不同数据集回答不同问题。共同出现TYMP或SPP1并不意味着样本、终点和统计单位可以合并；任何后续汇报仍需保持下列边界。')
    add_table(
        document,
        ['数据宇宙', '主要数据', '独立单位', '可回答问题', '不可替代'],
        [
            ('D01 原始单细胞', 'GSE156632/159115/178481/207493/242299', '患者；研究内分析', 'F1-TYMP、巨噬细胞状态、轨迹、调控、患者级桥接', '免疫治疗应答或空间定位'),
            ('D02 免疫治疗单细胞', 'GSE294109、GSE264586', '患者/配对患者', '反应分组和纵向免疫状态', 'D01机制的直接验证'),
            ('D03 三队列桥接', 'GSE207493/264586/294109', '患者内队列', 'SPP1细状态复现与情境差异', '把重叠队列当成同一分析'),
            ('D04-D06 空间', 'GLMF、TJ-RCC、GSE175540', '切片/患者；制备分层', '共定位、邻域、血管/TLS结构', '用spot数代替患者数'),
            ('D08 bulk临床', '外部ICI/PFS/OS队列', '患者', '预后与治疗关联', '单细胞机制证据'),
            ('D09-D12 公共代谢/细胞系', 'UnitedMet、NCI60、CCL180、CCLE', '肿瘤样本/细胞系', '代谢物与RNA的观察性迁移', '客户TYMP敲低实验'),
            ('D14 客户实验', '客户未公开/未完整移交数据', '客户定义生物学重复', '直接扰动与因果验证', '由公共数据库推测或替代'),
        ],
        [1150, 2050, 1300, 2650, 2250],
        font_size=8.1,
    )
    document.add_page_break()

    add_heading(document, '3. 完整证据链矩阵', 1)
    picture = document.add_picture(str(FIGURE), width=Cm(16.2))
    set_alt_text(picture, '当前证据链状态：F1-TYMP关联有支持；TYMP-MTA和TYMP/F1-SPP1不支持；MTA-SPP1未闭合；SPP1-内皮支持；免疫和临床转化不支持。')
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    crun = caption.add_run('图1  当前证据链状态。绿色为支持，黄色为未闭合，红色为当前数据不支持。')
    set_run_font(crun, size=9, color=MUTED, italic=True)

    evidence_rows = [
        ('E01', 'F1肿瘤程序-TYMP RNA', '支持但异质', '20例中17例正向；中位rho=0.115；BH q=0.0197；GSE207493 10/10正向', '可说相关；不可说F1驱动TYMP或TYMP功能已确认。'),
        ('E02', 'TYMP特异功能/代谢活性', '部分支持', '严格TYMP轴0/4；较宽嘧啶/代谢情境3/4为次要信号', '只能保留宽泛代谢情境，不能定义稳定TYMP特异功能轴。'),
        ('E03', 'TYMP-直接测量MTA', '不支持', 'UnitedMet 255例meta rho=0.0395，P=0.538；NCI60全谱系阴性；CCL180/CCLE不可估计', '公共基线数据不支持TYMP驱动MTA。'),
        ('E04', 'MTA-SPP1诱导', '未完成', '无合格MTA刺激+SPP1 RNA/蛋白+阻断/救援数据', '仍是扰动假设；实验可行性不等于公共数据闭环。'),
        ('E05', 'TYMP/F1-SPP1巨噬细胞', '不支持', '20例0/4丰度/极化终点；14例0/2肿瘤-SPP1通讯；空间TYMP邻域SPP1 q约0.39/0.38', '不能说TYMP一般性地通过SPP1巨噬细胞。'),
        ('E06-E08', 'SPP1-内皮/血管', '支持', 'SPP1->内皮14/14，内皮->SPP1 13/14，均q=0.00305；GSE175540两制备空间rho约0.21，q=0.0022/0.0044', '可说计算通讯兼容性和数据集特异空间关联；不可说配体结合或因果。'),
        ('E09', '血管-T细胞排斥/ICI/PFS/OS', '不支持', '空间方向异质；GSE294109应答P=0.561/0.724；bulk临床0个主要模型过FDR', '当前不能形成免疫排斥或临床预测结论。'),
        ('E10', '完整七特征迁移', '不支持', 'RC18 n=26 rho=-0.137，q=0.972；RC20 n=13 rho=0.011，q=0.972；meta P=0.594', '七特征阴性保留为补充，不再作为免疫主线门槛。'),
    ]
    add_table(document, ['编号', '证据边', '状态', '精确证据', '允许结论'], evidence_rows, [600, 1850, 1100, 3350, 2500], status_column=2, font_size=7.9)
    document.add_page_break()

    add_heading(document, '4. 现有证据逐层展开', 1)
    add_heading(document, '4.1 可以保留的上游起点：F1-TYMP相关，但不是TYMP功能证明', 2)
    add_body(document, '原始五研究恶性细胞患者级分析中，20例合格患者有17例为正向，研究内综合BH q=0.0197；GSE207493的10例均为正向，留一范围0.131-0.163。该结果说明F1程序与TYMP RNA在特定肿瘤状态中相关。')
    add_body(document, '限制：总体效应受GSE207493权重影响，研究间存在明显异质性；严格TYMP酶功能轴0/4未获支持。可以把F1/TYMP作为肿瘤状态标志，但不能把它直接升级为TYMP功能或因果起点。', bold_prefix='限制：')

    add_heading(document, '4.2 代谢桥：直接MTA与完整七特征均未形成公共复现', 2)
    add_table(
        document,
        ['分析', '独立n', '结果', '判定'],
        [
            ('UnitedMet TYMP-MTA', '255肿瘤样本', '固定meta rho=0.0395，P=0.538，队列方向分裂', '不支持观察性关联'),
            ('UnitedMet MTAP-MTA', '255肿瘤样本', '固定meta rho=-0.271，P=1.44e-05；随机P=0.00957', '情境对照支持，但不能替代TYMP结果'),
            ('NCI60 TYMP-MTA', '58细胞系；肾癌8', '全谱系rho=-0.064，P=0.634；肾癌rho=0.330，仅探索', '全谱系不支持；肾癌无推断'),
            ('CCL180/CCLE', '155/148细胞系', '直接MTA不可估计；相邻代谢物0/12过双FDR门槛', '不能用替代物补MTA'),
            ('七特征S7', 'RC18 26；RC20 13完整病例', 'rho=-0.137/0.011；两者q=0.972；meta P=0.594', '完整迁移不支持'),
            ('三特征S3', 'CPTAC 50；验证55', 'rho=-0.427，q=0.004；验证rho=-0.245，q=0.071', '候选子模块，不是七特征验证'),
        ],
        [2100, 1500, 3700, 2100],
        status_column=3,
        font_size=8.3,
    )
    add_body(document, '完整七特征正式判定代码：CURRENT_DATA_DO_NOT_SUPPORT_GENERAL_SEVEN_FEATURE_TRANSFER。', bold_prefix='完整七特征正式判定代码：')
    add_callout(document, '会议解释', '代谢分析并非“白做”：它明确回答了公共数据不能补上客户实验闭环。正确处理方式是保留阴性和候选三特征结果，把代谢降为补充层，而不是继续增加敏感性分析。', fill=AMBER, accent='7A5A00')

    add_heading(document, '4.3 SPP1巨噬细胞状态本身成立，但不等于TYMP驱动', 2)
    add_body(document, '在原始单细胞中，71,045个巨噬细胞池被划分为Other Macrophage 40,429、SPP1-state 15,322和C1QC-state 15,294。50,761个髓系细胞、40例患者、5个研究分图的DPT和Monocle3均支持Monocyte到SPP1-state的横断面后移。')
    add_body(document, '关键阴性：SPP1-versus-C1QC程序沿排序在5/5研究中都没有预期上升，因此不能描述为C1QC被SPP1统一替代，也不能写成真实时间的谱系转换。CEBPD/JUNB调控子为状态关联，不是直接结合或因果。', bold_prefix='关键阴性：')

    add_heading(document, '4.4 TYMP/F1-SPP1桥接是当前最主要断点', 2)
    add_body(document, '针对该断点的补充分析已经直接测试，而不是停留在“没有做”：20例患者中4个肿瘤轴-SPP1丰度/极化终点0/4获得支持，所有BH q至少为0.677；14例细状态通讯中，肿瘤->SPP1和SPP1->肿瘤均未获支持。GSE175540中肿瘤TYMP到邻近SPP1的空间相关在FFPE/冷冻中rho约0.009，q=0.393/0.380。')
    add_callout(document, '结论边界', '客户希望的“TYMP通过SPP1巨噬细胞”目前只能作为新免疫治疗队列中的待检验假设。既往阴性结果必须保留；新队列阳性也只能形成情境特异证据，不能抹去原始队列的阴性。', fill=RED, accent='9B1C1C')

    add_heading(document, '4.5 当前最稳定阳性：SPP1-内皮/血管界面', 2)
    add_body(document, '14例细状态CellChat重聚合显示，SPP1->内皮特异性14/14为正，中位0.21065，BH q=0.00305；内皮->SPP1为13/14正，中位0.14642，BH q=0.00305，留一患者后方向稳定。')
    add_body(document, 'GSE175540中，SPP1-state与血管重塑在12张FFPE和12张冷冻切片中均为正，rho=0.2121/0.2083，BH q=0.00220/0.00440。这两层证据可以三角互证，但不能把CellChat当成配体结合，也不能从空间共定位推出因果。')

    add_heading(document, '4.6 免疫排斥和治疗转化尚未闭合', 2)
    add_body(document, 'GSE294109中F1-TYMP与ICB获益比较样本极小（扩展5 vs 2；严格4 vs 1），P=0.561/0.724，不支持应答标志物。扩展定义下F1-TYMP与Treg比例有小样本单队列信号（7例，rho=0.929，q=0.0277），但严格层未过FDR，不能泛化。')
    add_body(document, 'GSE264586纵向分析每模型仅3-4对患者，54项特征没有一项过FDR。三队列细状态通讯没有主要终点在任一数据集正式支持。外部bulk ICI/PFS/OS中，预注册主要模型0个通过FDR。因此“治疗反应组/非反应组都有”并不等于现成结论已经存在。')
    document.add_page_break()

    add_heading(document, '5. 客户最新反馈：合理的需求纠偏与不可直接接受的结论', 1)
    add_table(
        document,
        ['客户表达', '合理部分', '证据风险', '建议转换'],
        [
            ('已有反应组/非反应组', '可以把治疗反应作为新的组织主线。', '各队列样本量、反应定义、时间点和细胞组成不同；现有分析并未形成稳定标志物。', '先做患者标签和可估计性审计，再分队列分析。'),
            ('代谢不是主线，只做特征性展示', '符合当前公共代谢阴性结果，也可减少无效敏感性扩展。', '单细胞只能显示转录通路代理，不能显示5-MTA等代谢物或代谢流。', 'UMAP仅展示TYMP及嘧啶/甲硫氨酸/SAM/多胺通路评分，并标注proxy。'),
            ('证明TYMP与免疫相关', '可以检验TYMP/F1与应答、SPP1状态、T细胞程序的患者级关联。', '“相关”需要跨队列方向；小样本细胞级P值不能替代患者n。', '冻结主要端点，患者为统计单位，队列分开。'),
            ('TYMP通过SPP1巨噬细胞', '是有生物学吸引力的桥接假设。', '既往0/4、0/2及空间阴性；目前不能预设为事实。', '设置三段门槛：TYMP-应答、TYMP-SPP1、SPP1-T细胞/应答。'),
            ('T细胞杀伤能力降低', '可分析细胞毒、效应和耗竭转录程序。', '转录组不能证明真实杀伤功能。', '表述为细胞毒性程序降低/耗竭增强；功能杀伤需实验。'),
            ('形成免疫屏障', '空间数据可做邻域、距离、边界和排斥分析。', '共定位不能自动等于物理屏障。', '预设空间几何指标，以切片/患者为单位。'),
            ('不要纠结敏感性', '同意不再做无目的模型扩展。', '必要QC和主要敏感性仍必须保留在内部审计。', '正文结果导向，完整QC放补充。'),
        ],
        [1750, 2200, 2900, 2550],
        font_size=7.9,
    )
    add_callout(document, '对主管的关键说明', '客户反馈改变的是“下一步分析主线”，不是让既往阴性结果失效。若转向反应分组免疫主线，应作为新的预注册模块启动，而不是把旧结果重新命名为支持。', fill=HEADER_FILL, accent=NAVY)

    add_heading(document, '6. 现在可以说什么、不能说什么', 1)
    add_table(
        document,
        ['可以说', '不能说'],
        [
            ('F1与TYMP RNA在原始恶性细胞中相关，但有研究异质性。', 'F1驱动TYMP，或TYMP特异功能已经由单细胞确认。'),
            ('SPP1巨噬细胞是可重复定义的状态，并呈现横断面状态排序。', 'SPP1状态由TYMP普遍诱导，或C1QC状态被统一替代。'),
            ('SPP1巨噬细胞与内皮/血管生态存在稳定计算和空间关联。', '已经证明配体结合、血管重塑因果或物理免疫屏障。'),
            ('完整七特征公共迁移当前不支持；三特征有候选线索。', '公共数据库复现了客户TYMP敲低实验，或三特征救回七特征。'),
            ('可在应答队列中重新检验TYMP相关免疫表型。', '已有反应分组就等于TYMP-应答或TYMP-SPP1已经被证明。'),
        ],
        [4700, 4700],
        font_size=8.8,
    )

    add_heading(document, '7. 主管可选路线', 1)
    option_rows = [
        ('A. 反应分组免疫主线', '以D02/D03反应组/非反应组为主，检验TYMP/F1、SPP1巨噬细胞、T细胞程序及通讯。', '最贴合客户；可以直接回答免疫治疗语境。', '样本量小；既往已有阴性；可能再次得不到TYMP-SPP1桥。', '可做，但必须允许阴性和情境特异结论。'),
        ('B. 下游锚点主线（推荐）', '以已支持的SPP1-内皮/血管生态为结果主轴；TYMP/F1作为独立肿瘤状态，客户实验作为独立上游验证接口。', '证据最稳；最容易形成可信结果。', '不能满足“TYMP通过SPP1”的强机制措辞。', '推荐作为报告骨架，并叠加A的最小验证。'),
        ('C. 三特征代谢子模块', '将5-MTA/Methionine/Thymine重新注册为候选S3，寻找外部验证或转录代理。', 'CPTAC方向较好；可作为补充特色。', '验证队列未过FDR；仍非免疫主线。', '只做补充，不建议优先。'),
        ('D. 继续补完整代谢闭环', '继续搜索或增加七特征/相邻代谢物敏感性。', '可能找到新数据。', '高时间成本；现有多平台阴性；偏离客户需求。', '不推荐。'),
        ('E. 暂停计算等待客户数据', '不启动新分析，等待客户实验或明确稿件结构。', '避免继续投入。', '客户已表示不再追加湿实验，项目可能长期停滞。', '只有主管不接受叙事降级时考虑。'),
    ]
    add_table(document, ['路线', '内容', '优点', '风险', '建议'], option_rows, [1500, 2600, 1750, 2100, 1450], font_size=7.7)
    add_callout(document, '综合建议', '采用“B作为报告骨架 + A作为最小新分析”。即：不再强求完整代谢链；先用应答队列检验TYMP/F1是否与免疫生态相关；无论TYMP-SPP1是否成立，都保留SPP1-内皮/血管作为独立、已支持的下游模块。', fill=GREEN, accent='274E13')

    add_heading(document, '8. 若批准继续：最小新分析块与停止规则', 1)
    add_table(
        document,
        ['步骤', '具体分析', '主要统计单位/端点', 'Go/No-Go规则'],
        [
            ('0. 标签审计', '逐队列核对Responder/Non-responder定义、治疗前后时间点、患者去重、细胞数和可用细胞类型。', '患者；每队列独立n', '任一组<3例只做描述；不合并不同定义。'),
            ('1. 肿瘤起点', '恶性细胞患者伪bulk：TYMP、F1及预设肿瘤状态分数比较应答组。', '患者级效应；队列内检验和FDR', '至少两个独立队列同向，才称应答相关；否则为情境依赖/不支持。'),
            ('2. 巨噬细胞与T细胞', 'SPP1/C1QC/Other比例和程序；CD8细胞毒、效应、耗竭模块；均按患者汇总。', '患者级比例/模块分数', '不以细胞数作n；小样本优先效应和一致性。'),
            ('3. 桥接检验', 'TYMP/F1-SPP1、SPP1-T细胞、TYMP/F1-T细胞三组预设关联；按队列分开。', '患者级Spearman/置换', '若TYMP/F1-SPP1仍不支持，禁止“通过SPP1”措辞，改为平行模块。'),
            ('4. 细胞通讯', '肿瘤->SPP1、SPP1->T细胞/内皮的预设通讯家族，患者级聚合。', '患者为n；配体受体家族FDR', 'CellChat只作通讯兼容性，不作因果。'),
            ('5. 空间屏障', 'TYMP/F1肿瘤区、SPP1巨噬细胞、CD8细胞毒性之间的距离、邻域、边界和排斥梯度。', '切片/患者为n；制备分层', 'SPP1-血管阳性不能替代TYMP-SPP1或T细胞排斥。'),
            ('6. 代谢展示', 'UMAP/空间图展示TYMP及嘧啶、甲硫氨酸/SAM、多胺转录通路分数。', '描述性；无因果检验', '必须标注transcriptional proxy，不显示为代谢物测量。'),
        ],
        [1050, 3100, 2200, 3050],
        font_size=7.8,
    )

    add_heading(document, '9. 主管会议需要明确的决策', 1)
    add_table(
        document,
        ['决策项', '推荐选择', '若不接受的后果'],
        [
            ('是否接受由“完整因果链”降级为“肿瘤状态-免疫生态关联链”', '接受', '现有证据无法诚实支持强因果主线。'),
            ('报告主锚点', 'SPP1-内皮/血管生态', '若必须以TYMP-SPP1为锚点，项目仍处于证据断裂状态。'),
            ('是否批准反应分组最小新分析块', '批准步骤0-3，结果后再决定通讯和空间扩展', '不做则无法回应客户最新需求。'),
            ('客户实验在报告中的位置', '独立扰动验证接口/研究前提', '不能由公共数据预测或代替。'),
            ('对阴性结果的容忍度', '允许保留并改变叙事', '若只接受预设阳性，继续分析存在明显选择性报告风险。'),
            ('是否继续公共代谢扩展', '暂停', '继续会增加投入但不解决客户当前免疫主线。'),
        ],
        [2800, 3800, 2800],
        font_size=8.5,
    )
    add_callout(document, '建议会议结论模板', '批准以SPP1-内皮/血管生态作为已验证锚点；批准在免疫治疗反应队列中开展患者级最小分析，TYMP-SPP1作为待检验桥而非预设事实；代谢降为补充性转录通路展示；若TYMP-SPP1再次阴性，则报告平行关联而不继续救链。', fill=GREEN, accent='274E13')
    document.add_page_break()

    add_heading(document, '附录A：关键统计结果速查', 1)
    add_table(
        document,
        ['模块', '数据/独立n', '效应与统计', '正式判定'],
        [
            ('F1-TYMP', 'D01 20患者', '17/20正向；中位rho=0.115；BH q=0.0197', '支持但异质'),
            ('GSE207493稳定性', '10患者', '10/10正向；留一rho=0.131-0.163', '单队列稳定'),
            ('SPP1细状态', '71,045细胞描述；40患者推断', 'Other 40,429；SPP1 15,322；C1QC 15,294', '操作性状态成立'),
            ('SPP1 vs C1QC轨迹', '5研究', '0/5预期正向', '不支持统一替代'),
            ('TYMP/F1-SPP1丰度/极化', '20患者', '0/4终点支持；所有q>=0.677', '一般性边不支持'),
            ('肿瘤-SPP1通讯', '14患者', '0/2端点支持', '情境依赖/不支持'),
            ('SPP1->内皮', '14患者', '14/14正；中位0.21065；q=0.00305', '支持计算兼容性'),
            ('内皮->SPP1', '14患者', '13/14正；中位0.14642；q=0.00305', '支持计算兼容性'),
            ('SPP1-血管空间', '12 FFPE+12冷冻', 'rho=0.2121/0.2083；q=0.00220/0.00440', '数据集特异支持'),
            ('肿瘤TYMP-邻近SPP1', '12 FFPE+12冷冻', 'rho=0.00894/0.00985；q=0.3928/0.3802', '不支持'),
            ('ICB反应', 'GSE294109 5vs2/4vs1', 'P=0.561/0.724', '不支持标志物'),
            ('完整S7', 'RC18 26；RC20 13', 'rho=-0.137/0.011；q=0.972/0.972；meta P=0.594', '不支持'),
            ('缩减S3', 'CPTAC 50；验证55', 'rho=-0.427/-0.245；q=0.004/0.071', '候选，未完成验证'),
        ],
        [2200, 1900, 3300, 2000],
        status_column=3,
        font_size=8.0,
    )

    add_heading(document, '附录B：权威结果路径', 1)
    path_rows = [
        ('数据集与结论分离', r'D:\project26210_external_joint_v1\phase106_dataset_separation_audit_local_v1\DATASET_SPECIFIC_CLAIM_MATRIX_v1.tsv'),
        ('链条证据矩阵', r'D:\project26210_external_joint_v1\phase107_evidence_escalation_audit_local_v1\CHAIN_EDGE_EVIDENCE_MATRIX_v1.tsv'),
        ('统一链条架构', r'D:\project26210_external_joint_v1\phase108_unified_chain_architecture_local_v1\UNIFIED_CHAIN_NODE_MATRIX_v1.tsv'),
        ('定向缺口分析', r'D:\project26210_external_joint_v1\phase109_targeted_gap_analysis_local_v1\final_synthesis_v1\PHASE109_CHAIN_EDGE_UPDATE_v1.tsv'),
        ('七特征分析总结', r'D:\project26210_external_joint_v1\phase114_unweighted_seven_feature_analysis_local_v1\PHASE114_ANALYSIS_SUMMARY_v1.md'),
        ('项目权威状态', r'D:\project26210_external_joint_v1\PROJECT_STATUS.md'),
    ]
    add_table(document, ['用途', '权威路径'], path_rows, [2200, 7200], font_size=8.2)
    add_body(document, '本简报为内部会议证据整理，不新增生物信息统计结果。所有数值均来自已经完成并通过各自验证门禁的项目模块。', italic=True, color=MUTED)

    document.core_properties.title = '项目26210现有证据总览与下一步决策简报'
    document.core_properties.subject = '主管会议内部证据与决策备忘录'
    document.core_properties.author = 'Codex'
    document.core_properties.keywords = '26210, TYMP, SPP1, macrophage, immunotherapy, evidence review, decision memo'
    document.save(DOCX)
    print(DOCX)


if __name__ == '__main__':
    build()
