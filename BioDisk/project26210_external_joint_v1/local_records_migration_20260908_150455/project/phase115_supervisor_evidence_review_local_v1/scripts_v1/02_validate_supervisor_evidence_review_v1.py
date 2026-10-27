import hashlib
import json
from pathlib import Path

import pdfplumber
from PIL import Image, ImageChops
from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r'D:\project26210_external_joint_v1\phase115_supervisor_evidence_review_local_v1')
DOCX = ROOT / 'output_v1' / '项目26210_现有证据总览与下一步决策简报_v1.docx'
PDF = ROOT / 'output_v1' / '项目26210_现有证据总览与下一步决策简报_v1.pdf'
RENDER = ROOT / 'render_final_v3'
OUT = ROOT / 'validation_v1' / 'DOCUMENT_VALIDATION_v1.json'


def sha256(path):
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest().upper()


def table_width(table):
    tblw = table._tbl.tblPr.find(qn('w:tblW'))
    return int(tblw.get(qn('w:w'))) if tblw is not None else None


def row_header_flag(row):
    trpr = row._tr.trPr
    return trpr is not None and trpr.find(qn('w:tblHeader')) is not None


def main():
    errors = []
    document = Document(str(DOCX))
    section = document.sections[0]
    geometry = {
        'page_width_cm': round(section.page_width.cm, 3),
        'page_height_cm': round(section.page_height.cm, 3),
        'top_margin_cm': round(section.top_margin.cm, 3),
        'bottom_margin_cm': round(section.bottom_margin.cm, 3),
        'left_margin_cm': round(section.left_margin.cm, 3),
        'right_margin_cm': round(section.right_margin.cm, 3),
    }
    expected = {
        'page_width_cm': 21.0, 'page_height_cm': 29.7,
        'top_margin_cm': 2.0, 'bottom_margin_cm': 2.0,
        'left_margin_cm': 2.2, 'right_margin_cm': 2.2,
    }
    for key, value in expected.items():
        if abs(geometry[key] - value) > 0.02:
            errors.append(f'geometry:{key}:{geometry[key]}')

    normal = document.styles['Normal']
    normal_east_asia = normal._element.rPr.rFonts.get(qn('w:eastAsia'))
    if normal.font.size.pt != 10.5 or normal_east_asia != 'Microsoft YaHei':
        errors.append(f'normal_style:{normal.font.size.pt}:{normal_east_asia}')

    content_tables = document.tables
    multirow_tables = [table for table in content_tables if len(table.rows) > 1]
    if len(multirow_tables) < 10:
        errors.append(f'multirow_table_count:{len(multirow_tables)}')
    for index, table in enumerate(content_tables):
        if table_width(table) != 9400:
            errors.append(f'table_width:{index}:{table_width(table)}')
        grid = table._tbl.tblGrid
        widths = [int(column.get(qn('w:w'))) for column in grid]
        if sum(widths) != 9400:
            errors.append(f'table_grid:{index}:{widths}')
        for row_index, row in enumerate(table.rows):
            cell_widths = []
            for cell in row.cells:
                tcw = cell._tc.get_or_add_tcPr().find(qn('w:tcW'))
                cell_widths.append(int(tcw.get(qn('w:w'))))
            if sum(cell_widths) != 9400:
                errors.append(f'table_cell_width:{index}:{row_index}:{cell_widths}')
        if len(table.rows) > 1 and not row_header_flag(table.rows[0]):
            errors.append(f'table_header_flag:{index}')

    drawings = document.inline_shapes
    if len(drawings) != 1:
        errors.append(f'inline_shape_count:{len(drawings)}')
    elif not drawings[0]._inline.docPr.get('descr'):
        errors.append('figure_alt_text_missing')

    full_text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    required = [
        '项目26210：现有证据总览与下一步决策简报',
        'CURRENT_DATA_DO_NOT_SUPPORT_GENERAL_SEVEN_FEATURE_TRANSFER',
        '17/20正向', '14/14正', '13/14正', 'q=0.00305',
        'rho=-0.137/0.011', 'P=0.561/0.724',
        'B作为报告骨架 + A作为最小新分析',
        'TYMP-SPP1作为待检验桥而非预设事实',
        'transcriptional proxy',
    ]
    all_text = full_text + '\n' + '\n'.join(cell.text for table in content_tables for row in table.rows for cell in row.cells)
    for term in required:
        if term not in all_text:
            errors.append(f'missing_term:{term}')
    for placeholder in ['TODO', 'TBD', '待补充', '[[', ']]']:
        if placeholder in all_text:
            errors.append(f'placeholder:{placeholder}')

    with pdfplumber.open(str(PDF)) as pdf:
        page_count = len(pdf.pages)
        page_text_lengths = [len((page.extract_text() or '').strip()) for page in pdf.pages]
    if page_count != 11:
        errors.append(f'page_count:{page_count}')
    if any(length < 100 for length in page_text_lengths):
        errors.append(f'sparse_page:{page_text_lengths}')

    pixel_audit = []
    pngs = sorted(RENDER.glob('page-*.png'))
    if len(pngs) != page_count:
        errors.append(f'png_count:{len(pngs)}')
    for png in pngs:
        image = Image.open(png).convert('RGB')
        white = Image.new('RGB', image.size, 'white')
        diff = ImageChops.difference(image, white).convert('L')
        mask = diff.point(lambda value: 255 if value > 12 else 0)
        bbox = mask.getbbox()
        if bbox is None:
            errors.append(f'blank_page:{png.name}')
            continue
        left, top, right, bottom = bbox
        margins = {
            'left_px': left, 'top_px': top,
            'right_px': image.width - right, 'bottom_px': image.height - bottom,
        }
        if min(margins.values()) < 28:
            errors.append(f'boundary_risk:{png.name}:{margins}')
        pixel_audit.append({'page': png.name, 'size': image.size, 'ink_bbox': bbox, **margins})

    result = {
        'status': 'PASS' if not errors else 'FAIL',
        'docx': str(DOCX), 'pdf': str(PDF),
        'docx_sha256': sha256(DOCX), 'pdf_sha256': sha256(PDF),
        'page_count': page_count, 'page_text_lengths': page_text_lengths,
        'geometry': geometry, 'table_count': len(content_tables),
        'multirow_table_count': len(multirow_tables),
        'pixel_audit': pixel_audit, 'errors': errors,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if not errors else 1)


if __name__ == '__main__':
    main()
