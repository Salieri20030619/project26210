from __future__ import annotations

import hashlib
import importlib.util
import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent
OUTPUT = ROOT / "PART1_EXTERNAL_SINGLE_CELL_COMPLETE_REPORT_TEXT_REVIEW_v3.docx"
STATUS = ROOT / "build_status_v3.json"
ASSETS = ROOT / "assets_v1"
FLOWCHART = ASSETS / "part1_workflow_v3.png"

helper_path = PROJECT / "phase23_part1_client_report_local_v1" / "03_build_part1_client_report_v1.py"
spec = importlib.util.spec_from_file_location("report_helper", helper_path)
h = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(h)


def font_path(bold=False):
    base = Path(r"C:\Windows\Fonts")
    candidates = ["msyhbd.ttc", "msyhbd.ttf"] if bold else ["msyh.ttc", "msyh.ttf"]
    for name in candidates:
        if (base / name).exists():
            return str(base / name)
    return str(base / "simhei.ttf")


def rounded_box(draw, xy, fill, outline, radius=24, width=4):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, p1, p2, color="#2E6F9E", width=8):
    draw.line([p1, p2], fill=color, width=width)
    angle = math.atan2(p2[1] - p1[1], p2[0] - p1[0])
    length = 22
    left = (p2[0] - length * math.cos(angle - math.pi / 6), p2[1] - length * math.sin(angle - math.pi / 6))
    right = (p2[0] - length * math.cos(angle + math.pi / 6), p2[1] - length * math.sin(angle + math.pi / 6))
    draw.polygon([p2, left, right], fill=color)


def draw_centered_multiline(draw, xy, text, font, fill, spacing=7):
    draw.multiline_text(xy, text, font=font, fill=fill, anchor="mm", align="center", spacing=spacing)


def build_flowchart(path):
    ASSETS.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1800, 1150), "white")
    draw = ImageDraw.Draw(img)
    f_title = ImageFont.truetype(font_path(True), 50)
    f_step = ImageFont.truetype(font_path(True), 30)
    f_body = ImageFont.truetype(font_path(False), 23)
    f_note = ImageFont.truetype(font_path(False), 22)
    draw_centered_multiline(draw, (900, 65), "第一部分分析流程", f_title, "#18324A")
    draw_centered_multiline(draw, (900, 125), "五个外部队列统一建库 · 患者作为统计独立单位 · 多方法交叉验证", f_note, "#64727D")

    xs = [70, 425, 780, 1135, 1490]
    box_w, box_h = 270, 265
    y1, y2 = 220, 620
    top = [
        ("01 数据整理", "5个队列\n75个文库\n40名患者", "#EAF3F8", "#2E6F9E"),
        ("02 质控注释", "275,150 → 258,120\n广义细胞类型\n肿瘤上皮识别", "#EAF3F8", "#2E6F9E"),
        ("03 恶性筛选", "inferCNV + CopyKAT\n保守恶性细胞\n11,950个", "#E9F5F3", "#2D8C87"),
        ("04 克隆与程序", "患者内CNV克隆\ncNMF秩7\nF1–F7", "#E9F5F3", "#2D8C87"),
        ("05 分子机制", "scDRS · pySCENIC\nmetaVIPER · MR\n代谢与酶功能轴", "#FBF4E4", "#B88A2A"),
    ]
    bottom = [
        ("06 发育轨迹", "Monocle3 · Slingshot\nCytoTRACE2 · scTour\nCellRank", "#FBF4E4", "#B88A2A"),
        ("07 细胞通讯", "CellChat · NicheNet\n受体可用性\n患者内汇总", "#E9F5F3", "#2D8C87"),
        ("08 临床验证", "TCGA-KIRC 510例\nOS / PFS\n多变量Cox", "#EAF3F8", "#2E6F9E"),
        ("09 证据分级", "4/4 高置信\n3/4 次级支持\n2/4 探索性", "#F3F5F7", "#64727D"),
        ("10 综合结论", "支持与反向结果\n按同一统计口径\n共同解释", "#F3F5F7", "#18324A"),
    ]
    for i, (title, body, fill, stroke) in enumerate(top):
        x = xs[i]
        rounded_box(draw, (x, y1, x + box_w, y1 + box_h), fill, stroke)
        draw_centered_multiline(draw, (x + box_w / 2, y1 + 62), title, f_step, stroke)
        draw_centered_multiline(draw, (x + box_w / 2, y1 + 165), body, f_body, "#202B33")
        if i < 4:
            arrow(draw, (x + box_w + 12, y1 + box_h / 2), (xs[i + 1] - 12, y1 + box_h / 2))
    arrow(draw, (xs[4] + box_w / 2, y1 + box_h + 18), (xs[4] + box_w / 2, y2 - 18))
    for i, (title, body, fill, stroke) in enumerate(bottom):
        x = xs[4 - i]
        rounded_box(draw, (x, y2, x + box_w, y2 + box_h), fill, stroke)
        draw_centered_multiline(draw, (x + box_w / 2, y2 + 62), title, f_step, stroke)
        draw_centered_multiline(draw, (x + box_w / 2, y2 + 165), body, f_body, "#202B33")
        if i < 4:
            arrow(draw, (x - 12, y2 + box_h / 2), (xs[3 - i] + box_w + 12, y2 + box_h / 2))
    draw_centered_multiline(
        draw,
        (900, 1010),
        "统计口径：n≥10报告效应与BH-FDR；n=3–9仅在名义P<0.05时报告显著性，其余只描述方向。",
        f_note,
        "#64727D",
    )
    img.save(path, dpi=(180, 180))


def configure_section(doc):
    section = doc.sections[0]
    h.configure_section(section)
    hp = section.header.paragraphs[0]
    hp.clear()
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    h.set_run(hp.add_run("project26210  |  第一部分外部单细胞验证"), size=8.5, color=h.MUTED)
    fp = section.footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.set_run(fp.add_run("—  "), size=9, color=h.MUTED)
    h.add_page_field(fp)
    h.set_run(fp.add_run("  —"), size=9, color=h.MUTED)


def add_title_page(doc):
    p = doc.add_paragraph()
    p.add_run("\n\n\n")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    h.set_run(p.add_run("EXTERNAL SINGLE-CELL VALIDATION"), size=10, bold=True, color=h.GOLD)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    h.set_run(p.add_run("第一部分 外部单细胞验证与机制分析报告"), size=25, bold=True, color=h.NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    h.set_run(p.add_run("TYMP相关恶性程序、调控网络、代谢轨迹与微环境通讯"), size=13.5, color=h.BLUE)
    h.add_callout(doc, "版本说明", "客户正文审阅稿：已纳入全部26项分析的目的、方法和结果；本版仅放分析流程图，数据结果图待正文确认后统一编排。", fill=h.LIGHT_BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(42)
    h.set_run(p.add_run("项目：project26210"), size=11, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.set_run(p.add_run("报告日期：2026年8月"), size=10.5, color=h.MUTED)
    doc.add_page_break()


def analysis(doc, code, title, why, method, result, conclusion=None):
    h.add_heading(doc, f"{code}  {title}", 2)
    h.add_labeled(doc, "为什么做", why)
    h.add_labeled(doc, "怎么做", method)
    h.add_labeled(doc, "得到什么", result, result=True)
    if conclusion:
        h.add_callout(doc, "本项结论", conclusion, fill=h.LIGHT_TEAL)


def add_picture(doc, path, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    h.set_image_alt(p, "第一部分分析流程", "五个外部单细胞队列从质控、恶性细胞筛选、克隆与程序、机制、轨迹、通讯到临床验证的分析流程。")
    cap = doc.add_paragraph("图1  第一部分总体分析流程。统计检验以患者为独立单位。", style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_document():
    build_flowchart(FLOWCHART)
    doc = Document()
    h.configure_styles(doc)
    configure_section(doc)
    add_title_page(doc)

    h.add_heading(doc, "执行摘要", 1)
    h.add_body(doc, "本报告整合5个外部肾癌单细胞队列，共75个文库、40名患者。统一对象含275,150个细胞，质控后保留258,120个细胞；其中38,670个肿瘤上皮细胞接受CNV分层，最终得到11,950个保守恶性细胞，覆盖全部5个队列。后续所有关联统计均先在患者内计算，再在患者层面汇总。")
    h.add_callout(doc, "核心正向结果", "cNMF的F1程序与TYMP表达总体同向；ILF2和HOXD3达到3/4次级调控证据；F1伴随缺氧/血管生成表达程序，但并不对应HIF1A蛋白活性增强；CASP9、CHMP1A和PRSS8活性与F1稳定同向；广义TYMP/嘧啶代谢轴达到3/4次级支持；内皮—恶性、巨噬—恶性和恶性—巨噬通讯，以及ICAM5、MST1配体活性获得患者层面支持。", fill=h.LIGHT_TEAL)
    h.add_callout(doc, "关键反向结果", "RCC遗传易感性与F1没有稳定关系；MR效应方向与可推断蛋白活性总体呈反向；HIF1A metaVIPER活性随F1降低；严格TYMP/RRM2B共同终末升高不稳定，广义RNR相关轴呈3/4反向；主要受体与连续TYMP/F1端点未见FDR阳性。反向结果与正向结果共同限定了机制解释。", fill=h.LIGHT_GOLD)
    h.add_body(doc, "综合来看，数据支持“TYMP相关恶性细胞程序与嘧啶代谢、特定调控因子及微环境通讯相伴”的较宽机制框架；不支持将该框架简化为HIF1A活性增强，或TYMP与RRM2B在所有轨迹末端共同升高。临床验证中，PRTG和IGDCC4相关双高组合与较差OS和PFS独立相关，提供了与患者结局相连的外部证据。")

    h.add_heading(doc, "1  数据范围与分析设计", 1)
    h.add_heading(doc, "1.1  正确数据范围", 2)
    h.add_table(
        doc,
        ["队列", "患者数", "定位"],
        [
            ["GSE156632", "7", "外部单细胞队列"],
            ["GSE159115", "7", "外部单细胞队列"],
            ["GSE178481", "5", "外部单细胞队列"],
            ["GSE207493", "12", "外部单细胞队列"],
            ["GSE242299", "9", "外部单细胞队列"],
            ["合计", "40", "75个文库"],
        ],
        [2800, 1500, 5060],
        numeric_cols=(1,),
    )
    h.add_body(doc, "本报告只使用上述五个队列。既往文档中出现的其他GSE编号不属于本次第一部分的证据范围。")
    h.add_heading(doc, "1.2  统计与证据分级", 2)
    h.add_body(doc, "患者是独立统计单位。细胞用于形成患者内效应，不将同一患者的细胞当作独立重复。患者数不少于10时，同时报告效应和BH校正后的FDR；患者数3–9时，名义双侧P值达到0.05可报告为小样本显著结果，否则只描述方向。")
    h.add_table(
        doc,
        ["跨方法通过数", "证据等级", "解释"],
        [["4/4", "高置信", "四类预设证据一致"], ["3/4", "次级支持", "保留为主要结果"], ["2/4", "探索性", "提示方向，需后续验证"], ["0–1/4", "未支持", "不作为正向机制依据"]],
        [1800, 2200, 5360],
        numeric_cols=(0,),
    )
    h.add_heading(doc, "1.3  总体流程", 2)
    add_picture(doc, FLOWCHART)

    doc.add_page_break()
    h.add_heading(doc, "2  数据基础与细胞图谱", 1)
    analysis(doc, "B01", "数据集与样本结构", "确认五个外部队列的样本单位、患者单位和可比范围，为全部患者层面统计建立共同底座。", "统一队列、文库、患者和细胞标识，核查队列来源及样本对应关系。", "共纳入5个队列、75个文库和40名患者；统一对象包含275,150个细胞。")
    analysis(doc, "B02", "质量控制与细胞注释", "去除低质量细胞并建立可用于肿瘤细胞与微环境分析的统一细胞图谱。", "完成细胞质控、scVI整合和广义细胞类型注释；注释同时保留队列与患者来源。", "质控后保留258,120个细胞。主要细胞群包括T细胞63,071、上皮细胞62,652、巨噬细胞42,546、内皮细胞30,025、NK细胞23,373以及周细胞/平滑肌细胞19,937；其余B细胞、浆细胞、单核细胞、树突细胞和肥大细胞等共同构成微环境。")

    doc.add_page_break()
    h.add_heading(doc, "3  恶性细胞、CNV克隆与cNMF程序", 1)
    analysis(doc, "M01", "肿瘤上皮细胞识别", "将肿瘤组织中的上皮细胞与其他微环境细胞分开，限定后续CNV和恶性程序分析对象。", "综合组织来源与统一细胞注释筛选肿瘤上皮细胞。", "共有38,670个肿瘤上皮细胞进入CNV分层。")
    analysis(doc, "M02", "保守恶性细胞筛选", "降低单一CNV工具造成的误判，形成更稳健的恶性细胞集合。", "联合inferCNV与CopyKAT构建CNV共识分层，并以保守规则选择恶性细胞。", "获得10,844个双方法一致恶性细胞和1,106个inferCNV单方法恶性细胞，合计11,950个保守恶性细胞；覆盖26名患者、29个文库及全部5个队列。")
    analysis(doc, "M03", "患者内CNV克隆树与程序映射", "观察同一患者内部CNV克隆差异及其与转录程序的对应关系。", "基于患者内HMM CNV轮廓构建确定性根锚定树，再将cNMF主导程序映射到克隆节点；根节点作为中性展示锚，不解释为时间祖先。", "14名患者形成43个主要克隆节点和43条根锚定边，227/227个HMM轮廓哈希复现。F1主导克隆见于5名患者、11个节点和3,682个细胞；其余F2–F7分别占据不同患者内克隆。", "CNV克隆结构与转录程序可在患者内联合描述，但树的边表示CNV距离关系，不直接代表真实演化时间。")
    analysis(doc, "M04", "cNMF恶性细胞程序发现", "从恶性细胞中提取可重复的共表达程序，避免只依赖单基因解释。", "在11,950个恶性细胞和2,500个高变基因上评估cNMF秩4–8，并以稳定性选择秩7。", "得到F1–F7共7个稳定程序。F1与TYMP的患者内Spearman相关中位数为0.115，85%的可评估患者方向为正，FDR=0.0381；留一队列敏感性分析的最低中位rho仍为0.1018。", "F1可作为TYMP相关恶性细胞程序的主要分析轴。F1主导细胞高度集中于GSE207493，因此跨患者结果与留一队列敏感性共同用于解释。")
    analysis(doc, "M05", "程序差异基因与功能富集", "为F1–F7提供可解释的基因和通路标签，并识别各程序的生物学差异。", "在14名合格患者中构建28个伪批量样本，对7个程序完成差异基因和功能富集分析。", "全部7个程序均完成注释，25/25项验证通过；这些结果用于后续调控、代谢和轨迹分析中的程序定位。")
    analysis(doc, "M06", "RCC遗传易感性与F1", "判断F1是否富集肾癌遗传易感相关表达信号。", "使用GWAS Catalog肾癌位点构建scDRS基因集，在患者内计算F1与scDRS得分相关，再进行患者层面符号检验和Bootstrap。", "20名患者的相关中位rho=-0.0490，Bootstrap 95%区间为-0.1292至0.0636，符号检验P=0.4833；8名患者为正、12名为负。", "F1与肾癌遗传易感表达分数没有稳定关联。")
    analysis(doc, "M07", "遗传易感性稳健性复核", "排查总体结果是否由单一队列或少数细胞驱动。", "比较队列内患者方向、细胞层面显著性和患者层面总体效应。", "未发现FDR<0.1或0.2的单细胞scDRS结果。GSE207493的10名患者中位rho=-0.1292，P=0.0303，但其方向未在全部队列中保持。", "单队列信号不足以改变总体不支持结论。")

    h.add_heading(doc, "4  调控网络、蛋白活性与MR方向", 1)
    analysis(doc, "M08", "双算法pySCENIC调控网络", "通过独立网络推断算法降低单一算法依赖，为调控因子证据分级提供基础。", "分别使用GENIE3和GRNBoost2重建网络并进行cisTarget/AUCell评分。", "GENIE3获得299个regulon，GRNBoost2获得314个regulon；11,950个恶性细胞全部进入评分，28/28项验证通过。")
    analysis(doc, "M09", "F1相关调控因子", "寻找随F1变化的上游调控因子，并按连续、分组和算法一致性分级。", "在两套pySCENIC矩阵中评估连续F1、F1高低组和四分位端点，使用4项证据规则综合。", "197个regulon可在两套矩阵中比较。没有4/4高置信因子；ILF2和HOXD3均达到3/4次级支持，HNF4A为2/4探索性。ILF2在n=9分组检验中P=0.00390625，HOXD3在n=9时P=0.03125。", "ILF2与HOXD3是当前最稳定的F1相关调控候选。")
    analysis(doc, "M10", "metaVIPER独立复核", "使用独立蛋白活性框架复核pySCENIC以外的调控关系。", "对5,566个调控因子和11,921个有完整输入的细胞运行metaVIPER，并按调控因子键对齐20个患者检查点。", "35/35项验证通过。20个患者检查点对齐后的最大绝对差均为0，说明修复只纠正了行列顺序，不改变任何已计算活性值。", "metaVIPER结果可以作为独立活性证据使用。")
    analysis(doc, "M11", "MR效应与蛋白活性方向一致性", "检验遗传因果推断方向是否在单细胞蛋白活性层面得到同向支持。", "在8个MR候选中，对可由metaVIPER推断的CASP9、CHMP1A、PRSS8和PXN比较MR效应与患者内活性。", "4个可推断蛋白的整体Pearson r=-0.6483。20/20名患者的相关均为负，中位r=-0.6637，FDR=3.81×10^-6。", "正向MR效应—蛋白活性一致性未得到支持，结果指向反向关系。")
    analysis(doc, "M12", "候选蛋白活性与F1", "直接评估候选蛋白活性是否随恶性程序F1变化。", "同时使用连续F1、F1高低组和Q4–Q1端点，在患者内估计方向并进行患者层面汇总。", "CASP9、CHMP1A和PRSS8在三个端点均为正。n=9分组检验中，三者分别有7/9、8/9和8/9患者为正，名义P分别为0.02734、0.01172和0.01172。PXN和HIF1A不支持正关联。", "CASP9、CHMP1A和PRSS8构成F1相关蛋白活性的主要正向信号。")
    analysis(doc, "M13", "终末共同激活与酶功能轴替代验证", "检验候选蛋白及TYMP/RRM2B是否在恶性轨迹末端共同上升；当调控网络无法覆盖TYMP/RRM2B时，以预先认可的酶功能轴完成替代验证。", "联合Monocle3、Slingshot和scTour终末窗口，另以连续F1、分组、四分位和轨迹四类证据评估严格及广义酶功能轴。", "CASP9、CHMP1A和PRSS8没有形成跨轨迹的稳定共同终末升高。严格TYMP轴与严格RRM2B轴均未支持；广义TYMP/嘧啶轴通过3/4证据，连续、四分位和分组端点FDR分别为0.000395、0.00121和0.00122；广义RNR相关轴则呈3/4反向，FDR分别为5.34×10^-5、0.000585和0.000488。", "酶功能替代验证支持TYMP相关的广义嘧啶代谢伴随关系，同时提示RRM2B/RNR方向与原先简单共同升高模型相反。")

    h.add_heading(doc, "5  代谢、酶矩阵与恶性轨迹", 1)
    analysis(doc, "M14", "代谢通路与scFEA评分", "从通路和代谢模块两个层面描绘恶性细胞代谢状态。", "对4条预设通路进行AUCell评分，并运行包含168个模块的scFEA。", "11,950个恶性细胞获得完整代谢评分，21/21项验证通过，为TYMP、RRM2B及其他代谢酶分析提供统一底座。")
    analysis(doc, "M15", "F1相关代谢程序", "确认F1是否对应一致的代谢重塑，而非单一基因波动。", "对连续F1、分组和四分位端点进行患者内代谢效应估计，并与轨迹端点共同分级。", "F1与广义TYMP/嘧啶代谢轴在三个横断面端点均稳定同向，达到3/4次级支持；轨迹端点没有形成一致的第四项证据。", "F1更适合解释为伴随广义嘧啶代谢状态，而不是固定的单一终末命运。")
    analysis(doc, "M16", "TYMP与RRM2B代谢解释", "区分两种酶的基因表达、通路状态和动态轨迹，避免把同一方向预设为必然结果。", "分别构建严格酶轴和广义功能轴，比较连续、分组、四分位和轨迹证据。", "TYMP严格轴未达标，但广义TYMP/嘧啶轴为3/4正向；RRM2B严格轴未达标，广义RNR相关轴为3/4反向。", "TYMP与RRM2B不构成稳定共同升高轴，二者应分别解释。")
    analysis(doc, "M17", "核苷酸与多胺代谢酶矩阵", "扩展到TYMP、TK1、ODC1、SRM、SMS、PCYT2等候选酶，寻找可重复的单基因和基因对关系。", "对12个候选基因进行连续F1、四分位和分组端点整合，并计算患者内基因对相关。", "ODC1达到2/3次级正向；TYMP、TK1、SRM、SMS和PCYT2为探索性。4对相关通过FDR：TYMP–SLC29A1、TYMP–SRM、TK1–SMS和SRM–PCYT2，FDR分别为0.0108、0.00352、0.01846和0.03185。", "单基因强证据有限，但TYMP相关核苷转运及多胺代谢存在可复核的协同结构。")
    analysis(doc, "M18", "五方法恶性轨迹", "以多算法判断恶性状态变化，降低单一伪时间方法的结构偏差。", "并行使用Monocle3、Slingshot、CytoTRACE2、scTour和CellRank，并统一到患者内可比较的轨迹与终末窗口。", "五种方法全部完成，13/13项验证通过。不同方法捕捉到的几何与方向信息用于M19和酶功能轴的交叉判定。")
    analysis(doc, "M19", "轨迹末端综合", "判断候选基因和蛋白是否在不同算法定义的末端共同变化。", "比较Monocle3、Slingshot和scTour的患者内终末变化，并结合其他两种轨迹方法评估稳定性。", "TYMP仅在scTour末端稳定升高，20名患者的P约为0.0238、FDR约为0.04765；RRM2B没有稳定末端变化。严格TYMP/RRM2B共同终末变化分别见于Monocle3 1/20、Slingshot 4/20和scTour 5/20名患者。", "轨迹结果支持TYMP的局部动态信号，但不支持跨方法统一的TYMP/RRM2B共同末端升高。")

    doc.add_page_break()
    h.add_heading(doc, "6  肿瘤微环境通讯与临床验证", 1)
    analysis(doc, "C01", "CellChat通讯方向", "识别与TYMP/F1恶性状态相伴的细胞间通讯方向。", "在14名合格患者内分别运行CellChat，再对内皮、巨噬与恶性细胞之间的通信差异进行患者层面汇总。", "内皮→恶性在13/14名患者为正，中位差0.1559、FDR=0.00244；巨噬→恶性在12/14名患者为正，中位差0.07994、FDR=0.00342；恶性→巨噬在11/14名患者为正，中位差0.2713、FDR=0.01139。恶性→内皮为8/14正，FDR=0.07849，未达支持标准。", "F1/TYMP相关恶性状态主要伴随内皮和巨噬细胞输入增强，以及恶性细胞对巨噬细胞的反馈增强。")
    analysis(doc, "C02", "受体可用性", "确认通讯候选是否在接收细胞中具备表达基础。", "按患者检查ITGB2、ITGAL、ITGAM、MST1R等受体或受体复合体相关特征的检测率和细胞分布。", "ITGB2/ITGAL相关特征在巨噬细胞中广泛可用；MST1R整体较稀疏。", "后续优先解释整合素相关轴，MST1R作为覆盖较窄的候选。")
    analysis(doc, "C03", "NicheNet配体活性", "从配体—靶基因角度验证哪些上游信号最可能解释恶性细胞响应。", "以患者内恶性细胞响应基因为目标，计算NicheNet配体活性并进行患者层面汇总。", "ICAM5在14/14名患者中为正，中位效应0.06761、FDR=0.00219；MST1在13/14名患者中为正，中位效应0.04025、FDR=0.00317。HIF1A加权响应多数为负，但未达支持标准。", "ICAM5和MST1是最稳定的上游配体候选。")
    analysis(doc, "C04", "受体与TYMP/F1状态关联", "把通讯候选连接到恶性细胞TYMP/F1状态，判断是否存在患者内剂量关系。", "以连续TYMP/F1为主要端点，并进行F1高低组小样本敏感性分析。", "主要连续端点没有受体关联通过FDR。n=9的敏感性分析中，ITGAM/ITGB2相关特征与TYMP在8/9名患者中同向，名义P=0.01172。", "整合素相关受体提供小样本敏感性信号，但连续主分析尚未形成稳定阳性。")
    analysis(doc, "C05", "TCGA-KIRC联合生存验证", "检验单细胞通讯候选与TYMP组合后是否关联患者结局。", "在510例TCGA-KIRC中使用预设中位数切点形成三组，分析OS和PFS；Cox模型调整年龄、性别和分期，不使用最优切点。", "12个全局三组检验中11个通过FDR。双高对双低的多变量结果显示：ICAM5–ITGAL/ITGB2组合与OS较差相关，HR=1.510，95%CI 1.019–2.236，P=0.03999；CD209组合OS HR=1.830，P=0.00749；PRTG组合OS HR=1.936、P=0.00976，PFS HR=2.200、P=0.00356；IGDCC4组合OS HR=3.190、P=3.20×10^-5，PFS HR=2.611、P=0.000660。MST1R调整后不显著。", "PRTG和IGDCC4相关组合同时获得OS与PFS支持，是临床关联最完整的候选。")

    doc.add_page_break()
    h.add_heading(doc, "7  综合解释", 1)
    h.add_heading(doc, "7.1  得到支持的主线", 2)
    h.add_body(doc, "第一，cNMF F1与TYMP在患者内总体同向，且留一队列分析保持正向。第二，ILF2和HOXD3为3/4次级调控候选，CASP9、CHMP1A和PRSS8蛋白活性与F1稳定同向。第三，广义TYMP/嘧啶代谢轴达到3/4次级支持，ODC1及四对代谢基因相关补充了核苷酸—多胺代谢结构。第四，内皮和巨噬细胞与恶性细胞之间的通讯增强，ICAM5和MST1配体活性稳定。第五，PRTG和IGDCC4相关双高组合在TCGA-KIRC中同时关联较差OS和PFS。")
    h.add_heading(doc, "7.2  反向与不支持结果", 2)
    h.add_body(doc, "F1不对应稳定的RCC遗传易感性富集。HIF1A靶向重算仍未在pySCENIC中恢复合格regulon；独立表达程序随F1升高，而metaVIPER HIF1A活性随F1降低，说明缺氧/血管生成表达伴随关系与HIF1A活性方向并不等价。MR效应与可推断蛋白活性呈患者一致的反向关系。严格TYMP/RRM2B共同末端升高没有跨轨迹复现，广义RNR相关轴也呈反向。")
    h.add_heading(doc, "7.3  最终结论", 2)
    h.add_callout(doc, "综合结论", "第一部分26项分析已经全部完成。现有证据支持一个以TYMP相关F1恶性程序为中心、连接广义嘧啶代谢、特定调控/蛋白活性和内皮—巨噬微环境通讯的模型；同时明确排除了HIF1A活性普遍增强、MR方向普遍同向以及TYMP/RRM2B稳定共同终末升高这三种过度简化解释。反向结果是机制边界的一部分，与正向结果具有同等解释价值。", fill=h.LIGHT_TEAL)

    h.add_heading(doc, "附录  分析完成清单", 1)
    rows = [
        ["B01–B02", "数据结构、质控与注释", "完成并验证"],
        ["M01–M05", "恶性筛选、CNV克隆、cNMF与功能注释", "完成并验证"],
        ["M06–M07", "scDRS遗传易感性", "完成并验证"],
        ["M08–M10", "pySCENIC与metaVIPER", "完成并验证"],
        ["M11–M13", "MR、蛋白活性、终末状态与酶功能替代", "完成并验证"],
        ["M14–M19", "代谢、酶矩阵与五方法轨迹", "完成并验证"],
        ["C01–C04", "CellChat、受体与NicheNet", "完成并验证"],
        ["C05", "TCGA-KIRC联合生存", "完成并验证"],
    ]
    h.add_table(doc, ["模块", "内容", "状态"], rows, [1700, 5260, 2400])
    h.add_body(doc, "本稿为正文审阅版，未插入数据结果图。正文确认后，可在不改变统计口径和结论的前提下加入CNV克隆树、cNMF、调控活性、代谢轨迹、通讯及生存图。")

    props = doc.core_properties
    props.title = "第一部分 外部单细胞验证与机制分析报告"
    props.subject = "project26210 客户正文审阅稿"
    props.author = "project26210 analysis team"
    props.keywords = "single-cell, TYMP, cNMF, CNV, pySCENIC, metaVIPER, CellChat, NicheNet, TCGA-KIRC"

    doc.save(OUTPUT)
    sha = hashlib.sha256(OUTPUT.read_bytes()).hexdigest()
    status = {
        "status": "PASS",
        "design_preset": "narrative_proposal with A4 6.5-inch content-width override",
        "output": str(OUTPUT),
        "size_bytes": OUTPUT.stat().st_size,
        "sha256": sha,
        "correct_cohorts": ["GSE156632", "GSE159115", "GSE178481", "GSE207493", "GSE242299"],
        "analysis_count": 26,
        "result_figures": 0,
        "process_flowcharts": 1,
    }
    STATUS.write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(status, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build_document()
