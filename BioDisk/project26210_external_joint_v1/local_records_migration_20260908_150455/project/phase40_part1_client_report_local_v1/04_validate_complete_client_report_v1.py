from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "PART1_EXTERNAL_SINGLE_CELL_COMPLETE_REPORT_TEXT_REVIEW_v3.docx"
RENDER = ROOT / "rendered_v3"
STATUS = ROOT / "validation_status_v1.json"

correct_cohorts = ["GSE156632", "GSE159115", "GSE178481", "GSE207493", "GSE242299"]
excluded_cohorts = ["GSE202374", "GSE264586", "GSE294109"]
analysis_ids = ["B01", "B02"] + [f"M{i:02d}" for i in range(1, 20)] + [f"C{i:02d}" for i in range(1, 6)]

doc = Document(DOCX)
text = "\n".join(p.text for p in doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        text += "\n" + "\t".join(cell.text for cell in row.cells)

checks = []


def add(name, passed, detail):
    checks.append({"check": name, "pass": bool(passed), "detail": detail})


for cohort in correct_cohorts:
    add(f"correct_cohort_{cohort}", cohort in text, "present" if cohort in text else "missing")
for cohort in excluded_cohorts:
    add(f"excluded_cohort_{cohort}", cohort not in text, "absent" if cohort not in text else "unexpectedly present")
for analysis_id in analysis_ids:
    add(f"analysis_{analysis_id}", analysis_id in text, "present" if analysis_id in text else "missing")

add("analysis_count", len(analysis_ids) == 26, str(len(analysis_ids)))
add("workflow_caption", "图1  第一部分总体分析流程" in text, "process flowchart caption")
add("statistical_policy", "患者是独立统计单位" in text and "患者数3–9" in text, "patient-level and small-n policy")
add("enzyme_axis", "广义TYMP/嘧啶轴" in text and "广义RNR相关轴" in text, "positive and reverse enzyme axes")
add("result_figure_policy", "未插入数据结果图" in text, "text-review edition")
add("render_status", (RENDER / "render_status.json").exists(), "render status exists")
pngs = sorted(RENDER.glob("page-*.png"))
add("rendered_pages", len(pngs) == 15, f"{len(pngs)} PNG pages")
add("pdf_exists", (RENDER / f"{DOCX.stem}.pdf").exists(), "rendered PDF")

with zipfile.ZipFile(DOCX) as zf:
    media = [n for n in zf.namelist() if n.startswith("word/media/")]
add("single_process_figure", len(media) == 1, f"{len(media)} embedded media files")

status = "PASS" if all(x["pass"] for x in checks) else "FAIL"
result = {
    "status": status,
    "docx": str(DOCX),
    "sha256": hashlib.sha256(DOCX.read_bytes()).hexdigest(),
    "analysis_count": len(analysis_ids),
    "rendered_pages": len(pngs),
    "checks_total": len(checks),
    "checks_passed": sum(x["pass"] for x in checks),
    "checks": checks,
}
STATUS.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
raise SystemExit(0 if status == "PASS" else 1)
