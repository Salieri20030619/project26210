from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document


PHASE = Path(__file__).resolve().parent
DOCX = PHASE / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_WITH_FIGURES_REVIEW_v9.docx"
RENDER = PHASE / "rendered_v3"
PDF = RENDER / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_WITH_FIGURES_REVIEW_v9.pdf"

doc = Document(DOCX)
text = "\n".join(p.text for p in doc.paragraphs)
caption_texts = [p.text.strip() for p in doc.paragraphs if p.style.name == "Caption"]

checks = []
def check(name, passed, detail):
    checks.append({"check": name, "pass": bool(passed), "detail": str(detail)})

check("docx_exists", DOCX.exists(), DOCX)
check("inline_shapes_22", len(doc.inline_shapes) == 22, len(doc.inline_shapes))
check("caption_count_22", len(caption_texts) == 22, len(caption_texts))
for i in range(1, 16):
    check(f"main_caption_{i}", f"结果图 {i} /" in text, f"结果图 {i}")
for i in range(1, 7):
    check(f"supplement_caption_{i}", f"补充结果图 S{i}" in text, f"补充结果图 S{i}")
check("version_statement_updated", "客户图文审阅稿" in text, "客户图文审阅稿")
check("stale_no_figure_statement_absent", "未插入数据结果图" not in text, "forbidden phrase absent")
check("module_mapping_statement", "B、M和C编号与正文分析模块一一对应" in text, "mapping statement")

pngs = sorted(RENDER.glob("page-*.png"))
check("rendered_pages_27", len(pngs) == 27, len(pngs))
check("pdf_exists", PDF.exists() and PDF.stat().st_size > 1000, PDF.stat().st_size if PDF.exists() else 0)
if PDF.exists():
    raw = PDF.read_bytes()
    check("pdf_header", raw.startswith(b"%PDF-"), raw[:8].decode("latin1", errors="replace"))
    check("pdf_eof", b"%%EOF" in raw[-2048:], "EOF in tail")

a11y = json.loads((PHASE / "a11y_audit_v1.json").read_text(encoding="utf-8"))
summary = a11y.get("summary", {})
check("a11y_zero_findings", sum(summary.get(k, 0) for k in ("high", "medium", "low")) == 0, summary)

failures = [c for c in checks if not c["pass"]]
status = {
    "status": "PASS_PHASE45_PART1_CLIENT_REPORT_WITH_FIGURES" if not failures else "FAIL",
    "checks_total": len(checks),
    "checks_passed": len(checks) - len(failures),
    "checks_failed": len(failures),
    "document_pages": len(pngs),
    "inline_shapes": len(doc.inline_shapes),
    "main_figures": 15,
    "supplementary_figures": 6,
    "docx_sha256": hashlib.sha256(DOCX.read_bytes()).hexdigest(),
    "pdf_sha256": hashlib.sha256(PDF.read_bytes()).hexdigest() if PDF.exists() else None,
    "checks": checks,
}
(PHASE / "validation_status_v1.json").write_text(json.dumps(status, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({k: v for k, v in status.items() if k != "checks"}, ensure_ascii=False, indent=2))
if failures:
    raise SystemExit(1)
