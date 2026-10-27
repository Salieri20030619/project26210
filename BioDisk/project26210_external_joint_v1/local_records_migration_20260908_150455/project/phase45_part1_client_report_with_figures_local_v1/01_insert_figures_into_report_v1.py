from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
PHASE = Path(__file__).resolve().parent
SOURCE_DOCX = ROOT / "phase42_part1_client_report_process_detail_local_v1" / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_PROCESS_DETAIL_REVIEW_v7.docx"
FIG_DIR = ROOT / "phase44_part1_client_figures_final_local_v1" / "figures_complete_v2"
OUTPUT_DOCX = PHASE / "PART1_EXTERNAL_SINGLE_CELL_CLIENT_REPORT_WITH_FIGURES_REVIEW_v9.docx"


FIGURES = [
    ("B02  质量控制与细胞注释", "Fig01_cell_landscape_composite_v1.png",
     "结果图 1 / Fig01（B01–B02）  五个外部队列的细胞组成、广义细胞类型 UMAP 与标志基因表达概览。"),
    ("M02  保守恶性细胞筛选", "Fig02_malignancy_CNV_composite_v1.png",
     "结果图 2 / Fig02（M01–M02）  代表病例 inferCNV 染色体图谱、CNV burden 分布以及 inferCNV 与 CopyKAT 判定组合。"),
    ("M03  患者内CNV克隆树与程序映射", "Fig03_CNV_clone_trees_by_cNMF_program_v1.png",
     "结果图 3 / Fig03（M03）  四个代表患者的 CNV 克隆结构，节点与分支按主导 cNMF 程序着色。"),
    ("M05  程序差异基因与功能富集", "Fig04_cNMF_rank7_TYMP_program_v1.png",
     "结果图 4 / Fig04（M04–M05）  cNMF rank 稳定性、TYMP 相关 F1 程序、功能富集及患者层证据。"),
    ("M07  遗传易感性稳健性复核", "Fig05_scDRS_Factor_TYMP_v1.png",
     "结果图 5 / Fig05（M06–M07）  RCC 遗传易感性 scDRS 评分在恶性细胞中的分布及其与 F1/TYMP 的患者内关联。"),
    ("M09  F1相关调控因子", "Fig06_pySCENIC_F1_driver_consensus_v1.png",
     "结果图 6 / Fig06（M08–M09）  GENIE3 与 GRNBoost2 两套 pySCENIC 结果对 F1 相关调控因子的共识证据。"),
    ("M09  F1相关调控因子", "Fig10_pySCENIC_selected_regulon_networks_v1.png",
     "结果图 7 / Fig10（M08–M09）  ILF2、HOXD3 与 HNF4A 的双算法 regulon 圆形网络；绿色节点为两种算法共同识别的高权重靶基因。"),
    ("M12  候选蛋白活性与F1", "Fig07_metaVIPER_MR_and_F1_activity_v1.png",
     "结果图 8 / Fig07（M10–M12）  可推断 master regulator 的 MR 效应、metaVIPER 活性及其与 F1 的患者层关联。"),
    ("M12  候选蛋白活性与F1", "Fig11_metaVIPER_candidate_activity_umaps_v1.png",
     "结果图 9 / Fig11（M12）  CASP9、CHMP1A、PRSS8 与 PXN 蛋白活性在恶性细胞 UMAP 上的连续分布。"),
    ("M17  核苷酸与多胺代谢酶矩阵", "Fig08_enzyme_transporter_synergy_matrix_v1.png",
     "结果图 10 / Fig08（M14–M17）  代谢酶与转运体的患者内协同相关矩阵，以及连续和分组端点的患者层方向。"),
    ("M17  核苷酸与多胺代谢酶矩阵", "Fig12_metabolic_pathway_activity_and_contrasts_v1.png",
     "结果图 11 / Fig12（M14–M16）  嘧啶、半胱氨酸/蛋氨酸、精氨酸/脯氨酸和甘油磷脂通路活性 UMAP、相关及配对差异。"),
    ("M19  轨迹末端综合", "Fig09_malignant_trajectory_five_methods_v1.png",
     "结果图 12 / Fig09（M18–M19）  Monocle3、Slingshot、CytoTRACE2、scTour 与 CellRank 五种方法的恶性细胞轨迹结果。"),
    ("C01  CellChat通讯方向", "Fig13_CellChat_communication_landscape_v1.png",
     "结果图 13 / Fig13（C01）  TYMP 分层后的 CellChat 通讯网络、收发角色、患者层端点差异与配体–受体对排序。"),
    ("C04  受体与TYMP/F1状态关联", "Fig14_NicheNet_ligand_receptor_evidence_v1.png",
     "结果图 14 / Fig14（C02–C04）  MST1/ICAM5 的 NicheNet 配体活性、配体–靶标权重、受体 UMAP 与患者内受体关联。"),
    ("C05  TCGA-KIRC联合生存验证", "Fig15_TCGA_TYMP_PRTG_IGDCC4_survival_v1.png",
     "结果图 15 / Fig15（C05）  TCGA-KIRC 中 TYMP 与 PRTG 或 IGDCC4 联合分组的 OS/PFS Kaplan–Meier 曲线。"),
]

SUPPLEMENT = [
    ("Fig01B_celltype_counts_horizontal_v1.png", "补充结果图 S1（B01）  广义细胞类型数量横向柱状图。"),
    ("Fig01C_celltype_counts_vertical_v1.png", "补充结果图 S2（B01）  广义细胞类型数量纵向柱状图。"),
    ("Fig01D_canonical_marker_dotplot_v1.png", "补充结果图 S3（B02）  完整 canonical marker DotPlot。"),
    ("Fig01D_one_marker_per_celltype_dotplot_v2.png", "补充结果图 S4（B02）  每类一个代表性 marker 的精简 DotPlot。"),
    ("Fig02B_CNV_burden_violin_v1.png", "补充结果图 S5（M01）  CNV 支持的恶性与非恶性细胞 CNV burden 分布。"),
    ("FigS01_TCGA_TYMP_receptor_survival_all_axes_v1.png", "补充结果图 S6（C05）  六条预设受体轴的 TYMP 联合分组 OS/PFS 完整敏感性结果。"),
]


def set_run_font(run, name="Arial", east_asia="Microsoft YaHei", size=Pt(9), color=RGBColor(78, 83, 88)):
    run.font.name = name
    run.font.size = size
    run.font.color.rgb = color
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def last_body_paragraph_after_heading(doc: Document, heading_text: str):
    paragraphs = doc.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == heading_text)
    end = len(paragraphs)
    for j in range(start + 1, len(paragraphs)):
        if paragraphs[j].style.name.startswith("Heading"):
            end = j
            break
    for j in range(end - 1, start, -1):
        if paragraphs[j].text.strip():
            return paragraphs[j]
    return paragraphs[start]


def add_figure_after(doc: Document, anchor, image_path: Path, caption_text: str):
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.keep_together = True
    run = caption.add_run(caption_text)
    set_run_font(run, size=Pt(8.5), color=RGBColor(55, 68, 82))

    image_p = doc.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.space_before = Pt(0)
    image_p.paragraph_format.space_after = Pt(8)
    image_p.paragraph_format.keep_together = True

    with Image.open(image_path) as im:
        aspect = im.width / im.height
    max_width = 6.45
    max_height = 7.35
    width = min(max_width, max_height * aspect)
    shape = image_p.add_run().add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("descr", caption_text)
    shape._inline.docPr.set("title", image_path.stem)

    anchor._p.addnext(caption._p)
    caption._p.addnext(image_p._p)
    return image_p


def main():
    PHASE.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE_DOCX)

    old = "版本说明  客户正文审阅稿：已纳入全部26项分析，正文补充了样本进入、算法处理、患者层面汇总和判定规则；本版仅放分析流程图，数据结果图待正文确认后统一编排。"
    new = "版本说明  客户图文审阅稿：已纳入全部26项分析，并按对应章节插入第一部分完整结果图；正文保留样本进入、算法处理、患者层面汇总和判定规则。"
    for p in doc.paragraphs:
        if p.text.strip() == old:
            p.text = new
            for run in p.runs:
                set_run_font(run, size=Pt(9), color=RGBColor(90, 96, 103))
            break

    appendix_old = "本稿为正文审阅版，未插入数据结果图。正文确认后，可在不改变统计口径和结论的前提下加入CNV克隆树、cNMF、调控活性、代谢轨迹、通讯及生存图。"
    appendix_new = "本稿已按对应分析章节插入15组正文结果图，并在下一附录提供6组备选版式与完整敏感性图。图注中的B、M和C编号与正文分析模块一一对应。"
    for p in doc.paragraphs:
        if p.text.strip() == appendix_old:
            p.text = appendix_new
            for run in p.runs:
                set_run_font(run, size=Pt(9.5), color=RGBColor(65, 72, 80))
            break

    grouped = {}
    for heading, filename, caption in FIGURES:
        grouped.setdefault(heading, []).append((filename, caption))

    for heading, items in grouped.items():
        anchor = last_body_paragraph_after_heading(doc, heading)
        for filename, caption in items:
            image_path = FIG_DIR / filename
            if not image_path.exists():
                raise FileNotFoundError(image_path)
            anchor = add_figure_after(doc, anchor, image_path, caption)

    doc.add_page_break()
    heading = doc.add_paragraph("附录  补充图件与敏感性结果", style="Heading 1")
    intro = doc.add_paragraph("以下图件提供正文复合图的备选版式、完整标志基因视图及联合生存敏感性结果，编号与第一部分分析代码对应。")
    intro.paragraph_format.space_after = Pt(8)
    anchor = intro
    for filename, caption in SUPPLEMENT:
        image_path = FIG_DIR / filename
        if not image_path.exists():
            raise FileNotFoundError(image_path)
        anchor = add_figure_after(doc, anchor, image_path, caption)

    doc.core_properties.title = "第一部分 外部单细胞验证与机制分析报告（图文审阅稿）"
    doc.core_properties.subject = "project26210 Part 1 client report with validated figures"
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)
    print(f"inline_shapes={len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
