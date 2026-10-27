#!/usr/bin/env python3
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import build_client_brief_v1 as b


ROOT = Path(r"C:\Users\31937\_codex_phase100_client_mta_mtap_brief_v1")
OUT = ROOT / "PROJECT26210_TYMP_MTA_IMMUNE_PUBLIC_VALIDATION_BRIEF_v3.docx"


def add_matrix(doc, headers, rows, widths, decision_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    b.mark_header_row(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        b.set_cell_shading(cell, b.LIGHT_GRAY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        b.set_run_font(r, size=9.1, bold=True, color=b.DEEP_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            r = p.add_run(str(value))
            color = b.DARK
            bold = False
            if decision_col is not None and i == decision_col:
                bold = True
                color = b.GREEN if str(value).startswith("支持") else b.RED
            b.set_run_font(r, size=8.7, bold=bold, color=color)
    b.set_table_geometry(table, widths, indent_dxa=120)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, b.BLUE, 16, 8),
        ("Heading 2", 13, b.BLUE, 12, 6),
        ("Heading 3", 12, b.DEEP_BLUE, 8, 4),
    ):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    hp = section.header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("项目26210  |  MTA与免疫机制公共数据验证")
    b.set_run_font(hr, size=8.5, bold=True, color=b.GRAY)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run("客户阅览稿  |  第 ")
    b.set_run_font(r, size=8.5, color=b.GRAY)
    b.add_field(fp, "PAGE")
    r = fp.add_run(" 页 / 共 ")
    b.set_run_font(r, size=8.5, color=b.GRAY)
    b.add_field(fp, "NUMPAGES")
    r = fp.add_run(" 页")
    b.set_run_font(r, size=8.5, color=b.GRAY)
    return doc


def build():
    doc = setup_document()

    b.add_para(doc, "PUBLIC AND IMMUNE DATA VALIDATION BRIEF", size=9, bold=True, color=b.BLUE, after=3)
    b.add_para(doc, "TYMP-MTA-SPP1机制链证据简报", size=24, bold=True, color=b.DEEP_BLUE, after=4, line=1.0)
    b.add_para(doc, "五队列转录代理、MTA相关免疫结果、四队列实测MTA及MTAP串联验证", size=12.5, color=b.GRAY, after=12)
    for label, value in (
        ("单细胞基础", "5个免疫/肿瘤队列，585,119个质控细胞，63名患者"),
        ("实测代谢验证", "4个配对RNA-代谢组队列，341个配对样本；255个肿瘤样本用于主分析"),
        ("质量控制", "实测MTA分析28/28检查通过；MTAP串联分析29/29检查通过"),
        ("日期", "2026年8月30日"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(label + "：")
        b.set_run_font(r1, size=10, bold=True, color=b.DEEP_BLUE)
        r2 = p.add_run(value)
        b.set_run_font(r2, size=10)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(10)
    b.paragraph_bottom_border(rule, color=b.BLUE, size="12", space="4")
    b.add_callout(
        doc,
        "综合结论",
        "免疫转录组中存在与TYMP和MTA反应代理相容的髓系代谢/免疫抑制背景，但最终CNV门控五队列分析未形成FDR支持；更关键的是，配对实测MTA数据不支持TYMP-MTA，插入MTAP后也不能闭合串联链。当前证据支持“广泛髓系背景”，不支持“TYMP经MTA选择性驱动SPP1巨噬细胞”的完整机制。",
    )

    doc.add_heading("1. 证据层级与阅读方式", level=1)
    b.add_para(doc, "本简报将不同数据层分开解释。早期五队列和免疫分析使用的是表达推导的polyamine/MTA context或文献来源MTA-response代理，反映转录反应相容性；UnitedMet分析使用的是同一样本中的实测MTA。代理分数可以支持生物学背景，但不能替代代谢物测量，也不能单独证明因果。")
    add_matrix(
        doc,
        ["证据层", "数据", "回答的问题", "主要限制"],
        [
            ("五队列转录代理", "5个单细胞队列", "TYMP是否与MTA相关表达背景同向", "仅2/5队列满足最终患者数门槛；不是实测MTA"),
            ("MTA相关免疫分析", "3个髓系队列及空间/扰动数据", "MTA反应代理是否连接SPP1/免疫抑制", "转录代理、跨细胞类型转移与计算扰动"),
            ("实测MTA", "4个配对RNA-代谢组队列", "TYMP/MTAP是否关联真实MTA", "观察性bulk数据"),
        ],
        [1800, 1900, 2940, 2720],
    )

    doc.add_heading("2. 五队列中的MTA转录代理结果", level=1)
    b.add_para(doc, "五个队列均进入统一单细胞流程，但最终CNV支持的恶性细胞门控后，TYMP与polyamine/MTA context的患者级检验只有GSE207493和GSE294109达到至少5名患者的预设门槛。其余三个队列不是阴性统计结果，而是样本量不足、不可检验。")
    add_matrix(
        doc,
        ["队列", "最终可检验患者", "TYMP-MTA代理结果", "证据判断"],
        [
            ("GSE178481", "n=2", "未达到n>=5门槛", "不可检验"),
            ("GSE202374", "n=3", "未达到n>=5门槛", "不可检验"),
            ("GSE207493", "n=11", "rho=0.318；P=0.340；BH q=0.510", "正向但未支持"),
            ("GSE264586", "n=2", "未达到n>=5门槛", "不可检验"),
            ("GSE294109", "n=6", "rho=0.371；P=0.468；BH q=0.951", "正向但未支持"),
            ("跨研究描述", "2个可检验队列；总n=17", "median rho=0.345；2/2正向；0个FDR通过", "描述性相容"),
        ],
        [1840, 1840, 3360, 2320],
        decision_col=3,
    )
    b.add_callout(doc, "五队列判断", "方向上存在TYMP与MTA转录背景同向的线索，但最终分析只有两个队列可检验，且均未通过名义P值或FDR门槛。因此只能作为代理层的描述性相容证据。", fill=b.PALE_BLUE)

    doc.add_heading("3. 免疫数据中与MTA相关的结果", level=1)
    doc.add_heading("3.1 髓系MTA-response代理桥接", level=2)
    b.add_para(doc, "在67,825个Myeloid/DC细胞和44个患者/时间点单位中，GSE207493用于发现与背景解释，GSE264586和GSE294109作为两个外部决策队列。分析在患者/时间点内进行，并校正S/G2M；细胞没有被当作独立重复。")
    add_matrix(
        doc,
        ["预设关系", "GSE264586", "GSE294109", "外部判断"],
        [
            ("TYMP -> MTA/A2B/STAT3代理", "rho=0.203；q=0.00147", "rho=0.117；q=0.00147", "支持于两个外部队列"),
            ("MTA代理 -> SPP1状态程序", "rho=0.096；q=0.0242", "rho=0.120；q=0.0105", "支持于两个外部队列"),
            ("TYMP -> SPP1状态程序", "rho=0.0175；q=0.0923", "rho=0.0441；q=0.0105", "仅一个外部队列"),
            ("MTA代理 -> 免疫抑制程序", "rho=0.0734；q=0.00488", "rho=0.0898；q=0.00928", "支持于两个外部队列"),
            ("MTAP -> MTA反应代理", "rho=-0.045；q=0.444", "rho=-0.032；q=0.178", "未支持"),
        ],
        [2500, 2260, 2260, 2340],
        decision_col=3,
    )
    b.add_callout(doc, "桥接判断", "预设三条主链边中只有2/3达到双外部队列规则，直接TYMP-SPP1边未完成复现；因此结论是“转录桥未完全支持”。该结果使用文献来源反应代理，不是实测5-MTA。", fill=b.PALE_BLUE)

    doc.add_heading("3.2 为什么从MTA引出SPP1？", level=2)
    b.add_para(doc, "这个连接有生物学依据，但目前属于跨研究拼接出的机制假说，不是已经被一项实验直接证明的MTA-SPP1通路。")
    b.add_bullet(doc, "MTA侧的直接实验：MTA可通过腺苷A2B受体重塑巨噬细胞，并依赖STAT3/C/EBP诱导ARG1、VEGFA、IL10和TIMP1等替代激活/血管生成程序；原研究没有检测SPP1。", "MTA侧的直接实验：")
    b.add_bullet(doc, "SPP1侧的直接实验：SPP1在缺氧巨噬细胞中升高；其他模型显示STAT3可结合并调控Spp1，IL10-STAT3与ERK等修复信号可促进产骨桥蛋白巨噬细胞形成。", "SPP1侧的直接实验：")
    b.add_bullet(doc, "本项目的连接证据：实测MTA与bulk SPP1呈名义正相关，同时MTA-response代理与SPP1状态程序在两个外部免疫队列中通过FDR，说明两端具有相容性。", "本项目的连接证据：")
    b.add_callout(doc, "准确写法", "MTA可能通过A2B-STAT3/C/EBP及缺氧/修复型巨噬细胞程序，促成与SPP1状态相容的免疫生态；但尚无直接证据证明MTA本身诱导SPP1。需要在MTA刺激后检测SPP1 RNA/蛋白，并用A2B或STAT3阻断进行验证。", fill="FFF4E5", label_color="8A5A00")

    doc.add_heading("3.3 TYMP相关免疫代谢背景", level=2)
    add_matrix(
        doc,
        ["免疫层", "外部复现结果", "证据等级", "与MTA机制的关系"],
        [
            ("TYMP-high Myeloid/DC", "9个氧化/缺氧、红氧与糖酵解代谢/应激签名在GSE264586和GSE294109复现", "A：双外部FDR核心", "支持广泛髓系免疫代谢背景；不是配对代谢组"),
            ("TYMP-high Treg", "6个缺氧、呼吸、红氧与糖酵解签名在GSE294109复现；GSE264586无可评估Treg注释", "B：单外部FDR核心", "支持情境性代谢适应；不是MTA特异信号"),
            ("细胞来源", "TYMP在髓系/DC总体最高；Treg相对其他T细胞富集，细胞周期校正后保留", "A/B背景证据", "说明TYMP免疫来源与状态，不证明MTA介导"),
        ],
        [1960, 3560, 1800, 2040],
    )

    doc.add_heading("3.4 空间与扰动压力测试", level=2)
    add_matrix(
        doc,
        ["分析", "主要结果", "证据判断"],
        [
            ("计算虚拟敲除", "TYMP、MTAP和STAT3针对MTA反应/生产程序的主要假设均未获得外部支持；可重复结果偏向单核/C1QC等广泛髓系网络", "未闭合MTA-SPP1链"),
            ("MTA-response空间转移", "24个切片、72,952个spots；0/6关系在FFPE和frozen双制备支持，0/4链乘积双制备支持", "未支持双制备空间链"),
            ("直接MTA扰动签名", "人源直接MTA暴露得到89基因签名；26基因跨情境共识因上下调方向不平衡而未通过预设资格门槛", "仅探索性"),
            ("MTAP自然扰动", "TCGA-KIRC MTAP缺失分子门通过，但TYMP交互和三个外部单细胞队列未形成机制支持", "未支持"),
        ],
        [2260, 5040, 2060],
        decision_col=2,
    )

    doc.add_heading("4. 四个配对队列的实测MTA验证", level=1)
    b.add_para(doc, "UnitedMet提供RC18、RC20、CPTAC和CPTAC_val同一样本的RNA与实测MTA。全部8个文件通过大小与MD5校验，341个配对样本映射成功；预设主分析使用255个肿瘤样本。")
    add_matrix(
        doc,
        ["检验节点", "合并效应", "统计结果", "证据判断"],
        [
            ("TYMP-MTA", "rho=0.0395", "95% CI -0.086至0.164；P=0.538；方向负、负、正、正", "未支持"),
            ("TYMP-MTA（校正MTAP）", "rho=0.0448", "P=0.485；方向仍不一致", "未支持"),
            ("MTAP-MTA", "rho=-0.2713", "95% CI -0.383至-0.151；P=1.44E-5；随机效应P=0.00957", "支持"),
            ("MTA-SPP1", "rho=0.1389", "名义P=0.0293；BH q=0.0975；校正TYMP/MTAP后消失", "弱且不稳健"),
            ("TYMP-SPP1", "rho=0.2173", "P=5.79E-4；描述性相关", "不能证明经MTA"),
            ("TYMP-MTA-SPP1中介", "8项bootstrap", "全部95% CI跨0", "未支持"),
        ],
        [2440, 1840, 3060, 2020],
        decision_col=3,
    )
    b.add_callout(doc, "实测MTA判断", "数据能稳定检出预期的MTAP-MTA负相关，说明代谢测量具有信息量；但TYMP-MTA接近零且方向冲突，所有MTA中介效应均跨0。", fill=b.PALE_BLUE)

    doc.add_heading("5. 插入MTAP后的串联验证", level=1)
    b.add_para(doc, "扩展链为TYMP -> MTAP -> 实测MTA -> SPP1 / SPP1相关程序。MTAP是合理的MTA清除因子，但只有在TYMP-MTAP关系和串联间接效应均成立时，才能作为桥梁。")
    add_matrix(
        doc,
        ["检验节点", "合并/间接效应", "统计结果", "证据判断"],
        [
            ("TYMP-MTAP", "rho=0.0051", "95% CI -0.120至0.130；P=0.937；方向正、正、负、负", "未支持"),
            ("TYMP-MTAP-MTA", "-0.00112", "95% CI -0.0353至0.0360；P=0.947", "未支持"),
            ("串联至SPP1", "-0.00010", "95% CI -0.00438至0.00411；BH q=0.948", "未支持"),
            ("串联至SPP1程序", "0.000005", "95% CI -0.00298至0.00208；BH q=0.948", "未支持"),
            ("队列层面复现", "12项bootstrap", "全部95% CI跨0；独立验证29/29通过", "未支持"),
        ],
        [2440, 2060, 2860, 2000],
        decision_col=3,
    )

    doc.add_heading("6. 合并解释：哪些结论成立，哪些尚未成立", level=1)
    b.add_bullet(doc, "成立：TYMP与广泛髓系免疫代谢/调控背景相关；MTA反应代理与SPP1状态及免疫抑制程序在两个外部队列中存在弱但FDR支持的细胞级关联。", "成立：")
    b.add_bullet(doc, "成立：实测MTA与MTAP呈稳定负相关，符合MTAP参与MTA清除的预期。", "成立：")
    b.add_bullet(doc, "尚未成立：最终五队列患者级MTA转录代理分析没有任何队列通过FDR，且只有2/5队列达到检验门槛。", "尚未成立：")
    b.add_bullet(doc, "尚未成立：实测TYMP-MTA、TYMP-MTAP以及所有短中介和完整串联间接效应。", "尚未成立：")
    b.add_bullet(doc, "不能外推：转录代理的正相关不能覆盖实测代谢物的阴性结果，也不能据此声称TYMP经MTA选择性诱导SPP1巨噬细胞、治疗抵抗或临床预后。", "不能外推：")
    b.add_callout(doc, "客户可用结论", "现有多层公共数据支持TYMP相关的广泛髓系免疫代谢背景，也支持MTAP与MTA的直接代谢联系；但缺少可重复的TYMP-MTA和TYMP-MTAP上游证据，因而完整TYMP-MTA-SPP1机制仍是待验证假说。", fill="FCEEEE", label_color=b.RED)

    doc.add_heading("7. 建议的下一步验证", level=1)
    b.add_numbered(doc, "TYMP扰动：", "在同一细胞/共培养模型中敲低或过表达TYMP。")
    b.add_numbered(doc, "同步测量：", "检测MTAP RNA、蛋白和酶活，并用靶向LC-MS/MS测定细胞内及上清MTA。")
    b.add_numbered(doc, "细胞来源：", "同时记录巨噬细胞SPP1状态、总巨噬细胞量和肿瘤细胞TYMP，避免bulk组成混杂。")
    b.add_numbered(doc, "桥梁救援：", "加入MTAP敲低/回补及外源MTA处理，检验是否出现可逆的MTA和SPP1表型。")

    doc.add_heading("数据来源与证据边界", level=2)
    b.add_para(doc, "单细胞队列：GSE178481、GSE202374、GSE207493、GSE264586、GSE294109。配对代谢组：UnitedMet Zenodo record 11286535（https://zenodo.org/records/11286535）。MTA重塑巨噬细胞：Hansen等，Scientific Reports 2022，PMC8907307。SPP1与缺氧巨噬细胞：Wei等，Frontiers in Immunology 2021，PMC8523849。STAT3调控Spp1及产骨桥蛋白巨噬细胞：PMC11081920、PMC7726992。MTAP-MTA反应参考Reactome R-HSA-1237160；TYMP反应参考Reactome R-HSA-112265。转录代理、观察性关联、空间转移和计算扰动均不能单独证明因果。", size=8.5, color=b.GRAY, after=0, line=1.0)

    props = doc.core_properties
    props.title = "TYMP-MTA-SPP1机制链证据简报"
    props.subject = "五队列MTA代理、MTA-SPP1机制依据、实测MTA与MTAP串联验证"
    props.author = "Project 26210 Analysis Team"
    props.keywords = "TYMP; MTAP; MTA; SPP1; immune; single-cell; metabolomics"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
