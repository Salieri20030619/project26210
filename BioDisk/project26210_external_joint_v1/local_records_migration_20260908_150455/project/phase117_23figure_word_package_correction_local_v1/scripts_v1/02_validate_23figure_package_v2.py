from __future__ import annotations

import csv
import hashlib
import json
import zipfile
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw
from docx import Document
from docx.oxml.ns import qn


BASE = Path(__file__).resolve().parent
PACKAGE = BASE / "package_v2"
DOCX = PACKAGE / "项目26210_2.1-2.2图件与统计说明_23图完整版_v2.docx"
RENDER = BASE / "render_word_v1"
OUT = BASE / "validation_v2"
OUT.mkdir(exist_ok=True)


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def contact_sheets(pages):
    outputs = []
    for batch_idx in range(0, len(pages), 7):
        batch = pages[batch_idx:batch_idx + 7]
        thumb_w, gap, label_h, cols = 350, 18, 22, 2
        thumbs = []
        for page in batch:
            with Image.open(page) as img:
                rgb = img.convert("RGB")
                h = round(rgb.height * thumb_w / rgb.width)
                thumbs.append((page.name, rgb.resize((thumb_w, h), Image.Resampling.LANCZOS)))
        thumb_h = max(x.height for _, x in thumbs)
        rows = (len(thumbs) + cols - 1) // cols
        sheet = Image.new("RGB", (cols * thumb_w + (cols + 1) * gap, rows * (thumb_h + label_h) + (rows + 1) * gap), "#D8DCE2")
        draw = ImageDraw.Draw(sheet)
        for idx, (name, img) in enumerate(thumbs):
            row, col = divmod(idx, cols)
            x = gap + col * (thumb_w + gap)
            y = gap + row * (thumb_h + label_h + gap)
            draw.text((x + 3, y + 3), name, fill="#1F2937")
            sheet.paste(img, (x, y + label_h))
        out = OUT / f"contact_sheet_{batch_idx // 7 + 1}_v2.jpg"
        sheet.save(out, quality=80, optimize=True)
        outputs.append(out)
    return outputs


def main():
    checks = []
    def add(name, passed, detail=""):
        checks.append({"check": name, "passed": bool(passed), "detail": str(detail)})

    pngs = sorted((PACKAGE / "01_图件PNG_23张").rglob("*.png"))
    pdfs = sorted((PACKAGE / "02_矢量PDF备份_18张").rglob("*.pdf"))
    add("package_png_count_23", len(pngs) == 23, len(pngs))
    add("package_pdf_backup_count_18", len(pdfs) == 18, len(pdfs))
    add("all_figure_files_nonzero", all(p.stat().st_size > 0 for p in pngs + pdfs))
    add("no_tsv_json_in_user_package", not any(p.suffix.lower() in {".tsv", ".json"} for p in PACKAGE.rglob("*") if p.is_file()))

    doc = Document(DOCX)
    add("word_inline_image_count_23", len(doc.inline_shapes) == 23, len(doc.inline_shapes))
    alt = [shape._inline.docPr.get("descr", "") for shape in doc.inline_shapes]
    add("all_23_images_have_alt_text", len(alt) == 23 and all(len(x) >= 18 for x in alt))
    add("all_23_image_alt_texts_unique", len(set(alt)) == 23)
    text_parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            text_parts.extend(cell.text for cell in row.cells)
    full = "\n".join(text_parts)
    add("word_says_23_png_embedded", "Word内嵌23张PNG" in full)
    add("word_removes_tsv_dependency", "无需另行打开TSV或JSON" in full)
    add("word_preserves_causal_boundary", "TYMP促使肿瘤进入更恶性状态" in full and "不支持" in full)
    add("figure_ordinals_5_to_23_present", all(f"图 {i}/23" in full for i in range(5, 24)))

    row_split_ok = True
    for table in doc.tables:
        for row in table.rows:
            row_split_ok &= row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
    add("all_table_rows_prevent_page_split", row_split_ok)

    render_status = json.loads((RENDER / "render_status.json").read_text(encoding="utf-8"))
    pages = sorted(RENDER.glob("page-*.png"))
    add("word_render_pass", render_status.get("status") == "PASS", render_status)
    add("rendered_pages_28", len(pages) == 28, len(pages))
    nonblank = True
    safe = True
    page_audit = []
    for page in pages:
        with Image.open(page) as img:
            rgb = img.convert("RGB")
            diff = ImageChops.difference(rgb, Image.new("RGB", rgb.size, "white"))
            bbox = diff.getbbox()
            count = sum(1 for px in diff.convert("L").getdata() if px > 8)
            frac = count / (rgb.width * rgb.height)
            is_nonblank = bool(bbox and frac > 0.003)
            is_safe = bool(bbox and bbox[0] >= 10 and bbox[1] >= 8 and bbox[2] <= rgb.width - 10 and bbox[3] <= rgb.height - 8)
            nonblank &= is_nonblank
            safe &= is_safe
            page_audit.append({"page": page.name, "bbox": bbox, "nonwhite_fraction": frac, "nonblank": is_nonblank, "safe": is_safe})
    add("no_blank_pages", nonblank)
    add("safe_pixel_boundaries_all_pages", safe)
    sheets = contact_sheets(pages)
    add("four_contact_sheets_created", len(sheets) == 4 and all(p.stat().st_size > 0 for p in sheets))

    with zipfile.ZipFile(DOCX) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
    add("docx_media_objects_23", len(media) == 23, len(media))

    result = {
        "status": "PASS" if all(x["passed"] for x in checks) else "FAIL",
        "n_checks": len(checks),
        "n_failed": sum(not x["passed"] for x in checks),
        "docx_sha256": sha(DOCX),
        "png_files": len(pngs),
        "pdf_backup_files": len(pdfs),
        "word_inline_images": len(doc.inline_shapes),
        "rendered_pages": len(pages),
        "checks": checks,
        "page_audit": page_audit,
    }
    (OUT / "PACKAGE_VALIDATION_v2.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    with (OUT / "validation_checks_v2.tsv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["check", "passed", "detail"], delimiter="\t")
        w.writeheader(); w.writerows(checks)
    print(json.dumps({k: result[k] for k in ("status", "n_checks", "n_failed", "word_inline_images", "rendered_pages")}, ensure_ascii=False))
    return 0 if result["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
