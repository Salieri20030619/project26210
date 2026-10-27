from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
PACKAGE = BASE / "package_v2"
SOURCE_DOCX = Path(r"D:\project26210_external_joint_v1\phase116_sections_2_1_2_2_figure_package_local_v1\deliverable_v1\项目26210_2.1-2.2图件证据说明_v1.docx")
OUTPUT_DOCX = PACKAGE / "项目26210_2.1-2.2图件与统计说明_23图完整版_v2.docx"
FONT_CJK = "Microsoft YaHei"
FONT_LATIN = "Calibri"


DETAILS = {
    "Fig02B_CNV_burden_violin_v1": ("共识恶性与非恶性细胞CNV burden", "比较患者阈值标准化CNV burden；细胞数仅作描述，不能作为独立生物学重复。"),
    "factor_contrast_threshold_sensitivity_v1": ("各cNMF因子克隆差值阈值敏感性", "7个rank-7因子在多个CNV克隆阈值下的TYMP差值；F1并非所有阈值下最强。"),
    "formal_primary_clone_tymp_landscape_v1": ("正式CNV克隆TYMP/cNMF景观", "展示患者内主要CNV克隆、主导cNMF程序和TYMP；F1预先由TYMP证据选择，不是独立验证。"),
    "infercnv_GSE156632_RCC2": ("代表病例inferCNV染色体热图", "展示GSE156632 RCC2的CNV结构；单一病例图不代表跨患者普遍性。"),
    "method_cross_table_v1": ("inferCNV与CopyKAT交叉判定", "展示两种方法的细胞层组合计数；需保留大量method-discordant细胞。"),
    "patient_malignancy_composition_v1": ("患者层恶性状态组成", "逐患者展示恶性共识、单方法恶性、非恶性和未决比例；不能解释为TYMP效应。"),
    "selected_F1_patient_tymp_deltas_v1": ("F1高克隆与其他克隆TYMP差值", "主阈值0.25下中位差0.200，9/14正向，P=0.0210，但7因子BH q=0.0735。"),
    "selected_F1_threshold_patient_sensitivity_v1": ("F1-TYMP克隆阈值敏感性", "0.225、0.25、0.275三个阈值的显著性不同；不能挑选最有利阈值。"),
    "CellRank_fate_confidence_v1": ("CellRank命运置信度", "用于展示模型不确定性，不是生物效应量。"),
    "CellRank_fate_Cycle_v1": ("CellRank增殖状态命运概率", "不能由增殖状态概率推出TYMP驱动恶性进展。"),
    "CellRank_fate_Hypoxia_v1": ("CellRank缺氧状态命运概率", "概率表示模型化状态兼容性，不是细胞真实迁移证据。"),
    "CellRank_fate_TYMP_state_v1": ("CellRank TYMP状态命运概率", "状态标签包含TYMP信息，不能作为TYMP作用的独立验证。"),
    "CytoTRACE2_potency_v1": ("CytoTRACE2细胞潜能", "F1方向与Slingshot/scTour相反；不能声称统一晚期F1状态。"),
    "Monocle3_pseudotime_v1": ("Monocle3恶性细胞拟时序", "F1和TYMP均未获支持；拟时序不是实际肿瘤时间或克隆祖先关系。"),
    "scTour_pseudotime_v1": ("scTour恶性细胞拟时序", "F1/TYMP仅名义方向信号，轨迹家族FDR未通过。"),
    "Slingshot_pseudotime_v2": ("Slingshot恶性细胞拟时序（正式v2）", "F1呈正向候选但FDR未通过；旧v1未纳入用户包。"),
    "F1_high_clone_pseudobulk_DE_volcano_v1": ("F1高克隆配对伪批量差异表达", "14名患者、28个伪批量样本；基因差异不等于F1驱动或TYMP因果效应。"),
    "F1_high_clone_top_DE_patient_heatmap_v1": ("F1高克隆差异基因患者热图", "展示主要差异基因的患者方向；异质性需保留。"),
    "rank7_program_pathway_ORA_v1": ("rank-7程序通路富集", "这是cNMF top-loading基因的程序注释；F1高克隆差异基因通路ORA本身FDR阴性。"),
}


ORDER = [
    ("01_图件PNG_23张/2.1_CNV与亚克隆", "infercnv_GSE156632_RCC2"),
    ("01_图件PNG_23张/2.1_CNV与亚克隆", "Fig02B_CNV_burden_violin_v1"),
    ("01_图件PNG_23张/2.1_CNV与亚克隆", "method_cross_table_v1"),
    ("01_图件PNG_23张/2.1_CNV与亚克隆", "patient_malignancy_composition_v1"),
    ("01_图件PNG_23张/2.1_CNV与亚克隆", "formal_primary_clone_tymp_landscape_v1"),
    ("01_图件PNG_23张/2.1_CNV与亚克隆", "selected_F1_patient_tymp_deltas_v1"),
    ("01_图件PNG_23张/2.1_CNV与亚克隆", "selected_F1_threshold_patient_sensitivity_v1"),
    ("01_图件PNG_23张/2.1_CNV与亚克隆", "factor_contrast_threshold_sensitivity_v1"),
    ("01_图件PNG_23张/2.1_恶性轨迹", "Monocle3_pseudotime_v1"),
    ("01_图件PNG_23张/2.1_恶性轨迹", "Slingshot_pseudotime_v2"),
    ("01_图件PNG_23张/2.1_恶性轨迹", "CytoTRACE2_potency_v1"),
    ("01_图件PNG_23张/2.1_恶性轨迹", "scTour_pseudotime_v1"),
    ("01_图件PNG_23张/2.1_恶性轨迹", "CellRank_fate_TYMP_state_v1"),
    ("01_图件PNG_23张/2.1_恶性轨迹", "CellRank_fate_Hypoxia_v1"),
    ("01_图件PNG_23张/2.1_恶性轨迹", "CellRank_fate_Cycle_v1"),
    ("01_图件PNG_23张/2.1_恶性轨迹", "CellRank_fate_confidence_v1"),
    ("01_图件PNG_23张/2.2_cNMF", "F1_high_clone_pseudobulk_DE_volcano_v1"),
    ("01_图件PNG_23张/2.2_cNMF", "F1_high_clone_top_DE_patient_heatmap_v1"),
    ("01_图件PNG_23张/2.2_cNMF", "rank7_program_pathway_ORA_v1"),
]


def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT_LATIN
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def replace_paragraph_text(paragraph, new_text):
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.add_run(new_text)
    set_font(run, size=9.5, color="667085")


def shade_paragraph(paragraph, fill, accent):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "16")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), accent)
    p_bdr.append(left)
    p_pr.append(p_bdr)


def image_width(path: Path):
    with Image.open(path) as img:
        ratio = img.width / img.height
    max_w, max_h = 6.15, 7.15
    return Inches(min(max_w, max_h * ratio))


def set_alt(shape, text):
    shape._inline.docPr.set("descr", text)
    shape._inline.docPr.set("title", text[:80])


def add_figure_page(doc, ordinal, path, title, note):
    doc.add_page_break()
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"图 {ordinal}/23｜{title}")
    set_font(r)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    shape = p.add_run().add_picture(str(path), width=image_width(path))
    set_alt(shape, f"图{ordinal}/23：{title}。{note}")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.22
    shade_paragraph(p, "FFF4D6", "7A5A00")
    r = p.add_run("结果与边界  ")
    set_font(r, size=10, bold=True, color="7A5A00")
    r = p.add_run(note)
    set_font(r, size=10, color="222222")


def main():
    if not SOURCE_DOCX.is_file():
        raise FileNotFoundError(SOURCE_DOCX)
    pngs = sorted((PACKAGE / "01_图件PNG_23张").rglob("*.png"))
    pdfs = sorted((PACKAGE / "02_矢量PDF备份_18张").rglob("*.pdf"))
    if len(pngs) != 23 or len(pdfs) != 18:
        raise RuntimeError((len(pngs), len(pdfs)))

    doc = Document(SOURCE_DOCX)
    if len(doc.inline_shapes) != 4:
        raise RuntimeError(f"source doc expected 4 figures, found {len(doc.inline_shapes)}")

    for table in doc.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip() == "图件组成":
                row.cells[1].text = "Word内嵌23张PNG；交付包另附18份对应矢量PDF备份"
                for p in row.cells[1].paragraphs:
                    for run in p.runs:
                        set_font(run, size=9, color="222222")

    for p in doc.paragraphs:
        if "逐文件格式、字节数、SHA-256和原始路径" in p.text:
            replace_paragraph_text(p, "本Word已完整嵌入23张PNG；下表用于快速查阅，随后按任务顺序逐图展示。用户交付包不再散放TSV/JSON。")
        elif "关键统计表位于本包“05_关键统计表”目录" in p.text:
            replace_paragraph_text(p, "主管需要的关键统计已经直接写入本Word的证据总表、逐图说明和建议表述中；无需另行打开TSV或JSON。")

    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("7. 其余19张支撑图逐图展示")
    set_font(r)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("前文已展示4张主管汇报主图。本节补齐其余19张正式支撑图，使Word中的图片对象总数达到23；每张图均附结果与解释边界。")
    set_font(r, size=10.5, color="222222")

    for ordinal, (folder, stem) in enumerate(ORDER, start=5):
        path = PACKAGE / folder / f"{stem}.png"
        if not path.is_file():
            raise FileNotFoundError(path)
        title, note = DETAILS[stem]
        add_figure_page(doc, ordinal, path, title, note)

    props = doc.core_properties
    props.title = "项目26210 2.1–2.2图件与统计说明（23图完整版）"
    props.subject = "主管讨论用；Word内嵌23张PNG"
    props.comments = "Non-overwriting v2 correction: all 23 PNG figures embedded; no TSV/JSON required for reading."
    doc.save(OUTPUT_DOCX)

    readme = """项目26210 2.1–2.2图件证据包｜23图完整版v2

1. 请先打开：项目26210_2.1-2.2图件与统计说明_23图完整版_v2.docx
2. Word内已完整嵌入23张PNG，并逐图附结果和解释边界。
3. 01_图件PNG_23张：23张独立图片，按2.1/2.2分组。
4. 02_矢量PDF备份_18张：仅供放大、排版或投稿使用；不是Word图片数量。
5. 用户包不包含TSV/JSON。主管需要的统计已直接写入Word。

核心证据：TYMP RNA与cNMF rank-7 F1转录程序相关；现有CNV、克隆和轨迹证据不支持“TYMP已被证明促使肿瘤进入更恶性状态”的因果结论。
"""
    (PACKAGE / "README_先看这里_v2.txt").write_text(readme, encoding="utf-8-sig")
    print(json.dumps({"docx": str(OUTPUT_DOCX), "embedded_images": 23, "png_files": len(pngs), "pdf_backups": len(pdfs)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
