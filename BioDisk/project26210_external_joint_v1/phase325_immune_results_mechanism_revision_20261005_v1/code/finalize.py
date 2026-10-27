from pathlib import Path
import sys,json,hashlib,csv
from datetime import datetime,timezone,timedelta
P=Path(__file__).resolve().parents[1];R=P.parent
sys.path.insert(0,str(R/'phase322_integrated_immune_report_20261005_v1/code/vendor'))
from docx import Document
b=json.loads((P/'inputs/DOCUMENT_CONTENT.json').read_text());f=next((P/'delivery').glob('*.docx'));d=Document(f)
text='\n'.join(x.text for x in d.paragraphs)
tables=[x for x in b if x['type']=='table']
checks={'all_paragraphs_present':all(x['text'] in text for x in b if x['type'] in ['p','h','title']),
'tables_exact':len(d.tables)==7 and all([[c.text for c in row.cells] for row in t.rows]==[x['headers']]+[[str(v) for v in row] for row in x['rows']] for t,x in zip(d.tables,tables)),
'source_unchanged':all(hashlib.sha256(Path(x['path']).read_bytes()).hexdigest()==x['sha256'] for x in json.loads((P/'inputs/SOURCE_MANIFEST.json').read_text())),
'twelve_pages':len(list((P/'render/v1').glob('page-*.png')))==12,
'render_exit_zero':(P/'logs/render.exit').read_text().strip()=='0'}
assert all(checks.values()),checks
now=datetime.now(timezone(timedelta(hours=8))).isoformat(timespec='seconds')
result={'status':'PASS','pages':12,'tables':7,'checks':checks,'visual_review':'Pages1-4 byte-identical to previously inspected Phase323 PNGs; pages5-12 individually inspected, all clean. Continued table repeats headers.','methods':'Complete Phase323 methods retained exactly in the same document','new_results_interpretation_paragraphs':13,'new_models':False}
(P/'qc/COMPLETE.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
sha=hashlib.sha256(f.read_bytes()).hexdigest()
(P/'qc/DELIVERY_MANIFEST.json').write_text(json.dumps({'file':str(f),'sha256':sha,'bytes':f.stat().st_size},ensure_ascii=False,indent=2))
record=f"""## Phase325（2026-10-05）：按用户纠正保留详细方法，仅增强结果机制解释
用户明确“不是方法介绍有问题，是结果部分缺少机制解释，方法详细介绍是有必要的”。本指令纠正此前将方法压缩和移至附录的执行理解。
以Phase323 v3为底稿，研究背景和完整方法逐项不变，保留在同一Word；六项结果补充13段具体功能解释：CTSB/TGFB1与溶酶体/调节信号、IFNγ与MHC-I/MHC-II和吞噬、共同糖酵解背景与免疫差异、组织混合信号与细胞特异结果、SPP1与成纤维/基质及T细胞、临床评分的不同生物学对象。原结果段、七表、全部统计值和关键阴性保留，表置于对应结果叙述之后；原讨论保留。新增文献解释与本项目实测证据区分，不新增因果结论或模型。
此前Phase324压缩方法的v4仅为未交付中间草稿，不采用、不作为当前入口。当前交付为{f}，SHA256 {sha}。12页渲染检查通过；前4页与已审v3逐字节一致，5–12页逐页检查；内容和七表数值核验PASS。
当前仍为免疫章节稿，未合并或覆盖Phase273完整报告。后续写作保持详细方法，在结果中解释生物学含义与机制联系，不用压缩方法代替结果重写。记录仅服务器，无新分析或下载。
"""
for n in ['CODEX_HANDOFF.md','PROJECT_STATUS.md']:
 q=R/n;q.write_text(record+'\n'+q.read_text())
with (R/'DECISIONS.md').open('a') as h:h.write('\n'+record)
m=R/'RESULTS_MANIFEST.tsv'
with m.open() as h:header=next(csv.reader(h,delimiter='\t'))
assert header==['result_id','stage','artifact_type','absolute_path','generated_at','validation_status','validation_evidence','notes']
with m.open('a',newline='') as h:
 w=csv.writer(h,delimiter='\t',lineterminator='\n')
 for rid,kind,path in [('phase325_result_mechanism_docx','report',f),('phase325_validation','validation',P/'qc/COMPLETE.json'),('phase325_sources','provenance',P/'inputs/MECHANISM_SOURCE_NOTES.json')]:
  w.writerow([rid,'phase325_immune_results_mechanism_revision',kind,str(path),now,'PASS',str(P/'qc/COMPLETE.json'),'Detailed methods retained; results biological explanation expanded; no new statistics'])
print(json.dumps({'status':'PASS','sha256':sha,'bytes':f.stat().st_size,'pages':12}))
