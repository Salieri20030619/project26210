from pathlib import Path
import sys,json,runpy,subprocess,os
P=Path(__file__).resolve().parents[1];sys.path.insert(0,'/mnt/BioDisk/wangmingyuan/project26210_external_joint_v1/phase322_integrated_immune_report_20261005_v1/code/vendor')
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
q=subprocess.run(['/mnt/BioDisk/ludingyu/project26210/check_memory_gate.sh'],capture_output=True,text=True);(P/'logs/memory_gate.log').write_text(q.stdout+q.stderr);assert q.returncode==0

blocks=json.loads((P/'inputs/DOCUMENT_CONTENT.json').read_text())
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2);sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
styles=doc.styles
def font_style(style,size,bold=False):
 s=styles[style];s.font.name='DejaVu Sans';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0);rp=s.element.get_or_add_rPr();rf=rp.rFonts
 if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
 for k,v in [('ascii','DejaVu Sans'),('hAnsi','DejaVu Sans'),('eastAsia','Noto Serif CJK SC'),('cs','DejaVu Sans')]:rf.set(qn('w:'+k),v)
 for el in list(rp):
  if el.tag==qn('w:u'):rp.remove(el)
 return s
normal=font_style('Normal',11);normal.paragraph_format.line_spacing=1.30;normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.widow_control=True
font_style('Title',23,True);styles['Title'].paragraph_format.space_after=Pt(13)
for st,size in [('Heading 1',16),('Heading 2',13)]:
 s=font_style(st,size,True);s.paragraph_format.space_before=Pt(16 if st=='Heading 1' else 11);s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=True
font_style('Caption',10,True);styles['Caption'].paragraph_format.keep_with_next=True;styles['Caption'].paragraph_format.space_before=Pt(7);styles['Caption'].paragraph_format.space_after=Pt(4)
def set_run(run,size=None,bold=None):
 run.font.name='DejaVu Sans';run.font.color.rgb=RGBColor(0,0,0)
 if size is not None:run.font.size=Pt(size)
 if bold is not None:run.font.bold=bold
 rp=run._element.get_or_add_rPr();rf=rp.rFonts
 if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
 rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
def para(t,style='Normal'):
 z=doc.add_paragraph(t,style=style)
 if style=='Normal':z.alignment=WD_ALIGN_PARAGRAPH.LEFT
 for run in z.runs:set_run(run)
 return z
header=sec.header.paragraphs[0];header.text=''
for run in header.runs:set_run(run,8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='1';rr.append(tt);field.append(rr);run._r.addnext(field)
appendix=False
for block in blocks:
 if block['type']=='h' and block['text'].startswith('\u9644\u5f55A'):appendix=True
 if block['type']=='title':para(block['text'],'Title')
 elif block['type']=='h':
  z=para(block['text'],'Heading '+str(block['level']))
  if block['level']==1 and block['text'] in ['1 \u65b9\u6cd5','2 \u7ed3\u679c','3 \u8ba8\u8bba','\u53c2\u8003\u8d44\u6599','\u9644\u5f55A \u5b8c\u6574\u7edf\u8ba1\u4e0e\u7248\u672c\u7d22\u5f15']:z.paragraph_format.page_break_before=False
 elif block['type']=='p':
  z=para(block['text'])
  if appendix:
   z.paragraph_format.line_spacing=1.15
   for run in z.runs:set_run(run,9.5)
  if block['text'].startswith('['):
   z.paragraph_format.line_spacing=1.25
   for run in z.runs:set_run(run,9.5)
 elif block['type']=='table':
  para(block['title'],'Caption')
  headers=block['headers'];rows=block['rows'];widths=block['widths'] or [16.6/len(headers)]*len(headers);scale=16.6/sum(widths);widths=[x*scale for x in widths]
  tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
  for col,w in zip(tb.columns,widths):col.width=Cm(w)
  pr=tb._tbl.tblPr;borders=OxmlElement('w:tblBorders')
  for tag in ['top','left','bottom','right','insideH','insideV']:
   e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
  pr.append(borders);mar=OxmlElement('w:tblCellMar')
  for tag,val in [('top','75'),('bottom','75'),('left','85'),('right','85')]:
   e=OxmlElement('w:'+tag);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');mar.append(e)
  pr.append(mar)
  for i,vals in enumerate([headers]+rows):
   row=tb.rows[0] if i==0 else tb.add_row();trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
   if i==0:rep=OxmlElement('w:tblHeader');trpr.append(rep)
   for j,(cell,text) in enumerate(zip(row.cells,vals)):
    cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=str(text)
    if i==0:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF5');cell._tc.get_or_add_tcPr().append(sh)
    for pp in cell.paragraphs:
     pp.paragraph_format.line_spacing=1.13;pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.space_before=Pt(1);pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j<2 or len(headers)<=3 else WD_ALIGN_PARAGRAPH.CENTER
     for run in pp.runs:set_run(run,9 if appendix else 9.5,i==0)
  if block.get('source'):
   z=para(block['source']);z.paragraph_format.space_after=Pt(7);z.paragraph_format.line_spacing=1.15
   for run in z.runs:set_run(run,9)
doc.core_properties.title='TYMP\u4e0eSPP1\u76f8\u5173\u514d\u75ab\u7279\u5f81\u7efc\u5408\u62a5\u544a';doc.core_properties.subject='Project26210 immune evidence synthesis';doc.core_properties.author='26210\u9879\u76ee\u5206\u6790';doc.core_properties.keywords='TYMP SPP1 macrophage immunity'
for node in list(doc.styles.element.iter(qn('w:pBdr')))+list(doc.element.iter(qn('w:pBdr'))):
 node.getparent().remove(node)
out=P/'delivery/26210_免疫章节_结果机制解释_v5.docx';doc.save(out)
runtime={'dependency_loader':'not available in this session','authoring_runtime':sys.executable,'python_docx':'1.2.0 project-local vendor','lxml':'5.4.0 project-local vendor','renderer':'bundled documents render_docx.py copied unchanged','libreoffice':'/usr/bin/libreoffice on sgs10; headless project render, no desktop app','operation_marker':'edit marker succeeded once for this revision before authoring'};(P/'qc/DOCUMENT_RUNTIME.json').write_text(json.dumps(runtime,indent=2))
print('DOCX_CREATED',out,out.stat().st_size,flush=True)
