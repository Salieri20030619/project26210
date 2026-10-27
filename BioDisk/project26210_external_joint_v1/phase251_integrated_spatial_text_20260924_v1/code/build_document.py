from pathlib import Path
import os,sys,re,json,hashlib,datetime,shutil,zipfile,subprocess,csv
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).resolve().parents[1];ROOT=P.parent
def now():return datetime.datetime.now().astimezone().isoformat()
def sha(f):return hashlib.sha256(Path(f).read_bytes()).hexdigest()
gate=subprocess.run(['/mnt/BioDisk/ludingyu/project26210/check_memory_gate.sh'],capture_output=True,text=True)
(P/'logs/document_memory.log').write_text(gate.stdout+gate.stderr);assert gate.returncode==0
md=P/'text/26210_TYMP_MTA_SPP1_spatial_integrated_text_v7.md'
text=md.read_text()
assert not re.search(r'!\[|<img|图\s*\d+[A-Z]?|Figure\s*\d+',text)
assert not any(line.startswith('|') for line in text.splitlines())
assert '0.493' in text and '0.559' in text and '一正两负' in text
assert text.count('### 2.')==14
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.4);sec.right_margin=Cm(2.4)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.9)
def setfont(style,latin,east,size,bold=False):
 style.font.name=latin;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
 rr=style.element.get_or_add_rPr();fonts=rr.find(qn('w:rFonts'))
 if fonts is None:fonts=OxmlElement('w:rFonts');rr.insert(0,fonts)
 for key,val in [('ascii',latin),('hAnsi',latin),('eastAsia',east),('cs',latin)]:fonts.set(qn('w:'+key),val)
 for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+key),None)
setfont(doc.styles['Normal'],'Times New Roman','SimSun',11)
np=doc.styles['Normal'].paragraph_format;np.line_spacing=1.4;np.space_after=Pt(5);np.first_line_indent=Pt(22);np.widow_control=True
setfont(doc.styles['Title'],'Times New Roman','SimHei',18,True)
tp=doc.styles['Title'].paragraph_format;tp.space_after=Pt(15);tp.first_line_indent=Pt(0);tp.line_spacing=1.2;tp.keep_with_next=True
for name,size,before,after in [('Heading 1',14,14,7),('Heading 2',11.8,10,5)]:
 setfont(doc.styles[name],'Times New Roman','SimHei',size,True)
 pp=doc.styles[name].paragraph_format;pp.first_line_indent=Pt(0);pp.space_before=Pt(before);pp.space_after=Pt(after);pp.keep_with_next=True;pp.keep_together=True;pp.line_spacing=1.2
ref=False;paragraphs=[]
for line in text.splitlines():
 line=line.strip()
 if not line:continue
 if line.startswith('# '):
  p=doc.add_paragraph(line[2:].replace('与SPP1', '\n与SPP1'),style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 elif line.startswith('## '):
  title=line[3:]
  if title=='参考文献与数据来源':ref=True
  p=doc.add_paragraph(title,style='Heading 1')
  if ref:p.paragraph_format.page_break_before=True
 elif line.startswith('### '):
  p=doc.add_paragraph(line[4:],style='Heading 2')
 else:
  p=doc.add_paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  if ref:
   p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(6)
   for run in p.runs:run.font.size=Pt(9.5)
 paragraphs.append(p.text)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Pt(0)
run=footer.add_run();run.font.name='Times New Roman';run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
doc.core_properties.title=paragraphs[0];doc.core_properties.subject='TYMP相关代谢与SPP1巨噬状态及空间组织';doc.core_properties.author='';doc.core_properties.last_modified_by=''
for root_element in [doc.styles.element, doc._element]:
 for border in root_element.xpath('.//w:pBdr'):
  border.getparent().remove(border)
out=P/'delivery/26210_TYMP_MTA_SPP1_空间整合正文_v7.docx';doc.save(out)
shutil.copy2(md,P/'delivery/26210_TYMP_MTA_SPP1_空间整合正文_v7.md')
# Detailed statistics are separate from the manuscript body.
source=(P/'source/previous_report_v6.md').read_text();lines=source.splitlines();tables=[];i=0
while i<len(lines):
 if lines[i].startswith('|'):
  start=i
  while i<len(lines) and lines[i].startswith('|'):i+=1
  prev=start-1
  while prev>=0 and not lines[prev].strip():prev-=1
  caption=lines[prev].strip() if prev>=0 else '原报告表'
  after=i
  while after<len(lines) and not lines[after].strip():after+=1
  note=lines[after].strip() if after<len(lines) and not lines[after].startswith('#') and not lines[after].startswith('|') else ''
  tables.append((caption,lines[start:i],note))
 else:i+=1
supp=['# 原报告统计表完整保留','', '本文件保留用户指定v6报告的全部表格数值。原表编号和校正口径按原稿保存；新增空间统计在配套TSV文件中分别提供。','']
for n,(caption,rows,note) in enumerate(tables,1):
 supp+=['## '+caption,'']+rows+(['',note] if note else [])+['']
(P/'text/原报告统计表完整保留.md').write_text('\n'.join(supp))
new=[
('phase234_part4_replan_execution_20260921_v1','tables/A1_COHORT_SUMMARY.tsv'),
('phase234_part4_replan_execution_20260921_v1','tables/A2_A4_COHORT_SUMMARY.tsv'),
('phase234_part4_replan_execution_20260921_v1','tables/A1_PATIENT_EFFECTS.tsv'),
('phase234_part4_replan_execution_20260921_v1','tables/A2_A4_PATIENT_EFFECTS.tsv'),
('phase232_legacy_part4_significance_audit_20260921_v1','tables/OLD_VASCULAR_ALL_RELATIONS.tsv'),
('phase246_spatial_marker_structure_20260924_v1','tables/MARKER_STRUCTURE_COHORT_SUMMARY_WITH_EFFECTIVE_N.tsv'),
('phase246_spatial_marker_structure_20260924_v1','tables/COMPOSITE_RECHECK_COHORT_SUMMARY_WITH_EFFECTIVE_N.tsv'),
('phase242_SpaMTP_12tumor_analysis_20260924_v1','tables/PATIENT_LEVEL_PRIMARY_AND_SENSITIVITY.tsv'),
('phase244_spatial_evidence_wholesection_20260924_v1','tables/WHOLE_SECTION_ASSOCIATIONS_ALL.tsv'),
('phase248_independent_cases_mainline_20260924_v1','tables/EXTERNAL_CASE_MAINLINE_EFFECTS.tsv'),
('phase250_independent_sixpatient_mainline_20260924_v1','tables/TWO_INDEPENDENT_COHORTS_MAINLINE.tsv'),
('phase250_independent_sixpatient_mainline_20260924_v1','tables/ALL_SECTION_EFFECTS.tsv'),
('phase250_independent_sixpatient_mainline_20260924_v1','tables/PATIENT_DIRECTION_SUMMARY.tsv'),
('phase250_independent_sixpatient_mainline_20260924_v1','tables/PAIRED_TUMOR_ADJACENT_DESCRIPTIVE.tsv'),
('phase250_independent_sixpatient_mainline_20260924_v1','tables/CASE_IDENTITY_FROZEN.tsv')]
index=[]
for stage,relative in new:
 f=ROOT/stage/relative;assert f.is_file()
 index.append(dict(stage=stage,path=str(f),archive_path=stage+'/'+f.name,bytes=f.stat().st_size,sha256=sha(f)))
(P/'text/STATISTICAL_SOURCE_INDEX.json').write_text(json.dumps(index,ensure_ascii=False,indent=2))
readme='''# 正文配套统计补充
原报告统计表完整保留.md：用户指定v6报告的全部32张原表，包含前两部分完整效应、P、FDR及模型口径。
phase234：空间主队列的全部患者效应与队列汇总，包括阴性、反向和不可估计端点。
phase232与phase246：原复合区域及去共享成分、标记邻域和区域距离检查。
phase242与phase244：局部空间质谱结果及整片患者对应的全部统计。
phase248与phase250：独立病例的逐片、患者方向及肿瘤旁组织配对结果。
各目录保留原定义和原校正家族。文件名中的阶段号仅用于溯源，不是正文叙事顺序。
TYMP与MTA原表中历史固定效应FDR与后续随机效应rho/P的口径说明保持原稿，不把二者拼成同一模型的显著性证据。
CD8为补充；新独立病例不计算P/FDR。所有原始输入和既有正式结果只读保留。
'''
(P/'text/统计补充说明.md').write_text(readme)
zipout=P/'delivery/26210_空间整合正文_统计补充_v7.zip'
with zipfile.ZipFile(zipout,'w',zipfile.ZIP_DEFLATED) as z:
 z.write(P/'text/原报告统计表完整保留.md','原报告统计表完整保留.md');z.write(P/'text/统计补充说明.md','统计补充说明.md');z.write(P/'text/STATISTICAL_SOURCE_INDEX.json','STATISTICAL_SOURCE_INDEX.json')
 for item in index:z.write(item['path'],item['archive_path'])
# OOXML structural checks.
with zipfile.ZipFile(out) as z:
 xml=z.read('word/document.xml').decode()
 assert '<w:drawing' not in xml and '<w:pict' not in xml and '<w:tbl>' not in xml
 assert not any(n.startswith('word/media/') for n in z.namelist())
 assert '•' not in xml and not re.search(r'图\s*\d+[A-Z]?',xml)
assert len(tables)==32,len(tables)
manifest=dict(created_at=now(),source_pdf_path=r'D:\电脑管家迁移文件\xwechat_files\wxid_913cjq0lvmfd32_3e57\msg\file\2026-09\26210_TYMP_MTA_SPP1_four_figures_narrative_v6.pdf',source_pdf_sha256='6c6a3840fab35d794f02268c5ea86c146a10c684ea09fd564aec072d9718ce64',source_md_sha256=sha(P/'source/previous_report_v6.md'),body_characters=len(text),paragraphs=len(paragraphs),main_body_tables=0,images=0,original_tables_preserved=len(tables),spatial_source_tables=len(new),output_docx=str(out),output_docx_sha256=sha(out),statistics_zip_sha256=sha(zipout),runtime_note='Workspace dependency loader unavailable in this Windows session. Authoring used a versioned server task directory with isolated pinned Python dependencies and canonical renderer; no global packages or desktop LibreOffice modified.')
(P/'qc/BUILD_VALIDATION.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps(manifest,ensure_ascii=False),flush=True)

