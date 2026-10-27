from pathlib import Path
from docx import Document
import re,json,hashlib
P=Path(__file__).resolve().parents[1];R=P.parent;old=R/'phase252_spatial_append_only_20260925_v1'
src=old/'delivery/26210_空间转录组追加正文_待审_v1.docx'
doc=Document(src)
mapping={'5 空间研究设计与分析方法':'1 研究设计与分析方法','6 空间分析结果':'2 分析结果','7 空间结果的生物学解释':'3 讨论'}
def edit(s):
 if s in mapping:return mapping[s]
 m=re.match(r'([567])\.(\d+) (.*)',s)
 if not m:return s
 group=int(m[1]);n=int(m[2]);base={5:11,6:13,7:7}[group]
 return str(group-4)+'.'+str(n+base)+' '+m[3]
before=[p.text for p in doc.paragraphs]
changed=[]
for p in doc.paragraphs:
 if p.style.name in ['Heading 1','Heading 2']:
  new=edit(p.text)
  if new!=p.text:changed.append([p.text,new]);p.text=new
assert len(changed)==16,len(changed)
assert [p.text for p in doc.paragraphs if p.style.name not in ['Heading 1','Heading 2']]==[p.text for p in Document(src).paragraphs if p.style.name not in ['Heading 1','Heading 2']]
out=P/'delivery/26210_空间追加正文_沿用原编号_v2.docx';doc.save(out)
md=(old/'text/26210_空间转录组追加正文_待审_v1.md').read_text()
newlines=[]
for line in md.splitlines():
 m=re.match(r'(#{2,3} )(.*)',line)
 newlines.append(m[1]+edit(m[2]) if m else line)
newmd='\n'.join(newlines)+'\n'
(P/'text/26210_空间追加正文_沿用原编号_v2.md').write_text(newmd)
assert [x for x in md.splitlines() if not x.startswith('#')]==[x for x in newmd.splitlines() if not x.startswith('#')]
source=R/'phase208_client_wording_SPP1_narrative_20260918_v6/report/26210_TYMP_MTA_SPP1_four_figures_narrative_v6.pdf'
sha=hashlib.sha256(source.read_bytes()).hexdigest();assert sha=='6c6a3840fab35d794f02268c5ea86c146a10c684ea09fd564aec072d9718ce64'
(P/'qc/NUMBERING_CHECK.json').write_text(json.dumps({'changes':changed,'body_text_unchanged':True,'original_pdf_sha256':sha,'scope':'Only headings and numbering changed. Methods1.12-1.15 Results2.14-2.19 Discussion3.8-3.10; physical append retained.'},ensure_ascii=False,indent=2))
print(json.dumps(changed,ensure_ascii=False))
