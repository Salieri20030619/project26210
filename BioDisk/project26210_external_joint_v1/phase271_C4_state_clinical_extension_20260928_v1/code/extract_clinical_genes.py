from pathlib import Path
import csv,gzip,json,hashlib,datetime
import pandas as pd,numpy as np
from openpyxl import load_workbook
p=Path(__file__).resolve().parents[1];h=p.parent.parent;b=h/'project26210_external_joint_v1_phase64_large_immune_validation_v1'
fr=json.loads((p/'qc/SIGNATURE_FREEZE_BEFORE_CLINICAL.json').read_text());genes=fr['genes'];want=set(genes)
files={}
cm=b/'inputs_v1/CheckMate_Braun2020/processed_v1/checkmate_rna_expression_v1.tsv.gz'
with gzip.open(cm,'rt') as f:
 rd=csv.reader(f,delimiter='\t');head=next(rd);rows=[row for row in rd if row[0] in want]
files['CheckMate']=(head,rows)
w=load_workbook(b/'downloads_v1/JAVELIN_Renal101/Supplementary_Tables.xlsx',read_only=True,data_only=True)
it=w['S13_Gene_expression_TPM'].iter_rows(values_only=True);next(it);head=next(it);rows=[]
for row in it:
 if row[0] in want:rows.append(row)
w.close();files['JAVELIN']=(head,rows)
common=set.intersection(*[set(row[0] for row in rows) for _,rows in files.values()])
assert len(common)>=16
coverage=[]
for ds,(head,rows) in files.items():
 assert len({row[0] for row in rows})==len(rows)
 with (p/f'inputs/{ds}_C4_genes.tsv').open('w') as f:
  wr=csv.writer(f,delimiter='\t',lineterminator='\n');wr.writerow(head);wr.writerows(rows)
 for g in genes:coverage.append({'dataset':ds,'gene':g,'present':g in {row[0] for row in rows},'common_used':g in common})
pd.DataFrame(coverage).to_csv(p/'tables/CLINICAL_GENE_COVERAGE.tsv',sep='\t',index=False)
scores=[]
for ds,(head,rows) in files.items():
 x=pd.DataFrame([row[1:] for row in rows],index=[row[0] for row in rows],columns=head[1:]).astype(float)
 x=x.loc[[g for g in genes if g in common]]
 assert x.columns.is_unique and np.isfinite(x.values).all()
 if ds=='JAVELIN':assert (x.values>=0).all();x=np.log2(x+1)
 v=x.mean(0)
 scores.extend({'dataset':ds,'sample_id':s,'C4_state':float(v[s]),'gene_n':len(common)} for s in x.columns)
pd.DataFrame(scores).to_csv(p/'tables/C4_CLINICAL_SCORES.tsv',sep='\t',index=False)
frozen={'time':datetime.datetime.now().astimezone().isoformat(),'genes_used':[g for g in genes if g in common],'n':len(common),'scores_sha256':hashlib.sha256((p/'tables/C4_CLINICAL_SCORES.tsv').read_bytes()).hexdigest(),'clinical_outcomes_not_loaded':True}
(p/'qc/CLINICAL_SCORE_FREEZE.json').write_text(json.dumps(frozen,indent=2));print(frozen)
