
from pathlib import Path
import re,json,hashlib,subprocess
from copy import deepcopy
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).resolve().parents[1];R=P.parent
for d in ['text','delivery','qc','logs']:(P/d).mkdir(exist_ok=True)
src=R/'phase253_spatial_continue_numbering_20260925_v2'
base=R/'phase208_client_wording_SPP1_narrative_20260918_v6/report/26210_TYMP_MTA_SPP1_four_figures_narrative_v6.pdf'
fig=R/'phase261_spatial_ternary_integrated_20260926_v5/delivery/26210_Spatial_A4_GSEA_Ternary_v5.png'
mec=R/'phase257_TYMP_lysosome_cholesterol_validation_20260926_v1/editorial_v2/delivery/Phase257_section_mainline_v2.md'
text=next((src/'text').glob('*.md')).read_text()
mt=mec.read_text()
def mechanism(n):
 m=re.search(r'^## '+re.escape(n)+r' ([^\n]+)\n\n(.*?)(?=^## |\Z)',mt,re.M|re.S)
 return '### '+n+' '+m[1]+'\n\n'+m[2].strip()+'\n\n'
def append_to_section(n,addition):
 global text
 pattern=r'(^### '+re.escape(n)+r' [^\n]+\n\n)(.*?)(?=^#{1,3} |\Z)'
 m=re.search(pattern,text,re.M|re.S);assert m,n
 text=text[:m.end(2)]+addition+'\n\n'+text[m.end(2):]
append_to_section('1.12','GLMF1/2的优势细胞成分与TYMP表达见图5A，代谢背景及SPP1-state q05丰度见图5B、图5C。高值重叠图使用各切片第80百分位阈值；代谢背景与去SPP1状态程序的组合见图5D，SPP1-state与内皮丰度的组合见图5E。')
append_to_section('2.14','GLMF1/2的组织组成和TYMP表达、代谢背景、SPP1-state丰度分别见图5A–C；代谢与去SPP1状态程序的高值区域重叠见图5D。区域图以实测spot为单位呈现各特征的分布，对应上述患者层分析。')
append_to_section('2.15','GLMF两例中SPP1-state与内皮高丰度区域的空间分布及重叠见图5E。')
append_to_section('2.18','六名独立患者的代表肿瘤切片见图5H；两例旁组织与肿瘤的SPP1表达对照见图5I，使用统一色标呈现病例间不同变化方向。')
methods=mechanism('1.16')+'''### 1.17 SPP1相关区域的功能富集分析
分别在GSE175540的23名患者和TJ-RCC的12名患者中，以每片SPP1-state q05丰度前20%的区域与其余区域比较逐基因标准化表达均值；同患者多制备等权汇总。以患者效应均值除以其标准误形成基因排序，排除TYMP、SPP1、GPNMB、LGALS3、FABP5、CD9和LPL后，使用冻结KEGG基因集进行GSEA，BH校正在每个来源的完整KEGG主分析家族内计算。图5F展示溶酶体和胆固醇代谢的原始区域比较曲线；组成及深度调整、完整通路和患者层统计保留于分析附表。

### 1.18 恶性富集区域的三元相对距离
在GLMF1/2中，以冻结恶性成分比例不低于0.20的spot为源位置，以SPP1-state、内皮和T细胞各自q05丰度前20%且满足阳性阈值的spot为三类目标。沿实测组织第一圈邻接图计算最短加权路径，边长由物理坐标确定，保留组织分量。三个距离各除以其总和得到总和为100%的距离份额，并绘制三元坐标（图5G）。保留真实零距离；仅排除任一目标不可达或三个距离均为零的点，不作抖动、平滑或抽样。两例分别作描述性展示。

'''
text=text.replace('## 2 分析结果',methods+'## 2 分析结果',1)
res=mechanism('2.20')
res+='''SPP1-state高丰度区域的基因排序进一步显示溶酶体与胆固醇代谢富集（图5F）。GSE175540中，两通路NES分别为1.99和1.73，完整KEGG家族FDR分别为1.8×10⁻⁵和1.2×10⁻³；TJ-RCC中NES分别为2.25和1.78，FDR分别为2.5×10⁻⁵和1.1×10⁻³。这些曲线呈现SPP1相关区域的功能背景，与本节TYMP区室内及空间相关分析采用不同的比较口径。

### 2.21 两例恶性富集区域呈现不同的三元距离分布
GLMF1和GLMF2分别有1,282和3,113个恶性富集spot具有可定义的三元距离（图5G）。GLMF1的点云更集中于T细胞距离份额较大的区域，GLMF2的分布更分散，显示两个病例的相对空间距离构成不同。三元坐标表达各类距离占总距离的份额，某轴份额较大表示该类目标相对更远；这一区域分布描述不等同于绝对距离大小或细胞比例。

[[FIGURE5]]

## 图5图注
[[CAPTION5]]

'''
caption='''A，GLMF1/2的组织优势成分及TYMP RNA表达。13类参考状态归并为恶性、髓系、淋巴、血管、正常上皮和肥大细胞六类；TYMP用log1p(CP10k)显示，共用0–2色标。GLMF1/2分别为GSE210041的Patient1/Patient7，代表两位独立患者。

B，多胺及MTA相关转录背景：MAT2A、MAT2B、AMD1、SRM、SMS的片内基因Z分数均值，红色高、蓝色低，色标−1.25至1.25。C，SPP1-state巨噬细胞估计丰度的后验第5百分位（q05），绿色越深丰度越高，色标0–3。

D，代谢背景与去SPP1状态程序的高值区域重叠。状态程序为GPNMB、LGALS3、FABP5、CD9、LPL；两评分分别取片内前20%，橙色仅代谢高、青色仅状态高、紫色两者高、灰色其余。E，SPP1-state与内皮q05丰度的高值重叠，同用前20%阈值；蓝色仅SPP1-state高、珊瑚色仅内皮高、紫色二者高、灰色其余。

F，SPP1-state高丰度区域的KEGG功能富集。从左至右为GSE175540（23位患者）的溶酶体、胆固醇代谢，以及TJ-RCC（12位患者）的相应两通路。横轴为完整排序基因列表中的位置比例，左端对应高丰度区域正向差异较强的基因；纵轴为累积富集分数ES。竖线从零基线连接到真实通路成员所在位置的曲线，曲线与竖线共用0–0.7的蓝—灰—紫ES渐变。NES为归一化富集分数，q为各来源完整KEGG主分析家族的BH FDR。展示原始区域比较，分析排除七个目标／状态定义基因；图中患者n与GLMF两例分开。

G，GLMF1/2恶性富集spot到SPP1-state、内皮和T细胞高丰度区域的三元距离份额。各点为一个源spot，三个物理组织图距离分别除以总和；顶部、左下、右下分别对应SPP1-state、内皮和T细胞。越靠近某顶点表示该目标的距离份额越大，即相对更远，不表示细胞比例或更邻近。GLMF1由1,314个源spot保留1,282个（11个不可达、21个全零），GLMF2由3,133个保留3,113个（20个全零）；真实零距离保留。蓝色与珊瑚色区分病例，两例独立描述，spot数不作为生物学样本量。

H，E-MTAB-12767六位独立患者的代表肿瘤地图，显示规则及颜色同D。每位患者选择质控spot数最多的肿瘤阵列，平局按名称排序。I，HG1/HG3的配对旁组织与肿瘤SPP1 RNA，四图共用0–5的log1p(CP10k)色标。

组织图均沿用真实影像坐标和实测spot六边形，无表达插值或平滑；比例尺500 μm，同病例A–E保持相同方向与范围。超出连续色标上限的值饱和显示。B为代谢相关转录程序，C/E为估计丰度，D为程序重叠；空间SPP1-state沿用既有参考定义。
'''
res=res.replace('[[CAPTION5]]','[[CAPTION5]]\n\n'+caption)
text=text.replace('## 3 讨论',res+'## 3 讨论',1)
disc=mechanism('3.11')
for a,b in [(1,16),(2,17),(3,18)]:disc=disc.replace(f'[{a}]',f'[{b}]')
text=text.replace('## 新增参考文献与数据来源',disc+'## 新增参考文献与数据来源',1)
for a,b in [(1,16),(2,17),(3,18)]:
 line=next(v for v in mt.splitlines() if v.startswith(f'[{a}] '))
 text+='\n'+line.replace(f'[{a}]',f'[{b}]',1)+'\n'
(P/'text/26210_spatial_figure5_body_v3.md').write_text(text)
doc=Document(next((src/'delivery').glob('*.docx')))
for elem in list(doc._element.body):
 if elem.tag!=qn('w:sectPr'):doc._element.body.remove(elem)
doc.styles['Caption'].font.color.rgb=RGBColor(0,0,0);doc.styles['Caption'].font.bold=True
doc.styles['Caption'].font.name='Times New Roman';doc.styles['Caption'].font.size=Pt(11)
doc.styles['Caption'].paragraph_format.first_line_indent=Pt(0)
doc.styles['Caption'].paragraph_format.keep_with_next=True
refcounts={}
def field(p,instr,value):
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),instr)
 r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=value;r.append(t);f.append(r);p._p.append(f)
def addtext(p,line,refs=True):
 # REF field has cached value for reliable headless rendering.
 parts=re.split(r'(图5(?:[A-I](?:[–-][A-I])?)?)',line) if refs else [line]
 for part in parts:
  if refs and re.fullmatch(r'图5(?:[A-I](?:[–-][A-I])?)?',part):
   p.add_run('图');field(p,' REF fig5_spatial \\h ','5');p.add_run(part[2:]);refcounts[part]=refcounts.get(part,0)+1
  else:p.add_run(part)
def normal_section(sec):
 sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.4);sec.right_margin=Cm(2.4)
 sec.page_width=Cm(21);sec.page_height=Cm(29.7)
caption_mode=False;refs_mode=False
for line in text.splitlines():
 if not line.strip():continue
 if line=='[[FIGURE5]]':
  sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
  sec.top_margin=Cm(.5);sec.bottom_margin=Cm(.5);sec.left_margin=Cm(.6);sec.right_margin=Cm(.6)
  sec.footer_distance=Cm(.1)
  p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
  p.add_run().add_picture(str(fig),width=Cm(19.8))
  p._p.xpath('.//wp:docPr')[0].set('descr','图5 空间组织分布 区域功能富集与两例三元距离 A至I')
  sec=doc.add_section(WD_SECTION_START.NEW_PAGE);normal_section(sec);sec.footer_distance=Cm(.9)
  continue
 if line=='## 图5图注':caption_mode=True;continue
 if line=='[[CAPTION5]]':
  p=doc.add_paragraph(style='Caption');p.add_run('图')
  start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),'500');start.set(qn('w:name'),'fig5_spatial');p._p.append(start)
  field(p,' SEQ Figure \\r 5 ','5')
  end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),'500');p._p.append(end)
  p.add_run(' 代谢背景与SPP1相关状态的组织分布及区域功能特征')
  continue
 if line.startswith('## 3 '):caption_mode=False;doc.add_page_break()
 if line.startswith('# '):
  p=doc.add_paragraph(line[2:],style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 elif line.startswith('### '):p=doc.add_paragraph(line[4:],style='Heading 2')
 elif line.startswith('## '):
  p=doc.add_paragraph(line[3:],style='Heading 1')
  refs_mode=line=='## 新增参考文献与数据来源'
 else:
  p=doc.add_paragraph();addtext(p,line,refs=not caption_mode and not refs_mode)
  p.alignment=WD_ALIGN_PARAGRAPH.LEFT if refs_mode else WD_ALIGN_PARAGRAPH.JUSTIFY
  if caption_mode or refs_mode:
   p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(5)
   for run in p.runs:run.font.size=Pt(9.5)
for sec in list(doc.sections)[1:]:
 pg=sec._sectPr.find(qn('w:pgNumType'))
 if pg is not None:sec._sectPr.remove(pg)
doc.core_properties.title='空间研究正文 图5与功能机制补充'
doc.core_properties.subject='完整报告的空间续稿 图5及正文交叉引用'
doc.save(P/'delivery/26210_空间正文_含图5与交叉引用_v3.docx')
inputs=[base,next((src/'text').glob('*.md')),next((src/'delivery').glob('*.docx')),fig,mec]
(P/'qc/BUILD.json').write_text(json.dumps(dict(inputs=[dict(path=str(f),sha256=hashlib.sha256(f.read_bytes()).hexdigest()) for f in inputs],references=refcounts,figure_number=5,source_base_pages=31,existing_results_retained=True,new_statistical_tests=0),ensure_ascii=False,indent=2))
print('BUILD COMPLETE',refcounts)
