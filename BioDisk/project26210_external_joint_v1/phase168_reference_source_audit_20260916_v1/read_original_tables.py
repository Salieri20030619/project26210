from pathlib import Path
import openpyxl,json
O=Path('/mnt/BioDisk/wangmingyuan/project26210_external_joint_v1/phase168_reference_source_audit_20260916_v1')
result={}
for f in [O/'original_mmc2.xlsx',O/'original_mmc3.xlsx']:
 wb=openpyxl.load_workbook(f,read_only=True,data_only=True);out={}
 for ws in wb:
  rows=[[v for v in row] for row in ws.iter_rows(values_only=True)]
  out[ws.title]=rows
  print('FILE',f.name,'SHEET',ws.title,'ROWS',ws.max_row,'COLS',ws.max_column)
  for row in rows[:12]:print(row)
 result[f.name]=out
(O/'ORIGINAL_TABLES_EXTRACT.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str))
