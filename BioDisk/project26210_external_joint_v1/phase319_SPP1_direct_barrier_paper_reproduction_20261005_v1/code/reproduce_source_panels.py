from pathlib import Path
import openpyxl,numpy as np,pandas as pd,json,hashlib,subprocess,re
P=Path(__file__).resolve().parents[1]
q=subprocess.run(['/mnt/BioDisk/ludingyu/project26210/check_memory_gate.sh'],capture_output=True,text=True);(P/'logs/memory_gate.log').write_text(q.stdout+q.stderr);assert q.returncode==0
books={i:openpyxl.load_workbook(P/'downloads'/f'41588_2025_2380_MOESM{i}_ESM.xlsx',data_only=True) for i in [9,10,20]}
numeric=[];units=[];fields=[];issues=[];inventory=[]
def num(v):return isinstance(v,(int,float)) and not isinstance(v,bool) and np.isfinite(v)
def label(s,row,col):
 v=s.cell(row,col).value
 if v is None:
  for m in s.merged_cells.ranges:
   if m.min_row<=row<=m.max_row and m.min_col<=col<=m.max_col:v=s.cell(m.min_row,m.min_col).value;break
 return str(v).replace('\xa0',' ').strip() if v is not None else ''
def add(s,endpoint,group,uid,vals,kind,loc,batch='',source_col=None):
 vals=[float(x) for x in vals if num(x)]
 if not vals:return
 units.append(dict(panel=s.title,endpoint=endpoint,group=group,unit=uid,unit_kind=kind,n_observations=len(vals),mean=float(np.mean(vals)),median=float(np.median(vals)),source_cells=loc,batch=batch,source_col=source_col))
 for k,v in enumerate(vals):fields.append(dict(panel=s.title,endpoint=endpoint,group=group,unit=uid,unit_kind=kind,observation=k+1,value=v,source_col=source_col))
for bi,w in books.items():
 for s in w:
  for row in s:
   for c in row:
    if num(c.value):numeric.append(dict(workbook=bi,sheet=s.title,cell=c.coordinate,value=c.value))
  name=s.title;mode=''
  if name in ['Fig7D','Fig7O','Fig8B CD3','Fig8B Ki67 CD8','Fig8C CD3','Fig8C Ki67 CD8','Ext Fig10C','Fig8J','Fig8K']:
   mode='microscopy';start=2 if name in ['Fig8J','Fig8K'] else 3;seen={}
   for c in range(1,s.max_column+1):
    group=label(s,1,c);uid=label(s,2,c) if start==3 else f'column_{c}';vals=[s.cell(i,c).value for i in range(start,s.max_row+1)]
    if not any(num(v) for v in vals):continue
    kind='identified_mouse' if start==3 else 'anonymous_column_unverified';uid=f'{group}:{uid}' if uid.startswith('Mouse') else uid
    h=hashlib.sha256(repr(vals).encode()).hexdigest();key=(group,uid)
    if key in seen:
     identical=h==seen[key];issues.append(dict(panel=name,type='duplicate_identifier',group=group,unit=uid,identical=identical,column=c))
     if identical:continue
    seen[key]=h
    add(s,name,group,uid,vals,kind,f'{s.cell(start,c).coordinate}:{s.cell(s.max_row,c).coordinate}',source_col=c)
  elif name in ['Fig7F','Fig7G','Fig7H','Fig7P','Ext Fig10G']:
   mode='animal_rows'
   for row in range(2,s.max_row+1):
    endpoint=label(s,row,1)
    for c in range(2,s.max_column+1):add(s,endpoint,label(s,1,c),f'column_{c}',[s.cell(row,c).value],'caption_declared_mouse',s.cell(row,c).coordinate,source_col=c)
  elif name=='Fig7C':
   mode='animal_size_histogram'
   for c in range(2,s.max_column+1):
    vals=[s.cell(i,c).value for i in range(2,s.max_row+1)]
    if not all(num(v) for v in vals):continue
    total=sum(vals)
    for row,v in enumerate(vals,2):add(s,'fraction_'+label(s,row,1),label(s,1,c),f'column_{c}',[v/total if total else 0],'caption_declared_mouse',s.cell(row,c).coordinate,source_col=c)
  else:
   mode='flat_animal_or_assay'
   for c in range(1,s.max_column+1):
    if name in ['Fig8B Mets','Fig8C Mets']:
     if c==1:continue
     group='Recipient '+label(s,1,c)+' Donor '+label(s,2,c);start=3
    elif name=='Fig7N':group=' | '.join(label(s,k,c) for k in [1,2,3]);start=4
    else:group=label(s,1,c);start=2
    batch=''
    for row in range(start,s.max_row+1):
     v=s.cell(row,c).value
     if isinstance(v,str) and 'Experiment' in v:batch=v
     if name in ['Fig7J','Fig7L'] and isinstance(s.cell(row,1).value,str) and 'Experiment' in s.cell(row,1).value:batch=s.cell(row,1).value
     kind='assay_replicate_unresolved' if name=='Ext Fig10F' else 'caption_declared_mouse'
     add(s,name,group,f'{c}_{row}',[v],kind,s.cell(row,c).coordinate,batch=batch,source_col=c)
  inventory.append(dict(workbook=bi,panel=name,mode=mode,numeric_cells=sum(1 for z in numeric if z['workbook']==bi and z['sheet']==name)))
raw=pd.DataFrame(fields);u=pd.DataFrame(units)
# Collapse all fields within the same declared animal. Duplicate identifier conflicts remain explicit.
keys=['panel','endpoint','group','unit','unit_kind']
agg=raw.groupby(keys,dropna=False)['value'].agg(['count','mean','median']).reset_index().rename(columns={'count':'n_observations'})
summary=agg.groupby(['panel','endpoint','group','unit_kind']).agg(n_units=('unit','size'),mean=('mean','mean'),median=('mean','median'),sd=('mean','std'),min=('mean','min'),max=('mean','max')).reset_index();summary['sem']=summary.sd/np.sqrt(summary.n_units)
# Fixed within-panel contrasts. Every eligible group pair retained, no significance screening.
contrasts=[]
from itertools import combinations
for (panel,endpoint),df in summary.groupby(['panel','endpoint'],sort=False):
 for _,a in df.iterrows():
  for _,b in df.iterrows():
   if a['group']>=b['group']:continue
   contrasts.append(dict(panel=panel,endpoint=endpoint,group_a=a['group'],group_b=b['group'],n_a=a.n_units,n_b=b.n_units,mean_a=a['mean'],mean_b=b['mean'],delta_b_minus_a=b['mean']-a['mean'],ratio_b_over_a=b['mean']/a['mean'] if a['mean'] else np.nan,unit_kind=a.unit_kind,inference='descriptive; see separate qualified model table'))
for name,df in [('ALL_NUMERIC_SOURCE_CELLS',pd.DataFrame(numeric)),('ALL_OBSERVATIONS',raw),('SOURCE_COLUMN_UNITS',u),('BIOLOGICAL_OR_DECLARED_UNIT_SUMMARY',agg),('ALL_GROUP_SUMMARIES',summary),('ALL_DESCRIPTIVE_CONTRASTS',pd.DataFrame(contrasts)),('PANEL_INVENTORY',pd.DataFrame(inventory))]:df.to_csv(P/'tables'/f'{name}.tsv',sep='\t',index=False)
# Cross-panel identifiers that encode incompatible host/donor/treatment; source labels are preserved.
identified=u[u.unit_kind=='identified_mouse'];
for uid,df in identified[identified.panel.str.startswith('Fig8')].groupby('unit'):
 gs=set(df.group)
 if len(gs)>1:issues.append(dict(type='cross_panel_condition_conflict',unit=uid,conditions=sorted(gs),panels=sorted(set(df.panel))))
(P/'qc/SOURCE_IDENTITY_ISSUES.json').write_text(json.dumps(issues,indent=2))
print('PANELS',len(inventory),'NUMERIC_CELLS',len(numeric),'UNITS',len(agg),'GROUPS',len(summary),'CONTRASTS',len(contrasts),'ISSUES',len(issues))
print(summary[summary.panel.isin(['Fig7F','Fig7G','Ext Fig10C','Ext Fig10H','Fig8J','Fig8K','Ext Fig10F'])][['panel','endpoint','group','n_units','mean','unit_kind']].to_string(index=False))
