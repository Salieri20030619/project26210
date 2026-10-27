from pathlib import Path
import re,subprocess,hashlib,json
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).resolve().parents[1]; R=P.parent
old=R/'phase251_integrated_spatial_text_20260924_v1'
source=R/'phase208_client_wording_SPP1_narrative_20260918_v6/report/26210_TYMP_MTA_SPP1_four_figures_narrative_v6.pdf'
t=(old/'text/26210_TYMP_MTA_SPP1_spatial_integrated_text_v7.md').read_text()
def body(num):
 m=re.search(r'^### '+re.escape(num)+r' [^\n]+\n\n(.*?)(?=^#{1,3} |\Z)',t,re.M|re.S)
 assert m and len(m.group(1).strip())>100,num
 return m.group(1).strip()
parts=['# 空间转录组补充正文',
'前述分析已将TYMP相关代谢特征、MTA候选联系与SPP1富集巨噬状态连接起来。空间转录组进一步回答这些特征在真实肿瘤组织中如何分布，以及它们与血管和免疫区域有何联系。结果显示，代谢转录背景与SPP1相关巨噬及其状态程序在组织区域中共同出现，独立病例提供了方向性支持。本部分围绕这一组织层面的发现展开。',
'## 5 空间研究设计与分析方法']
for i,(n,title) in enumerate([('1.7','空间队列与细胞状态读出'),('1.8','空间区域与邻域比较'),('1.9','独立病例与空间质谱对照'),('1.10','患者层汇总与统计方法')],1):
 b=body(n).replace('[11-13]','[11-13]').replace('生物学样本量以独立患者为单位。单细胞结果沿用各冻结端点的Wilcoxon、相关及配对模型；空间队列先计算片内效应，再在患者层进行精确符号检验及既定伴随检验。','生物学样本量以独立患者为单位。空间队列先计算片内效应，再在患者层进行精确符号检验及既定伴随检验。')
 parts += ['### 5.'+str(i)+' '+title,b]
parts += ['## 6 空间分析结果']
for i,(n,title) in enumerate([('2.9','代谢背景与SPP1相关巨噬在组织区域中共同分布'),('2.10','SPP1相关区域伴随血管绝对丰度增加'),('2.11','SPP1相关区域伴随免疫调节和侵袭程序'),('2.12','SPP1相关区域保留T细胞相对分布差异'),('2.13','独立病例支持代谢背景与SPP1相关程序的空间对应'),('2.14','空间质谱提供精胺与SPP1表达的候选联系')],1):
 parts += ['### 6.'+str(i)+' '+title,body(n)]
parts += ['## 7 空间结果的生物学解释',
'### 7.1 空间共分布将代谢背景与巨噬表型连接到组织环境',
'前述单细胞分析确定了SPP1富集、脂质处理和溶酶体程序增强的巨噬状态。空间结果进一步表明，相关代谢背景与SPP1巨噬丰度及状态程序在组织区域中共同出现，使前述细胞状态发现获得组织分布层面的对应。独立病例中，代谢背景与去SPP1状态程序的同向变化支持以多基因表型理解这一联系。',
'巨噬绝对丰度、巨噬内部状态比例及SPP1单基因表达并不完全一致，表明组织中细胞数量和细胞状态是两个不同层次。现有结果支持代谢背景与巨噬相关区域的联系，并为研究局部代谢环境与巨噬功能适应提供方向。',
'单细胞C4与既有空间SPP1-state采用不同定义：前者来自本项目髓系状态发现，后者来自既有空间参考。当前空间结果提供SPP1相关巨噬背景的对应，尚未完成C4本身的独立空间定位。',
'### 7.2 血管与免疫特征描绘SPP1相关区域的功能背景',
'血管绝对丰度、免疫调节和侵袭程序的共同分布，说明SPP1相关区域包含值得进一步研究的微环境特征。受体表达对应、T细胞相对分布和CD8区域距离则分别提供信号兼容性、组织组成及局部结构的补充线索。',
'这些联系随细胞组成、来源和尺度变化。当前数据不支持将其概括为已确认的CD8排斥、CAF屏障或SPP1特异血管重塑。空间部分建立的是组织区域联系；这些区域特征与临床结局或治疗反应的对应仍由后续免疫临床板块另行回答。',
'### 7.3 空间证据的适用范围',
'多胺及MTA相关转录程序描述合成背景，不等同于组织内实测MTA浓度；空间共变化也不直接证明TYMP、MTA与SPP1之间的诱导顺序。独立病例的调整结果和两例配对对照保留了状态程序与单基因表达的差别，因此不能将分段对应写成已完成的串联机制。',
'空间映射、混合spot程序和表达区域距离不能替代独立染色、细胞分割及病理审核。既有病理边界仍缺少独立审核，C4空间混合恢复尚未达到全部预设门槛。部分区域关联在组成、深度或共同基因处理后减弱，这限定了细胞特异机制解释，但不否定已观察到的区域共变化。',
'空间质谱中的精胺及亚精胺仍为质量相容候选，缺少对应MS/MS或标准品确证。形态配准尚未取得独立解剖标志或作者变换确认，因此整片患者关联、局部共定位和代谢物身份分别解释。新增独立来源按实际患者数给出方向性结果，未将spot、区域或重复切片扩大为生物学样本量。',
'## 新增参考文献与数据来源']
refs=t.split('## 参考文献与数据来源\n',1)[1]
parts += [v.strip() for v in refs.split('\n\n') if re.match(r'\[(11|12|13|14|15)\]',v.strip())]
text='\n\n'.join(parts)+'\n'
md=P/'text/26210_空间转录组追加正文_待审_v1.md';md.write_text(text)
assert not re.search(r'!\[|图\s*\d|Figure\s*\d',text)
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.4);sec.right_margin=Cm(2.4)
sec.footer_distance=Cm(.9)
def font(name,size,east,bold=False):
 st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
 rr=st.element.get_or_add_rPr();fs=rr.find(qn('w:rFonts'))
 if fs is None:fs=OxmlElement('w:rFonts');rr.insert(0,fs)
 for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia',east)]:fs.set(qn('w:'+k),v)
 for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fs.attrib.pop(qn('w:'+k),None)
 return st.paragraph_format
pp=font('Normal',11,'SimSun');pp.line_spacing=1.4;pp.first_line_indent=Pt(22);pp.space_after=Pt(5);pp.widow_control=True
pp=font('Title',18,'SimHei',True);pp.line_spacing=1.2;pp.first_line_indent=Pt(0);pp.space_after=Pt(15);pp.keep_with_next=True
for name,size in [('Heading 1',14),('Heading 2',11.8)]:
 pp=font(name,size,'SimHei',True);pp.first_line_indent=Pt(0);pp.space_before=Pt(12);pp.space_after=Pt(6);pp.keep_with_next=True;pp.keep_together=True;pp.line_spacing=1.2
ref=False
for line in text.splitlines():
 if not line.strip():continue
 if line.startswith('# '):
  p=doc.add_paragraph(line[2:],style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 elif line.startswith('## '):
  p=doc.add_paragraph(line[3:],style='Heading 1')
  if line[3:]=='新增参考文献与数据来源':ref=True
 elif line.startswith('### '):p=doc.add_paragraph(line[4:],style='Heading 2')
 else:
  p=doc.add_paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ref else WD_ALIGN_PARAGRAPH.JUSTIFY
  if ref:
   p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.2
   for run in p.runs:run.font.size=Pt(9.5)
for root in [doc.styles.element,doc._element]:
 for b in root.xpath('.//w:pBdr'):b.getparent().remove(b)
pages=int(re.search(r'Pages:\s+(\d+)',subprocess.check_output(['pdfinfo',str(source)],text=True)).group(1))
pg=sec._sectPr.find(qn('w:pgNumType'))
if pg is None:pg=OxmlElement('w:pgNumType');sec._sectPr.append(pg)
pg.set(qn('w:start'),str(pages+1))
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
doc.core_properties.title='空间转录组补充正文';doc.core_properties.author='';doc.core_properties.subject='接续原报告的空间部分待审稿'
out=P/'delivery/26210_空间转录组追加正文_待审_v1.docx';doc.save(out)
(P/'qc/BUILD.json').write_text(json.dumps({'original_pdf':str(source),'original_pages':pages,'original_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'new_characters':len(text),'docx':str(out),'scope':'append only; original untouched; immune clinical pending; not an integrated rewrite'},ensure_ascii=False,indent=2))
print('Original pages',pages,'New chars',len(text))
