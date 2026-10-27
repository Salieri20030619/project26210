from pathlib import Path
import csv,hashlib,json,time
from openpyxl import load_workbook
h=Path('/mnt/BioDisk/wangmingyuan'); p=h/'project26210_external_joint_v1/phase270_SPP1_literature_clinical_revalidation_20260927_v1'
b=h/'project26210_external_joint_v1_phase64_large_immune_validation_v1'
x=b/'downloads_v1/JAVELIN_Renal101/Supplementary_Tables.xlsx'
w=load_workbook(x,read_only=True,data_only=True); sheet=w['S13_Gene_expression_TPM']
it=sheet.iter_rows(values_only=True); title=next(it); header=next(it)
genes={'SPP1','FN1','CLEC5A','OLR1','IL1RN','TYMP'}; rows=[]
for i,row in enumerate(it):
 if row[0] in genes: rows.append(row);print('found',row[0],i,flush=True)
assert len(rows)==6 and len({r[0] for r in rows})==6
with (p/'inputs/javelin_genes.tsv').open('w') as f:
 z=csv.writer(f,delimiter='\t',lineterminator='\n');z.writerow(header);z.writerows(rows)
print('JAVELIN DONE',len(header)-1,flush=True)
w.close()
x=b/'downloads_v1/CheckMate_Braun2020/Table_S4_RNA.xlsx'
w=load_workbook(x,read_only=True,data_only=True)
print('CM title',next(w['RNA_Expression'].iter_rows(values_only=True)),flush=True)
rows=[];it=w['RNA_Expression'].iter_rows(values_only=True);next(it);header=next(it)
for row in it:
 if row[0] in genes:rows.append(row)
with (p/'inputs/checkmate_genes_raw_workbook.tsv').open('w') as f:
 z=csv.writer(f,delimiter='\t',lineterminator='\n');z.writerow(header);z.writerows(rows)
print('CHECKMATE RAW DONE',len(rows),flush=True)
w.close()
(p/'qc/EXTRACTION_COMPLETE.json').write_text(json.dumps({'time':time.time(),'javelin_genes':6,'checkmate_genes':len(rows)}))
