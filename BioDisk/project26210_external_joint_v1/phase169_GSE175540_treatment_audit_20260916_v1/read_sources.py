from pathlib import Path
import openpyxl,json,xml.etree.ElementTree as ET,re
O=Path('/mnt/BioDisk/wangmingyuan/project26210_external_joint_v1/phase169_GSE175540_treatment_audit_20260916_v1')
wb=openpyxl.load_workbook(O/'original_mmc4.xlsx',read_only=True,data_only=True)
out={}
for ws in wb:
 rows=list(ws.values);out[ws.title]=rows
 print('SHEET',ws.title,ws.max_row,ws.max_column)
 for row in rows[:15]:print(row)
(O/'MMC4_EXTRACT.json').write_text(json.dumps(out,default=str,indent=2))
root=ET.parse(O/'reanalysis_fulltext.xml').getroot()
for el in root.iter():
 if el.tag in ['p','table-wrap']:
  txt=' '.join(el.itertext())
  if re.search('nivolumab|treatment|Table 1|immunological|immunotherapy|pretreat',txt,re.I):print(txt[:6500])
trial=json.loads((O/'bionikk_trial.json').read_text())
p=trial['protocolSection'];print('ELIGIBILITY',json.dumps(p.get('eligibilityModule',{})));print('DESCRIPTION',json.dumps(p.get('descriptionModule',{})))
