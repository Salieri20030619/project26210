from pathlib import Path
import json,re,subprocess
import anndata as ad,numpy as np,pandas as pd
from scipy import sparse
import openpyxl
P=Path(__file__).resolve().parents[1]
q=subprocess.run(['bash','/mnt/BioDisk/ludingyu/project26210/check_memory_gate.sh'],capture_output=True,text=True)
(P/'logs/public_inspect_memory.log').write_text(q.stdout+q.stderr);assert q.returncode==0
a=ad.read_h5ad(P/'inputs/xenium_rcc_core.h5ad')
targets=['TYMP','SPP1','CD68','CD163','C1QA','C1QB','C1QC','LST1','TYROBP','DCN','LUM','COL1A1','COL1A2','COL3A1','PDGFRA','FAP','POSTN','CTHRC1','RGS5','ACTA2','MYH11','CD3D','CD3E','CD8A','CD8B','EPCAM','PAX8','CA9','KRT8','KRT18','PECAM1','VWF']
vals=a.X.data if sparse.issparse(a.X) else a.X.ravel()
coverage=[dict(gene=g,present=g in a.var_names,positive_cells=int(np.asarray((a[:,g].X>0).sum()).item()) if g in a.var_names else None) for g in targets]
pd.DataFrame(coverage).to_csv(P/'tables/NICHEVERSE_PANEL_QUALIFICATION.tsv',sep='\t',index=False)
info=dict(shape=list(a.shape),obs=list(a.obs.columns),samples=a.obs.sample_id.astype(str).unique().tolist(),var_head=a.var_names[:5].tolist(),source=str(a.uns.get('source')),finite=bool(np.isfinite(vals).all()),integer=bool(np.equal(vals,np.floor(vals)).all()),nonnegative=bool((vals>=0).all()),coordinates_finite=bool(np.isfinite(a.obsm['spatial']).all()),targets=coverage)
(P/'qc/NICHEVERSE_INPUT_QC.json').write_text(json.dumps(info,indent=2));print(json.dumps(info))
sheets=[]
for f in sorted((P/'inputs').glob('*.xlsx')):
 w=openpyxl.load_workbook(f,read_only=True,data_only=True)
 for s in w:
  top=[list(r) for r in s.iter_rows(min_row=1,max_row=3,values_only=True)]
  rec=dict(file=f.name,sheet=s.title,rows=s.max_row,columns=s.max_column,top=top);sheets.append(rec)
 print(f.name,json.dumps([r for r in sheets if r['file']==f.name],default=str))
 w.close()
(P/'resources/XLSX_SCHEMA.json').write_text(json.dumps(sheets,indent=2,default=str))
j=json.loads((P/'resources/S-BIAD2470.json').read_text())
strings=[]
def walk(o):
 if isinstance(o,dict):
  for v in o.values():walk(v)
 elif isinstance(o,list):
  for v in o:walk(v)
 elif isinstance(o,str):strings.append(o)
walk(j)
urls=sorted({u for s in strings for u in re.findall(r'https?[^\s<>"]+',s)})
(P/'resources/BIOIMAGE_URLS.json').write_text(json.dumps(urls,indent=2))
print('BIOIMAGE_URLS',json.dumps(urls))
