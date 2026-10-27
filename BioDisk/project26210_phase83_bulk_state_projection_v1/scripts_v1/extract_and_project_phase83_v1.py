#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import gzip
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import nnls


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value).replace("\t", " ").replace("\r", " ").replace("\n", " ")


def collapse_rows(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.copy()
    frame.iloc[:, 0] = frame.iloc[:, 0].astype(str).str.upper()
    frame = frame.groupby(frame.columns[0], sort=True).mean(numeric_only=True)
    frame.index.name = "gene"
    return frame


def read_checkmate(path: Path, target: set[str]) -> pd.DataFrame:
    selected = []
    for chunk in pd.read_csv(path, sep="\t", chunksize=2000, compression="gzip"):
        genes = chunk.iloc[:, 0].astype(str).str.upper()
        hit = chunk.loc[genes.isin(target)].copy()
        if len(hit):
            selected.append(hit)
    if not selected:
        raise RuntimeError("No CheckMate markers found")
    return collapse_rows(pd.concat(selected, ignore_index=True))


def extract_javelin(workbook_path: Path, target: set[str]) -> pd.DataFrame:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["S13_Gene_expression_TPM"]
    header = [clean(cell.value) for cell in sheet[2]]
    rows = []
    for values in sheet.iter_rows(min_row=3, values_only=True):
        gene = clean(values[0]).upper()
        if gene in target:
            rows.append([gene, *[np.nan if value is None else float(value) for value in values[1:]]])
    workbook.close()
    if not rows:
        raise RuntimeError("No JAVELIN markers found")
    return collapse_rows(pd.DataFrame(rows, columns=header))


def project(expression: pd.DataFrame, reference: pd.DataFrame, dataset: str):
    genes = reference.index.tolist()
    matrix = expression.loc[genes]
    if not np.isfinite(matrix.to_numpy()).all():
        raise RuntimeError(f"{dataset}: nonfinite expression")
    ranked = matrix.rank(axis=0, method="average", pct=True)
    a = reference.to_numpy(dtype=float)
    weights = []
    diagnostics = []
    for sample_id in ranked.columns:
        b = ranked[sample_id].to_numpy(dtype=float)
        coefficient, residual = nnls(a, b)
        coefficient_sum = coefficient.sum()
        if not np.isfinite(coefficient_sum) or coefficient_sum <= 0:
            raise RuntimeError(f"{dataset}/{sample_id}: invalid NNLS coefficient sum")
        fitted = a @ coefficient
        relative_error = float(np.linalg.norm(b - fitted) / np.linalg.norm(b))
        coefficient = coefficient / coefficient_sum
        for state, value in zip(reference.columns, coefficient):
            weights.append({"dataset": dataset, "sample_id": sample_id, "state": state, "weight": value})
        diagnostics.append({
            "dataset": dataset,
            "sample_id": sample_id,
            "relative_reconstruction_error": relative_error,
            "active_states": int((coefficient > 1e-12).sum()),
            "raw_coefficient_sum": float(coefficient_sum),
            "weight_sum": float(coefficient.sum()),
            "minimum_weight": float(coefficient.min()),
        })
    return pd.DataFrame(weights), pd.DataFrame(diagnostics)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--reference", required=True, type=Path)
    parser.add_argument("--markers", required=True, type=Path)
    parser.add_argument("--checkmate-expression", required=True, type=Path)
    parser.add_argument("--javelin-workbook", required=True, type=Path)
    parser.add_argument("--allow-common-intersection", action="store_true")
    parser.add_argument("--amendment", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError(args.output)
    args.output.mkdir(parents=True)

    marker_table = pd.read_csv(args.markers, sep="\t")
    target = sorted(marker_table["gene"].astype(str).str.upper().unique())
    reference = pd.read_csv(args.reference, sep="\t", compression="gzip")
    reference.iloc[:, 0] = reference.iloc[:, 0].astype(str).str.upper()
    reference = reference.set_index(reference.columns[0])
    if target != sorted(reference.index.tolist()):
        raise RuntimeError("Marker/reference gene mismatch")
    if "TYMP" in target or "SPP1" in target:
        raise RuntimeError("Target leakage into marker union")

    checkmate = read_checkmate(args.checkmate_expression, set(target))
    javelin = extract_javelin(args.javelin_workbook, set(target))
    coverage_rows = []
    missing_by_dataset = {}
    for dataset, frame in (("CheckMate", checkmate), ("JAVELIN", javelin)):
        missing = sorted(set(target) - set(frame.index))
        missing_by_dataset[dataset] = missing
        for gene in target:
            coverage_rows.append({"dataset": dataset, "gene": gene, "present": gene in frame.index})
    pd.DataFrame(coverage_rows).to_csv(args.output / "phase83_marker_coverage_v1.tsv", sep="\t", index=False)

    status = {
        "status": "PASS" if not any(missing_by_dataset.values()) else "COVERAGE_FAIL",
        "frozen_marker_genes": len(target),
        "checkmate_samples": checkmate.shape[1],
        "javelin_samples": javelin.shape[1],
        "checkmate_present": checkmate.shape[0],
        "javelin_present": javelin.shape[0],
        "checkmate_missing": missing_by_dataset["CheckMate"],
        "javelin_missing": missing_by_dataset["JAVELIN"],
    }
    (args.output / "PHASE83_EXTRACTION_STATUS_v1.json").write_text(
        json.dumps(status, indent=2) + "\n", encoding="utf-8"
    )
    if status["status"] != "PASS" and not args.allow_common_intersection:
        print(json.dumps(status, indent=2))
        raise SystemExit(2)

    if status["status"] != "PASS":
        if args.amendment is None or not args.amendment.exists():
            raise RuntimeError("Common-intersection mode requires a frozen amendment file")
        common = sorted(set(target) & set(checkmate.index) & set(javelin.index))
        if len(common) < 0.90 * len(target):
            raise RuntimeError(f"Common marker coverage below 90 percent: {len(common)}/{len(target)}")
        dropped = sorted(set(target) - set(common))
        reference = reference.loc[common]
        norms = np.sqrt((reference ** 2).sum(axis=0))
        if (norms <= 0).any():
            raise RuntimeError("Zero reference column after platform intersection")
        reference = reference / norms
        checkmate = checkmate.loc[common]
        javelin = javelin.loc[common]
        status.update({
            "status": "PASS_WITH_PREREG_AMENDMENT",
            "common_marker_genes": len(common),
            "common_marker_fraction": len(common) / len(target),
            "dropped_from_common_projection": dropped,
            "amendment_file": str(args.amendment),
        })

    checkmate.to_csv(args.output / "checkmate_phase83_marker_expression_v1.tsv.gz", sep="\t", compression="gzip")
    javelin.to_csv(args.output / "javelin_phase83_marker_expression_v1.tsv.gz", sep="\t", compression="gzip")
    all_weights = []
    all_diagnostics = []
    for dataset, expression in (("CheckMate", checkmate), ("JAVELIN", javelin)):
        weights, diagnostics = project(expression, reference, dataset)
        all_weights.append(weights)
        all_diagnostics.append(diagnostics)
    weight_table = pd.concat(all_weights, ignore_index=True)
    diagnostic_table = pd.concat(all_diagnostics, ignore_index=True)
    weight_table.to_csv(args.output / "phase83_projection_weights_v1.tsv.gz", sep="\t", index=False, compression="gzip")
    diagnostic_table.to_csv(args.output / "phase83_projection_diagnostics_v1.tsv", sep="\t", index=False)
    status.update({
        "projected_samples": int(diagnostic_table.shape[0]),
        "states": int(reference.shape[1]),
        "max_weight_sum_error": float((diagnostic_table["weight_sum"] - 1).abs().max()),
        "minimum_weight": float(weight_table["weight"].min()),
        "median_relative_reconstruction_error": float(diagnostic_table["relative_reconstruction_error"].median()),
    })
    (args.output / "PHASE83_EXTRACTION_STATUS_v1.json").write_text(
        json.dumps(status, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(status, indent=2))


if __name__ == "__main__":
    main()
