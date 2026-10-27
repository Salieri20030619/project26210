#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import gzip
import json
from pathlib import Path

from openpyxl import load_workbook


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value).replace("\t", " ").replace("\r", " ").replace("\n", " ")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--members", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError(args.output)
    args.output.mkdir(parents=True)

    with args.members.open(encoding="utf-8", newline="") as handle:
        members = list(csv.DictReader(handle, delimiter="\t"))
    target = sorted({row["gene"].upper() for row in members})
    if len(target) != 31 or "TYMP" not in target:
        raise RuntimeError(f"Unexpected target set: n={len(target)}")

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    sheet = workbook["S13_Gene_expression_TPM"]
    header = [clean(cell.value) for cell in sheet[2]]
    found = {}
    output_path = args.output / "javelin_phase81_expression_v1.tsv.gz"
    with gzip.open(output_path, "wt", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(header)
        for values in sheet.iter_rows(min_row=3, values_only=True):
            gene = clean(values[0]).upper()
            if gene in target:
                writer.writerow([gene, *[clean(value) for value in values[1:]]])
                found[gene] = found.get(gene, 0) + 1
    workbook.close()

    with (args.output / "javelin_phase81_gene_coverage_v1.tsv").open(
        "w", encoding="utf-8", newline=""
    ) as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["gene", "present", "row_count"])
        for gene in target:
            writer.writerow([gene, str(gene in found).upper(), found.get(gene, 0)])
    status = {
        "status": "PASS" if len(header) - 1 == 726 and len(found) == len(target) else "FAIL",
        "samples": len(header) - 1,
        "target_genes": len(target),
        "present_genes": len(found),
        "duplicated_rows": sum(value - 1 for value in found.values()),
    }
    (args.output / "javelin_phase81_extraction_status_v1.json").write_text(
        json.dumps(status, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(status, indent=2))
    if status["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
